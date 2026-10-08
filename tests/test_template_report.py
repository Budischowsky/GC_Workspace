"""Writing a Template report: the red-band workbook, the Word document, the payload and the batch document."""
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from gcws.report import table as TB
from gcws.report import template as TP
from gcws.report import template_report as TR


def _table(**changes):
    from test_report_table import _data, _row
    rows = [_row(12.0, "Bisphenol A", "80-05-7", (0.2,), sml=0.05, ref="10/2011", footnote="Group SML"),
            _row(14.0, "Hydrocarbons C20", "", (0.05,)), _row(15.0, "Hydrocarbons C22", "", (0.07,))]
    tpl = TP.preset("NIAS")
    for k, v in changes.items():
        tpl["extras"][k] = v
    return TB.build(_data(rows), tpl)


def _job(tmp_path, table, **kw):
    target = tmp_path / "out" / "26016605_NIAS_Template_Report.xlsx"
    return SimpleNamespace(kind="template", table=table, target=target, word=target.with_suffix(".docx"),
                           audit=[SimpleNamespace(timestamp="2026-10-08 10:00", run="run0", action="Edit",
                                                  before="a", after="b", detail="")],
                           notes=["no blank"], record_seen=False, sample_key="26016605", **kw)


def test_the_workbook_has_the_nias_look(tmp_path):
    res = TR.generate(_job(tmp_path, _table()))
    assert res.target.exists() and res.word is not None and res.word.exists()
    wb = load_workbook(res.target)
    assert wb.sheetnames[0] == "Result" and TR.PAYLOAD_SHEET in wb.sheetnames
    assert wb[TR.PAYLOAD_SHEET].sheet_state == "veryHidden"
    sh = wb["Result"]
    assert sh["A1"].value == "GC-MS/FID – NIAS-Screening –" and sh["A1"].fill.fgColor.rgb.endswith("BA0C2F")
    assert sh["A2"].value == "PA 26.007, single determination"
    assert sh["A3"].value == "Sample:" and sh["B3"].value == "26016605"
    head = [sh.cell(5, c).value for c in range(1, 9)]
    assert head == ["RT (min)", "Name", "CAS-No.", "% match", "Conc. mg/dm²", "Conc. mg/kg", "SML (mg/kg)", "Ref."]
    assert sh["B6"].value == "Bisphenol A" and sh["F6"].number_format == "0.000" and sh["E6"].number_format == "0.0000"
    assert sh["F6"].font.bold                                    # above its SML
    assert "(a)" in str(sh["H6"].value)                          # footnote marker on Ref.
    assert str(sh["B7"].value).startswith("Sum of hydrocarbons") and sh["B7"].font.bold
    assert sh["F7"].value == pytest.approx(0.12)
    texts = [sh.cell(r, c).value for r in range(8, sh.max_row + 1) for c in (1, 2)]
    assert "Group SML" in texts and any(str(t).startswith("**") for t in texts if t)
    assert sh.page_setup.orientation == "landscape" and sh.freeze_panes == "A6"
    assert res.rows == 1 and "no blank" in res.warnings
    assert res.combined[0]["name"] == "Bisphenol A"


def test_the_word_document_carries_links_markers_and_the_audit_page(tmp_path):
    from docx import Document
    res = TR.generate(_job(tmp_path, _table()))
    doc = Document(res.word)
    text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "GC-MS/FID – NIAS-Screening –" in text and "Bisphenol A" in text and "0.200" in text
    xml = doc.element.xml
    assert "w:hyperlink" in xml and "80-05-7" in xml
    assert any(r.font.superscript for p in doc.paragraphs for r in p.runs)
    assert any(p.text == "Audit Trail" for p in doc.paragraphs)
    assert doc.sections[0].orientation == 1                      # landscape


def test_payload_round_trip_and_batch_word(tmp_path):
    job = _job(tmp_path, _table(word=False))
    res = TR.generate(job)
    assert res.word is None and TR.is_template_workbook(res.target)
    table, audit = TR.read_payload(res.target)
    assert table == job.table and audit[0][2] == "Edit"
    other = tmp_path / "other.xlsx"
    TR.generate(SimpleNamespace(**dict(vars(job), target=other, word=other.with_suffix(".docx"))))
    out = TR.combined_word([res.target, other], tmp_path / "batch.docx")
    from docx import Document
    doc = Document(out)
    assert len(doc.tables) >= 2 and len(doc.sections) == 2
    assert not TR.is_template_workbook(Path(__file__))


def test_an_empty_template_is_refused(tmp_path):
    with pytest.raises(ValueError, match="no column"):
        TR.generate(_job(tmp_path, TB.ReportTable()))


def test_the_preview_html():
    from gcws.report.template_html import to_html
    html = to_html(_table())
    assert "NIAS-Screening" in html and "<b>0.200</b>" in html and "<sup>(a)</sup>" in html
    assert "Sum of hydrocarbons" in html
    assert "more rows" in to_html(_table(), max_rows=0)


def test_batch_word_takes_template_reports_and_refuses_a_mix(tmp_path):
    from openpyxl import Workbook
    from gcws.automation import batch as BA
    job = _job(tmp_path, _table())
    TR.generate(job)
    out = BA.combined_word([job.target], tmp_path / "b" / "batch.docx")
    assert out.exists()
    other = tmp_path / "nias.xlsx"
    Workbook().save(other)
    with pytest.raises(ValueError, match="cannot be combined"):
        BA.combined_word([job.target, other], tmp_path / "mixed.docx")

"""Focused checks for copying a result sheet into a Word table."""

from docx import Document
from docx.table import Table
from openpyxl import Workbook

from gcws.report.legacy_api import main_script


def test_word_report_reuses_row_cells_after_merging(tmp_path, monkeypatch):
    book = Workbook()
    sheet = book.active
    sheet.title = "NIAS Result"
    sheet.merge_cells("A1:H1")
    sheet["A1"] = "Example result"
    for column in range(1, 9):
        sheet.cell(5, column, f"Column {column}")
    for row in range(6, 46):
        for column in range(1, 9):
            sheet.cell(row, column, f"{row}:{column}")
    source = tmp_path / "input.xlsx"
    target = tmp_path / "output.docx"
    book.save(source)

    original_cell = Table.cell
    large_table_calls = 0

    def counted_cell(self, row_index, col_index):
        nonlocal large_table_calls
        if len(self._tbl.tr_lst) >= 40:
            large_table_calls += 1
        return original_cell(self, row_index, col_index)

    monkeypatch.setattr(Table, "cell", counted_cell)
    main_script().create_combined_word([source], target)
    calls_during_generation = large_table_calls

    table = Document(target).tables[0]
    assert table.cell(0, 0).text == "Example result"
    assert table.cell(5, 0).text == "6:1"
    assert table.cell(44, 7).text == "45:8"
    # Merge handling still uses Table.cell; body cells are row-cached.
    assert calls_during_generation <= 3

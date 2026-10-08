"""End-to-end reports from our own integration (identifications from LIB as test data)."""
from pathlib import Path

import pytest

from gcws.core.ident import Identification


def _ws_with(samples, prefixes, qapp):
    from gcws.core.model import FID
    from gcws.io.run_loader import load_run
    from gcws.signal.delay import estimate_delay
    from gcws.ui.workspace import Workspace
    from gcws.integration.engine import integrate
    from gcws.quant.nias_bridge import engine
    ws = Workspace()
    eng = engine()
    for prefix in prefixes:
        d = next(samples.glob(prefix + "*.D"))
        run = load_run(d)
        st = ws.add_run(run, {FID: integrate(run.fid, ws.methods.get(ws.methods.default_name(FID)))},
                        delay=estimate_delay(run.fid, run.signal("TIC")))
        if not (d / "RESULTS.CSV").exists():
            continue
        # test oracle: library hits of the old ChemStation LIB, bound by RT
        _tic, _fid, pbm = eng.parse_results(str(d / "RESULTS.CSV"))
        lib = eng.parse_library_report(str(d / "LIB"))
        res = st.results[FID]
        for p in pbm:
            hits = lib.get(p.peak) or p.hits
            if not hits:
                continue
            peak = min(res.peaks, key=lambda x: abs(x.apex_rt - (p.rt + st.delay_value)))
            if abs(peak.apex_rt - (p.rt + st.delay_value)) > 0.02:
                continue
            st.ident_set(FID).set(Identification(
                apex_rt=peak.apex_rt, name=hits[0].name, cas=hits[0].cas, score=hits[0].quality,
                hits=[{"name": h.name, "cas": h.cas, "score": h.quality} for h in hits], source="LIB (test)"))
    ws.recompute_quant()
    return ws


@pytest.fixture(scope="module")
def qapp():
    # a QApplication, not a QCoreApplication: the GUI tests run in the same
    # process and cannot create widgets on a core application
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


@pytest.mark.parametrize("kind", ["nias", "fingerprint", "total_extraction"])
def test_duplicate_reports(samples, qapp, tmp_path, kind):
    from openpyxl import load_workbook
    from gcws.report.service import ReportJob, SUFFIXES, generate
    from gcws.quant.service import cas_lookup
    from gcws import paths
    ws = _ws_with(samples, ["06_", "07_", "08_", "11_"], qapp)
    ids = [s.id for s in ws.states() if s.role == "sample"]
    assert len(ids) == 2
    for rid in ids:
        sample = ws.nias_sample(rid)
        assert sample is not None and sample.mean_factor
        assert ws.runs[rid].blanks and ws.runs[rid].blanks_istd
    from gcws.quant.nias_bridge import make_settings
    settings = make_settings(ws.quant.get("settings"))
    target = tmp_path / f"26016606{SUFFIXES[kind]}.xlsx"
    job = ReportJob(kind=kind, samples=[ws.nias_sample(r) for r in ids], names=[ws.runs[r].name for r in ids],
                    settings=settings, target=target, word=target.with_suffix(".docx"),
                    cas_path=paths.RESOURCES / "CASINFO.xlsx",
                    migration={"analyst": "Test", "migration_cell": "cell", "occupancy": "1",
                               "simulant": "EtOH 95 %", "temperature": "60 C", "duration": "10 d",
                               "cell_area_dm2": 0.51, "occupancy_factor": 1, "volume_ml": 100,
                               "ov_ratio": 6.0},
                    batch_target=tmp_path / "26016606_Doppelbestimmung.xlsx" if kind == "nias" else None,
                    keep_middle=tmp_path / "middle.xlsx")
    res = generate(job)
    assert res.target.exists() and res.rows > 0
    assert res.word is not None and res.word.exists()
    wb = load_workbook(tmp_path / "middle.xlsx")
    if kind == "nias":
        assert "Bestimmung_2" in wb.sheetnames
        assert (tmp_path / "26016606_Doppelbestimmung.xlsx").exists()


def test_triplicate_nias_report(samples, qapp, tmp_path):
    from openpyxl import load_workbook
    from gcws.report.service import ReportJob, generate
    from gcws import paths
    from gcws.quant.nias_bridge import make_settings
    ws = _ws_with(samples, ["06_", "07_", "08_", "09_", "11_"], qapp)
    ids = [s.id for s in ws.states() if s.role == "sample"]
    assert len(ids) == 3
    target = tmp_path / "triple_NIAS_Report.xlsx"
    job = ReportJob(kind="nias", samples=[ws.nias_sample(r) for r in ids], names=[ws.runs[r].name for r in ids],
                    settings=make_settings(ws.quant.get("settings")), target=target,
                    word=target.with_suffix(".docx"), cas_path=paths.RESOURCES / "CASINFO.xlsx",
                    policy="majority", keep_middle=tmp_path / "middle.xlsx")
    res = generate(job)
    assert res.target.exists()
    wb = load_workbook(tmp_path / "middle.xlsx")
    ws3 = wb["Doppelbestimmung"]
    headers = [c.value for c in ws3[1]]
    assert "Concentration 3 [mg/kg]" in headers and "SD [mg/kg]" in headers
    assert headers[3] == "mg/kg (mean)"
    means = [r[3] for r in ws3.iter_rows(min_row=2, values_only=True) if isinstance(r[3], (int, float))]
    assert means, "N-fold mean must be written as values"
    assert "Bestimmung_3" in wb.sheetnames


def _small_pdf(path, pages=2):
    from PySide6.QtGui import QPageSize, QPainter, QPdfWriter
    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.A5))
    painter = QPainter(writer)
    for i in range(pages):
        if i:
            writer.newPage()
        painter.drawText(200, 400, f"page {i + 1}")
    painter.end()
    return path


def test_render_pages_gives_qimages(qapp, tmp_path):
    pytest.importorskip("pypdfium2")
    from PySide6.QtGui import QImage
    from gcws.report.service import render_pages
    pages = render_pages(_small_pdf(tmp_path / "x.pdf"), scale=0.5)
    assert len(pages) == 2
    assert all(isinstance(p, QImage) and not p.isNull() and p.width() > 50 for p in pages)
    # the page is white with some dark text pixels
    img = pages[0]
    assert img.pixelColor(2, 2).lightness() > 240


def test_preview_dialog_shows_warnings(qapp, tmp_path):
    from gcws.report.service import render_pages
    from gcws.ui.dialogs.report_preview import ReportPreview
    pages = render_pages(_small_pdf(tmp_path / "x.pdf", 1), scale=0.5)
    dlg = ReportPreview("t", {"xlsx": tmp_path / "r.xlsx"}, pages, warnings=["Register not updated: x"])
    assert "Register not updated" in dlg.notes.text() and not dlg.notes.isHidden()
    assert dlg.col.count() == 1
    empty = ReportPreview("t", {"xlsx": tmp_path / "r.xlsx"}, [], warnings=[])
    assert "No page images" in empty.notes.text()


MIGRATION = {"analyst": "Test", "migration_cell": "cell", "occupancy": "1", "simulant": "EtOH 95 %",
             "temperature": "60 C", "duration": "10 d", "cell_area_dm2": 0.51, "occupancy_factor": 1,
             "volume_ml": 100, "ov_ratio": 6.0}


def _report_table(path):
    """(RT, name, mg/kg) of the rows of a NIAS report."""
    from openpyxl import load_workbook
    sh = load_workbook(path, data_only=True).worksheets[0]
    out, started = [], False
    for r in sh.iter_rows(values_only=True):
        if r and r[0] == "RT (min)":
            started = True
            continue
        if started and isinstance(r[0], (int, float)):
            out.append((float(r[0]), r[1], r[5]))
    return out


def test_duplicate_edits_reach_nias_report(samples, qapp, tmp_path):
    from gcws.report.service import ReportJob, combined_rows, generate
    from gcws.quant.nias_bridge import make_settings
    from gcws.quant.replicates import combine, engine_peaks
    from gcws import paths
    ws = _ws_with(samples, ["06_", "07_", "08_", "11_"], qapp)
    ids = [s.id for s in ws.states() if s.role == "sample"]
    settings = make_settings(ws.quant.get("settings"))

    def job(name, edits):
        t = tmp_path / f"{name}_NIAS_Report.xlsx"
        return ReportJob(kind="nias", samples=[ws.nias_sample(r) for r in ids], names=[ws.runs[r].name for r in ids],
                         settings=settings, target=t, word=t.with_suffix(".docx"),
                         cas_path=paths.RESOURCES / "CASINFO.xlsx", migration=MIGRATION, record_seen=False,
                         edits=edits)
    raw = combine([engine_peaks(ws.nias_sample(r)) for r in ids], 0.035)
    plain = generate(job("plain", {}))
    base = _report_table(plain.target)
    # no edits: every reported value is the merged mean, and no artefact is reported
    for rt, _name, mgkg in base:
        src = min(raw, key=lambda r: abs(r["rt"] - rt))
        assert not str(src.get("status", "")).startswith("Artefact")
        assert mgkg == pytest.approx(src["mean"], rel=1e-3, abs=1e-3)
    confirmed = [r for r in raw if r.get("mean") and r["mean"] > 0.05 and r["c1"] and r["c2"]]
    drop, change = confirmed[0], confirmed[1]
    artefact = max((r for r in raw if str(r.get("status", "")).startswith("Artefact")),
                   key=lambda r: max(v for v in (r.get("c1"), r.get("c2")) if v is not None))
    edits = {f"{drop['rt']:.3f}": {"rt": drop["rt"], "report": False},
             f"{change['rt']:.3f}": {"rt": change["rt"], "mean": 1.234},
             f"{artefact['rt']:.3f}": {"rt": artefact["rt"], "report": True, "c1": 0.5, "c2": 0.5}}
    j = job("edited", edits)
    assert len(combined_rows(j)) == len(combined_rows(job("n", {}))) + 0   # one dropped, one kept artefact
    res = generate(j)
    rep = _report_table(res.target)
    near = lambda rt: [r for r in rep if abs(r[0] - rt) < 0.02]
    assert not near(drop["rt"])
    assert near(change["rt"]) and near(change["rt"])[0][2] == pytest.approx(1.234, abs=1e-3)
    assert near(artefact["rt"]) and near(artefact["rt"])[0][2] == pytest.approx(0.5, abs=1e-3)
    assert any("analyst" in w for w in res.warnings)
    # everything else is as without edits
    others = [r for r in base if all(abs(r[0] - x["rt"]) > 0.02 for x in (drop, change, artefact))]
    assert all(near(r[0]) and near(r[0])[0][2] == pytest.approx(r[2]) for r in others)


def test_control_characters_in_names_do_not_break_the_report(samples, qapp, tmp_path):
    """A Shimadzu library name with a NUL byte ("\x00n-Tridecan-1-ol") made openpyxl refuse the
    workbook ("... cannot be used in worksheets")."""
    from openpyxl import load_workbook
    from gcws.core.text import excel_safe
    from gcws.report.service import ReportJob, SUFFIXES, generate
    from gcws.quant.nias_bridge import make_settings
    from gcws import paths
    ident = Identification(apex_rt=1.0, name="\x00n-Tridecan-1-ol\x02", cas="112-70-9\x1f",
                           hits=[{"name": "\x08Tridecanol", "cas": "112-70-9"}])
    assert (ident.name, ident.cas, ident.hits[0]["name"]) == ("n-Tridecan-1-ol", "112-70-9", "Tridecanol")
    assert Identification.from_dict({"apex_rt": 1.0, "name": "a\x0bb"}).name == "ab"
    assert excel_safe("x\x00y") == "xy" and excel_safe(3.5) == 3.5
    ws = _ws_with(samples, ["06_", "07_", "08_", "11_"], qapp)
    ids = [s.id for s in ws.states() if s.role == "sample"]
    samples_ = [ws.nias_sample(r) for r in ids]
    named = [r for r in samples_[0].rows if r.name]
    assert named
    named[0].name = "\x00n-Tridecan-1-ol"            # a project saved before names were cleaned
    target = tmp_path / f"26016606{SUFFIXES['nias']}.xlsx"
    job = ReportJob(kind="nias", samples=samples_, names=[ws.runs[r].name for r in ids],
                    settings=make_settings(ws.quant.get("settings")), target=target,
                    word=target.with_suffix(".docx"), cas_path=paths.RESOURCES / "CASINFO.xlsx",
                    keep_middle=tmp_path / "middle.xlsx")
    res = generate(job)
    assert res.target.exists()
    assert named[0].name == "n-Tridecan-1-ol"
    texts = [c.value for sh in load_workbook(tmp_path / "middle.xlsx").worksheets for row in sh.iter_rows()
             for c in row if isinstance(c.value, str)]
    assert not any("\x00" in t for t in texts)


def test_nias_report_without_reportable_substance(samples, qapp, tmp_path):
    """Nothing above the reporting limit: the report is still made, with one line instead of the
    substances (it used to stop: 'Spalte D enthält nur unberechnete Formeln')."""
    from openpyxl import load_workbook
    from gcws.report.service import ReportJob, combined_rows, generate
    from gcws.quant.nias_bridge import make_settings
    from gcws.quant.replicates import combine, engine_peaks
    from gcws import paths
    ws = _ws_with(samples, ["06_", "07_", "08_", "11_"], qapp)
    ids = [s.id for s in ws.states() if s.role == "sample"]
    raw = combine([engine_peaks(ws.nias_sample(r)) for r in ids], 0.035)
    target = tmp_path / "empty_NIAS_Report.xlsx"
    job = ReportJob(kind="nias", samples=[ws.nias_sample(r) for r in ids], names=[ws.runs[r].name for r in ids],
                    settings=make_settings(ws.quant.get("settings")), target=target,
                    word=target.with_suffix(".docx"), cas_path=paths.RESOURCES / "CASINFO.xlsx",
                    migration=MIGRATION, record_seen=False,
                    edits={f"{r['rt']:.3f}": {"rt": r["rt"], "report": False} for r in raw})
    assert combined_rows(job) == []
    res = generate(job)
    assert res.rows == 0 and res.word is not None and res.word.exists()
    texts = [c for row in load_workbook(target).worksheets[0].iter_rows(values_only=True) for c in row if c]
    assert "No substance above 10 ppb detected." in texts
    assert not any(isinstance(t, str) and t.startswith("=") for t in texts)

"""Quantification report (extraction method): RT, Name, CAS, Qual, Conc. 1, Conc. 2 for single and double
determination."""
from types import SimpleNamespace

import pytest

from gcws.report import service as RS
from gcws.report import simple as S


def _job(n, overrides=None):
    return SimpleNamespace(samples=[object()] * n, overrides=overrides or {})


def _row(c1, c2=None, mean=None, **kw):
    sources = [{"score": 91}, {"quality": 85}] if c2 is not None else [{"score": 91}]
    cs = [c1, c2] if c2 is not None else [c1]
    return dict(rt=10.0, name="Substance", cas="1-2-3", c1=c1, c2=c2, cs=cs,
                mean=mean if mean is not None else sum(cs) / len(cs), sources=sources, **kw)


# A: 1 g, B: 2 g, both in 10 mL; Conc. 1 mg/kg, Conc. 2 µg/L (µg/L = mg/kg × g × 1000 ÷ mL ÷ ... per run)
RATIOS = {0: 100.0, 1: 200.0}
ratio = lambda k, unit: 1.0 if unit == "mg/kg" else RATIOS[k]


def test_single_determination_reports_its_own_values():
    rows, dets = S.report_rows(_job(1), [_row(10.0)], ["mg/kg", "µg/L"], ratio)
    assert rows == [[10.0, "Substance", "1-2-3", 91.0, 10.0, 1000.0]]
    assert dets[0][3:6] == [10.0, 1000.0, ""]


def test_double_determination_means_each_unit_with_its_own_amount():
    rows, _ = S.report_rows(_job(2), [_row(10.0, 6.0)], ["mg/kg", "µg/L"], ratio)
    rt, name, cas, qual, c1, c2 = rows[0]
    assert qual == 85.0                                   # the lower match of the two determinations
    assert c1 == pytest.approx(8.0)
    assert c2 == pytest.approx((10.0 * 100 + 6.0 * 200) / 2)      # the mean of the converted A and B


def test_edits_and_a_dismissed_outlier_reach_both_units():
    rows, dets = S.report_rows(_job(2, {0: {"c2": 8.0, "mean": 9.0}}), [_row(10.0, 6.0)], ["mg/kg", "µg/L"], ratio)
    assert rows[0][4] == pytest.approx(9.0)
    assert rows[0][5] == pytest.approx((10.0 * 100 + 8.0 * 200) / 2)
    rows, dets = S.report_rows(_job(2, {0: {"mean": 6.0, "dismissed": 1}}), [_row(10.0, 6.0)],
                               ["mg/kg", "µg/L"], ratio)
    assert rows[0][4] == pytest.approx(6.0) and rows[0][5] == pytest.approx(1200.0)
    assert dets[0][5] == "dismissed"


def test_reporting_limit_and_rows_without_result():
    combined = [_row(10.0), dict(_row(0.001)), dict(_row(5.0), mean=None)]
    combined[2]["mean"] = None
    rows, _ = S.report_rows(_job(1), combined, ["mg/kg", "µg/L"], ratio, limit=0.01)
    assert [r[4] for r in rows] == [10.0]


def test_the_report_kind_follows_the_mode():
    assert RS.default_kind({"mode": "extraction"}) == "quant"
    assert RS.default_kind({"mode": "hs_screening"}) == "hs_screening"
    assert RS.default_kind({"mode": "nias_mgkg", "detector": "TIC"}) == "nias"       # not by the detector
    assert RS.kind_fits("quant", {"mode": "extraction"}) and not RS.kind_fits("nias", {"mode": "extraction"})
    assert RS.kind_fits("fingerprint", {"mode": "area_pct"}) and not RS.kind_fits("quant", {"mode": "nias_mgkg"})


@pytest.fixture(scope="module")
def pair(samples, qapp):
    from gcws.io.run_loader import load_run
    from gcws.ui.workspace import Workspace
    from tests.conftest import run_dir
    ws = Workspace()
    ids = [ws.add_run(load_run(run_dir(p))).id for p in ("07_", "11_")]
    ws.process_runs(ids)
    q = dict(ws.quant)
    q["mode"] = "extraction"
    q["method"] = {"name": "Test", "sample_type": "solid", "amount": 1.0, "extract_volume_ml": 10.0,
                   "units": ["mg/kg", "µg/L"], "reporting_limit": 0.5}
    q["method_samples"] = {ids[1]: {"amount": 2.0}}
    ws.quant = q
    ws.recompute_quant()
    return ws, ids


def test_quantification_report_single_and_double(pair, tmp_path):
    from openpyxl import load_workbook
    from docx import Document
    from gcws.report import assemble as AS
    ws, ids = pair
    single = {"id": "", "name": "A", "members": ids[:1], "policy": "all"}
    job = AS.build_job(ws, "quant", single, tmp_path / "single_Quantification_Report.xlsx", record_seen=False)
    res = RS.generate(job)
    sh = load_workbook(res.target)["Result"]
    assert [sh.cell(5, c).value for c in range(1, 7)] == ["RT (min)", "Name", "CAS-No.", "Qual",
                                                         "Conc. 1 [mg/kg]", "Conc. 2 [µg/L]"]
    assert "Single determination" in sh["A2"].value and "(1 g)" in sh["B4"].value
    assert res.rows > 0 and len(res.combined) >= res.rows
    first = [sh.cell(6, c).value for c in range(1, 7)]
    assert first[4] >= 0.5                                # reporting limit in Conc. 1
    assert first[5] == pytest.approx(first[4] * 1.0 / 10.0 * 1000.0, rel=1e-9)   # mg/kg × 1 g ÷ 10 mL → µg/L
    doc = Document(res.word)
    assert doc.tables[0].cell(4, 5).text == "Conc. 2 [µg/L]"

    double = {"id": "g", "name": "AB", "members": ids, "policy": "all", "extraction_edits:mg/kg": {}}
    job = AS.build_job(ws, "quant", double, tmp_path / "double_Quantification_Report.xlsx", record_seen=False)
    res = RS.generate(job)
    wb = load_workbook(res.target)
    sh = wb["Result"]
    assert "Double determination" in sh["A2"].value and "A 1 g, B 2 g" in sh["B4"].value
    det = wb["Determinations"]
    head = [c.value for c in det[1]]
    assert head[3:9] == ["A [mg/kg]", "A [µg/L]", "A outlier", "B [mg/kg]", "B [µg/L]", "B outlier"]
    calc = wb["Calculation"]
    assert [calc.cell(r, 5).value for r in (2, 3)] == [1.0, 2.0]          # B's own sample mass
    for r in range(2, det.max_row + 1):
        a1, a2, b1, b2 = (det.cell(r, c).value for c in (4, 5, 7, 8))
        if None not in (a1, b1):
            assert a2 == pytest.approx(a1 * 100.0) and b2 == pytest.approx(b1 * 200.0)
            break
    else:
        pytest.fail("no substance found in both determinations")

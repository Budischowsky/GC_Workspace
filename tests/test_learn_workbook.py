"""gcws.learn: parsing human NIAS evaluation workbooks (synthetic workbooks, no client data)."""
from learn_fixtures import FID_HEADER, FID_PEAKS, TIC_HEADER, TIC_PEAKS


def _rohdaten_rows():
    return [("[contents]",), ("count=2",), (r"[INT TIC: 05_X_A.D\data.ms]",), ("Time=", "x"), tuple(TIC_HEADER),
            *map(tuple, TIC_PEAKS), (r"[INT 05_X_A.D\FID1A.ch]",), tuple(FID_HEADER), *map(tuple, FID_PEAKS)]


def test_parse_rohdaten_sections():
    from gcws.learn.model import RawPeak
    from gcws.learn.workbook import parse_rohdaten
    peaks = parse_rohdaten(_rohdaten_rows())
    tic = [p for p in peaks if p.signal == "TIC"]
    fid = [p for p in peaks if p.signal == "FID"]
    assert len(tic) == 2 and len(fid) == len(FID_PEAKS)
    assert fid[3] == RawPeak(signal="FID", number=4, rt=7.07, area=119966715, height=3238076, start=6.938,
                             end=7.313, peak_type="M", manual=True)
    t = tic[1]
    assert (t.number, t.rt, t.area, t.start, t.end, t.peak_type, t.manual) == (2, 8.203, 40534814, None, None,
                                                                                "VV", False)


def test_parse_rohdaten_keeps_column_gaps():
    """Empty cells inside a row (e.g. a blank peak type) must not shift the columns."""
    from gcws.learn.workbook import parse_rohdaten
    rows = [(r"[INT 05_X_A.D\FID1A.ch]",), tuple(FID_HEADER), ("1=", 1, 5.0, 4.9, 5.1, None, 100, 2000, 0, 0)]
    (p,) = parse_rohdaten(rows)
    assert (p.peak_type, p.height, p.area, p.manual) == ("", 100, 2000, False)


def test_normalise_cas():
    from gcws.learn.model import normalise_cas
    assert normalise_cas("000097-88-1") == "97-88-1"
    assert normalise_cas(" 128-37-0 ") == "128-37-0"
    assert normalise_cas("-") == ""
    assert normalise_cas(None) == ""


def _cells(path, sheet="Auswertung", data_only=True):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=data_only)
    ws = wb[sheet]
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}


def test_header_v2(tmp_path):
    from learn_fixtures import make_workbook
    from gcws.learn.workbook import detect_template, parse_header
    cells = _cells(make_workbook(tmp_path / "w.xlsx"))
    h = parse_header(cells)
    assert detect_template(cells) == "v2"
    assert (h.simulant, h.temperature, h.duration, h.volume, h.sv_ratio) == ("EtOH 95%", 40, "10 d", 10, 6)
    assert (h.gc_method, h.inj_volume, h.syn_summary, h.syn_id, h.evaluator, h.operator) == (
        "PA 26.009", 30, "25011662", "2501675", "BlM", "BlM")
    assert h.sample_name == "05_X_A" and h.data_file.endswith("05_X_A.D")
    assert [i.name for i in h.istd] == ["C17", "BBP", "DnNP", "DBP-d4"]
    assert [i.area for i in h.istd] == [9871653, 7095674, 7904784, 742832]
    assert h.istd[0].conc == 0.84 and h.istd[3].conc == 0.0752
    assert abs(h.istd_mean_area - 8290703.67) < 0.01
    assert len(h.alkanes) == 21
    c10 = next(a for a in h.alkanes if a.name == "C10")
    assert (round(c10.rt, 3), c10.ri) == (6.203, 1000)
    assert h.conc_units == ["mg/dm²", "mg/kg"]


def test_header_v1_no_alkanes(tmp_path):
    from learn_fixtures import make_workbook
    from gcws.learn.workbook import detect_template, parse_header
    cells = _cells(make_workbook(tmp_path / "w.xlsx", template="v1"))
    assert parse_header(cells).alkanes == []
    assert detect_template(cells) == "v1"


def test_header_div0(tmp_path):
    from learn_fixtures import make_workbook
    from gcws.learn.workbook import parse_header
    h = parse_header(_cells(make_workbook(tmp_path / "w.xlsx", istd_area_c17="#DIV/0!")))
    assert h.istd[0].name == "C17" and h.istd[0].area is None


def _table(path):
    from gcws.learn.workbook import classify_rows, parse_eval_table
    rows = parse_eval_table(_cells(path), _cells(path, data_only=False))
    classify_rows(rows)
    return rows


def test_eval_table(tmp_path):
    from learn_fixtures import make_workbook
    rows = _table(make_workbook(tmp_path / "w.xlsx"))
    by_rt = {r.rt: r for r in rows if r.rt is not None}
    bma = by_rt[7.07]
    assert (bma.label, bma.cas, bma.library, bma.match) == ("Butyl methacrylate", "97-88-1", "NIST05", 90)
    assert (bma.area_formula, bma.area_original, bma.row_class) == ("=L29-F30", 119966715, "named")
    assert (bma.sml, bma.reference, bma.sheet_row) == ("6", "[1],[2]", 29)
    assert bma.conc == {"mg/dm²": 0.243, "mg/kg": 1.458}
    assert by_rt[6.917].row_class == "unnamed" and by_rt[6.917].area == 395154
    assert by_rt[6.917].area_formula == "" and by_rt[6.917].area_original is None
    assert by_rt[14.33].row_class == "istd" and by_rt[16.935].row_class == "istd"
    assert by_rt[21.048].row_class == "group" and by_rt[21.449].row_class == "group"
    assert by_rt[26.5].row_class == "named_no_cas"
    assert by_rt[20.439].row_class == "unknown"
    assert by_rt[17.878].row_class == "coelution"
    assert by_rt[28.862].row_class == "derivative"
    sums = [r for r in rows if r.row_class == "sum"]
    assert len(sums) == 1 and sums[0].label.startswith("Sum of styrene") and sums[0].rt is None
    assert not any(r.label == "Ende" for r in rows)


def test_eval_table_other_units(tmp_path):
    from learn_fixtures import make_workbook
    rows = _table(make_workbook(tmp_path / "w.xlsx", conc_headers=("Conc. µg/L", "Conc. µg/Zipper",
                                                                    "Conc. µg/0.25gGranulat"),
                                final_rows=[{"A": 7.5, "B": "X", "C": "50-00-0", "F": 10, "G": 1, "H": 2,
                                             "I": 3}]))
    assert rows[0].conc == {"µg/L": 1, "µg/Zipper": 2, "µg/0.25gGranulat": 3}
    assert rows[0].sml == ""


def test_parse_report(tmp_path):
    from learn_fixtures import FOOTNOTES, make_workbook
    from gcws.learn.workbook import parse_report
    rows, notes = parse_report(_cells(make_workbook(tmp_path / "w.xlsx"), "externerBericht"))
    assert [r.rt for r in rows] == [7.07, 10.846, 17.878, 21.048, 21.449]
    bma = rows[0]
    assert (bma.label, bma.cas, bma.library, bma.match, bma.sml, bma.reference) == (
        "Butyl methacrylate", "97-88-1", "NIST05", 90, "6", "[1],[2]")
    assert bma.conc == {"mg/dm²": 0.243, "mg/kg": 1.458}
    assert rows[1].sml == "CC I(a)"
    assert [r.row_class for r in rows] == ["named", "named", "coelution", "group", "group"]
    assert notes == FOOTNOTES


def test_parse_report_lowercase_headers(tmp_path):
    import openpyxl
    from learn_fixtures import make_workbook
    from gcws.learn.workbook import parse_report
    path = make_workbook(tmp_path / "w.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["externerBericht"]
    ws["B25"], ws["G25"], ws["J25"] = "name", "migration conc.", "ref."
    wb.save(path)
    rows, _ = parse_report(_cells(path, "externerBericht"))
    assert rows[0].label == "Butyl methacrylate" and rows[0].reference == "[1],[2]"
    assert list(rows[0].conc) == ["mg/dm²", "mg/kg"]


def test_parse_workbook_synthetic(tmp_path):
    from learn_fixtures import make_workbook
    from gcws.learn.workbook import parse_workbook
    ev = parse_workbook(make_workbook(tmp_path / "w.xlsx"))
    assert ev.problems == []
    assert ev.template == "v2" and ev.header.simulant == "EtOH 95%"
    assert {p.signal for p in ev.raw} == {"TIC", "FID"}
    assert ev.pre_clean is not None and len(ev.pre_clean) == len(ev.final) + 2
    assert len(ev.report) == 5 and ev.footnotes
    assert parse_workbook(make_workbook(tmp_path / "v.xlsx", pre_clean=False)).pre_clean is None


def test_parse_workbook_bad_file(tmp_path):
    from gcws.learn.workbook import parse_workbook
    bad = tmp_path / "bad.xlsm"
    bad.write_text("not a workbook")
    ev = parse_workbook(bad)
    assert len(ev.problems) == 1 and ev.problems[0].startswith("cannot open: ")
    assert (ev.raw, ev.final, ev.report) == ([], [], [])


def test_parse_workbook_div0_is_a_problem(tmp_path):
    from learn_fixtures import make_workbook
    from gcws.learn.workbook import parse_workbook
    ev = parse_workbook(make_workbook(tmp_path / "w.xlsx", istd_area_c17="#DIV/0!"))
    assert ev.problems == ["ISTD area missing: C17"]


def test_removed_peaks(tmp_path):
    from learn_fixtures import make_workbook
    from gcws.learn.workbook import parse_workbook, removed_peaks
    removed = removed_peaks(parse_workbook(make_workbook(tmp_path / "w.xlsx")))
    assert [p.rt for p in removed] == [5.132, 6.409]

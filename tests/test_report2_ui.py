"""P62 / Report² rework: one list with filter chips, To do and Archive, one-click accept and reject
with a reason (undo for a few seconds), right-click menus, delete (hide) and restore, keys, and the
batch report of a folder from GC Workspace. Dialogs are never left open."""
import json
import time
from pathlib import Path

import pytest

from test_ui import _load, win  # noqa: F401  (fixture)


@pytest.fixture()
def data(tmp_path, monkeypatch):
    from gcws import paths
    d = tmp_path / "data"
    d.mkdir()
    monkeypatch.setattr(paths, "DATA", d)
    return d


class NoWatcher:
    def status(self, journal=None):
        return {"state": "stopped", "heartbeat": None, "current": ""}

    def send(self, cmd, timeout=0):
        return None


def _seed(data, tmp_path):
    """A workflow and a journal with samples in every state."""
    from gcws.automation import journal as J, templates
    watch = tmp_path / "watch"
    batch = watch / "26016605_GIOSUN1635"
    batch.mkdir(parents=True)
    wf = templates.make("nias", "Lab", source=str(watch), method="NIAS", folder_a=str(tmp_path / "A"),
                        folder_b=str(tmp_path / "B"))
    wf.save()
    rep = wf.by_type("report")[0]
    jr = J.Journal()
    b = jr.batch(wf.id, batch)
    m = wf.methods()[0].id
    out = tmp_path / "jobs"
    out.mkdir()
    ids = {}
    for name, state in (("S-control", J.CONTROL), ("S-auto", J.ACCEPTED_AUTO), ("S-wait", J.WAITING),
                        ("S-noblank", J.NOT_PROCESSED), ("S-failed", J.FAILED)):
        j = jr.ensure_job(wf.id, m, b["id"], name, name, [name + "_A.D"], {}, "fp", state=J.QUEUED
                          if state not in (J.WAITING,) else J.WAITING, reason="waiting for B" if state == J.WAITING
                          else "")
        files = {}
        for fmt in ("xlsx", "docx"):
            p = out / f"{name}_NIAS_Report.{fmt}"
            p.write_text(fmt)
            files[fmt] = str(p)
        if state != J.WAITING:
            jr.transition(j.id, J.QUEUED, J.PROCESSING)
            fields = {"files": {rep.id: files}, "job_dir": str(out)}
            if state == J.CONTROL:
                fields["findings"] = [{"rule": "sml_exceeded", "level": "control", "text": "0.5 mg/kg > SML 0.05",
                                       "member": "07_A", "substance": "Bisphenol A", "cas": "80-05-7", "rt": 17.2}]
            if state == J.NOT_PROCESSED:
                fields["reason"] = "07_A: no Blank from the same batch"
            jr.transition(j.id, J.PROCESSING, state, **fields)
        ids[name] = j.id
    return wf, jr, ids, batch


def _names(dock, bucket=None):
    """The samples in the list (of one chip filter)."""
    if bucket is not None and dock.filter != bucket:
        dock.set_filter(bucket)
    tree = dock.tree
    out = []
    for i in range(tree.topLevelItemCount()):
        top = tree.topLevelItem(i)
        out += [top.child(k).text(0) for k in range(top.childCount())]
    return out


def _no_word(docx, pdf):
    raise RuntimeError("no Word in the tests")


@pytest.fixture()
def settings(tmp_path):
    """QSettings of the test only (the preview switch is remembered there)."""
    from PySide6.QtCore import QCoreApplication, QSettings
    QCoreApplication.setOrganizationName("GCWorkspaceTest")
    QCoreApplication.setApplicationName("pytest")
    QSettings.setDefaultFormat(QSettings.IniFormat)
    QSettings.setPath(QSettings.IniFormat, QSettings.UserScope, str(tmp_path))
    QSettings().clear()


def _dock(jr, monkeypatch, qtbot, word=_no_word):
    """Report² on ``jr``; Microsoft Word (the preview of Word reports) is replaced by ``word``."""
    from gcws.report import service as RS
    from gcws.ui.docks.report2 import Report2Dock
    monkeypatch.setattr("gcws.automation.control.WatcherControl", lambda *a, **k: NoWatcher())
    monkeypatch.setattr(RS, "docx_to_pdf", word)
    dock = Report2Dock(journal=jr, poll_ms=60000)
    qtbot.addWidget(dock)
    return dock


def test_report2_list_counters_and_review(qtbot, data, tmp_path, monkeypatch):
    from gcws.automation import journal as J
    wf, jr, ids, batch = _seed(data, tmp_path)
    dock = _dock(jr, monkeypatch, qtbot)
    assert set(_names(dock)) == {"S-control", "S-auto", "S-wait", "S-noblank", "S-failed"}
    assert _names(dock, "control") == ["S-control"] and _names(dock, "accepted") == ["S-auto"]
    assert _names(dock, "waiting") == ["S-wait"]
    assert _names(dock, "not_processed") == ["S-noblank"] and _names(dock, "failed") == ["S-failed"]
    dock.set_filter("failed")                                  # clicked again: all
    assert dock.filter == "all"
    assert dock.chips["all"].text() == "All 5"
    assert dock.chips["control"].text() == "Control needed 1"
    assert dock.chips["accepted"].text() == "Accepted 1" and "1 automatic" in dock.chips["accepted"].toolTip()
    assert dock.chips["rejected"].text() == ""                 # none: hidden
    assert dock.b_todo.text() == "To do (1)" and dock.b_archive.text() == "Archive (0)"
    assert dock.watcher.text() == "● Watcher stopped" and not dock.start_btn.isHidden()
    batch_row = dock.tree.topLevelItem(0)
    assert batch_row.text(0) == batch.name and batch_row.text(1) == "1/5 accepted · 1 control · 2 failed/rejected · 1 waiting"
    dock.select(ids["S-control"])
    assert dock.reasons.text() == "" and dock.b_accept.isEnabled() and dock.b_open.isEnabled()   # no labels
    assert dock._items[f"j:{ids['S-control']}"].childCount() == 0     # no double determination: nothing below
    jr.update_job(ids["S-control"], evidence={"features": [
        {"light": "red", "name": "Bisphenol A", "rt": 17.2}, {"light": "red", "name": "unknown", "rt": 9.1},
        {"light": "yellow", "name": "BHT", "rt": 11.87}, {"light": "green", "name": "x"}, {"light": "grey"}]})
    dock.refresh()
    row = dock._items[f"j:{ids['S-control']}"]
    assert [(row.child(k).text(0), row.child(k).text(1)) for k in range(row.childCount())] ==         [("Red - to decide", "2"), ("Yellow - to check", "1")]                # counts, not every finding
    assert "Bisphenol A" in row.child(0).toolTip(0)
    # accept: one click, no comment; delivered once the undo time is over (no watcher running)
    assert dock.review(True)
    j = jr.job(ids["S-control"])
    assert j.state == J.ACCEPTED_MANUAL and j.comment == "" and j.reviewer
    assert dock.bar.isVisibleTo(dock) and dock.b_undo.isVisibleTo(dock)
    assert not (tmp_path / "A" / batch.name / "S-control_NIAS_Report.xlsx").is_file()
    dock.dismiss_message()
    assert (tmp_path / "A" / batch.name / "S-control_NIAS_Report.xlsx").is_file()
    assert (tmp_path / "B" / batch.name / "S-control_NIAS_Report.docx").is_file()
    assert _names(dock, "control") == [] and set(_names(dock, "accepted")) == {"S-control", "S-auto"}
    assert dock._items[f"j:{ids['S-control']}"].text(3) == "✓"
    # reject with a reason, process again, process without blank
    dock.set_filter("all")
    dock.select(ids["S-auto"])
    assert dock.reject("Wrong sample / mix-up") and jr.job(ids["S-auto"]).state == J.REJECTED
    assert jr.job(ids["S-auto"]).comment == "Wrong sample / mix-up"
    assert dock.reprocess("full") and jr.job(ids["S-auto"]).state == J.QUEUED
    assert "start the watcher" in dock.bar_text.text()
    dock.select(ids["S-noblank"])
    assert dock.a_noblank.isEnabled()
    assert dock.process_without_blank(confirm=False)
    j = jr.job(ids["S-noblank"])
    assert j.state == J.QUEUED and j.override["allow_no_blank"] and j.revision == 2


def test_open_in_gc_workspace(qtbot, data, tmp_path, monkeypatch):
    wf, jr, ids, batch = _seed(data, tmp_path)
    proj = tmp_path / "x.gcws"
    proj.write_text("{}")
    jr.update_job(ids["S-control"], project_path=str(proj))
    dock = _dock(jr, monkeypatch, qtbot)
    got = []
    dock.openProject.connect(got.append)
    dock.select(ids["S-control"])
    assert dock.b_project.isEnabled()
    dock.open_project()
    assert got == [str(proj)]


def test_accept_and_reject_comments_are_optional(qtbot):
    from PySide6.QtWidgets import QDialog
    from gcws.ui.dialogs.report2 import ReviewDialog

    for accept in (True, False):
        dlg = ReviewDialog(accept, "S-control", 2)
        qtbot.addWidget(dlg)
        assert not dlg.required
        dlg._ok()
        assert dlg.result() == QDialog.Accepted and dlg.text() == ""


def test_processed_pairs_are_nested_and_available_to_switch(qtbot, data, tmp_path, monkeypatch):
    wf, jr, ids, batch = _seed(data, tmp_path)
    project = tmp_path / "pair.gcws"
    project.write_text("{}")
    for name in ("S-control", "S-auto"):
        jr.update_job(ids[name], members=[f"{name}_A.D", f"{name}_B.D"], project_path=str(project))
    dock = _dock(jr, monkeypatch, qtbot)
    child = dock._items[f"j:{ids['S-control']}"]
    assert batch.name == dock.tree.topLevelItem(0).text(0)
    assert "S-control_A" in child.text(0) and "S-control_B" in child.text(0)
    assert {j.id for j in dock.visible_pairs()} == {ids["S-control"], ids["S-auto"]}
    opened = []
    dock.openDetermination.connect(opened.append)
    dock.open_determination(ids["S-control"])
    assert opened == [ids["S-control"]]
    opened.clear()
    menu = dock._tree_context_menu(dock.tree.visualItemRect(child).center())
    texts = [a.text() for a in menu.actions()]
    assert "Open in Replicates / results" in texts
    next(a for a in menu.actions() if a.text() == "Open in Replicates / results").trigger()
    assert opened == [ids["S-control"]]
    dock.search.setText("S-auto")
    assert [j.id for j in dock.visible_pairs()] == [ids["S-auto"]]


class FakeLauncher:
    """The job process of GC Workspace's LocalJobs, completed by the test."""

    def __init__(self, parent=None):
        from PySide6.QtCore import QObject, Signal

        class _S(QObject):
            finished = Signal(str, int, str)
        self._s = _S()
        self.finished = self._s.finished
        self.started, self.job, self.killed = [], None, 0

    def running(self):
        return self.job is not None

    def start(self, job_id, spec_path, kind="job"):
        self.started.append((job_id, json.loads(Path(spec_path).read_text(encoding="utf-8")), kind))
        self.job = job_id

    def kill(self):
        self.killed += 1
        self.job = None

    def complete(self, result):
        job_id, spec, _kind = self.started[-1]
        (Path(spec["out_dir"]) / "result.json").write_text(json.dumps(result), encoding="utf-8")
        self.job = None
        self.finished.emit(job_id, 0, "")


def test_edited_accept_makes_the_report_again_here(qtbot, data, tmp_path, monkeypatch):
    """Accepting an edited report makes it again in GC Workspace at once (no watcher needed); it stays
    accepted and is delivered."""
    import os
    from gcws.automation import journal as J
    wf, jr, ids, batch = _seed(data, tmp_path)
    project = tmp_path / "pair.gcws"
    project.write_text("{}")
    jr.update_job(ids["S-control"], project_path=str(project), edited=1)
    launcher = FakeLauncher()
    monkeypatch.setattr("gcws.automation.watcher.ProcessLauncher", lambda parent=None: launcher)
    dock = _dock(jr, monkeypatch, qtbot)
    dock.select(ids["S-control"])
    assert dock.review(True)
    job = jr.job(ids["S-control"])
    assert job.state == J.PROCESSING and job.review_pending and job.pid == os.getpid()
    assert not dock.b_undo.isVisibleTo(dock)                   # regenerating: nothing to undo
    assert _names(dock, "control") == ["S-control"] and "S-control" not in _names(dock, "waiting")
    job_id, spec, kind = launcher.started[-1]
    assert job_id == job.id and spec["mode"] == "rereport" and spec["project_path"] == str(project)
    rep = wf.by_type("report")[0]
    files = {}
    for fmt in ("xlsx", "docx", "pdf"):
        p = Path(spec["out_dir"]) / f"S-control_NIAS_Report.{fmt}"
        p.write_text(fmt)
        files[fmt] = str(p)
    new = {"rule": "manual_check", "level": "control", "text": "new", "member": "B", "substance": "Y", "rt": 2.0}
    launcher.complete({"state": "control", "reason": "", "files": {rep.id: files}, "project": str(project),
                       "evidence": {}, "findings": list(job.findings) + [new], "warnings": [], "timings": {}})
    after = jr.job(job.id)
    assert after.state == J.ACCEPTED_MANUAL and after.revision == 2 and after.review_pending is None
    assert "1 new finding" in after.reason and after.export_pending == 0 and after.export_state == "done"
    assert (tmp_path / "A" / batch.name / "S-control_NIAS_Report.xlsx").is_file()
    assert (tmp_path / "B" / batch.name / "S-control_NIAS_Report.docx").is_file()
    assert "updated report is made and accepted" in dock.bar_text.text()


def test_edited_report_left_when_gc_workspace_closes_goes_back_to_the_queue(qtbot, data, tmp_path, monkeypatch):
    from gcws.automation import journal as J
    wf, jr, ids, batch = _seed(data, tmp_path)
    project = tmp_path / "pair.gcws"
    project.write_text("{}")
    jr.update_job(ids["S-control"], project_path=str(project), edited=1)
    launcher = FakeLauncher()
    monkeypatch.setattr("gcws.automation.watcher.ProcessLauncher", lambda parent=None: launcher)
    dock = _dock(jr, monkeypatch, qtbot)
    dock.select(ids["S-control"])
    assert dock.review(True) and jr.job(ids["S-control"]).state == J.PROCESSING
    dock._stop_local()
    job = jr.job(ids["S-control"])
    assert launcher.killed == 1 and job.state == J.QUEUED and job.review_pending and job.pid is None


def test_undo_and_next_report_needing_control(qtbot, data, tmp_path, monkeypatch):
    from gcws.automation import journal as J
    wf, jr, ids, batch = _seed(data, tmp_path)
    second = jr.ensure_job(wf.id, wf.methods()[0].id, jr.batch(wf.id, batch)["id"], "S-control2", "S-control2",
                           ["x.D"], {}, "fp", state=J.QUEUED)
    jr.transition(second.id, J.QUEUED, J.PROCESSING)
    jr.transition(second.id, J.PROCESSING, J.CONTROL)
    dock = _dock(jr, monkeypatch, qtbot)
    dock.select(ids["S-control"])
    assert dock.review(True)
    assert dock.current == second.id                           # the next one needing control is selected
    assert dock.undo() and jr.job(ids["S-control"]).state == J.CONTROL
    assert not dock.b_undo.isVisibleTo(dock) and dock._deliver == set()
    dock.dismiss_message()
    assert not (tmp_path / "A" / batch.name / "S-control_NIAS_Report.xlsx").is_file()   # undone: not delivered
    # a reject is undone the same way
    dock.select(second.id)
    assert dock.reject("Bad chromatography") and jr.job(second.id).state == J.REJECTED
    assert dock.undo() and jr.job(second.id).state == J.CONTROL
    # several at once
    dock.set_filter("control")
    dock.tree.selectAll()
    assert len(dock.selected_jobs()) == 2 and "2 reports selected" in dock.title.text()
    assert dock.review(True) and {jr.job(i).state for i in (ids["S-control"], second.id)} == {J.ACCEPTED_MANUAL}
    assert "and 1 more" in dock.bar_text.text()


def test_reject_reasons_menu_and_other(qtbot, data, tmp_path, monkeypatch):
    from gcws.automation import journal as J, rules as RU
    from gcws.ui.dialogs.report2 import ReviewDialog
    wf, jr, ids, batch = _seed(data, tmp_path)
    RU.save_reject_reasons(["Too dilute", "Bad chromatography"])
    dock = _dock(jr, monkeypatch, qtbot)
    dock.select(ids["S-control"])
    dock.reject_menu.aboutToShow.emit()
    texts = [a.text() for a in dock.reject_menu.actions() if not a.isSeparator()]
    assert texts == ["Too dilute", "Bad chromatography", "Other..."]
    with monkeypatch.context() as m:
        m.setattr(ReviewDialog, "exec", lambda self: (self.comment.setPlainText("smells"), 1)[1])
        dock.reject_menu.actions()[-1].trigger()
    assert jr.job(ids["S-control"]).state == J.REJECTED and jr.job(ids["S-control"]).comment == "smells"
    # a rejected report can be accepted after all
    assert dock.b_accept.isEnabled() and dock.review(True)
    assert jr.job(ids["S-control"]).state == J.ACCEPTED_MANUAL


def test_context_menus_of_a_sample_and_a_batch(qtbot, data, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    from gcws.automation import journal as J
    wf, jr, ids, batch = _seed(data, tmp_path)
    dock = _dock(jr, monkeypatch, qtbot)
    row = dock._items[f"j:{ids['S-control']}"]
    menu = dock._tree_context_menu(dock.tree.visualItemRect(row).center())
    texts = [a.text() for a in menu.actions() if not a.isSeparator()]
    for t in ("Open report", "Save as Word...", "Save as Excel...", "Edit in GC Workspace", "Accept", "Accept with comment...", "Reject",
              "Process again", "Remove from the queue...", "Deliver to the target folders now",
              "Show history...", "Open the job folder", "Copy sample name", "Delete..."):
        assert t in texts, t
    assert dock.current == ids["S-control"]                    # right-click selects the row
    top = dock.tree.topLevelItem(0)
    menu = dock._tree_context_menu(dock.tree.visualItemRect(top).center())
    texts = [a.text() for a in menu.actions() if not a.isSeparator()]
    assert texts == ["Open batch report", "Save batch report as Word...", "Save batch report as Excel...",
                     "Open batch folder", 'Accept all "control needed" (1)...', "Reject batch",
                     "Process batch again...", "Delete batch..."]
    bid = top.data(0, 0x0100)
    with monkeypatch.context() as m:
        m.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
        assert dock.accept_batch(bid)
        assert jr.job(ids["S-control"]).state == J.ACCEPTED_MANUAL
        assert dock.reject_batch(bid, "Repeat measurement")
        assert {jr.job(ids[n]).state for n in ("S-control", "S-auto")} == {J.REJECTED}
        assert dock.reprocess_batch(bid)
        assert jr.job(ids["S-failed"]).state == J.QUEUED and jr.job(ids["S-control"]).state == J.QUEUED


def test_delete_hides_and_show_deleted_restores(qtbot, data, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    wf, jr, ids, batch = _seed(data, tmp_path)
    dock = _dock(jr, monkeypatch, qtbot)
    dock.select(ids["S-failed"])
    asked = []
    with monkeypatch.context() as m:
        m.setattr(QMessageBox, "question", lambda *a, **k: asked.append(a[2]) or QMessageBox.Yes)
        assert dock.delete_selected() == [ids["S-failed"]]
    assert "Nothing on disk is deleted" in asked[0]
    assert "S-failed" not in _names(dock) and dock.chips["all"].text() == "All 4"
    assert dock.undo() and "S-failed" in _names(dock)            # Undo restores it
    dock.select(ids["S-failed"])
    assert dock.delete_selected(confirm=False)
    dock.a_show_deleted.setChecked(True)
    row = dock._items[f"j:{ids['S-failed']}"]
    assert row.text(1).startswith("Deleted") and row.font(0).italic()
    dock.select(ids["S-failed"])
    assert dock.a_restore.isEnabled() and not dock.a_delete.isEnabled()
    assert dock.restore_selected() == [ids["S-failed"]] and not jr.job(ids["S-failed"]).deleted
    # a whole batch
    bid = jr.batch(wf.id, batch)["id"]
    assert dock.delete_batch(bid, confirm=False)
    dock.a_show_deleted.setChecked(False)
    assert dock.tree.topLevelItemCount() == 0 and "Nothing to do" in dock.empty.text()
    dock.a_show_deleted.setChecked(True)
    assert dock.tree.topLevelItem(0).text(1) == "Deleted"
    assert dock.restore_batch(bid) and len(_names(dock)) == 5


def test_archive_and_reopen(qtbot, data, tmp_path, monkeypatch):
    from gcws.automation import journal as J
    wf, jr, ids, batch = _seed(data, tmp_path)
    for name in ("S-wait", "S-noblank", "S-failed"):
        jr.delete([ids[name]])
    dock = _dock(jr, monkeypatch, qtbot)
    dock.select(ids["S-control"])
    assert dock.review(True)
    dock.dismiss_message()                                     # delivered: the whole batch is done
    jr.update_job(ids["S-auto"], export_state="done")
    dock.refresh()
    assert dock.tree.topLevelItemCount() == 0 and "Archive" in dock.empty.text()
    assert dock.b_archive.text() == "Archive (1)"
    dock.set_mode("archive")
    top = dock.tree.topLevelItem(0)
    assert top.text(0) == batch.name and top.text(1) == "2/2 accepted" and top.text(3) == "2/2"
    assert not top.isExpanded()                                # the archive lists batches
    bid = top.data(0, 0x0100)
    dock.search.setText("S-auto")                              # a sample name finds its batch
    assert dock.tree.topLevelItem(0).isExpanded() and _names(dock) == ["S-auto"]
    dock.search.setText("")
    menu = dock._tree_context_menu(dock.tree.visualItemRect(dock.tree.topLevelItem(0)).center())
    assert "Reopen" in [a.text() for a in menu.actions()]
    assert dock.reopen_batch(bid) and dock.mode == "todo" and dock.b_todo.isChecked()
    assert set(_names(dock)) == {"S-control", "S-auto"}
    dock.select(ids["S-auto"])
    assert dock.reject("Repeat measurement") and jr.job(ids["S-auto"]).state == J.REJECTED


def test_to_do_counts_only_batches_with_something_to_show(qtbot, data, tmp_path, monkeypatch):
    """A batch folder without samples, or whose reports were all deleted, is neither listed nor counted."""
    from gcws.automation import journal as J
    wf, jr, ids, batch = _seed(data, tmp_path)
    jr.batch(wf.id, tmp_path / "watch" / "26016700_EMPTY")                  # seen, no sample
    gone = jr.batch(wf.id, tmp_path / "watch" / "26016701_GONE")
    j = jr.ensure_job(wf.id, wf.methods()[0].id, gone["id"], "g", "S-gone", ["g.D"], {}, "fp", state=J.QUEUED)
    jr.delete([j.id])
    dock = _dock(jr, monkeypatch, qtbot)
    assert dock.b_todo.text() == "To do (1)" and dock.tree.topLevelItemCount() == 1
    dock.a_show_deleted.setChecked(True)
    assert dock.b_todo.text() == "To do (2)" and "S-gone" in _names(dock)


def test_keys_history_and_watcher(qtbot, data, tmp_path, monkeypatch):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    from gcws.automation import journal as J
    wf, jr, ids, batch = _seed(data, tmp_path)
    dock = _dock(jr, monkeypatch, qtbot)
    dock.select(ids["S-control"])
    dlg = dock.show_history()
    try:
        assert dlg.isVisible() and not dlg.isModal()
        assert dlg.files.rowCount() == 2 and dlg.history.rowCount() >= 1
    finally:
        dlg.close()
    dock.tree.setFocus()
    QTest.keyClick(dock.tree, Qt.Key_A)
    assert jr.job(ids["S-control"]).state == J.ACCEPTED_MANUAL
    QTest.keyClick(dock.tree, Qt.Key_Z, Qt.ControlModifier)
    assert jr.job(ids["S-control"]).state == J.CONTROL
    dock.select(ids["S-failed"])
    QTest.keyClick(dock.tree, Qt.Key_A)                        # not possible: nothing happens
    assert jr.job(ids["S-failed"]).state == J.FAILED
    with monkeypatch.context() as m:
        m.setattr("PySide6.QtWidgets.QMessageBox.question", lambda *a, **k: 0x00010000)   # No
        QTest.keyClick(dock.tree, Qt.Key_Delete)
    assert not jr.job(ids["S-failed"]).deleted


class RunningWatcher(NoWatcher):
    def status(self, journal=None):
        return {"state": "running", "heartbeat": 1, "current": ""}


def test_with_a_running_watcher_the_accept_is_left_to_it(qtbot, data, tmp_path, monkeypatch):
    from gcws.automation import journal as J
    from gcws.ui.docks.report2 import Report2Dock
    wf, jr, ids, batch = _seed(data, tmp_path)
    monkeypatch.setattr("gcws.automation.control.WatcherControl", lambda *a, **k: RunningWatcher())
    monkeypatch.setattr("gcws.report.service.docx_to_pdf", _no_word)
    dock = Report2Dock(journal=jr, poll_ms=60000)
    qtbot.addWidget(dock)
    assert dock.watcher.text() == "● Watcher running" and dock.start_btn.isHidden()
    dock.select(ids["S-control"])
    assert dock.review(True) and dock._deliver == set()
    job = jr.job(ids["S-control"])
    assert job.export_pending and job.deliver_after > time.time() and job.state == J.ACCEPTED_MANUAL
    assert dock._items[f"j:{ids['S-control']}"].text(3) == "pending"


def _pdf(path):
    from PySide6.QtGui import QPainter, QPdfWriter
    w = QPdfWriter(str(path))
    p = QPainter(w)
    p.drawText(200, 200, "NIAS report")
    p.end()
    return path


def test_pdf_preview(qtbot, data, tmp_path, monkeypatch, settings):
    from PySide6.QtCore import QSettings
    wf, jr, ids, batch = _seed(data, tmp_path)
    job = jr.job(ids["S-control"])
    files = dict(job.files)
    node = next(iter(files))
    pdf = _pdf(tmp_path / "S-control_NIAS_Report.pdf")
    files[node] = dict(files[node], pdf=str(pdf))
    jr.update_job(job.id, files=files)
    dock = _dock(jr, monkeypatch, qtbot)
    assert dock.b_preview.isChecked()
    dock.select(ids["S-auto"])                                  # Word and Excel only, and no Word here
    qtbot.waitUntil(lambda: "no Word in the tests" in dock.preview_note.text())
    dock.select(ids["S-control"])
    assert dock.preview.currentWidget() is dock.pdf_view and dock._pdf_doc.pageCount() == 1
    pdf.unlink()                                                # read into memory: the file is not held open
    dock.b_preview.setChecked(False)
    assert not dock.preview.isVisibleTo(dock) and QSettings().value("report2/preview", type=bool) is False
    dock.b_preview.setChecked(True)
    assert dock.preview.isVisibleTo(dock)


def test_word_report_is_previewed_through_a_pdf_made_once(qtbot, data, tmp_path, monkeypatch, settings):
    """No PDF among the report's formats: Word converts the Word report once (in a thread); the PDF is
    kept beside it in the job folder, made again only when the Word report is newer."""
    import os
    import shutil
    wf, jr, ids, batch = _seed(data, tmp_path)
    made = _pdf(tmp_path / "made_by_word.pdf")
    calls = []

    def word(docx, pdf):
        calls.append(Path(docx).name)
        shutil.copyfile(made, pdf)
        return Path(pdf)
    dock = _dock(jr, monkeypatch, qtbot, word)
    assert dock.b_preview.isChecked()
    dock.select(ids["S-control"])
    assert "Making a preview" in dock.preview_note.text()
    qtbot.waitUntil(lambda: dock.pdf_view is not None and dock.preview.currentWidget() is dock.pdf_view)
    docx = Path(jr.job(ids["S-control"]).files[next(iter(jr.job(ids["S-control"]).files))]["docx"])
    kept = dock.preview_pdf(docx)
    assert kept == docx.with_name(docx.stem + ".preview.pdf") and kept.is_file()
    assert calls == [docx.name]
    dock.select(ids["S-auto"])                                  # another report: converted as well
    qtbot.waitUntil(lambda: len(calls) == 2 and dock.preview.currentWidget() is dock.pdf_view)
    dock.select(ids["S-control"])                               # kept: no Word again
    assert dock.preview.currentWidget() is dock.pdf_view and len(calls) == 2
    later = kept.stat().st_mtime + 10
    os.utime(docx, (later, later))                              # the Word report was made again
    dock.select(ids["S-auto"])
    dock.select(ids["S-control"])
    qtbot.waitUntil(lambda: len(calls) == 3 and dock.preview.currentWidget() is dock.pdf_view)


def test_word_preview_failure_is_shown_once(qtbot, data, tmp_path, monkeypatch, settings):
    wf, jr, ids, batch = _seed(data, tmp_path)
    calls = []

    def no_word(docx, pdf):
        calls.append(docx)
        raise RuntimeError("Word is not installed")
    dock = _dock(jr, monkeypatch, qtbot, no_word)
    dock.select(ids["S-control"])
    qtbot.waitUntil(lambda: "Word is not installed" in dock.preview_note.text())
    dock.select(ids["S-auto"])
    qtbot.waitUntil(lambda: len(calls) == 2 and dock._converting is None)
    dock.select(ids["S-control"])                               # not tried again and again
    assert len(calls) == 2 and "Word could not convert" in dock.preview_note.text()


def test_docks_and_menus(qtbot, win):
    assert "report2" in win.docks and "automation" in win.docks
    labels = [a.text() for a in win.report_menu.actions()]
    assert "Batch report of this folder..." in labels
    assert "Report²" in [a.text() for a in win.automation_menu.actions()]


def test_replicates_has_report2_pair_picker(qtbot, win, monkeypatch):
    page = win.replicates.duplicate
    assert page.b_report2.text() == "Load from Report²…"
    page.resize(1000, 600)
    page.show()
    from PySide6.QtWidgets import QApplication
    QApplication.processEvents()
    assert page.b_report2.y() == page.b_compare.y()
    assert page.b_report2.x() + page.b_report2.width() < 388  # fits the usual dock width
    win.replicates.report2Requested.disconnect(win._show_report2_pair_menu)
    selected = []
    win.replicates.report2Requested.connect(lambda: selected.append(True))
    page.b_report2.click()
    assert selected == [True]


def test_processed_pair_display_does_not_apply_automatic_changes(qtbot, win, monkeypatch):
    page = win.replicates.duplicate
    calls = []
    monkeypatch.setattr(page, "compare", lambda sync=False: calls.append(sync))
    page.set_pair("A", "B", processed=True)
    assert calls == [False] and page._view_processed


def test_switch_saves_accepted_project_and_returns_it_to_control(qtbot, win, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    from gcws.automation import journal as J

    jr = J.Journal(tmp_path / "journal.sqlite")
    batch = jr.batch("wf", tmp_path)
    job = jr.ensure_job("wf", "method", batch["id"], "pair", "Pair", ["A.D", "B.D"], {}, "fp",
                        state=J.QUEUED)
    project = tmp_path / "pair.gcws"
    project.write_text("original")
    jr.transition(job.id, J.QUEUED, J.PROCESSING)
    jr.transition(job.id, J.PROCESSING, J.ACCEPTED_AUTO, project_path=str(project))
    win.report2._journal = jr
    win._report2_job = (job.id, 1, project)
    win._report2_baseline = {"quant": {"limit": 1}}
    win.ws.project_path = project
    win.ws.dirty = True
    monkeypatch.setattr("gcws.ui.main_window.P.to_dict", lambda ws, path: {"quant": {"limit": 2}})
    monkeypatch.setattr("gcws.ui.main_window.P.save", lambda ws, path: (project.write_text("edited"), project)[1])

    assert win._save_report2_project()
    assert project.read_text() == "edited"
    assert jr.job(job.id).state == J.CONTROL and jr.job(job.id).edited == 1
    assert jr.accept_edited(job.id, user="analyst")
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: None)
    assert win._save_report2_project()  # switching after Accept must not trap the open project
    writes = []
    monkeypatch.setattr(win.ws, "states", lambda: [object()])
    monkeypatch.setattr("gcws.ui.main_window.P.save", lambda ws, path: (writes.append(path), path)[1])
    win.save_project()
    assert writes == []  # Ctrl+S on an old revision must not overwrite the pending source


def test_run_rename_marks_report2_project_edited(qtbot, win, tmp_path, monkeypatch):
    from gcws.automation import journal as J
    from gcws.ui.main_window import _report_inputs

    jr = J.Journal(tmp_path / "journal.sqlite")
    batch = jr.batch("wf", tmp_path)
    job = jr.ensure_job("wf", "method", batch["id"], "pair", "Pair", ["A.D", "B.D"], {}, "fp",
                        state=J.QUEUED)
    project = tmp_path / "pair.gcws"
    project.write_text("original")
    jr.transition(job.id, J.QUEUED, J.PROCESSING)
    jr.transition(job.id, J.PROCESSING, J.ACCEPTED_AUTO, project_path=str(project))
    win.report2._journal = jr
    win._report2_job = (job.id, 1, project)
    original = {"runs": [{"id": "a", "name": "Old"}]}
    renamed = {"runs": [{"id": "a", "name": "New"}]}
    win._report2_baseline = _report_inputs(original)
    win.ws.project_path = project
    win.ws.dirty = False
    monkeypatch.setattr("gcws.ui.main_window.P.to_dict", lambda ws, path: renamed)
    monkeypatch.setattr("gcws.ui.main_window.P.save", lambda ws, path: (project.write_text("edited"), project)[1])

    assert win._save_report2_project()
    assert project.read_text() == "edited"
    assert jr.job(job.id).state == J.CONTROL and jr.job(job.id).edited == 1


def test_rejected_report2_project_can_be_saved_and_requeued(qtbot, win, tmp_path, monkeypatch):
    from gcws.automation import journal as J
    from gcws.ui.main_window import _report_inputs

    jr = J.Journal(tmp_path / "journal.sqlite")
    batch = jr.batch("wf", tmp_path)
    job = jr.ensure_job("wf", "method", batch["id"], "pair", "Pair", ["A.D", "B.D"], {}, "fp",
                        state=J.QUEUED)
    project = tmp_path / "pair.gcws"
    project.write_text("original")
    jr.transition(job.id, J.QUEUED, J.PROCESSING)
    jr.transition(job.id, J.PROCESSING, J.CONTROL, project_path=str(project))
    assert jr.review(job.id, False, "bad data", user="reviewer")
    win.report2._journal = jr
    win._report2_job = (job.id, 1, project)
    win._report2_baseline = _report_inputs({"quant": {"limit": 1}})
    win.ws.project_path = project
    win.ws.dirty = False
    monkeypatch.setattr("gcws.ui.main_window.P.to_dict", lambda ws, path: {"quant": {"limit": 2}})
    monkeypatch.setattr("gcws.ui.main_window.P.save", lambda ws, path: (project.write_text("edited"), project)[1])
    monkeypatch.setattr("gcws.ui.main_window.QMessageBox.warning", lambda *args: None)

    assert win._save_report2_project()
    saved = jr.job(job.id)
    assert project.read_text() == "edited"
    assert saved.state == J.REJECTED and saved.edited == 1
    assert saved.reviewer == "reviewer" and saved.comment == "bad data"
    assert jr.accept_edited(job.id, user="analyst")
    queued = jr.job(job.id)
    assert queued.state == J.QUEUED and queued.review_pending["user"] == "analyst"
    assert queued.reviewer is None and queued.comment is None and queued.reviewed_at is None


def test_stale_report2_project_refuses_unsaved_display_changes(qtbot, win, tmp_path, monkeypatch):
    from gcws.automation import journal as J
    from gcws.ui.main_window import _report_inputs

    jr = J.Journal(tmp_path / "journal.sqlite")
    batch = jr.batch("wf", tmp_path)
    job = jr.ensure_job("wf", "method", batch["id"], "pair", "Pair", ["A.D", "B.D"], {}, "fp",
                        state=J.QUEUED)
    project = tmp_path / "pair.gcws"
    project.write_text("original")
    win.report2._journal = jr
    win._report2_job = (job.id, 1, project)
    win._report2_baseline = _report_inputs({"runs": [{"id": "a", "color": "#111111"}]})
    win.ws.project_path = project
    win.ws.dirty = True
    jr.update_job(job.id, revision=2, project_path=str(project))
    monkeypatch.setattr("gcws.ui.main_window.P.to_dict", lambda ws, path: {"runs": [{"id": "a", "color": "#222222"}]})
    monkeypatch.setattr("gcws.ui.main_window.P.save", lambda ws, path: pytest.fail("stale source overwritten"))
    warnings = []
    monkeypatch.setattr("gcws.ui.main_window.QMessageBox.warning", lambda *args: warnings.append(args[-1]))

    assert not win._save_report2_project()
    assert warnings and project.read_text() == "original"


def test_incomplete_report2_project_is_not_attached_or_saved(qtbot, win, tmp_path, monkeypatch):
    from gcws.automation import journal as J

    jr = J.Journal(tmp_path / "journal.sqlite")
    batch = jr.batch("wf", tmp_path)
    job = jr.ensure_job("wf", "method", batch["id"], "pair", "Pair", ["A.D", "B.D"], {}, "fp",
                        state=J.QUEUED)
    project = tmp_path / "pair.gcws"
    project.write_text("original")
    jr.transition(job.id, J.QUEUED, J.PROCESSING)
    jr.transition(job.id, J.PROCESSING, J.CONTROL, project_path=str(project))
    win.report2._journal = jr
    monkeypatch.setattr("gcws.ui.main_window.P.read", lambda path: {"runs": [{"id": "a"}, {"id": "b"}]})
    monkeypatch.setattr("gcws.ui.main_window.P.resolve_run_path",
                        lambda entry, path: project if entry["id"] == "a" else None)
    warnings = []
    monkeypatch.setattr("gcws.ui.main_window.QMessageBox.warning", lambda *args: warnings.append(args[-1]))

    assert not win.open_report2_job(job.id)
    assert win._report2_job is None and project.read_text() == "original"
    assert warnings


def test_async_report2_load_failure_does_not_bind_partial_workspace(qtbot, win, tmp_path, monkeypatch):
    from gcws.automation import journal as J

    jr = J.Journal(tmp_path / "journal.sqlite")
    batch = jr.batch("wf", tmp_path)
    job = jr.ensure_job("wf", "method", batch["id"], "pair", "Pair", ["A.D", "B.D"], {}, "fp",
                        state=J.QUEUED)
    project = tmp_path / "pair.gcws"
    project.write_text("original")
    jr.transition(job.id, J.QUEUED, J.PROCESSING)
    jr.transition(job.id, J.PROCESSING, J.CONTROL, project_path=str(project))
    win.report2._journal = jr
    monkeypatch.setattr("gcws.ui.main_window.P.read", lambda path: {"runs": [{"id": "a"}, {"id": "b"}]})
    monkeypatch.setattr("gcws.ui.main_window.P.resolve_run_path",
                        lambda entry, path: tmp_path / f"{entry['id']}.D")
    monkeypatch.setattr(win, "load_runs", lambda *args, **kwargs: win._finish_project_load())
    warnings = []
    monkeypatch.setattr("gcws.ui.main_window.QMessageBox.warning", lambda *args: warnings.append(args[-1]))

    assert win.open_report2_job(job.id)
    assert win._report2_job is None and project.read_text() == "original"
    assert any("did not all load" in text for text in warnings)


def test_batch_report_of_a_folder(qtbot, win, samples, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    from test_pipeline import MIGRATION, lib_oracle
    _load(qtbot, win, samples, ["06_", "07_", "08_", "09_", "10_", "11_", "12_", "13_"])
    ws = win.ws
    lib_oracle(ws, [s.id for s in ws.states() if s.role == "sample"])
    ws.quant["migration"] = dict(MIGRATION)
    ws.recompute_quant()
    out = tmp_path / "reports"
    out.mkdir()
    with monkeypatch.context() as m:
        m.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.No)
        win.batch_report("nias", out_dir=str(out))
        qtbot.waitUntil(lambda: getattr(win, "last_batch", None) is not None, timeout=240000)
    entries, files = win.last_batch
    assert [e["name"] for e in entries] == ["26016606_130m_min_GIOSUN1635", "26016607_170m_min_GIOSUN1635"]
    assert all(Path(e["xlsx"]).is_file() for e in entries)
    assert Path(files["batch_docx"]).is_file() and Path(files["batch_xlsx"]).is_file()
    assert all(e["state"] in ("accepted_auto", "control") for e in entries)
    assert any(r.action == "Batch report" for r in ws.audit.records)
    assert sorted(p.name for p in samples.iterdir())                         # nothing written into the batch
    assert not any(p.name.endswith("_Batch_Report.docx") for p in samples.iterdir())


def test_accepted_pair_is_listed_in_report2(qtbot, win, samples, tmp_path, monkeypatch):
    """A pair no workflow processed: Accept lists it in Report² (its report and a project copy, accepted,
    in the batch of its folder). Accepted again under a name Report² has, the analyst renames the pair or
    overwrites the entry; a workflow's report that is overwritten is hidden."""
    from types import SimpleNamespace
    from gcws.automation import journal as J
    from gcws.automation import manual as MA
    from gcws.automation import store
    from gcws.core import project as P
    from gcws.report import service as RS
    from gcws.ui.dialogs import report2 as RD
    from test_pipeline import MIGRATION, lib_oracle
    _load(qtbot, win, samples, ["07_", "08_", "11_"])
    ws = win.ws
    lib_oracle(ws, [s.id for s in ws.states() if s.role == "sample"])
    ws.quant["migration"] = dict(MIGRATION)
    ws.recompute_quant()
    a = next(s.id for s in ws.states() if s.name.startswith("07_"))
    b = next(s.id for s in ws.states() if s.name.startswith("11_"))
    jr = J.Journal(tmp_path / "journal.sqlite")
    win.report2._journal = jr
    monkeypatch.setattr(store, "jobs_dir", lambda: tmp_path / "jobs")
    monkeypatch.setattr(RS, "docx_to_pdf", _no_word)                 # the Report² preview needs no Word
    win.loaded_samples.pairRequested.emit(a, b)
    page = win.replicates.duplicate
    said = []
    ws.message.connect(said.append)

    def listed(n):
        qtbot.waitUntil(lambda: sum("listed in Report²" in t and "not listed" not in t for t in said) >= n,
                        timeout=240000)

    page.b_accept.click()
    listed(1)
    (job,) = jr.jobs(include_batch=False)
    g = page.group()
    assert MA.is_manual(job) and job.state == J.ACCEPTED_MANUAL and job.group_name == g["name"]
    assert job.reviewer == g["signoff"]["by"] and not job.export_pending
    assert Path(jr.batch_by_id(job.batch_id)["folder"]) == ws.runs[a].run.path.parent
    files = job.files[MA.MANUAL]
    assert Path(files["xlsx"]).is_file() and Path(files["docx"]).is_file()
    assert win.report2._is_pair(job) and P.read(job.project_path)["automation"]["workflow_id"] == MA.MANUAL
    assert win.report2.current == job.id and ws.project_path is None and not ws.automation
    # accepted again: the name is listed - cancelled, overwritten, renamed (the report is not made again here)
    def fake(rjob, progress=None):
        rjob.target.parent.mkdir(parents=True, exist_ok=True)
        rjob.target.write_text("report")
        return SimpleNamespace(target=rjob.target, word=None, batch=None, combined=[], summary={}, warnings=[],
                               reported=[])
    monkeypatch.setattr(RS, "generate", fake)
    asked = []
    with monkeypatch.context() as m:
        m.setattr(RD, "ask_name_clash", lambda *args: asked.append(args[1]))
        assert not win._accept_pair(g["id"])
    assert asked == [g["name"]] and len(jr.jobs(include_batch=False)) == 1
    with monkeypatch.context() as m:
        m.setattr(RD, "ask_name_clash", lambda *args: ("overwrite", args[1]))
        assert win._accept_pair(g["id"])
        listed(2)
    assert [(j.id, j.revision) for j in jr.jobs(include_batch=False)] == [(job.id, 2)]
    assert Path(jr.job(job.id).files[MA.MANUAL]["xlsx"]).parent.name == "r2"
    with monkeypatch.context() as m:
        m.setattr(RD, "ask_name_clash", lambda *args: ("rename", args[4]))      # the suggested free name
        assert win._accept_pair(g["id"])
        listed(3)
    second = next(j for j in jr.jobs(include_batch=False) if j.id != job.id)
    assert second.group_name == f"{g['name']} (2)" == page.group()["name"]
    # a workflow's report of that name: overwritten = hidden (restorable), the pair listed by hand
    wf = jr.ensure_job("wf", "m", job.batch_id, "other", "Other", ["X.D"], {}, "fp", state=J.QUEUED)
    win.replicates.rename_group(g["id"], "Other")
    with monkeypatch.context() as m:
        m.setattr(RD, "ask_name_clash", lambda *args: ("overwrite", args[1]))
        assert win._accept_pair(g["id"])
        listed(4)
    assert jr.job(wf.id).deleted
    third = MA.clash(jr, job.batch_id, "Other")
    assert MA.is_manual(third) and third.state == J.ACCEPTED_MANUAL
    # Report² never processes or delivers such an entry again
    from gcws.ui.docks.report2 import can_reprocess
    assert not can_reprocess(third)
    ws.dirty = False


def test_edit_an_accepted_report_and_update_it(qtbot, win, data, tmp_path, monkeypatch):
    """Edit in GC Workspace works for any report (here a single determination, accepted automatically);
    Update report in the status bar makes it again: accepted by the analyst, delivered again."""
    from gcws.automation import journal as J
    wf, jr, ids, batch = _seed(data, tmp_path)
    job = jr.job(ids["S-auto"])
    project = tmp_path / "s-auto.gcws"
    project.write_text("original")
    jr.update_job(job.id, project_path=str(project))
    win.report2._journal = jr
    launcher = FakeLauncher()
    monkeypatch.setattr("gcws.automation.watcher.ProcessLauncher", lambda parent=None: launcher)
    monkeypatch.setattr("gcws.automation.control.WatcherControl", lambda *a, **k: NoWatcher())
    routed, bind = [], win.open_report2_job
    monkeypatch.setattr(win, "open_report2_job", lambda jid: routed.append(jid) or True)
    win.report2.select(job.id)
    assert win.report2.b_project.text() == "Edit in GC Workspace" and win.report2.b_project.isEnabled()
    win.report2.open_project()
    assert routed == [job.id]                              # bound to the report, not just opened
    monkeypatch.setattr(win, "open_report2_job", bind)
    assert not win.r2_bar.isVisibleTo(win)
    win._report2_job = (job.id, job.revision, project)
    win._report2_baseline = {"quant": {"limit": 1}}
    win.ws.project_path, win.ws.dirty = project, True
    assert win.r2_bar.isVisibleTo(win) and "S-auto" in win.r2_label.text() and "accepted" in win.r2_label.text()
    monkeypatch.setattr("gcws.ui.main_window.P.to_dict", lambda ws, path: {"quant": {"limit": 2}})
    monkeypatch.setattr("gcws.ui.main_window.P.save", lambda ws, path: (Path(path).write_text("edited"), path)[1])
    try:
        assert win.update_report2()
        assert project.read_text() == "edited"
        queued = jr.job(job.id)
        assert queued.state == J.PROCESSING and queued.review_pending and queued.revision == 2
        assert not win.b_r2_update.isEnabled()             # being made
        job_id, spec, _kind = launcher.started[-1]
        assert spec["mode"] == "rereport" and spec["project_path"] == str(project)
        rep = wf.by_type("report")[0]
        files = {}
        for fmt in ("xlsx", "docx", "pdf"):
            p = Path(spec["out_dir"]) / f"S-auto_NIAS_Report.{fmt}"
            p.write_text("updated " + fmt)
            files[fmt] = str(p)
        newer = Path(spec["out_dir"]) / "S-auto.gcws"
        newer.write_text("regenerated")
        launcher.complete({"state": "accepted_auto", "reason": "", "files": {rep.id: files}, "project": str(newer),
                           "evidence": {}, "findings": [], "warnings": [], "timings": {}})
        after = jr.job(job.id)
        assert after.state == J.ACCEPTED_MANUAL and after.revision == 2 and after.export_state == "done"
        assert (tmp_path / "B" / batch.name / "S-auto_NIAS_Report.docx").read_text() == "updated docx"
        assert win._report2_job == (job.id, 2, newer) and win.ws.project_path == newer   # still editing it
        assert win.b_r2_update.isEnabled()
        assert not win.update_report2()                    # nothing changed since: nothing to update
        asked = []
        from PySide6.QtWidgets import QMessageBox
        monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: asked.append(a) or QMessageBox.No)
        win.ws.dirty = True
        assert win.stop_report2_edit()
        assert asked and win._report2_job is None and not win.r2_bar.isVisibleTo(win)
        assert win.ws.project_path is None                 # never saved into the report's project by accident
        assert jr.job(job.id).state == J.ACCEPTED_MANUAL
    finally:
        win._report2_job = None
        win.ws.dirty = False


def test_edit_is_not_offered_for_a_batch_report_or_a_missing_project(qtbot, win, data, tmp_path, monkeypatch):
    from gcws.automation import journal as J
    wf, jr, ids, batch = _seed(data, tmp_path)
    win.report2._journal = jr
    assert not win.open_report2_job(ids["S-auto"])         # no project file
    report = jr.ensure_job(wf.id, wf.methods()[0].id, jr.batch(wf.id, batch)["id"], J.BATCH_KEY, "Batch", [], {},
                           "fp", state=J.QUEUED)
    assert not win.open_report2_job(report.id)


# -- save as Word / Excel (Oct 2026) ---------------------------------------------------------------------

def _workbook(path: Path, title: str, width: float = 30.0):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    sh = wb.active
    sh.title = "NIAS Result"
    sh["A1"] = title
    sh["A1"].font = Font(bold=True, size=14)
    sh.merge_cells("A1:D1")
    sh["A3"], sh["B3"] = "Bisphenol A", 0.5
    sh["B3"].number_format = "0.000"
    sh.column_dimensions["A"].width = width
    hidden = wb.create_sheet("_AuditData")
    hidden.sheet_state = "hidden"
    wb.save(path)
    return path


def test_combined_workbook_keeps_each_report(tmp_path):
    from openpyxl import load_workbook
    from gcws.automation import batch as BA
    a = _workbook(tmp_path / "a.xlsx", "Sample A", 31.5)
    b = _workbook(tmp_path / "b.xlsx", "Sample B")
    out = BA.combined_workbook([("26016606_x [A/B]", a), ("26016606_x [A/B]", b)], tmp_path / "all.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["26016606_x _A_B_", "26016606_x _A_B_ (2)"]
    first = wb.worksheets[0]
    assert first["A1"].value == "Sample A" and first["A1"].font.bold and first["B3"].number_format == "0.000"
    assert "A1:D1" in {str(r) for r in first.merged_cells.ranges} and first.column_dimensions["A"].width == 31.5


def test_save_a_report_as_word_and_excel(qtbot, data, tmp_path, monkeypatch):
    from gcws.automation import batch as BA
    wf, jr, ids, batch = _seed(data, tmp_path)
    dock = _dock(jr, monkeypatch, qtbot)
    dock.select(ids["S-control"])
    assert dock.a_save_word.isEnabled() and dock.b_save.isEnabled()
    word = dock.save_as("docx", tmp_path / "out" / "mine.docx") if (tmp_path / "out").mkdir() is None else None
    assert word.read_text() == "docx"
    assert dock.save_as("xlsx", tmp_path / "out" / "mine.xlsx").read_text() == "xlsx"
    # a report delivered without Word: its Word report is made from the Excel report
    job = jr.job(ids["S-auto"])
    node = next(iter(job.files))
    Path(job.files[node]["docx"]).unlink()
    made = []
    monkeypatch.setattr(BA, "combined_word", lambda xlsx, target: made.append(xlsx) or Path(target).write_text("w"))
    dock.select(ids["S-auto"])
    target = dock.save_as("docx", tmp_path / "out" / "auto.docx")
    qtbot.waitUntil(lambda: target.is_file())
    assert [Path(p).name for p in made[0]] == ["S-auto_NIAS_Report.xlsx"]
    qtbot.waitUntil(lambda: "Saved" in dock.bar_text.text())
    asked = []
    monkeypatch.setattr("gcws.ui.docks.report2.QFileDialog.getSaveFileName",
                        lambda *a, **k: asked.append(a[2]) or ("", ""))
    assert dock.save_as("xlsx") is None and asked[0].endswith("S-auto_NIAS_Report.xlsx")   # cancelled


def test_save_the_batch_report_as_word_and_excel(qtbot, data, tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from gcws.automation import batch as BA
    wf, jr, ids, batch = _seed(data, tmp_path)
    for name in ("S-control", "S-auto"):                   # real Excel reports
        job = jr.job(ids[name])
        _workbook(Path(job.files[next(iter(job.files))]["xlsx"]), name)
    dock = _dock(jr, monkeypatch, qtbot)
    bid = jr.job(ids["S-auto"]).batch_id
    menu = dock._batch_menu(bid)
    labels = [a.text() for a in menu.actions()]
    assert "Save batch report as Word..." in labels and "Save batch report as Excel..." in labels
    out = tmp_path / "out"
    out.mkdir()
    summary = dock.save_batch(bid, "xlsx", out / "batch.xlsx")
    combined = out / f"{BA.batch_stem(batch.name)}_Sample_Reports.xlsx"
    qtbot.waitUntil(lambda: "Saved" in dock.bar_text.text() or "could not" in dock.bar_text.text(), timeout=10000)
    assert combined.name in dock.bar_text.text(), dock.bar_text.text()
    assert load_workbook(combined).sheetnames == ["S-control", "S-auto"]
    rows = [r[0] for r in load_workbook(summary).active.iter_rows(min_row=4, values_only=True)]
    assert rows == ["S-control", "S-auto", "S-wait", "S-noblank", "S-failed"]
    made = []
    monkeypatch.setattr(BA, "combined_word", lambda xlsx, target: made.append(xlsx) or Path(target).write_text("w"))
    word = dock.save_batch(bid, "docx", out / "batch.docx")   # no batch report yet: made now
    qtbot.waitUntil(lambda: word.is_file() and "Saved" in dock.bar_text.text())
    assert [Path(p).name for p in made[0]] == ["S-control_NIAS_Report.xlsx", "S-auto_NIAS_Report.xlsx"]


# -- the register (Oct 2026) -----------------------------------------------------------------------------

def test_done_means_accepted_and_delivered(tmp_path):
    from gcws.automation import journal as J
    from gcws.ui.docks.report2 import is_done
    jr = J.Journal(tmp_path / "j.sqlite")
    b = jr.batch("wf", tmp_path)
    job = jr.ensure_job("wf", "m", b["id"], "s", "S", ["S.D"], {}, "fp", state=J.QUEUED)
    jr.transition(job.id, J.QUEUED, J.PROCESSING)
    jr.transition(job.id, J.PROCESSING, J.ACCEPTED_AUTO, export_pending=1)
    assert not is_done(jr.job(job.id))                         # accepted, not delivered yet
    jr.update_job(job.id, export_pending=0, export_state="done")
    assert is_done(jr.job(job.id))
    jr.review(job.id, False, "bad")
    assert not is_done(jr.job(job.id))                         # rejected: not done
    manual = jr.record_manual("jm", workflow_id=J.MANUAL_WORKFLOW, method_node="manual", batch_id=b["id"],
                              group_key="manual:p", group_name="P", members=[], files={}, project_path="", job_dir="",
                              evidence={}, findings=[], summary={}, reviewer="a")
    assert is_done(manual)                                     # no workflow delivers it: accepted is done


def test_register_lists_every_sample_done_or_not(qtbot, data, tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from gcws.automation import journal as J
    wf, jr, ids, batch = _seed(data, tmp_path)
    jr.update_job(ids["S-auto"], export_state="done")
    dock = _dock(jr, monkeypatch, qtbot)
    dock.set_mode("register")
    assert dock.b_register.text() == "Register (5)" and dock.list_stack.currentWidget() is dock.register
    assert dock.done_filter.isVisibleTo(dock) and dock.b_export_register.isVisibleTo(dock)
    rows = {dock.register.item(r, 0).text(): [dock.register.item(r, c).text() for c in range(9)]
            for r in range(dock.register.rowCount())}
    assert set(rows) == {"S-control", "S-auto", "S-wait", "S-noblank", "S-failed"}
    assert rows["S-auto"][1] == batch.name and rows["S-auto"][2] == "Lab" and rows["S-auto"][4] == "Done"
    assert rows["S-control"][4] == "Not done" and rows["S-control"][3] == "Control needed"
    dock.done_filter.setCurrentIndex(dock.done_filter.findData("open"))
    dock.refresh()
    assert dock.register.rowCount() == 4
    dock.done_filter.setCurrentIndex(dock.done_filter.findData("done"))
    dock.refresh()
    assert [dock.register.item(r, 0).text() for r in range(dock.register.rowCount())] == ["S-auto"]
    dock.done_filter.setCurrentIndex(0)
    dock.set_filter("control")                                 # the chips filter the register too
    assert [dock.register.item(r, 0).text() for r in range(dock.register.rowCount())] == ["S-control"]
    dock.select(ids["S-control"])                              # the same actions as in the list
    assert dock.current == ids["S-control"] and dock.a_accept.isEnabled() and dock.review(True)
    assert jr.job(ids["S-control"]).state == J.ACCEPTED_MANUAL
    dock.set_filter("control")
    out = dock.export_register(tmp_path / "register.xlsx")
    sheet = load_workbook(out).active
    assert [c.value for c in sheet[1]][:5] == ["Sample", "Batch", "Workflow", "Status", "Done"]
    assert sheet.max_row == 6                                  # header + every sample
    dock.set_mode("todo")
    assert not dock.done_filter.isVisibleTo(dock) and dock.list_stack.currentWidget() is not dock.register

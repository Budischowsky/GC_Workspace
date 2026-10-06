import copy
import struct
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest

from gcws.core.ident import Identification, IdentificationSet
from gcws.core.model import Baseline, Peak
from gcws.quant import hs
from gcws.quant.service import compute, mode_unit


def workspace():
    peaks = [Peak(i + 1, i + 1.2, i + 1.1, Baseline("line", i + 1, 0, i + 1.2, 0),
                  area=100 if i < 7 else 200) for i in range(8)]
    ids = IdentificationSet([Identification(p.apex_rt, name=f"STD{i + 1}" if i < 7 else "Analyte",
                                            status="Accepted", score=95) for i, p in enumerate(peaks)])
    defs = hs.default_defs()
    for i, d in enumerate(defs):
        d.update(name=f"STD{i + 1}", target_rt=peaks[i].apex_rt)
    run = SimpleNamespace(path=Path("sample.qgd"), ms=SimpleNamespace(rt=np.arange(10.), tic=lambda: np.ones(10)))
    st = SimpleNamespace(id="sample", name="Sample", role="sample", run=run, results={"TIC": SimpleNamespace(peaks=peaks)},
                         ident_set=lambda _: ids, blanks=[], blanks_istd=[])
    cfg = dict(istd_defs=defs, unit="µg/HS", samples={"sample": {"area_dm2": 0.5, "mass_g": 2}})
    return SimpleNamespace(quant={"mode": "hs_screening", "hs": cfg}, runs={st.id: st}, states=lambda: [st])


@pytest.mark.parametrize("unit,expected", [("µg/HS", 2), ("µg/dm²", 4), ("µg/g", 1)])
def test_hs_units_and_average(unit, expected):
    ws = workspace()
    ws.quant["hs"]["unit"] = unit
    result = compute(ws)
    assert not result.errors
    assert result.rows["sample"][7]["conc"] == pytest.approx(expected)
    assert mode_unit(ws.quant) == unit


def test_single_and_missing_standards():
    ws = workspace()
    for d in ws.quant["hs"]["istd_defs"]:
        d["quantify"] = d["code"] == "HS2"
    ws.quant["hs"]["use_mean_area"] = False
    ws.runs["sample"].results["TIC"].peaks[1].area = 50
    assert compute(ws).rows["sample"][7]["conc"] == 4
    ws.quant["hs"]["istd_bindings"] = {"sample": {"HS2": None}}
    result = compute(ws)
    assert "sample" in result.errors
    assert result.rows["sample"][7]["conc"] is None


def test_blank_correction_and_istd_protection():
    ws = workspace()
    st = ws.runs["sample"]
    blank = copy.copy(st)
    blank.name = "Blank"
    blank.results = copy.deepcopy(st.results)
    for p in blank.results["TIC"].peaks:
        p.area = 50
    st.blanks = ["blank"]
    ws.runs["blank"] = blank
    result = compute(ws)
    assert result.rows[st.id][7]["conc"] == pytest.approx(1.5)
    assert result.rows[st.id][0]["blank_area"] == 0
    assert result.rows[st.id][0]["conc"] == 1


def external_workspace(*areas):
    """Samples without ISTD plus one Standard run per entry of ``areas`` (the area of every HS standard)."""
    ws = workspace()
    ws.quant["hs"]["calibration"] = "external"
    sample = ws.runs["sample"]
    sample.results = {"TIC": SimpleNamespace(peaks=[copy.copy(sample.results["TIC"].peaks[7])])}
    sample.ident_set = lambda _: IdentificationSet([Identification(8.1, name="Analyte", status="Accepted")])
    states = [sample]
    for n, area in enumerate(areas):
        st = copy.copy(workspace().runs["sample"])
        st.id, st.name, st.role = f"std{n}", f"Std {n + 1}", "standard"
        for p in st.results["TIC"].peaks[:7]:
            p.area = area
        ws.runs[st.id] = st
        states.append(st)
    ws.states = lambda: states
    return ws


def test_external_calibration_averages_standard_runs():
    ws = external_workspace(80, 120)
    result = compute(ws)
    assert not result.errors
    sample = result.samples["sample"]
    assert sample.mean_factor == pytest.approx(7 / 700)            # Σ 1 µg ÷ Σ mean area 100
    assert result.rows["sample"][0]["conc"] == pytest.approx(2)     # area 200 → 2 µg/HS
    assert result.rows["sample"][0]["istd"] == ""
    assert [s["area"] for s in sample.standards] == pytest.approx([100] * 7)
    assert sample.standards[0]["status"] == "Mean of 2 Standard runs"
    assert sample.meta["calculation"].startswith("External 1-point calibration")
    assert "Std 1, Std 2" in sample.meta["calculation"]
    # a Standard run shows its own standards and is quantified with the common factor
    std = result.samples["std0"]
    assert [s["area"] for s in std.standards] == [80] * 7
    assert result.rows["std0"][0]["istd"] == "HS1"
    assert result.rows["std0"][7]["conc"] == pytest.approx(2)


def test_external_calibration_single_standard_and_weights():
    ws = external_workspace(100)
    ws.runs["std0"].results["TIC"].peaks[1].area = 50
    ws.quant["hs"]["istd_defs"][1]["concentration"] = 2
    for d in ws.quant["hs"]["istd_defs"]:
        d["quantify"] = d["code"] == "HS2"
    ws.quant["hs"]["use_mean_area"] = False
    result = compute(ws)
    assert result.rows["sample"][0]["conc"] == pytest.approx(8)       # 200 × 2 µg ÷ 50


def test_external_calibration_needs_every_standard_in_every_run():
    ws = external_workspace(100, 100)
    ws.quant["hs"]["istd_bindings"] = {"std1": {"HS3": None}}
    result = compute(ws)
    assert "Identify or bind HS3 in Standard run Std 2" in result.errors["sample"]
    assert result.rows["sample"][0]["conc"] is None
    assert result.samples["sample"].standards[2]["status"] == "Not found in Std 2"
    ws.quant["hs"]["istd_defs"][2]["quantify"] = False                # an inactive standard may be missing
    assert not compute(ws).errors


def test_external_calibration_without_standard_run():
    ws = external_workspace()
    result = compute(ws)
    assert result.errors["sample"] == "Mark at least one run as Standard for external calibration"
    assert result.rows["sample"][0]["conc"] is None
    ws.quant["hs"]["calibration"] = "internal"
    assert "Identify or bind" in compute(ws).errors["sample"]          # the internal mode looks in the sample


def test_external_calibration_blank_and_report(tmp_path):
    from gcws.report.service import ReportJob, generate
    from gcws.quant.nias_bridge import make_settings
    from openpyxl import load_workbook
    ws = external_workspace(100)
    blank = copy.copy(ws.runs["sample"])
    blank.id, blank.name, blank.role = "blank", "Blank", "blank"
    blank.results = copy.deepcopy(blank.results)
    blank.results["TIC"].peaks[0].area = 50
    ws.runs["sample"].blanks = ["blank"]
    ws.runs["blank"] = blank
    result = compute(ws)
    assert result.rows["sample"][0]["conc"] == pytest.approx(1.5)
    job = ReportJob("hs_screening", [result.samples["sample"]], ["Sample"], make_settings(),
                    tmp_path / "HS.xlsx", tmp_path / "HS.docx", None, record_seen=False)
    wb = load_workbook(generate(job).target)
    assert wb["HS Result"]["E6"].value == pytest.approx(1.5)
    assert wb["HS calculation"]["G2"].value.startswith("External 1-point calibration")
    assert wb["HS standards"]["H2"].value == "Mean of 1 Standard run"


@pytest.mark.parametrize("bad", [None, 0, -1, float("nan"), float("inf")])
def test_invalid_normalization_never_produces_result(bad):
    ws = workspace()
    ws.quant["hs"].update(unit="µg/g", samples={"sample": {"mass_g": bad}})
    result = compute(ws)
    assert "sample" in result.errors
    assert result.rows["sample"][7]["conc"] is None


@pytest.mark.parametrize("width", [1, 2, 3, 4])
def test_qgd_binary_scan(width):
    from gcws.io.shimadzu import decode_scan
    header = bytearray(32)
    struct.pack_into("<I", header, 4, 150000)
    struct.pack_into("<HH", header, 20, width, 2)
    payload = b"".join(struct.pack("<H", m * 20) + n.to_bytes(width, "little") for m, n in [(43, 12), (57, 99)])
    mz, ab = decode_scan(bytes(header) + payload, 150000)
    assert mz.tolist() == [43, 57]
    assert ab.tolist() == [12, 99]
    with pytest.raises(ValueError):
        decode_scan(bytes(header) + payload[:-1], 150000)


def test_actual_qgd_sample():
    from gcws.io.run_loader import load_run
    from gcws.io.folders import list_runs, sources
    p = Path(__file__).resolve().parents[2] / "NIAS Working/samples/26012850_Sample1_A.qgd"
    if not p.exists():
        pytest.skip("Local QGD fixture unavailable")
    run = load_run(p)
    assert run.fid is None and run.role == "sample"
    assert run.ms.n_scans == 4420
    assert run.ms.rt[[0, -1]].tolist() == pytest.approx([2.5, 24.595])
    assert np.array_equal(run.signal("TIC").y, run.ms.stored_tic)      # the instrument TIC
    assert p in list_runs(p.parent)
    assert sources(p) == {"MS": p.name}


@pytest.mark.parametrize("unit,expected", [("µg/HS", 2), ("µg/dm²", 4), ("µg/g", 1)])
def test_report_quantities_replace_area(tmp_path, unit, expected):
    from gcws.report.service import ReportJob, generate
    from gcws.quant.nias_bridge import make_settings
    from openpyxl import load_workbook
    from docx import Document
    ws = workspace()
    ws.quant["hs"]["unit"] = unit
    result = compute(ws)
    job = ReportJob("hs_screening", list(result.samples.values()), ["Sample"], make_settings(),
                    tmp_path / "HS.xlsx", tmp_path / "HS.docx", None, record_seen=False)
    report = generate(job)
    wb = load_workbook(report.target)
    assert wb["HS Result"]["E5"].value == unit
    assert wb["HS Result"]["E6"].value == expected
    assert report.rows == 1
    assert wb["HS standards"].max_row == 8
    doc = Document(report.word)
    assert doc.tables[0].cell(4, 4).text == unit
    assert doc.tables[0].cell(5, 4).text == f"{expected:.4f}"
    # Header and body cell widths must agree or Word expands the table off-page.
    assert [c.width for c in doc.tables[0].rows[4].cells] == [c.width for c in doc.tables[0].rows[5].cells]


def test_hs_ui_switch_and_project(qtbot, tmp_path):
    from gcws.ui.workspace import Workspace
    from gcws.ui.docks.quant import QuantDock
    from gcws.core import project
    ws = Workspace()
    original = copy.deepcopy(ws.quant)
    dock = QuantDock(ws)
    qtbot.addWidget(dock)
    dock.mode.setCurrentIndex(dock.mode.findData("hs_screening"))
    dock._mode_changed()
    assert ws.panel_keys[0] == "TIC" and ws.signal_key == "TIC"
    assert dock.hs_panel.defs.rowCount() == 7
    assert ws.quant.get("settings") == original.get("settings")
    saved = project.read(project.save(ws, tmp_path / "hs.gcws"))
    assert saved["quant"]["mode"] == "hs_screening"
    ws.project_undo.undo()
    assert ws.quant == original
    assert ws.panel_keys == ["FID", "TIC"]


def test_hs_panel_calibration_choice(qtbot):
    from gcws.ui.workspace import Workspace
    from gcws.ui.docks.quant import QuantDock
    ws = Workspace()
    dock = QuantDock(ws)
    qtbot.addWidget(dock)
    dock.mode.setCurrentIndex(dock.mode.findData("hs_screening"))
    dock._mode_changed()
    panel = dock.hs_panel
    assert panel.calibration.currentData() == "internal"
    panel.calibration.setCurrentIndex(panel.calibration.findData("external"))
    panel.save_inputs()
    assert ws.quant["hs"]["calibration"] == "external"
    panel.refresh()
    assert "Standard runs" in panel.note.text()
    assert not panel.bind_buttons[1].isEnabled()           # no run loaded: nothing to bind
    ws.project_undo.undo()
    panel.refresh()
    assert panel.calibration.currentData() == "internal"


def test_hs_full_window_qgd(qtbot, monkeypatch, tmp_path):
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QMessageBox
    from gcws.ui.main_window import MainWindow
    from gcws.core import project, proc_method
    p = Path(__file__).resolve().parents[2] / "NIAS Working/samples/26012850_Sample1_A.qgd"
    if not p.exists():
        pytest.skip("Local QGD fixture unavailable")
    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.Yes))
    notices = []
    monkeypatch.setattr(QMessageBox, "information", staticmethod(lambda *a, **k: notices.append(a[2])))
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: notices.append(a[2])))
    QSettings().setValue("integration/solvent_cut", False)
    win = MainWindow()
    qtbot.addWidget(win)
    win.show()
    win.load_runs([str(p)])
    qtbot.waitUntil(lambda: win.loading == 0 and win.ws.active is not None, timeout=30000)
    win.ws.process_runs()
    ws, st = win.ws, win.ws.active
    win.quant.mode.setCurrentIndex(win.quant.mode.findData("hs_screening"))
    win.quant._mode_changed()
    assert ws.panel_keys[0] == "TIC"
    assert st.run.fid is None
    assert len(st.results["TIC"].peaks) >= 7
    original = copy.deepcopy(ws.quant.get("settings", {}))
    ws.set_solvent_cut(True, 3.0)
    assert ws.solvent_cut(st, "TIC") == 3.0
    assert ws.quant.get("settings", {}) == original
    q = copy.deepcopy(ws.quant)
    defs = hs.default_defs()
    # Test-only bindings validate plumbing, not this sample's chemical identities.
    peaks = st.results["TIC"].peaks
    bindings = {}
    for i, d in enumerate(defs):
        d["name"] = f"TEST STANDARD {i + 1}"
        bindings[d["code"]] = peaks[i].apex_rt
    q.setdefault("hs", {}).update(istd_defs=defs, istd_bindings={st.id: bindings},
                                   samples={st.id: {"area_dm2": 1, "mass_g": 2}})
    ws.push_quant("Test bindings", q)
    assert not ws.quant_result.errors
    assert ws.quant_rows(st.id)
    win.replicates.show_pair(st.id)
    win.replicates.duplicate.compare()
    assert win.replicates.duplicate.quant_signal() == "TIC"
    method = proc_method.collect(win, "HS test")
    assert len(method["sections"]["quant"]["hs"]["istd_defs"]) == 7
    assert "samples" not in method["sections"]["quant"]["hs"]
    saved = project.save(ws, tmp_path / "qgd.gcws")
    data = project.read(saved)
    assert data["runs"][0]["fingerprint"][p.name][0] == p.stat().st_size
    win.open_project(saved)
    qtbot.waitUntil(lambda: win.loading == 0 and win._pending_project is None, timeout=30000)
    win.ws.recompute_quant()
    assert win.ws.active is not None
    assert win.ws.panel_keys[0] == "TIC"
    assert not win.ws.quant_result.errors
    assert not notices
    from gcws.ui.workers import load_and_integrate
    run, results, delay = load_and_integrate(p, "sample", ws.methods)
    extra = ws.add_run(run, results, delay=delay)
    assert all(p.apex_rt >= 3 for p in extra.results["TIC"].peaks)
    win.ws.dirty = False
    win.close()


def test_ambiguous_or_duplicate_istd_is_blocked():
    ws = workspace()
    cfg = ws.quant["hs"]
    cfg["istd_defs"][1]["target_rt"] = cfg["istd_defs"][0]["target_rt"]
    result = compute(ws)
    assert "sample" in result.errors
    assert result.rows["sample"][7]["conc"] is None


def test_qgd_saturated_width_and_wrong_time():
    from gcws.io.shimadzu import decode_scan
    header = bytearray(32)
    struct.pack_into("<I", header, 4, 150000)
    struct.pack_into("<HH", header, 20, 1, 1)
    block = bytes(header) + struct.pack("<HI", 860, 0x8000007b)
    assert decode_scan(block, 150000)[1].tolist() == [123]
    with pytest.raises(ValueError, match="time"):
        decode_scan(block, 150001)


def test_derived_tic_uses_same_quantities_once():
    from gcws.ui.workspace import Workspace
    ws = workspace()
    ws.quant_result = compute(ws)
    ws.signal_key = "TIC - Blank"
    ws.effective_key = lambda st, key: key
    st = ws.runs["sample"]
    st.results[ws.signal_key] = copy.deepcopy(st.results["TIC"])
    st.results[ws.signal_key].peaks[-1].area = 10
    ws.result = lambda rid, key: ws.runs[rid].results.get(key)
    assert Workspace.quant_rows(ws, st.id)[7]["conc"] == 2
    assert Workspace.quant_rows(ws, st.id, "FID") == {}

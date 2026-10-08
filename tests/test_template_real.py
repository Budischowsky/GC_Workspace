"""The Template report on real runs: the NIAS layout gives the NIAS Report's substances and numbers, and a
template works in every quantification mode."""
import copy

import pytest

from test_report import _ws_with, qapp  # noqa: F401  (fixture)
from test_report_assemble import MIGRATION


@pytest.fixture(scope="module")
def ws(samples, qapp):
    w = _ws_with(samples, ["06_", "07_", "08_", "11_"], qapp)
    w.quant["migration"] = dict(MIGRATION)
    return w


def _group(ws):
    ids = [s.id for s in ws.states() if s.role == "sample"]
    return {"id": "g", "name": "26016606", "members": ids, "policy": "all"}


def _build(ws, tpl, group=None):
    from gcws.report import assemble as AS
    from gcws.report import table as TB
    from gcws.report import template as TP
    from gcws.report import template_data as TD
    g = group or _group(ws)
    data = TD.collect(ws, g["members"], g, fields=AS.template_fields(TP.normalise(tpl)))
    return data, TB.build(data, tpl)


def _legacy_rows(path):
    from openpyxl import load_workbook
    sh = load_workbook(path)["NIAS Result"]
    out = []
    for r in range(6, sh.max_row + 1):
        name = sh.cell(r, 2).value
        if not name or sh.cell(r, 6).value is None:
            continue
        out.append((str(name), round(float(sh.cell(r, 6).value), 3), round(float(sh.cell(r, 5).value or 0), 4)))
    return out


def test_the_nias_layout_matches_the_nias_report(ws, tmp_path):
    from gcws.report import assemble as AS
    from gcws.report import service as RS
    from gcws.report import template as TP
    g = _group(ws)
    members, samples = AS.prepare(ws, "nias", g)
    job = AS.build_job(ws, "nias", g, tmp_path / "x_NIAS_Report.xlsx", members=members, samples=samples,
                       preview=True)
    RS.generate(job)
    legacy = _legacy_rows(job.target)
    data, table = _build(ws, TP.preset("NIAS"))
    assert data.n == 2 and table.subtitle == "PA 26.007, double determination"
    head = [c.header for c in table.columns]
    i_name, i_kg, i_dm2 = head.index("Name"), head.index("Conc. mg/kg"), head.index("Conc. mg/dm²")
    ours = [(r.cells[i_name].value, round(r.cells[i_kg].value, 3), round(r.cells[i_dm2].value or 0, 4))
            for r in table.rows]
    assert legacy, "the NIAS report has substances"
    # the NIAS Report names a CAS-only row from PubChem; the Template report does not look names up
    named = {n for n, _kg, _dm2 in ours if n}
    theirs = sorted(x for x in legacy if x[0] in named or x[0].startswith("Sum of"))
    mine = sorted(x for x in ours if x[0])
    assert [x[:2] for x in theirs] == [x[:2] for x in mine]
    for a, b in zip(theirs, mine):
        # the NIAS Report rounds the mg/dm² of a repeated-substance sum to 3 decimals
        assert a[2] == b[2] or (a[0].startswith("Sum of") and abs(a[2] - b[2]) < 0.001), (a, b)
    assert [r.kind for r in table.rows].count("sum") == sum(1 for x in legacy if x[0].startswith("Sum of"))


def test_the_determinations_sheet_and_the_sml_check_need_no_further_column(ws):
    """The Determinations sheet lists each determination's Conc. 1 / 2, and the SML check compares the
    mg/kg, whether or not the template shows those columns."""
    from gcws.report import template as TP
    data, table = _build(ws, TP.preset("NIAS"))               # columns mg/dm² and mg/kg, not Conc. 1 / 2
    assert set(data.values) == set(TP.PLACEHOLDERS)           # every header value can be inserted
    head, lines = table.determinations[0], table.determinations[1:]
    assert head[3:8] == ["A [mg/kg]", "A [mg/dm²]", "A outlier", "B [mg/kg]", "B [mg/dm²]"]
    assert lines and all(line[3] is not None or line[6] is not None for line in lines)
    assert all(isinstance(line[4], float) for line in lines if line[3] is not None)
    tpl = TP.preset("Empty")
    tpl["columns"] = [{"field": "name"}, {"field": "sml_check"}]
    _data, table = _build(ws, tpl)
    assert table.rows and all(r.cells[1].value in ("no SML", "≤ SML", "> SML") for r in table.rows)


def test_every_quantification_mode_reports(ws, tmp_path):
    from gcws.report import template as TP
    from gcws.report import template_report as TR
    from types import SimpleNamespace
    saved = copy.deepcopy(ws.quant)
    tpl = TP.preset("Empty")
    tpl["columns"] = [{"field": f} for f in ("rt", "name", "cas", "conc", "report_unit_1", "report_unit_2",
                                             "area", "corr_area", "height", "sn", "reldiff")]
    tpl["columns"][3]["view"] = "each_mean"
    try:
        for mode, extra in (("nias_mgkg", {}), ("istd_conc", {"istd_conc_value": 10.0, "unit": "µg/mL"}),
                            ("total_ugl", {}), ("area_pct", {}),
                            ("extraction", {"method": {"sample_type": "solid", "amount": 1.0,
                                                       "extract_volume_ml": 10.0, "spike_ul": 10.0,
                                                       "units": ["mg/kg", "µg/L"]}})):
            q = copy.deepcopy(saved)
            q.update(mode=mode, **extra)
            ws.quant = q
            ws.recompute_quant()
            data, table = _build(ws, tpl)
            head = [c.header for c in table.columns]
            assert "Area" in head and "Diff. %" in head, mode
            assert any(h.startswith("Conc.") and h.endswith(" A") for h in head), (mode, head)
            if mode in ("istd_conc", "area_pct"):
                assert not any(h.startswith("Conc. 2") for h in head) and table.warnings
            values = [r.cells[head.index("Area")].value for r in table.rows]
            assert table.rows and any(v for v in values), mode
            target = tmp_path / mode / "x.xlsx"
            res = TR.generate(SimpleNamespace(kind="template", table=table, target=target,
                                              word=target.with_suffix(".docx"), audit=[], notes=[],
                                              record_seen=False, sample_key="x"))
            assert res.target.exists() and res.word.exists()
    finally:
        ws.quant = saved
        ws.recompute_quant()


def test_extraction_matches_the_quantification_report(ws, tmp_path):
    from openpyxl import load_workbook
    from gcws.report import assemble as AS
    from gcws.report import service as RS
    from gcws.report import template as TP
    saved = copy.deepcopy(ws.quant)
    try:
        q = copy.deepcopy(saved)
        q.update(mode="extraction", method={"sample_type": "solid", "amount": 1.0, "extract_volume_ml": 10.0,
                                            "spike_ul": 10.0, "units": ["mg/kg", "µg/L"]},
                 method_samples={_group(ws)["members"][1]: {"amount": 2.0}})
        ws.quant = q
        ws.recompute_quant()
        g = _group(ws)
        members, samples = AS.prepare(ws, "quant", g)
        job = AS.build_job(ws, "quant", g, tmp_path / "q.xlsx", members=members, samples=samples, preview=True)
        RS.generate(job)
        sh = load_workbook(job.target)["Result"]
        fixed = sorted((str(sh.cell(r, 2).value), round(sh.cell(r, 5).value, 4), round(sh.cell(r, 6).value, 2))
                       for r in range(6, sh.max_row + 1) if sh.cell(r, 5).value is not None)
        _data, table = _build(ws, TP.preset("Quantification"))
        head = [c.header for c in table.columns]
        assert head[-2:] == ["Conc. 1 [mg/kg]", "Conc. 2 [µg/L]"]
        ours = sorted((r.cells[1].value, round(r.cells[4].value, 4), round(r.cells[5].value, 2)) for r in table.rows)
        assert fixed and ours == fixed
    finally:
        ws.quant = saved
        ws.recompute_quant()


def test_the_template_report_through_the_report_job(ws, tmp_path):
    from gcws.report import assemble as AS
    from gcws.report import service as RS
    from gcws.report import template as TP
    g = _group(ws)
    saved = dict(ws.quant)
    try:
        ws.quant.pop(TP.QUANT_KEY, None)
        with pytest.raises(AS.ReportNotPossible) as exc:
            AS.prepare(ws, "template", g)
        assert exc.value.code == "no_template"
        with pytest.raises(AS.ReportNotPossible) as exc:
            AS.prepare(ws, "template", g, template=TP.preset("Empty"))
        assert exc.value.code == "empty_template"
        tpl = TP.stamped(TP.preset("NIAS"), "Customer A")
        ws.quant[TP.QUANT_KEY] = tpl
        assert RS.kind_fits("template", {"mode": "hs_screening"}) and RS.preview_kind(ws.quant) == "template"
        assert AS.default_target(ws, "template", g["members"]).name == "26016606_Customer_A_Report.xlsx"
        assert RS.suffix("template", None, TP.stamped(tpl, "NIAS")) == "_NIAS_Template_Report"
        members, samples = AS.prepare(ws, "template", g, recompute=False)
        job = AS.build_job(ws, "template", g, tmp_path / "t.xlsx", members=members, samples=samples, preview=True,
                           method_name="NIAS EtOH")
        assert job.template["name"] == "Customer A" and job.table.rows and job.features is None
        res = RS.generate(job)
        assert res.target.exists() and res.word.exists() and res.rows == len(job.table.reported)
        assert res.combined and "name" in res.combined[0]
    finally:
        ws.quant.clear()
        ws.quant.update(saved)

"""P59: one sample processed unattended from a copied batch folder (the job process's work), the
batch report and the delivery to the target folders. Real data, only ever on copies."""
import json
import shutil
from pathlib import Path

import pytest

from conftest import run_dir

MIGRATION = {"analyst": "Test", "migration_cell": "cell", "occupancy": "1", "simulant": "EtOH 95 %",
             "temperature": "60 C", "duration": "10 d", "cell_area_dm2": 0.51, "occupancy_factor": 1,
             "volume_ml": 100, "ov_ratio": 6.0}
METHOD = {"format": "gcws-processing-method", "version": 1, "name": "Test NIAS", "sections": {
    "quant": {"mode": "nias_mgkg", "detector": "FID", "ms_solvent": {"enabled": False, "end": 0.0}},
    "migration": MIGRATION}}
A, B = "07_26016606_130m_min_GIOSUN1635_A.D", "11_26016606_130m_min_GIOSUN1635_B.D"


@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


@pytest.fixture()
def data(tmp_path, monkeypatch):
    from gcws import paths
    d = tmp_path / "data"
    d.mkdir()
    shutil.copy2(paths.DATA / "settings.json", d / "settings.json")
    monkeypatch.setattr(paths, "DATA", d)
    return d


def copy_batch(samples: Path, dst: Path, prefixes) -> Path:
    b = dst / samples.name
    b.mkdir(parents=True)
    for p in prefixes:
        src = run_dir(p)
        shutil.copytree(src, b / src.name)
    for f in samples.glob("*Sequence Log*"):
        shutil.copy2(f, b / f.name)
    return b


def snapshot(folder: Path) -> dict:
    return {str(p.relative_to(folder)): (p.stat().st_size, p.stat().st_mtime_ns)
            for p in folder.rglob("*") if p.is_file()}


def lib_oracle(ws, run_ids):
    """Library search stand-in: the hits of the old ChemStation LIB, bound by RT (as test_report)."""
    from gcws.core.ident import Identification
    from gcws.core.model import FID
    from gcws.quant.nias_bridge import engine
    eng = engine()
    for rid in run_ids:
        st = ws.runs[rid]
        d = st.run.path
        _tic, _fid, pbm = eng.parse_results(str(d / "RESULTS.CSV"))
        lib = eng.parse_library_report(str(d / "LIB"))
        res = ws.result(rid, FID)
        for p in pbm:
            hits = lib.get(p.peak) or p.hits
            if not hits:
                continue
            peak = min(res.peaks, key=lambda x: abs(x.apex_rt - (p.rt + st.delay_value)))
            if abs(peak.apex_rt - (p.rt + st.delay_value)) > 0.02:
                continue
            st.ident_set(FID).set(Identification(
                apex_rt=peak.apex_rt, name=hits[0].name, cas=hits[0].cas, score=hits[0].quality,
                hits=[{"name": h.name, "cas": h.cas, "score": h.quality} for h in hits], source="LIB (test)"))


def spec(batch: Path, out: Path, **kw) -> dict:
    s = {"job_id": "j1", "revision": 1, "mode": "full", "method": METHOD, "batch_folder": str(batch),
         "group": {"key": "26016606:x", "name": "26016606_130m_min_GIOSUN1635", "members": [A, B]},
         "blanks": {"blank": ["08_EtOH.D", "13_EtOH.D"], "blank_istd": ["06_EtOH_ISTD.D"]},
         "reports": [{"node": "r1", "kind": "nias", "formats": ["xlsx", "docx", "dd"]}],
         "rules": None, "auto_accept": True, "has_review": True, "out_dir": str(out), "search": True,
         "istd_detect": True, "min_confidence": "high", "require_blank": "blank"}
    s.update(kw)
    return s


def test_process_one_sample(qapp, samples, data, tmp_path):
    from gcws.automation import pipeline as PL
    from gcws.automation import headless as H
    batch = copy_batch(samples, tmp_path / "watch", ["06_", "07_", "08_", "11_", "13_"])
    before = snapshot(batch)
    res = PL.run_job(spec(batch, tmp_path / "job"), identify=lib_oracle)
    assert res.state in (PL.ACCEPTED_AUTO, PL.CONTROL), res.reason
    files = res.files["r1"]
    assert set(files) == {"xlsx", "docx", "dd"} and all(Path(p).is_file() for p in files.values())
    assert Path(res.project).is_file() and Path(res.project).parent == tmp_path / "job"
    ev = res.evidence
    assert [m["name"] for m in ev["members"]] == [A[:-2], B[:-2]]
    assert all(m["mean_factor"] and m["blank_ok"] for m in ev["members"])
    assert ev["rows"] and "sml_exceedances" in ev["summary"]
    det = ev["members"][1]["istd_detection"]
    assert det["IS1"]["applied"] and det["IS1"]["confidence"] == "high"
    json.dumps(res.to_dict())
    assert snapshot(batch) == before                                   # the raw data are untouched
    # the analyst's project opens without the main window
    ws = H.new_workspace()
    notes = H.open_project(ws, res.project)
    assert not [n for n in notes if "not found" in n]
    assert len(ws.runs) == 5 and ws.replicate_groups[0]["members"] == [s.id for s in ws.states()
                                                                        if s.role == "sample"]
    # the analyst unbinds IS2 in the project: reported again, the ISTD rule finds it
    import copy
    q = copy.deepcopy(ws.quant)
    for rid in ws.replicate_groups[0]["members"]:
        q.setdefault("istd_bindings", {}).setdefault(rid, {})["IS2"] = None
    ws.push_quant("unbind IS2", q)
    from gcws.core import project as P
    P.save(ws, res.project)
    again = PL.run_job(spec(batch, tmp_path / "job2", mode="rereport", project_path=res.project))
    assert again.state == PL.CONTROL
    assert any(f["rule"] == "istd_qc" and "IS2" in f["text"] for f in again.findings)
    assert snapshot(batch) == before


def test_no_blank_in_the_batch(qapp, samples, data, tmp_path):
    from gcws.automation import pipeline as PL
    batch = copy_batch(samples, tmp_path / "watch", ["06_", "07_", "11_"])
    res = PL.run_job(spec(batch, tmp_path / "job", blanks={"blank_istd": ["06_EtOH_ISTD.D"]}), identify=lib_oracle)
    assert res.state == PL.NOT_PROCESSED and "no Blank from the same batch" in res.reason
    allowed = PL.run_job(spec(batch, tmp_path / "job2", blanks={"blank_istd": ["06_EtOH_ISTD.D"]},
                              override={"allow_no_blank": True}), identify=lib_oracle)
    assert allowed.state == PL.CONTROL and any(f["rule"] == "no_blank" for f in allowed.findings)


def test_batch_report_and_delivery(qapp, samples, data, tmp_path):
    from gcws.automation import batch as BA
    from gcws.automation import export, journal as J, pipeline as PL, templates
    batch = copy_batch(samples, tmp_path / "watch", ["06_", "07_", "08_", "11_", "13_"])
    res = PL.run_job(spec(batch, tmp_path / "job", reports=[{"node": "r1", "kind": "nias",
                                                             "formats": ["xlsx", "docx"]}]), identify=lib_oracle)
    xlsx = res.files["r1"]["xlsx"]
    entries = [{"name": "S1", "state": "accepted_auto", "xlsx": xlsx, "findings": []},
               {"name": "S2", "state": "control", "xlsx": xlsx, "findings": [{"text": "SML", "substance": "X"}]}]
    files, warnings = BA.batch_report("26016605_GIOSUN1635", entries, tmp_path / "batch",
                                      ["batch_docx", "batch_xlsx"])
    assert not warnings and Path(files["batch_docx"]).is_file() and Path(files["batch_xlsx"]).is_file()
    from openpyxl import load_workbook
    sh = load_workbook(files["batch_xlsx"]).active
    assert sh.cell(4, 1).value == "S1" and sh.cell(5, 2).value == "Control needed"
    # delivery: Excel to A, Word to B, nothing twice, never into the raw data
    wf = templates.make("nias", "T", source=str(tmp_path / "watch"), method="Test NIAS",
                        folder_a=str(tmp_path / "A"), folder_b=str(tmp_path / "B"))
    rep = wf.by_type("report")[0]
    jr = J.Journal(tmp_path / "j.sqlite")
    b = jr.batch(wf.id, batch)
    job = jr.ensure_job(wf.id, wf.methods()[0].id, b["id"], "k", "26016606", [A, B], {}, "fp", state=J.QUEUED)
    jr.transition(job.id, J.QUEUED, J.PROCESSING)
    jr.transition(job.id, J.PROCESSING, J.CONTROL, files={rep.id: res.files["r1"]}, export_pending=1)
    assert export.deliver(jr, wf, jr.job(job.id)) == []                 # control: held back
    assert jr.review(job.id, True, "ok", user="analyst")
    lines = export.deliver(jr, wf, jr.job(job.id))
    assert len(lines) == 2
    assert (tmp_path / "A" / batch.name / Path(xlsx).name).is_file()
    assert (tmp_path / "B" / batch.name / Path(res.files["r1"]["docx"]).name).is_file()
    assert not (tmp_path / "B" / batch.name / Path(xlsx).name).exists()
    assert export.deliver(jr, wf, jr.job(job.id)) == []                 # delivered once
    wf.by_type("folder")[0].params["path"] = str(batch / "reports")
    wf.by_type("folder")[0].params["overwrite"] = "overwrite"
    jr.request(job.id)
    jr.transition(job.id, J.QUEUED, J.PROCESSING)
    jr.transition(job.id, J.PROCESSING, J.ACCEPTED_AUTO, files={rep.id: res.files["r1"]})
    export.deliver(jr, wf, jr.job(job.id))
    assert not (batch / "reports").exists()                            # refused: inside the raw data
    assert any(e["state"] == "error" for e in jr.exports(job.id))
    jr.close()


def test_job_process(samples, data, tmp_path):
    """``python -m gcws --process-job spec.json``: no window, the outcome in result.json."""
    import os
    import subprocess
    import sys
    from gcws import paths
    batch = copy_batch(samples, tmp_path / "watch", ["06_", "07_", "08_"])
    s = spec(batch, tmp_path / "job", group={"key": "k", "name": "07", "members": [A]},
             blanks={"blank": ["08_EtOH.D"], "blank_istd": ["06_EtOH_ISTD.D"]}, search=False,
             reports=[{"node": "r1", "kind": "fingerprint", "formats": ["xlsx"]}])
    (tmp_path / "job").mkdir()
    path = tmp_path / "job" / "spec.json"
    path.write_text(json.dumps(s), encoding="utf-8")
    env = dict(os.environ, GCWS_DATA=str(data), QT_QPA_PLATFORM="offscreen")
    done = subprocess.run([sys.executable, "-m", "gcws", "--process-job", str(path)], cwd=str(paths.ROOT), env=env,
                          timeout=600, capture_output=True)
    result = json.loads((tmp_path / "job" / "result.json").read_text(encoding="utf-8"))
    assert done.returncode == 0, (done.stdout[-500:], (tmp_path / "job" / "job.log").read_text()[-1500:])
    assert result["state"] in ("accepted_auto", "control") and Path(result["files"]["r1"]["xlsx"]).is_file()
    assert (tmp_path / "job" / "job.log").is_file()
    # a spec that cannot be processed still ends in result.json
    bad = tmp_path / "bad"
    bad.mkdir()
    (bad / "spec.json").write_text(json.dumps(dict(s, out_dir=str(bad), batch_folder=str(tmp_path / "gone"))),
                                   encoding="utf-8")
    done = subprocess.run([sys.executable, "-m", "gcws", "--process-job", str(bad / "spec.json")],
                          cwd=str(paths.ROOT), env=env, timeout=600, capture_output=True)
    assert json.loads((bad / "result.json").read_text(encoding="utf-8"))["state"] in ("retry", "failed")


def test_job_makes_the_feature_double_determination(qapp, samples, data, tmp_path):
    from gcws.automation import headless as H
    from gcws.automation import pipeline as PL
    batch = copy_batch(samples, tmp_path / "watch", ["06_", "07_", "08_", "11_", "13_"])
    res = PL.run_job(spec(batch, tmp_path / "job"), identify=lib_oracle)
    assert res.state in (PL.ACCEPTED_AUTO, PL.CONTROL), res.reason
    assert "features" in res.timings
    feats = res.evidence["features"]
    assert feats and all(f["light"] in ("green", "yellow", "red", "grey") for f in feats)
    red = [f for f in feats if f["light"] == "red"]
    found = [f for f in res.findings if f["rule"] == "feature_review"]
    assert len(found) == len(red)
    if red:
        assert res.state == PL.CONTROL
    # the analyst's project holds the automatic gap fills and the record of what was made automatically
    ws = H.new_workspace()
    H.open_project(ws, res.project)
    g = ws.replicate_groups[0]
    assert g["features"]["ids"] and g["features"]["auto_done"]
    gap = [e for st in ws.states() for e in st.events("FID") if e.option == "gapfill"]
    assert gap and all("gap fill" in e.comment for e in gap)
    # P73: the two determinations were found, harmonised and compared; Report² tells what was done
    dd = res.evidence["double_determination"]
    assert len(dd["members"]) == 2 and dd["pairing"] == "features" and dd["features"] == len(feats)
    info = [f for f in res.findings if f["rule"] == "double_determination"]
    assert [f["level"] for f in info] == ["info"] and "boundaries harmonised" in info[0]["text"]
    moved = [e for st in ws.states() for e in st.events("FID") if "as in" in (e.comment or "")]
    assert dd["boundaries"] == len({(round(e.ref_rt or 0, 4), st.id) for st in ws.states()
                                    for e in st.events("FID") if "as in" in (e.comment or "")})
    assert dd["boundaries"] > 0 and moved
    assert any(s.startswith("boundary|") for s in g["features"]["auto_done"])


def test_single_determination_rule():
    from gcws.automation import rules as RU
    ev = {"double_determination": {"members": ["07_x_A.D"], "pairing": "single"}}
    res = RU.evaluate(RU.default_rules(), ev)
    assert res.status == RU.CONTROL and "one determination only" in res.findings[0].text
    ev = {"double_determination": {"members": ["a", "b"], "pairing": "features", "text": "a + b: 3 features"}}
    res = RU.evaluate(RU.default_rules(), ev)
    assert res.status == RU.ACCEPTED_AUTO and res.findings[0].level == "info"
    # a single determination is information, an incomplete double determination needs control
    sd = {"members": ["07_x"], "pairing": "single", "determination": "single", "reported": ["07_x"]}
    res = RU.evaluate(RU.default_rules(), {"double_determination": sd})
    assert res.status == RU.ACCEPTED_AUTO and [(f.level, f.text) for f in res.findings] ==         [("info", "single determination")]
    res = RU.evaluate(RU.default_rules(), {"double_determination": dict(sd, determination="incomplete")})
    assert res.status == RU.CONTROL


def test_template_report_in_the_automation(qapp, samples, data, tmp_path):
    import copy
    from gcws.automation import headless as H
    from gcws.automation import pipeline as PL
    from gcws.report import template as TP
    batch = copy_batch(samples, tmp_path / "watch", ["06_", "07_", "08_", "11_", "13_"])
    method = copy.deepcopy(METHOD)
    tpl = TP.stamped(TP.preset("NIAS"), "Customer A")
    tpl["columns"].append({"field": "reldiff"})
    method["sections"]["report_template"] = tpl
    res = PL.run_job(spec(batch, tmp_path / "job", method=method,
                          reports=[{"node": "r1", "kind": "template", "formats": ["xlsx", "docx"]}]),
                     identify=lib_oracle)
    assert res.state in (PL.ACCEPTED_AUTO, PL.CONTROL), res.reason
    files = res.files["r1"]
    assert Path(files["xlsx"]).name.endswith("_Customer_A_Report.xlsx") and Path(files["docx"]).is_file()
    assert res.evidence["kind"] == "template" and res.evidence["rows"]
    # the analyst's project keeps the template: a re-report makes the same report
    ws = H.new_workspace()
    H.open_project(ws, res.project)
    assert TP.of(ws.quant)["name"] == "Customer A"
    again = PL.run_job(spec(batch, tmp_path / "job2", mode="rereport", project_path=res.project,
                            reports=[{"node": "r1", "kind": "template", "formats": ["xlsx"]}]))
    assert Path(again.files["r1"]["xlsx"]).name.endswith("_Customer_A_Report.xlsx")


PLAIN = "07_26016606_130m_min_GIOSUN1635.D"


def label_of(xlsx) -> str:
    """The determination line of a NIAS report (the second header line of its first sheet)."""
    from openpyxl import load_workbook
    return str(load_workbook(xlsx).worksheets[0]["A2"].value)


def test_a_run_without_a_suffix_is_a_single_determination(qapp, samples, data, tmp_path):
    """Issue 02: a lone run without a replicate suffix is a single determination, reported in the double
    determination layout; Report² gets an information line, not "Double determination incomplete". The
    batch report lists it with a double determination."""
    from gcws.automation import batch as BA
    from gcws.automation import pipeline as PL
    batch = copy_batch(samples, tmp_path / "watch", ["06_", "07_", "08_", "11_", "13_"])
    shutil.move(str(batch / A), str(batch / PLAIN))
    shutil.copytree(batch / PLAIN, batch / A)
    blanks = {"blank": ["08_EtOH.D"], "blank_istd": ["06_EtOH_ISTD.D"]}
    import copy
    from gcws.report import template as TP
    method = copy.deepcopy(METHOD)
    method["sections"]["report_template"] = TP.stamped(TP.preset("NIAS"), "Customer A")
    sd = PL.run_job(spec(batch, tmp_path / "sd", group={"key": "k", "name": "07", "members": [PLAIN]},
                         blanks=blanks, method=method,
                         reports=[{"node": "r1", "kind": "nias", "formats": ["xlsx"]},
                                  {"node": "r2", "kind": "template", "formats": ["xlsx"]}]), identify=lib_oracle)
    assert sd.state in (PL.ACCEPTED_AUTO, PL.CONTROL), sd.reason
    assert label_of(sd.files["r1"]["xlsx"]) == "PA 26.007, single determination"
    from openpyxl import load_workbook
    head = [str(c.value) for row in load_workbook(sd.files["r2"]["xlsx"]).worksheets[0].iter_rows(max_row=6)
            for c in row if c.value]
    assert any("single determination" in v for v in head), head            # {determination} of the template
    dd = sd.evidence["double_determination"]
    assert dd["determination"] == "single" and dd["reported"] == dd["members"] and len(dd["members"]) == 1
    found = [f for f in sd.findings if f["rule"] == "double_determination"]
    assert [(f["level"], f["text"]) for f in found] == [("info", "single determination")]
    # a lone _A whose partner never came stays an incomplete double determination
    inc = PL.run_job(spec(batch, tmp_path / "inc", group={"key": "k", "name": "07", "members": [A]},
                          blanks=blanks, reports=[]), identify=lib_oracle)
    assert inc.evidence["double_determination"]["determination"] == "incomplete"
    found = [f for f in inc.findings if f["rule"] == "double_determination"]
    assert [f["level"] for f in found] == ["control"] and "one determination only" in found[0]["text"]
    # the batch report lists the single and the double determination together
    pair = PL.run_job(spec(batch, tmp_path / "dd", reports=[{"node": "r1", "kind": "nias", "formats": ["xlsx"]}]),
                      identify=lib_oracle)
    assert label_of(pair.files["r1"]["xlsx"]) == "PA 26.007, double determination"
    entries = [{"name": "07", "state": sd.state, "xlsx": sd.files["r1"]["xlsx"], "findings": sd.findings},
               {"name": "26016606", "state": pair.state, "xlsx": pair.files["r1"]["xlsx"], "findings": []}]
    files, warnings = BA.batch_report("26016605_GIOSUN1635", entries, tmp_path / "batch",
                                      ["batch_docx", "batch_xlsx"])
    assert not warnings and Path(files["batch_docx"]).is_file()
    from openpyxl import load_workbook
    sh = load_workbook(files["batch_xlsx"]).active
    assert [sh.cell(4, 1).value, sh.cell(5, 1).value] == ["07", "26016606"]


def test_single_determination_only_rules():
    """Issue 03: in a single determination the red A/B differences are information, unless the rule is set
    to control; the info line names the control determination."""
    from gcws.automation import rules as RU
    red = [{"feature_id": "F1", "rt": 10.0, "name": "X", "light": "red", "verdict": "found in A only"}]
    sd = {"members": ["a", "b"], "pairing": "features", "determination": "single", "reported": ["a"],
          "control": "b"}
    res = RU.evaluate(RU.default_rules(), {"double_determination": sd, "features": red})
    assert res.status == RU.ACCEPTED_AUTO
    assert [(f.rule, f.level) for f in res.findings] == [("feature_review", "info"), ("double_determination", "info")]
    assert "b is the control determination" in res.findings[1].text
    rules = RU.default_rules()
    next(r for r in rules if r.id == "feature_review").params["single_control"] = True
    assert RU.evaluate(rules, {"double_determination": sd, "features": red}).status == RU.CONTROL
    dd = dict(sd, determination="double", reported=["a", "b"], control="")
    assert RU.evaluate(RU.default_rules(), {"double_determination": dd, "features": red}).status == RU.CONTROL


def test_single_determination_only_job(qapp, samples, data, tmp_path):
    """Issue 03: with Single determination only, B is processed and compared, the report has A's values
    only and is labelled single determination; Report² gets the SD info line."""
    import copy
    from gcws.automation import pipeline as PL
    batch = copy_batch(samples, tmp_path / "watch", ["06_", "07_", "08_", "11_", "13_"])
    method = copy.deepcopy(METHOD)
    method["sections"]["quant"]["single_only"] = True
    sd = PL.run_job(spec(batch, tmp_path / "sd", method=method,
                         reports=[{"node": "r1", "kind": "nias", "formats": ["xlsx"]}]), identify=lib_oracle)
    assert sd.state in (PL.ACCEPTED_AUTO, PL.CONTROL), sd.reason
    assert label_of(sd.files["r1"]["xlsx"]) == "PA 26.007, single determination"
    dd = sd.evidence["double_determination"]
    assert dd["determination"] == "single" and dd["reported"] == [Path(A).stem] and dd["control"] == Path(B).stem
    assert dd["pairing"] == "features" and sd.evidence["features"]          # B processed and compared
    found = [f for f in sd.findings if f["rule"] == "double_determination"]
    assert found[0]["level"] == "info" and found[0]["text"].startswith("single determination")
    assert all(f["level"] == "info" for f in sd.findings if f["rule"] == "feature_review")
    from openpyxl import load_workbook
    sheet = " | ".join(str(c.value) for ws in load_workbook(sd.files["r1"]["xlsx"]).worksheets
                      for row in ws.iter_rows() for c in row if c.value is not None)
    assert Path(A).stem in sheet and Path(B).stem not in sheet
    # issue 05: the analyst reports B instead (Replicates), and the report made again has B's values
    from gcws.automation import headless as H
    from gcws.core import project as P
    from gcws.quant import double_determination as DD
    ws = H.new_workspace()
    H.open_project(ws, sd.project)
    g = ws.replicate_groups[0]
    b = next(m for m in g["members"] if ws.runs[m].run.path.name == B)
    ws.replicate_groups = DD.set_reported(ws.replicate_groups, g["id"], b)
    P.save(ws, sd.project)
    again = PL.run_job(spec(batch, tmp_path / "sd2", mode="rereport", project_path=sd.project,
                            reports=[{"node": "r1", "kind": "nias", "formats": ["xlsx"]}]))
    assert again.state in (PL.ACCEPTED_AUTO, PL.CONTROL), again.reason
    dd = again.evidence["double_determination"]
    assert dd["reported"] == [Path(B).stem] and dd["control"] == Path(A).stem
    sheet = " | ".join(str(c.value) for ws in load_workbook(again.files["r1"]["xlsx"]).worksheets
                      for row in ws.iter_rows() for c in row if c.value is not None)
    assert Path(B).stem in sheet and Path(A).stem not in sheet
    assert label_of(again.files["r1"]["xlsx"]) == "PA 26.007, single determination"


def test_single_determination_only_b_alone(qapp, samples, data, tmp_path):
    """Issue 04: with Single determination only and only B in the batch, B is reported as a single
    determination (not incomplete) and Report² shows the info line "A missing, B reported"."""
    import copy
    from gcws.automation import pipeline as PL
    batch = copy_batch(samples, tmp_path / "watch", ["06_", "08_", "11_", "13_"])
    method = copy.deepcopy(METHOD)
    method["sections"]["quant"]["single_only"] = True
    group = {"key": "26016606:x", "name": "26016606_130m_min_GIOSUN1635", "members": [B]}
    res = PL.run_job(spec(batch, tmp_path / "sd", method=method, group=group), identify=lib_oracle)
    assert res.state in (PL.ACCEPTED_AUTO, PL.CONTROL), res.reason
    dd = res.evidence["double_determination"]
    assert dd["determination"] == "single" and dd["reported"] == [Path(B).stem] and not dd["control"]
    found = [f for f in res.findings if f["rule"] == "double_determination"]
    assert [(f["level"], f["text"]) for f in found] == [("info", "A missing, B reported")]

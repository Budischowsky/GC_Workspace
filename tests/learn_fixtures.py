"""Synthetic NIAS evaluation workbooks for the gcws.learn tests (same cell layout as the real
NIAS-Screening-*.xlsm, no client data)."""
from __future__ import annotations

from pathlib import Path

import openpyxl

TIC_HEADER = ["Header=", "Peak", "R.T.", "First", "Max", "Last", "PK  TY", "Height", "Area", "Pct Max", "Pct Total"]
FID_HEADER = ["Header=", "Peak", "R.T.", "Start", "End", "PK TY", "Height", "Area", "Pct Max", "Pct Total"]

TIC_PEAKS = [
    ["1=", 1, 8.051, 231, 236, 244, "BV  ", 84245, 1487790, 0.5, 0.084],
    ["2=", 2, 8.203, 244, 256, 272, "VV  ", 2697716, 40534814, 13.57, 2.301],
]
FID_PEAKS = [
    ["1=", 1, 5.132, 5.1, 5.16, "BV  ", 50000, 7497857, 0, 0],
    ["2=", 2, 6.409, 6.38, 6.44, "VV  ", 30000, 1241851, 0, 0],
    ["3=", 3, 6.917, 6.873, 6.938, "BV  ", 19706, 395154, 0, 0],
    ["4=", 4, 7.07, 6.938, 7.313, "  M ", 3238076, 119966715, 0.23, 0.062],
    ["5=", 5, 7.189, 7.15, 7.22, "VV  ", 20000, 988201, 0, 0],
    ["6=", 6, 14.33, 14.299, 14.385, "  M ", 636187, 9871653, 0.02, 0.005],
    ["7=", 7, 21.048, 21.0, 21.09, "VV  ", 90000, 2563570, 0, 0],
    ["8=", 8, 21.449, 21.4, 21.48, "VV  ", 50000, 1167300, 0, 0],
]

ALKANES = [(f"C{8 + 2 * i}", 5.003 + i * 1.2, 800 + 200 * i) for i in range(21)]

# (RT, name, CAS, DB, match, area, conc1, conc2, SML, Ref) ; area may be a formula; extra keys via dict
FINAL_ROWS = [
    {"A": 6.917, "F": 395154, "G": 0.0008, "H": 0.0048},
    {"A": 7.07, "B": "Butyl methacrylate", "C": "000097-88-1", "D": "NIST05", "E": 90, "F": "=L29-F30",
     "G": 0.243, "H": 1.458, "I": 6, "J": "[1],[2]", "L": 119966715},
    {"A": 7.189, "B": "α-Methylstyrene", "C": "000098-83-9", "D": "NIST05", "E": 93, "F": 988201,
     "G": 0.002, "H": 0.012, "I": 0.05, "J": "[1],[2]"},
    {"A": 14.33, "B": "IS1", "F": 9871653, "G": 0.02, "H": 0.12},
    {"A": 16.935, "B": "IS4", "F": 742832, "G": 0.0015, "H": 0.009},
    {"A": 17.878, "B": "mehrere Verbindungen", "F": 932973, "G": 0.0019, "H": 0.011},
    {"A": 20.439, "B": "unknown m/z 145/167/270", "F": 854508, "G": 0.0017, "H": 0.010},
    {"A": 21.048, "B": "Styrene Oligomer", "F": 2563570, "G": 0.0052, "H": 0.031},
    {"A": 21.449, "B": "Styrene Oligomer", "F": 1167300, "G": 0.0024, "H": 0.014},
    {"A": 26.5, "B": "Some Additive", "F": 100000, "G": 0.0002, "H": 0.001},
    {"A": 28.862, "B": "possible derivative of an antioxidant (Irganox 1076) m/z 528/57", "F": 2569060.5,
     "G": 0.0052, "H": 0.031},
    {"A": "Sum of styrene oligomers (estimated)**", "F": 3730870, "G": 0.0076, "H": 0.045},
    {"A": "Ende"},
]

PRE_CLEAN_ROWS = [{"A": 5.132, "F": 7497857, "G": 0.015, "H": 0.09},
                  {"A": 6.409, "F": 1241851, "G": 0.0025, "H": 0.015}] + FINAL_ROWS

REPORT_ROWS = [
    {"A": 7.07, "B": "Butyl methacrylate", "C": "97-88-1", "D": "NIST05", "E": 90, "F": 118978514, "G": 0.243,
     "H": 1.458, "I": 6, "J": "[1],[2]"},
    {"A": 10.846, "B": "Benzaldehyde, 2,4,6-trimethyl-", "C": "487-68-3", "D": "NIST05", "E": 94, "F": 1538244,
     "G": 0.0031, "H": 0.0188, "I": "CC I(a)"},
    {"A": 17.878, "B": "mehrere Verbindungen", "F": 932973, "G": 0.0019, "H": 0.011},
    {"A": 21.048, "B": "Styrene Oligomer", "F": 2563570, "G": 0.0052, "H": 0.031},
    {"A": 21.449, "B": "Styrene Oligomer", "F": 1167300, "G": 0.0024, "H": 0.014},
]
FOOTNOTES = ["[1] Regulation (EU) No 10/2011", "** estimated as sum"]


def _header(ws, template: str, conc_headers, istd_area_c17) -> None:
    put = lambda coord, v: ws.__setitem__(coord, v)  # noqa: E731
    for coord, v in {
        "A1": "10ppb-Auswertung", "A2": "Probenname", "B2": " 05_X_A", "I2": "Auswerter:", "J2": "BlM",
        "A3": "Syn-Proben-ID", "B3": 2501675, "A4": "Syn-Summary", "B4": 25011662, "I4": "GC-Operator:",
        "J4": " BlM", "A5": "Datenfile", "B5": r"X:\Daten\05_X_A.D", "A6": "Probenvorbereitung",
        "D8": "Temperatur:", "E8": 40, "D9": "Dauer:", "E9": "10 d", "A10": "Simulans:", "B10": "EtOH 95%",
        "A12": "Volumen:", "B12": 10, "D12": "O/V-Ratio", "E12": 6, "E20": "Conc", "F20": "Fläche",
        "E21": "mg/mL", "F21": "cts", "A22": "GC-Methode:", "B22": "PA 26.009", "D22": "C17", "E22": 0.84,
        "F22": istd_area_c17, "A23": "Inj-Vol:", "B23": 30, "D23": "BBP", "E23": 0.85, "F23": 7095674,
        "I23": "Conc mg/mL", "J23": "Fläche cts", "A24": "Interner Std.", "B24": 10, "D24": "DnNP",
        "E24": 0.85, "F24": 7904784, "H24": "DBP-d4", "I24": 0.0752, "J24": 742832, "F25": 8290703.666666667,
    }.items():
        put(coord, v)
    if template == "v2":
        ws["M1"] = "n-Alkane"
        ws["M4"], ws["N4"], ws["O4"] = "n-Alkane", "RT [min]", "RI []"
        for i, (name, rt, ri) in enumerate(ALKANES):
            ws[f"M{5 + i}"], ws[f"N{5 + i}"], ws[f"O{5 + i}"] = name, round(rt, 3), ri
    cols = ["RT / min", "Name", "CAS-#", "DB", "%match", "Fläche"] + list(conc_headers)
    for i, h in enumerate(cols):
        ws.cell(row=27, column=i + 1, value=h)


def _rows(ws, rows, start: int) -> None:
    for i, row in enumerate(rows):
        for col, v in row.items():
            ws[f"{col}{start + i}"] = v


def make_workbook(path: Path, *, template: str = "v2", pre_clean: bool = True, final_rows=None,
                  report_rows=None, conc_headers=("Conc. mg/dm²", "Conc. mg/kg", "SML", "Reference"),
                  istd_area_c17=9871653, fid_peaks=None) -> Path:
    wb = openpyxl.Workbook()
    raw = wb.active
    raw.title = "Rohdaten"
    lines = [["[contents]"], ["count=2"], ["Name=", r"X:\Daten\05_X_A.D"], ["1=", r"INT TIC: 05_X_A.D\data.ms"],
             ["2=", r"INT 05_X_A.D\FID1A.ch"], [r"[INT TIC: 05_X_A.D\data.ms]"], ["Time=", "Tue Jun 17 2025"],
             TIC_HEADER, *TIC_PEAKS, [r"[INT 05_X_A.D\FID1A.ch]"], ["Time=", "Tue Jun 17 2025"], FID_HEADER,
             *(FID_PEAKS if fid_peaks is None else fid_peaks)]
    for line in lines:
        raw.append(line)
    if pre_clean:
        ws2 = wb.create_sheet("Auswertung (2)")
        _header(ws2, template, conc_headers, istd_area_c17)
        _rows(ws2, PRE_CLEAN_ROWS, 28)
    ws = wb.create_sheet("Auswertung")
    _header(ws, template, conc_headers, istd_area_c17)
    _rows(ws, FINAL_ROWS if final_rows is None else final_rows, 28)
    rep = wb.create_sheet("externerBericht")
    rep["A21"], rep["A22"] = "GC–MS/FID – NIAS-Screening –", "PA 26.007"
    rep["A23"], rep["B23"] = "Sample name:", " 05_X_A"
    for col, v in zip("ABCEGIJ", ["RT", "Name", "CAS-No.", "%\nmatch", "Migration concentration", "SML", "Ref"]):
        rep[f"{col}25"] = v
    for col, v in zip("AGHI", ["min", "mg/dm²", "mg/kg", "mg/kg"]):
        rep[f"{col}26"] = v
    rows = REPORT_ROWS if report_rows is None else report_rows
    _rows(rep, rows, 28)
    for i, note in enumerate(FOOTNOTES):
        rep[f"A{28 + len(rows) + 2 + i}"] = note
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path

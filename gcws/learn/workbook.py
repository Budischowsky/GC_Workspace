"""Parse one human NIAS-Screening evaluation workbook into a HumanEvaluation."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from gcws.learn.model import (AlkanePoint, EvalRow, Header, HumanEvaluation, IstdEntry, RawPeak, ReportRow,
                              normalise_cas)


def _num(v) -> Optional[float]:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().replace(",", "."))
    except (TypeError, ValueError):
        return None


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _norm_col(v) -> str:
    return re.sub(r"\s+", " ", _text(v)).upper()


def _section_signal(title: str) -> str:
    inner = title.strip()[1:-1].strip()
    if inner.upper().startswith("INT TIC"):
        return "TIC"
    if "FID" in inner.upper():
        return "FID"
    return inner


def parse_rohdaten(rows: list[tuple]) -> list[RawPeak]:
    """The integration tables of the 'Rohdaten' sheet (INI-like '[INT …]' sections, a 'Header=' row and
    'N=' peak rows). Columns are located from each section's own header, so TIC (scan numbers) and FID
    (start/end in minutes) layouts both work and empty cells do not shift values."""
    peaks: list[RawPeak] = []
    signal = ""
    cols: dict[str, int] = {}
    for row in rows:
        cells = list(row)
        first_i = next((i for i, v in enumerate(cells) if v not in (None, "")), None)
        if first_i is None:
            continue
        first = _text(cells[first_i])
        if first.startswith("[") and first.endswith("]"):
            signal = "" if first.lower() == "[contents]" else _section_signal(first)
            cols = {}
            continue
        if not signal:
            continue
        if first == "Header=":
            cols = {_norm_col(v): i for i, v in enumerate(cells) if i > first_i and v not in (None, "")}
            continue
        if not cols or not re.fullmatch(r"\d+=", first):
            continue

        def get(*names):
            for n in names:
                i = cols.get(n)
                if i is not None and i < len(cells):
                    return cells[i]
            return None

        rt, area = _num(get("R.T.", "RT")), _num(get("AREA"))
        if rt is None or area is None:
            continue
        ptype = _text(get("PK TY", "PK  TY", "TYPE"))
        peaks.append(RawPeak(
            signal=signal, number=int(_num(get("PEAK")) or first[:-1]), rt=rt, area=area,
            height=_num(get("HEIGHT")),
            start=_num(get("START")) if "START" in cols else None,
            end=_num(get("END")) if "END" in cols else None,
            peak_type=ptype, manual="M" in ptype))
    return peaks


# --- 'Auswertung' sheet -------------------------------------------------------------------------------------------

ROW_CLASSES = ("istd", "named", "named_no_cas", "group", "unknown", "coelution", "derivative", "background",
               "unnamed", "sum")

_RE_ISTD = re.compile(r"^IS\d+$", re.IGNORECASE)
_RE_UNKNOWN = re.compile(r"^unknown\b", re.IGNORECASE)
_RE_COELUTION = re.compile(r"mehrere verb|several compounds|co-?elut", re.IGNORECASE)
_RE_DERIVATIVE = re.compile(r"derivat|transformation product|^(possible\s+)?degradation product", re.IGNORECASE)
# remarks by which the analyst marks a listed peak as not coming from the sample
_RE_BACKGROUND = re.compile(r"\bblank\b|nicht aus (der )?probe|\bseptum\b|\b(im|in) (std|standard)\b",
                            re.IGNORECASE)
_RE_SUM = re.compile(r"^(sum|summe)\b", re.IGNORECASE)
_RE_COORD = re.compile(r"^([A-Z]{1,3})(\d+)$")
_RE_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")


def _col_index(col: str) -> int:
    n = 0
    for ch in col:
        n = n * 26 + ord(ch) - 64
    return n


def _grid(cells: dict[str, object]) -> dict[int, dict[int, object]]:
    """{row: {column index: value}} without empty cells."""
    grid: dict[int, dict[int, object]] = {}
    for coord, v in cells.items():
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        m = _RE_COORD.match(coord)
        if m:
            grid.setdefault(int(m.group(2)), {})[_col_index(m.group(1))] = v
    return grid


def _plain(v) -> str:
    """Cell value as text; whole floats without '.0'."""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return _text(v)


def _table_header_row(grid) -> Optional[int]:
    for r in sorted(grid):
        a, b = _text(grid[r].get(1)), _text(grid[r].get(2))
        if a.upper().startswith("RT") and b.casefold() == "name":
            return r
    return None


def _alkane_anchor(grid) -> Optional[tuple[int, int]]:
    for r in sorted(grid):
        for c, v in grid[r].items():
            if _text(v).casefold() == "n-alkane" and _text(grid[r].get(c + 1)).upper().startswith("RT"):
                return r, c
    return None


def detect_template(cells: dict[str, object]) -> str:
    """'v2' = worksheet with an n-Alkane RT/RI block, 'v1' = without."""
    return "v2" if _alkane_anchor(_grid(cells)) else "v1"


_HEADER_LABELS = {
    "probenname": "sample_name", "syn-proben-id": "syn_id", "syn-summary": "syn_summary",
    "auswerter:": "evaluator", "gc-operator:": "operator", "datenfile": "data_file", "simulans:": "simulant",
    "temperatur:": "temperature", "dauer:": "duration", "volumen:": "volume", "o/v-ratio": "sv_ratio",
    "s/v-ratio": "sv_ratio", "gc-methode:": "gc_method", "inj-vol:": "inj_volume",
}
_NUMERIC_HEADER = {"temperature", "volume", "sv_ratio", "inj_volume"}


def parse_header(cells: dict[str, object]) -> Header:
    """Header values found by their label: the next non-empty cell to the right in the same row."""
    grid = _grid(cells)
    stop = _table_header_row(grid) or (max(grid) + 1 if grid else 0)
    h = Header()
    for r in sorted(grid):
        if r >= stop:
            break
        row = grid[r]
        for c in sorted(row):
            key = _HEADER_LABELS.get(_text(row[c]).casefold())
            if not key or getattr(h, key) not in ("", None):
                continue
            right = [row[k] for k in sorted(row) if k > c]
            if right:
                setattr(h, key, _num(right[0]) if key in _NUMERIC_HEADER else _plain(right[0]))
                if key == "temperature" and h.temperature is None:
                    h.temperature_text = _plain(right[0])        # a condition, e.g. "USB" (ultrasonic bath)
    h.istd, h.istd_mean_area = _istd(grid, stop)
    h.alkanes = _alkanes(grid)
    h.conc_units = [unit for unit, _ in _table_columns(grid)["conc"]]
    return h


def _istd(grid, stop: int) -> tuple[list[IstdEntry], Optional[float]]:
    """ISTD blocks: a 'Conc' header with 'Fläche' to its right; the ISTD name is left of the Conc column.
    The mean area is the number below the first block's areas (the sheet's AVERAGE cell)."""
    entries: list[IstdEntry] = []
    mean: Optional[float] = None
    for r in sorted(grid):
        if r >= stop:
            break
        for c, v in sorted(grid[r].items()):
            t = _text(v)
            if not t.startswith("Conc") or t.startswith("Conc."):
                continue
            if not _text(grid[r].get(c + 1)).startswith("Fläche"):
                continue
            block: list[IstdEntry] = []
            last = r
            for rr in range(r + 1, stop):
                row = grid.get(rr, {})
                name = row.get(c - 1)
                if name is None or _num(name) is not None or "/" in _text(name):   # numbers, units
                    continue
                block.append(IstdEntry(_text(name), _num(row.get(c)), _num(row.get(c + 1))))
                last = rr
            if block and mean is None:
                below = _num(grid.get(last + 1, {}).get(c + 1))
                areas = [e.area for e in block if e.area is not None]
                mean = below if below is not None else (sum(areas) / len(areas) if areas else None)
            entries.extend(block)
    return entries, mean


def _alkanes(grid) -> list[AlkanePoint]:
    anchor = _alkane_anchor(grid)
    if not anchor:
        return []
    r0, c = anchor
    out: list[AlkanePoint] = []
    for r in range(r0 + 1, max(grid) + 1):
        row = grid.get(r, {})
        name, ri = _text(row.get(c)), _num(row.get(c + 2))
        if not re.fullmatch(r"C\d+", name) or ri is None:
            break
        out.append(AlkanePoint(name, _num(row.get(c + 1)), ri))
    return out


def _table_columns(grid) -> dict:
    """Column indexes of the worksheet table, found by header text."""
    hr = _table_header_row(grid)
    cols: dict = {"conc": []}
    if hr is None:
        return cols
    cols["row"] = hr
    for c, v in sorted(grid[hr].items()):
        t = _text(v)
        tl = t.casefold()
        if c == 1:
            cols["rt"] = c
        elif tl == "name":
            cols["label"] = c
        elif tl.startswith("cas"):
            cols["cas"] = c
        elif tl == "db":
            cols["library"] = c
        elif "match" in tl:
            cols["match"] = c
        elif tl in ("fläche", "area"):
            cols["area"] = c
        elif tl.startswith("conc."):
            cols["conc"].append((t[5:].strip(), c))
        elif tl == "sml":
            cols["sml"] = c
        elif tl.startswith("ref"):
            cols["reference"] = c
        elif tl.startswith("ri"):
            cols["ri"] = c
    return cols


def _original_area(formula: str, row: int, area_col: int, grid) -> Optional[float]:
    """The same-row cell an area formula starts from (e.g. L29 in '=L29-F30')."""
    for col, r in _RE_REF.findall(formula):
        ci = _col_index(col)
        if int(r) == row and ci != area_col:
            value = _num(grid.get(row, {}).get(ci))
            if value is not None:
                return value
    return None


def parse_eval_table(values: dict[str, object], formulas: dict[str, object]) -> list[EvalRow]:
    """The worksheet table below the 'RT / min | Name | …' header, up to 'Ende'."""
    grid = _grid(values)
    fgrid = _grid({k: getattr(v, "text", v) for k, v in formulas.items()})
    cols = _table_columns(grid)
    if "row" not in cols:
        return []
    known = {c for k, c in cols.items() if isinstance(c, int) and k not in ("row", "ri")}
    known |= {c for _, c in cols["conc"]}
    area_col = cols.get("area", 0)
    last = max([*grid, *fgrid])
    rows: list[EvalRow] = []
    for r in range(cols["row"] + 1, last + 1):
        vrow, frow = grid.get(r, {}), fgrid.get(r, {})
        first = vrow.get(1)
        if _text(first).casefold() == "ende":
            break
        rt = _num(first)
        label = _text(vrow.get(cols.get("label", 0)))
        if rt is None and isinstance(first, str):
            label = _text(first)
        if rt is None and not label:
            continue
        formula = frow.get(area_col)
        formula = formula if isinstance(formula, str) and formula.startswith("=") else ""
        rows.append(EvalRow(
            sheet_row=r, rt=rt, label=label, cas=normalise_cas(vrow.get(cols.get("cas", 0))),
            library=_text(vrow.get(cols.get("library", 0))), match=_num(vrow.get(cols.get("match", 0))),
            area=_num(vrow.get(area_col)), area_formula=formula,
            area_original=_original_area(formula, r, area_col, grid) if formula else None,
            conc={unit: _num(vrow.get(c)) for unit, c in cols["conc"]},
            sml=_plain(vrow.get(cols["sml"])) if "sml" in cols else "",
            reference=_plain(vrow.get(cols["reference"])) if "reference" in cols else "",
            note=" ".join(_text(v) for c, v in sorted(vrow.items())
                          if c not in known and c != 1 and isinstance(v, str))))
    return rows


def classify_rows(rows: list) -> None:
    """Set row_class from what the analyst wrote (parsing only, no decision rules). A label without CAS
    that the analyst used on two or more rows is a group label (e.g. an oligomer family)."""
    no_cas: dict[str, int] = {}
    for r in rows:
        if r.label and not r.cas:
            no_cas[r.label.casefold()] = no_cas.get(r.label.casefold(), 0) + 1
    for r in rows:
        label = r.label
        if r.rt is None and _RE_SUM.match(label):
            r.row_class = "sum"
        elif _RE_ISTD.match(label):
            r.row_class = "istd"
        elif _RE_BACKGROUND.search(label):
            r.row_class = "background"
        elif _RE_UNKNOWN.match(label):
            r.row_class = "unknown"
        elif _RE_COELUTION.search(label):
            r.row_class = "coelution"
        elif _RE_DERIVATIVE.search(label):
            r.row_class = "derivative"
        elif r.cas:
            r.row_class = "named"
        elif not label:
            r.row_class = "unnamed"
        elif no_cas.get(label.casefold(), 0) >= 2:
            r.row_class = "group"
        else:
            r.row_class = "named_no_cas"


# --- 'externerBericht' sheet (client report) ----------------------------------------------------------------------

def parse_report(values: dict[str, object]) -> tuple[list[ReportRow], list[str]]:
    """Substance lines of the client report and the footnote lines below them. The concentration units are
    in the row under the header, from the concentration column up to the SML column."""
    grid = _grid(values)
    hr = next((r for r in sorted(grid) if _text(grid[r].get(1)).upper() == "RT"
               and _text(grid[r].get(2)).casefold() == "name"), None)
    if hr is None:
        return [], []
    cols: dict[str, int] = {}
    for c, v in sorted(grid[hr].items()):
        tl = _text(v).casefold()
        if tl == "name":
            cols["label"] = c
        elif tl.startswith("cas"):
            cols["cas"] = c
        elif "match" in tl:
            cols["match"] = c
        elif "conc" in tl:
            cols["conc"] = c
        elif tl == "sml":
            cols["sml"] = c
        elif tl.startswith("ref"):
            cols["reference"] = c
    if "cas" in cols and cols["cas"] + 1 not in cols.values():
        cols["library"] = cols["cas"] + 1                    # the DB column has no header text
    units = grid.get(hr + 1, {})
    conc_cols = []
    if "conc" in cols:
        end = cols.get("sml", max(units, default=cols["conc"]) + 1)
        conc_cols = [(_text(units[c]), c) for c in sorted(units) if cols["conc"] <= c < end]
    rows: list[ReportRow] = []
    notes: list[str] = []
    for r in sorted(grid):
        if r <= hr + 1:
            continue
        row = grid[r]
        rt = _num(row.get(1))
        if rt is None:
            text = " ".join(_text(v) for _, v in sorted(row.items()) if isinstance(v, str))
            if text and rows:
                notes.append(text)
            continue
        rows.append(ReportRow(
            rt=rt, label=_text(row.get(cols.get("label", 0))), cas=normalise_cas(row.get(cols.get("cas", 0))),
            library=_text(row.get(cols.get("library", 0))), match=_num(row.get(cols.get("match", 0))),
            conc={unit: _num(row.get(c)) for unit, c in conc_cols},
            sml=_plain(row.get(cols["sml"])) if "sml" in cols else "",
            reference=_plain(row.get(cols["reference"])) if "reference" in cols else ""))
    classify_rows(rows)
    return rows, notes


# --- whole workbook ----------------------------------------------------------------------------------------------

def _sheet_cells(ws) -> dict[str, object]:
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row
            if getattr(c, "value", None) is not None}


def parse_workbook(path: Path) -> HumanEvaluation:
    """One evaluation workbook. Content problems are recorded in .problems; nothing raises."""
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")      # openpyxl: unsupported Excel extensions in these templates
        return _parse_workbook(path)


def _parse_workbook(path: Path) -> HumanEvaluation:
    import openpyxl

    ev = HumanEvaluation(path=str(path))
    opened = []
    try:
        for data_only in (True, False):
            opened.append(openpyxl.load_workbook(path, data_only=data_only, read_only=True, keep_vba=False))
    except Exception as exc:  # noqa: BLE001 - any unreadable file is a corpus problem, not a crash
        for wb in opened:
            wb.close()
        ev.problems.append(f"cannot open: {type(exc).__name__}: {exc}")
        return ev
    wv, wf = opened
    try:
        names = set(wv.sheetnames)
        if "Rohdaten" in names:
            ev.raw = parse_rohdaten(list(wv["Rohdaten"].iter_rows(values_only=True)))
        else:
            ev.problems.append("sheet missing: Rohdaten")
        if "Auswertung" in names:
            cells = _sheet_cells(wv["Auswertung"])
            ev.template = detect_template(cells)
            ev.header = parse_header(cells)
            ev.final = parse_eval_table(cells, _sheet_cells(wf["Auswertung"]))
            classify_rows(ev.final)
            ev.problems += [f"ISTD area missing: {i.name}" for i in ev.header.istd if i.area is None]
            if not ev.final:
                ev.problems.append("no worksheet table in Auswertung")
        else:
            ev.problems.append("sheet missing: Auswertung")
        if "Auswertung (2)" in names:
            ev.pre_clean = parse_eval_table(_sheet_cells(wv["Auswertung (2)"]), _sheet_cells(wf["Auswertung (2)"]))
            classify_rows(ev.pre_clean)
        if "externerBericht" in names:
            ev.report, ev.footnotes = parse_report(_sheet_cells(wv["externerBericht"]))
        else:
            ev.problems.append("sheet missing: externerBericht")
    except Exception as exc:  # noqa: BLE001 - an unexpected sheet layout is a corpus problem, not a crash
        ev.problems.append(f"parse error: {type(exc).__name__}: {exc}")
    finally:
        wv.close()
        wf.close()
    return ev


def removed_peaks(ev: HumanEvaluation, rt_tol: float = 0.01) -> list[RawPeak]:
    """FID peaks of the raw integration that have no row in the final worksheet: what the analyst removed."""
    kept = [r.rt for r in ev.final if r.rt is not None]
    return [p for p in ev.raw if p.signal == "FID" and not any(abs(p.rt - k) <= rt_tol for k in kept)]

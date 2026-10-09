"""The program side of the comparison: each evaluated run processed by the production job
(gcws.automation.pipeline.run_job) as a single determination with its batch blanks, and the per-peak evidence
of its FID peaks."""
from __future__ import annotations

import hashlib
import json
import os
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path, PureWindowsPath
from typing import Any, Optional

from gcws.learn.corpus import CorpusEntry
from gcws.learn.model import Header

#: per-peak values taken from the Peaks panel's getters (gcws.quant.peak_values.VALUES)
PEAK_KEYS = ("num", "rt", "ms_rt", "start", "end", "area", "height", "w50", "sym", "sn", "name", "cas", "score",
             "status", "library", "ri", "istd", "blank_area", "blank_ratio", "in_blank", "area_minus_blank",
             "class_hint", "origin")
N_HITS = 3


def _plain(v: Any):
    if v is None or isinstance(v, (bool, str)):
        return v
    if isinstance(v, (int, float)):
        return float(v)
    return str(v)


def collect_peaks(ws, run_id: str) -> list[dict]:
    """The FID peaks of ``run_id`` with PEAK_KEYS and the top library hits, sorted by RT."""
    from gcws.core.model import FID
    from gcws.quant.peak_values import VALUES, rows_for
    out = []
    for row in rows_for(ws, run_id, FID):
        d = {}
        for k in PEAK_KEYS:
            try:
                d[k] = _plain(VALUES[k](row, ws, run_id, FID))
            except Exception:  # noqa: BLE001 - a value the panel cannot give for this peak
                d[k] = None
        hits = getattr(row.ident, "hits", None) or []
        d["hits"] = [{"name": str(h.get("name", "")), "cas": str(h.get("cas", "")), "score": _plain(h.get("score"))}
                     for h in hits[:N_HITS]]
        out.append(d)
    return sorted(out, key=lambda p: p["rt"] if p["rt"] is not None else float("inf"))


def collect(ws, members: list[str]) -> dict:
    """``run_job`` inspect hook: the evidence of the single evaluated run."""
    return {"peaks": collect_peaks(ws, members[0])}


@dataclass
class ProgramResult:
    run_dir: str
    method: str
    state: str
    reason: str = ""
    peaks: list[dict] = field(default_factory=list)
    reported: list[dict] = field(default_factory=list)      # the NIAS report's reported rows {"name","cas","rt"}
    rows: list[dict] = field(default_factory=list)          # run_job evidence rows (slim combined rows)
    report_lines: list[dict] = field(default_factory=list)  # the client report's lines {rt, name, cas, kind}
    warnings: list[str] = field(default_factory=list)
    timings: dict = field(default_factory=dict)


def run_key(run_dir: str, method: str) -> str:
    name = PureWindowsPath(run_dir).name
    digest = hashlib.sha1(f"{run_dir}|{method}".encode("utf-8")).hexdigest()[:8]
    return re.sub(r"_+", "_", re.sub(r"[^A-Za-z0-9_.-]", "_", f"{name}-{method}"))[:100] + f"-{digest}"


def migration_from_header(h: Header) -> dict:
    """The migration conditions the analyst entered (worksheet header and 'Berechnungen'), in the form of a
    processing method's migration section. Inputs of the evaluation, not decisions."""
    from gcws.quant.migration import _fmt, complete
    m: dict = {"analyst": h.evaluator, "simulant": h.simulant,
               "temperature": f"{_fmt(h.temperature)} °C" if h.temperature is not None else h.temperature_text,
               "duration": h.duration}
    for key, value in (("cell_area_dm2", h.cell_area_dm2), ("occupancy_factor", h.occupancy_factor),
                       ("volume_ml", h.volume), ("ov_ratio", h.sv_ratio)):
        if value is not None:
            m[key] = value
    return complete(m)


def build_spec(entry: CorpusEntry, method: dict, out_dir: Path, migration: Optional[dict] = None) -> dict:
    """The automation job for one evaluated run: a single determination with the blanks of its batch, the
    library search, ISTD detection and the NIAS report, as the user's automation workflow runs it.
    ``migration`` (the analyst's conditions) goes into a copy of the method."""
    import copy
    from gcws.io.sequence import BLANK_ISTD, classify_role
    if migration is not None:
        method = copy.deepcopy(method)
        method.setdefault("sections", {})["migration"] = dict(migration)
    run = Path(entry.run_dir).name
    blanks = {"blank": [b for b in entry.blanks if classify_role(b) != BLANK_ISTD],
              "blank_istd": [b for b in entry.blanks if classify_role(b) == BLANK_ISTD]}
    return {"job_id": "learn", "revision": 1, "mode": "full", "method": method, "batch_folder": entry.batch_dir,
            "group": {"key": run, "name": run[:-2] if run.lower().endswith(".d") else run, "members": [run]},
            "blanks": blanks, "override": {"allow_no_blank": True},
            "reports": [{"node": "learn", "kind": "nias", "formats": ["xlsx"]}], "rules": None,
            "auto_accept": True, "has_review": True, "out_dir": str(out_dir), "search": True, "istd_detect": True,
            "min_confidence": "high", "require_blank": "auto"}


def load_program(path: Path) -> ProgramResult:
    return ProgramResult(**json.loads(Path(path).read_text(encoding="utf-8")))


def parse_program_report(path: Path) -> list[dict]:
    """The lines of the generated NIAS report ('NIAS Result' sheet) as the client sees them: substance lines
    with an RT, and 'Sum of …' lines (kind "sum", no RT). Footnotes below the table are not lines."""
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb["NIAS Result"] if "NIAS Result" in wb.sheetnames else wb.worksheets[0]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    lines: list[dict] = []
    header = next((i for i, r in enumerate(rows) if r and str(r[0] or "").strip().upper().startswith("RT")), None)
    if header is None:
        return lines
    for r in rows[header + 1:]:
        cells = [v for v in r if v is not None and str(v).strip()]
        if not cells:
            if lines:
                break                      # the empty row before the footnotes
            continue
        first = r[0]
        name = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
        if isinstance(first, (int, float)):
            cas = str(r[2]).strip() if len(r) > 2 and isinstance(r[2], str) else ""
            lines.append({"rt": float(first), "name": name, "cas": cas, "kind": "line"})
        else:                              # 'Sum of …' lines: no RT, the text in the name (or first) column
            text = name if first is None else str(first).strip()
            if text.casefold().startswith("sum"):
                lines.append({"rt": None, "name": text, "cas": "", "kind": "sum"})
    return lines


def _report_lines(folder: Path) -> list[dict]:
    files = sorted((folder / "job" / "learn").glob("*.xlsx"))
    try:
        return parse_program_report(files[0]) if files else []
    except Exception:  # noqa: BLE001 - an unreadable report gives no lines; the register list is the fallback
        return []


def _save(prog: ProgramResult, cache: Path) -> None:
    cache.parent.mkdir(parents=True, exist_ok=True)
    tmp = cache.with_suffix(".tmp")
    tmp.write_text(json.dumps(asdict(prog), ensure_ascii=False, default=str), encoding="utf-8")
    os.replace(tmp, cache)


def process(entry: CorpusEntry, out_root: Path, method_name: str = "NIAS", *, method: Optional[dict] = None,
            migration: Optional[dict] = None, force: bool = False, identify=None) -> ProgramResult:
    """Run (or read from the cache) the production job for ``entry``'s run. ``method`` replaces the stored
    processing method ``method_name`` (tests). A job that raises is a failed result, not an exception."""
    from gcws.automation import pipeline as PL
    if method is None:
        from gcws.core import proc_method as PM
        method = PM.load(method_name)
    name = method.get("name") or method_name
    folder = Path(out_root) / "runs" / run_key(entry.run_dir, name)
    cache = folder / "program.json"
    if cache.is_file() and not force:
        prog = load_program(cache)
        if not prog.report_lines and prog.state != "failed":       # cached before report lines were kept
            prog.report_lines = _report_lines(folder)
            if prog.report_lines:
                _save(prog, cache)
        return prog
    try:
        res = PL.run_job(build_spec(entry, method, folder / "job", migration), identify=identify, inspect=collect)
        ev = res.evidence or {}
        prog = ProgramResult(run_dir=entry.run_dir, method=name, state=res.state, reason=res.reason,
                             peaks=(ev.get("inspect") or {}).get("peaks", []),
                             reported=((ev.get("reported") or {}).get("learn") or {}).get("rows", []),
                             rows=ev.get("rows", []), warnings=list(res.warnings), timings=dict(res.timings),
                             report_lines=_report_lines(folder))
    except Exception as exc:  # noqa: BLE001 - one run failing must not stop a corpus run
        prog = ProgramResult(run_dir=entry.run_dir, method=name, state="failed", reason=f"{type(exc).__name__}: {exc}")
    _save(prog, cache)
    return prog

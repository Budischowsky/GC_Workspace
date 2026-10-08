"""Double determination at a glance.

Pick determination A and B (the partner is suggested from the run names),
press *Compare*: chips counting the colours, a verdict per substance in plain language and
a mirror plot (A up, B down) show whether the two determinations agree. The
difference limit is the report parameter ``duplicate_max_reldiff``, so what
is flagged here is exactly what the report flags. Choosing A and B keeps the
replicate group in step, so the reports use the same pair.

The analyst decides what goes out: the *Report* box of each substance (default:
AutoLib's rule), the name and CAS (they become the identification of the peak
in both determinations), areas, concentrations and the mean, and a comment.
Every change is undoable, marked in the table and written to the audit trail;
the NIAS report uses exactly these rows and values.

With the feature pairing (``gcws.features``, the default) *Compare* pairs the peaks
by retention time and spectrum, fills gaps and sets one name per substance
automatically (one undo step), and every row gets a traffic light: green is
taken over, yellow was made consistent automatically, red needs the analyst.
The chips filter by colour and F3 (next open red) leads through the exceptions; the right-click menu
chooses a candidate name, removes a gap fill or takes over harmonised
boundaries.
"""
from __future__ import annotations

import copy
import uuid

import numpy as np
import pyqtgraph as pg
from PySide6.QtCore import QTimer, Qt, Signal as QtSignal
from PySide6.QtGui import QBrush, QColor, QKeySequence, QShortcut
from PySide6.QtWidgets import (QComboBox, QDoubleSpinBox, QFileDialog,
                               QHBoxLayout, QHeaderView, QLabel, QMenu, QPushButton, QSplitter, QStyle,
                               QTableWidgetItem, QToolButton, QVBoxLayout, QWidget)

from gcws.core.model import FID
from gcws.quant import duplicate_view as DV
from gcws.quant import service as QS
from gcws.ui import theme, workers
from gcws.ui.icons import color_chip
from gcws.ui.widgets.cell_marks import EDITED_ROLE, LEVEL_ROLE, CheckDelegate, DiffGaugeDelegate, EditedDelegate
from gcws.ui.widgets.chips import Chip, ElidedLabel

#: SNIP window (min) of the baseline removed from the mirror plot's traces: wider than any peak
BASELINE_WINDOW = 1.0
#: wider views (min) of the mirror plot scale to the substances without the internal standards
WIDE_VIEW = 3.0
WIDE_PERCENTILE = 90
KEYS_NOTE = ("Enter: report · Backspace: not reported · a click or Space: switch the Report box · Delete: delete the "
             "row (Restore under the Deleted chip) · type or F2: edit · Ctrl+C / Ctrl+V · Ctrl+D or drag the small "
             "square of the marking: copy down. Changes are marked, undoable and logged.")
ICON = {"ok": "✔", "warn": "⚠", "bad": "✖", "info": "ℹ", "neutral": "·", "decided": "◉"}
#: icon cell: True on a red row that still waits for the analyst (F3 goes there)
OPEN_ROLE = Qt.UserRole + 1
SORT_ROLE = Qt.UserRole + 2
#: the order of the icon column: open red first, then decided, yellow, green, grey (then by RT)
SEVERITY = {"bad": 0, "decided": 1, "warn": 2, "info": 2, "ok": 3, "neutral": 4}
#: table column -> edited field
C_ICON, C_REPORT, C_RT, C_NAME, C_CAS, C_A1, C_A2, C_C1, C_C2, C_MEAN, C_DIFF, C_VERDICT, C_NOTES, C_COMMENT = range(14)
#: feature pairing only (appended, so the columns above keep their places)
C_FEATURE, C_SIM, C_HIT_A, C_HIT_B = range(14, 18)
FEATURE_COLUMNS = {C_FEATURE, C_SIM, C_HIT_A, C_HIT_B}
#: the concentrations in the further units of the NIAS modes: A, B and the mean of each unit
UNIT_SIDES = (("a", "c1"), ("b", "c2"), ("mean", "mean"))
C_UNIT0 = 18
UNIT_DECIMALS = {"mg_dm2": 4, "ug_dm2": 3, "ug_l": 2, "mg_l": 4, "mg_ml": 6, "mg_g": 6, "mg_kg": 4, "ug_g": 4,
                 "ug_kg": 2}
#: every column's name by index: the column choice is remembered by name
COLUMN_KEYS = ["icon", "report", "rt", "name", "cas", "area_a", "area_b", "conc_a", "conc_b", "mean", "diff",
               "verdict", "notes", "comment", "feature", "similarity", "hit_a", "hit_b"] + \
    [f"{unit}_{side}" for unit in QS.CONC_UNITS for side, _field in UNIT_SIDES]
UNIT_OF = {C_UNIT0 + i: key.rsplit("_", 1)[0] for i, key in enumerate(COLUMN_KEYS[C_UNIT0:])}
#: columns that cannot be hidden, and the columns hidden until the analyst shows them
FIXED_COLUMNS = {C_ICON, C_REPORT, C_NAME}
HIDDEN_BY_DEFAULT = (C_A1, C_A2, C_NOTES, C_HIT_A, C_HIT_B) + tuple(range(C_UNIT0, len(COLUMN_KEYS)))
COLUMNS_SETTING = "replicates/columns"
#: the filter chips: key -> (label, level, tooltip)
FILTERS = {"all": ("All", "info", "Every substance"),
           "check": ("To check", "accent", "Red and yellow rows and the rows you changed"),
           "red": ("Red", "bad", "Your decision: found in one determination only, spectra differ, difference too "
                                 "large. F3: next open red row."),
           "yellow": ("Yellow", "warn", "Made consistent automatically: a quick look"),
           "green": ("Green", "ok", "Confirmed in both determinations"),
           "grey": ("Grey", "neutral", "Not reported anyway: below the reporting limit or at blank level"),
           "deleted": ("Deleted", "neutral", "The rows you deleted: not in the list, the counts or the report. "
                                             "Right-click: Restore row.")}
FIELD_OF = {C_REPORT: "report", C_NAME: "name", C_CAS: "cas", C_A1: "a1", C_A2: "a2", C_C1: "c1", C_C2: "c2",
            C_MEAN: "mean", C_COMMENT: "comment"}


class _AbsAxis(pg.AxisItem):
    """Intensity axis of the mirror plot: B is drawn downwards but has positive values too."""

    def tickStrings(self, values, scale, spacing):
        return super().tickStrings([abs(v) for v in values], scale, spacing)


def _stale(group: dict) -> None:
    """An analyst change after the acceptance: the acceptance no longer holds (kept, to say who accepted)."""
    if group.get("signoff"):
        group["signoff"] = dict(group["signoff"], stale=True)


class _SeverityItem(QTableWidgetItem):
    """A cell that sorts by ``SORT_ROLE`` instead of its text: the icon by severity, the Report box
    by its state, each then by RT."""

    def __lt__(self, other):
        a, b = self.data(SORT_ROLE), other.data(SORT_ROLE)
        if a is None or b is None:
            return super().__lt__(other)
        return a < b


class DuplicatePage(QWidget):
    reportRequested = QtSignal(str, str)        # kind, group id
    previewRequested = QtSignal(str, str)       # kind, group id
    report2Requested = QtSignal()
    summaryChanged = QtSignal()                 # the counts of the rows changed (compared or edited)
    acceptRequested = QtSignal(str)             # group id: the analyst accepted this double determination

    def __init__(self, ws, set_groups, parent=None):
        super().__init__(parent)
        self.ws = ws
        self.set_groups = set_groups            # callable(groups, text): undoable group change
        self.rows: list[dict] = []               # with the analyst's edits applied
        self.base_rows: list[dict] = []          # as merged from the integrations
        self.verdicts: list = []
        self._filling = False
        self.members: list[str] = []
        self._loading = False
        self._view_processed = False

        self.a = QComboBox()
        self.b = QComboBox()
        for c in (self.a, self.b):
            c.setMinimumContentsLength(26)
            c.setSizeAdjustPolicy(QComboBox.AdjustToMinimumContentsLengthWithIcon)
        self.a.activated.connect(lambda *_: self._a_picked())
        self.b.activated.connect(lambda *_: self.compare(sync=True))
        swap = QToolButton()
        swap.setText("⇄")
        swap.setToolTip("Swap A and B")
        swap.clicked.connect(self._swap)
        self.b_compare = QPushButton("Compare")
        theme.set_primary(self.b_compare)
        self.b_compare.clicked.connect(lambda: self.compare(sync=True))
        self.b_report2 = QPushButton("Load from Report²…")
        self.b_report2.setToolTip("Switch to a processed double determination shown in Report²")
        self.b_report2.clicked.connect(self.report2Requested.emit)
        self.b_more = more = QToolButton()
        more.setText("More")
        more.setToolTip("Settings, three or more determinations, reset")
        more.setPopupMode(QToolButton.InstantPopup)
        menu = QMenu(more)
        menu.addAction("Settings…", self.edit_settings).setToolTip(
            "Pairing (features or classic), gap filling, consensus name")
        menu.addAction("3+ determinations…", self._to_groups_tab)
        menu.addSeparator()
        menu.addAction("Reset all", self.reset_all).setToolTip("Undo every change made in this double determination")
        menu.setToolTipsVisible(True)
        more.setMenu(menu)
        # Compare and Report² first: they stay in view in a narrow dock (choosing A or B compares anyway)
        pick = QHBoxLayout()
        for w in (self.b_compare, self.b_report2, QLabel("A"), self.a, swap, QLabel("B"), self.b):
            pick.addWidget(w)
        pick.addStretch(1)
        pick.addWidget(more)

        self.limit = QDoubleSpinBox()
        self.limit.setRange(0.0, 200.0)
        self.limit.setDecimals(1)
        self.limit.setSuffix(" %")
        self.limit.setToolTip("Maximum relative difference |A−B| / mean. This is the report parameter "
                              "'Duplicate difference limit': changing it changes the report too.")
        self.limit.editingFinished.connect(self._limit_changed)
        # the chips count the rows of each colour and filter the list (clicked again: all rows)
        self.filter = "all"
        self.counts: dict[str, int] = {}
        self.chips: dict[str, Chip] = {}
        lim = QHBoxLayout()
        lim.setSpacing(4)
        for key, (label, level, tip) in FILTERS.items():
            chip = Chip()
            chip.setToolTip(tip)
            chip.clicked.connect(lambda k=key: self.set_filter(k))
            theme.set_chip(chip, label, level)
            self.chips[key] = chip
            lim.addWidget(chip)
        lim.addStretch(1)
        lim.addWidget(QLabel("Difference limit"))
        lim.addWidget(self.limit)
        self.limit_chip = theme.chip("report parameter", "info")
        lim.addWidget(self.limit_chip)

        self.banner = ElidedLabel()               # one line; the whole text is its tooltip
        self.banner.setObjectName("chip")

        from gcws.ui.widgets.sheet_table import SheetTable
        self.table = SheetTable(0, 0)            # Excel-like: keys, Ctrl+C/V/D, fill handle
        self.table.verticalHeader().setVisible(False)
        self.table.itemChanged.connect(self._cell_edited)
        self.table.markRequested.connect(self._mark_rows)
        self.table.toggleRequested.connect(self._toggle_rows)
        self.table.deleteRequested.connect(self._delete_rows)
        self.table.bulkEdit.connect(self._bulk_edit)
        self.table.setAlternatingRowColors(True)
        self.table.setItemDelegate(EditedDelegate(self.table))     # changed cells: a corner mark
        self.table.setItemDelegateForColumn(C_DIFF, DiffGaugeDelegate(lambda: DV.limits(self.ws)[0], self.table))
        self.table.setItemDelegateForColumn(C_REPORT, CheckDelegate(self._report_clicked, self.table))
        self._sort = (C_ICON, Qt.AscendingOrder)                   # most severe first, until a header is clicked
        # the severity order as of the last Compare, chip or header click: an edited row keeps its place
        self._frozen: dict | None = None
        hh = self.table.horizontalHeader()
        hh.sortIndicatorChanged.connect(self._sort_changed)
        hh.setContextMenuPolicy(Qt.CustomContextMenu)
        hh.customContextMenuRequested.connect(self._column_menu)
        self._nav = QTimer(self)                 # arrow keys: jump to the peak once the cursor rests
        self._nav.setSingleShot(True)
        self._nav.setInterval(150)
        self._nav.timeout.connect(self._row_selected)
        self.table.currentItemChanged.connect(lambda *_: self._nav.start())
        self.table.cellDoubleClicked.connect(self._open_row)
        self.table.setContextMenuPolicy(Qt.CustomContextMenu)
        self.table.customContextMenuRequested.connect(self._context_menu)
        QShortcut(QKeySequence(Qt.Key_F3), self, activated=self.next_red)
        self.table_features = None                   # the FeatureTable of the rows (feature pairing)
        self._search = None                          # the running consensus search
        self._comparing = False
        self._refresh = QTimer(self)                 # a comparison asked for while one is running
        self._refresh.setSingleShot(True)
        self._refresh.setInterval(0)
        self._refresh.timeout.connect(self._refresh_now)

        self.mirror = pg.PlotWidget(axisItems={"left": _AbsAxis("left")})
        self.mirror.setMenuEnabled(False)
        self.mirror.showGrid(x=True, y=True, alpha=theme.PLOT["grid_alpha"])
        self.mirror.setLabel("bottom", "RT (FID)", units="min")
        self.mirror.setLabel("left", "A  ↑   FID   ↓  B")
        self.mirror.getAxis("left").setWidth(62)
        self.mirror.getAxis("bottom").enableAutoSIPrefix(False)
        self._traces, self._marks = [], []
        vb = self.mirror.getViewBox()
        vb.setMouseEnabled(x=True, y=False)          # wheel / drag: time only, the intensity follows
        self._fit_timer = QTimer(self)
        self._fit_timer.setSingleShot(True)
        self._fit_timer.setInterval(0)
        self._fit_timer.timeout.connect(self._fit_y)
        vb.sigXRangeChanged.connect(lambda *_: self._fit_timer.start())
        theme.register_plot(self.mirror, lambda: self._draw_mirror() if self.members else None)
        self.mirror.scene().sigMouseClicked.connect(lambda ev: self.full_view() if ev.double() else None)
        self.mirror.setToolTip("A up, B down, in FID signal units. Wheel or drag: time; double-click: "
                               "whole chromatogram")
        self.cursor = pg.InfiniteLine(angle=90, movable=False,
                                      pen=pg.mkPen(theme.ACCENT, width=1, style=Qt.DashLine))
        self.spec = pg.PlotWidget(axisItems={"left": _AbsAxis("left")})
        self.spec.setMenuEnabled(False)
        self.spec.showGrid(x=False, y=True, alpha=theme.PLOT["grid_alpha"])
        self.spec.setLabel("bottom", "m/z")
        self.spec.setLabel("left", "A  ↑   %   ↓  B")
        self.spec.getAxis("left").setWidth(62)
        self.spec.setToolTip("The spectra of the selected substance: A up, B down (co-eluting ions, "
                             "background subtracted), each scaled to its base peak")
        theme.register_plot(self.spec, lambda: self._draw_spectra())
        self.plots = plots = QSplitter(Qt.Horizontal)
        plots.addWidget(self.mirror)
        plots.addWidget(self.spec)
        plots.setSizes([520, 300])
        self._bounds = []                            # the shaded peaks of the selected substance
        self._signed = {}                            # determination index -> (rt, signed signal) as drawn
        split = QSplitter(Qt.Vertical)
        split.addWidget(self.table)
        split.addWidget(plots)
        split.setSizes([320, 220])

        from gcws.ui.widgets.report_button import report_button
        buttons = QHBoxLayout()
        self.b_report = report_button(
            self, lambda: self._report("hs_screening" if self.quant_signal() == "TIC" else "nias", preview=True),
            lambda kind: self._report(kind), lambda: self.export())
        buttons.addWidget(self.b_report)
        buttons.addStretch(1)
        self.b_bounds = QPushButton("Harmonise boundaries")
        self.b_bounds.setToolTip("Take over every proposed integration boundary (one undo step)")
        self.b_bounds.clicked.connect(lambda: self.apply_boundaries(None))
        self.b_bounds.setVisible(False)
        buttons.addWidget(self.b_bounds)
        self.b_accept = QPushButton("Accept double determination")
        theme.set_primary(self.b_accept)
        self.b_accept.clicked.connect(self.accept)
        self.b_accept.setEnabled(False)
        buttons.addWidget(self.b_accept)

        lay = QVBoxLayout(self)
        lay.setContentsMargins(4, 4, 4, 4)
        lay.setSpacing(6)
        self.edit_note = theme.hint("", False)
        from PySide6.QtCore import QSettings
        self.b_plots = QToolButton()
        self.b_plots.setText("Plots")
        self.b_plots.setCheckable(True)
        self.b_plots.setToolTip("Show or hide the chromatograms and spectra below the list")
        self.b_plots.toggled.connect(self.set_plots_visible)
        self.b_columns = QToolButton()
        self.b_columns.setText("Columns...")
        self.b_columns.setToolTip("Choose and order the columns of the list, e.g. the concentrations in µg/dm², "
                                  "µg/L, mg/L or mg/mL")
        self.b_columns.clicked.connect(self.choose_columns)
        keys = QToolButton()
        keys.setText("?")
        keys.setAutoRaise(True)
        keys.setToolTip(KEYS_NOTE)
        status = QHBoxLayout()
        status.addWidget(self.banner, 1)
        status.addWidget(self.edit_note)
        status.addWidget(self.b_columns)
        status.addWidget(self.b_plots)
        status.addWidget(keys)
        lay.addLayout(pick)
        lay.addLayout(lim)
        lay.addLayout(status)
        lay.addWidget(split, 1)
        lay.addLayout(buttons)

        self.b_plots.setChecked(QSettings().value("replicates/plots", True, type=bool))
        self.set_plots_visible(self.b_plots.isChecked())
        ws.runAdded.connect(lambda *_: self.refresh_choices())
        ws.runRemoved.connect(lambda *_: self.refresh_choices())
        ws.runChanged.connect(lambda *_: self.refresh_choices())
        ws.quantChanged.connect(self._quant_changed)
        ws.replicatesChanged.connect(self._reapply)
        self.refresh_choices()

    # -- choices ---------------------------------------------------------------------

    def _to_groups_tab(self):
        p = self.parent()
        while p is not None and not hasattr(p, "show_nfold"):
            p = p.parent()
        if p is not None:
            p.show_nfold()

    def _candidates(self):
        return [s for s in self.ws.states() if s.role in ("sample", "standard")]

    def refresh_choices(self):
        cur_a, cur_b = self.a.currentData(), self.b.currentData()
        self._loading = True
        for combo in (self.a, self.b):
            combo.clear()
            for st in self._candidates():
                combo.addItem(color_chip(st.color), st.name, st.id)
        self.b.addItem("— single determination —", "")
        self._loading = False
        ids = [s.id for s in self._candidates()]
        if cur_a in ids:
            self._select(self.a, cur_a)
            self._select(self.b, cur_b if (cur_b in ids or cur_b == "") else DV.suggest_partner(self.ws, cur_a))
        elif ids:
            self.set_pair(self._default_a())
        self.show_limit()

    def _default_a(self):
        act = self.ws.active
        if act is not None and act.role in ("sample", "standard"):
            return act.id
        c = self._candidates()
        return c[0].id if c else None

    @staticmethod
    def _select(combo, data):
        i = combo.findData(data if data is not None else "")
        combo.setCurrentIndex(max(0, i))

    def set_pair(self, a, b=None, compare=True, processed=False):
        """Show the double determination of ``a`` (partner ``b`` or the suggested one)."""
        if a is None:
            return
        self._view_processed = processed
        self._select(self.a, a)
        partner = b if b is not None else DV.suggest_partner(self.ws, a)
        self._select(self.b, partner or "")
        if compare:
            self.compare(sync=b is not None and not processed)

    def _a_picked(self):
        a = self.a.currentData()
        self._select(self.b, DV.suggest_partner(self.ws, a) or "")
        self.compare(sync=True)

    def _swap(self):
        a, b = self.a.currentData(), self.b.currentData()
        if not b:
            return
        self._select(self.a, b)
        self._select(self.b, a)
        self.compare(sync=True)

    # -- groups -------------------------------------------------------------------------

    def group(self):
        """The replicate group holding exactly the chosen determinations, if any."""
        want = [m for m in (self.a.currentData(), self.b.currentData()) if m]
        for g in self.ws.replicate_groups:
            if [m for m in g["members"] if m in self.ws.runs] == want:
                return g
        return None

    def _sync_group(self, members):
        if not members or self.group() is not None:
            return
        groups = copy.deepcopy(self.ws.replicate_groups)
        target = next((g for g in groups if members[0] in g["members"] and len(g["members"]) <= 2), None)
        if target is None:
            from gcws.io.sequence import replicate_stem
            name = replicate_stem(self.ws.runs[members[0]].run.path.name) or self.ws.runs[members[0]].name
            target = {"id": uuid.uuid4().hex[:8], "name": name, "members": [], "policy": "all"}
            groups.append(target)
        for g in groups:                        # a run belongs to one pair or single; groups of 3+ keep theirs
            if g is not target and len(g["members"]) <= 2:
                g["members"] = [m for m in g["members"] if m not in members]
        groups = [g for g in groups if g["members"] or g is target]
        target["members"] = list(members)
        target["policy"] = "all"
        names = " / ".join(self.ws.runs[m].name for m in members)
        self.set_groups(groups, f"double determination {names}")

    # -- compute ---------------------------------------------------------------------------

    def _quant_changed(self):
        """A new quantification: compare again -- never inside a running comparison (the building
        of the feature table recomputes the quantification itself), and once for several changes."""
        if self.isVisible() and self.members:
            self._refresh.start()
        self.show_limit()

    def labels(self):
        from gcws.io.sequence import replicate_label
        out = []
        for i, m in enumerate(self.members):
            lab = replicate_label(self.ws.runs[m].run.path.name) if m in self.ws.runs else ""
            out.append(lab if lab and lab not in out else "AB"[i] if i < 2 else str(i + 1))
        return tuple(out) if len(out) == 2 else ("A", "B")

    def edits_key(self):
        return DV.edits_key(self.ws.quant, self.ws.quant_unit())

    def quant_signal(self):
        from gcws.quant.service import quant_detector
        return quant_detector(self.ws.quant)

    def _refresh_now(self):
        if self._comparing:
            self._refresh.start(50)
        elif self.isVisible() and self.members:
            self.compare(sync=False)

    def compare(self, sync: bool = False):
        if sync:
            self._view_processed = False
        if self._comparing:                         # re-entered through a signal: once more afterwards
            self._refresh.start()
            return
        self._comparing = True
        window = self.window()
        activity = f"compare:{id(self)}"
        if hasattr(window, "begin_activity"):
            window.begin_activity(activity, "Comparing determinations…")
        try:
            self._compare(sync)
        finally:
            if hasattr(window, "end_activity"):
                window.end_activity(activity)
            self._comparing = False

    def _compare(self, sync: bool = False):
        key = self.quant_signal()
        self.mirror.setLabel("bottom", f"RT ({key})", units="min")
        self.mirror.setLabel("left", f"A  ↑   {key}   ↓  B")
        self.mirror.setToolTip(f"A up, B down, in {key} signal units. Wheel or drag: time; double-click: full view.")
        if self._loading:
            return
        a, b = self.a.currentData(), self.b.currentData()
        self.members = [m for m in (a, b) if m]
        if not self.members:
            self._show([], [], ["load the determinations and give them role Sample"])
            return
        if sync:
            self._sync_group(self.members)
        self.table_features = None
        if len(self.members) >= 2 and self.features_mode() and not DV.member_problems(self.ws, self.members):
            from gcws.features import service as SV
            window = self.window()
            progress = (lambda message: window.update_activity(f"compare:{id(self)}", message)) \
                if hasattr(window, "update_activity") else None
            if self.ws.quant_result is None:
                self.ws.recompute_quant()
            # a deliberate Compare makes the automatic changes (each only once); a refresh only shows
            stack = self._stack()
            before = (stack.count(), stack.index())
            table = SV.run(self.ws, self.members, apply_auto=None if sync else False, stack=stack,
                           progress=progress, search=False)
            if (stack.count(), stack.index()) != before:
                self.ws.recompute_quant()              # the gap-filled peaks get their concentrations now
            self.table_features = DV.features_table(self.ws, self.members, table)
        window = self.window()
        if hasattr(window, "update_activity"):
            window.update_activity(f"compare:{id(self)}", "Preparing comparison table…")
        rows, problems = DV.compute(self.ws, self.members, "all")
        limit, rl = DV.limits(self.ws)
        labels = self.labels()
        unit = self.ws.quant_unit()
        verdicts = [DV.plain_verdict(r, limit, rl, labels, unit) for r in rows]
        self.base_rows = rows
        if hasattr(window, "update_activity"):
            window.update_activity(f"compare:{id(self)}", "Drawing comparison results…")
        self._show(DV.apply_edits(rows, verdicts, self.edits(), self._tol()), verdicts, problems)
        if self.table_features is not None and not self._view_processed:
            self._start_consensus_search()

    # -- analyst edits ------------------------------------------------------------------------

    def _tol(self) -> float:
        from gcws.quant.nias_bridge import make_settings
        return float(getattr(make_settings(self.ws.quant.get("settings")), "rt_tolerance", 0.035) or 0.035)

    def edits(self) -> dict:
        g = self.group()
        return dict((g or {}).get(self.edits_key()) or {})

    def _reapply(self):
        if not self.base_rows or self._filling:
            return
        self.rows = DV.apply_edits(self.base_rows, self.verdicts, self.edits(), self._tol())
        if self.rows and self.rows[0].get("light"):
            self._show_lights(self.rows)
        elif self.rows:
            self._show_summary(self.rows, self.labels())
        theme._repolish(self.banner)
        self._fill_table()
        self._update_report_note()

    def _update_report_note(self):
        rows = [r for r in self.rows if not r.get("deleted")]
        n = sum(1 for r in rows if r.get("report"))
        changed = sum(1 for r in rows if r.get("edited"))
        deleted = len(self.rows) - len(rows)
        text = f"{n} of {len(rows)} in the report"
        if changed:
            text += f" · {changed} changed"
        if deleted:
            text += f" · {deleted} deleted"
        self.edit_note.setText(text)
        self.edit_note.setToolTip(f"{n} of {len(rows)} substances go into the report"
                                  + (f"; {changed} changed by the analyst" if changed else "")
                                  + (f"; {deleted} deleted (the Deleted chip shows them)" if deleted else ""))

    def _live(self):
        """``(rows, verdicts)`` without the rows the analyst deleted."""
        pairs = [(r, v) for r, v in zip(self.rows, self.verdicts) if not r.get("deleted")]
        return [r for r, _ in pairs], [v for _, v in pairs]

    def _names(self) -> str:
        return " / ".join(self.ws.runs[m].name for m in self.members if m in self.ws.runs)

    def set_edit(self, row: dict, field: str, value) -> None:
        """Store (or with ``None`` remove) one analyst change of ``row`` in the replicate group."""
        self.set_edits([(row, field, value)])

    def set_edits(self, changes: list, label: str | None = None) -> None:
        """Several ``(row, field, value)`` changes as one undo step (``label``: its text); each is logged."""
        from datetime import datetime
        from gcws.core.audit import current_user
        if not changes:
            return
        self._sync_group(self.members)
        g = self.group()
        if g is None:
            return
        groups = copy.deepcopy(self.ws.replicate_groups)
        tg = next(x for x in groups if x["id"] == g["id"])
        edits = tg.setdefault(self.edits_key(), {})
        logs = []
        for row, field, value in changes:
            key = row.get("edit_key") or DV.edit_key(row["rt"])
            e = edits.setdefault(key, {"rt": float(row["rt"])})
            before = row.get(field) if field not in DV.NUMERIC_EDITS else \
                row.get("edited", {}).get(field, row.get(field))
            if value is None:
                e.pop(field, None)
            else:
                e[field] = value
            e["by"], e["at"] = current_user(), datetime.now().isoformat(timespec="seconds")
            if not any(k in e for k in DV.NUMERIC_EDITS + DV.ROW_FLAGS):
                edits.pop(key, None)
            name = row.get("name") or f"RT {row['rt']:.3f}"
            if field == "deleted":
                logs.append((f"{name} (RT {row['rt']:.3f})", "", "row deleted" if value else "row restored"))
            else:
                logs.append((f"{name} (RT {row['rt']:.3f}): {field}", "" if before is None else str(before),
                             "reset" if value is None else str(value)))
        if label is None and len(changes) == 1:
            label = f"double determination: {changes[0][1]} of {name}"
        elif label is None:
            label = f"double determination: {len(changes)} cells changed"
        _stale(tg)
        self.set_groups(groups, label)
        for what, before, after in logs:
            self.ws.log("Double determination changed", self._names(), what, before, after)

    def set_identity(self, row: dict, name: str | None = None, cas: str | None = None) -> None:
        """Name / CAS of a substance: the identification of its FID peak in both determinations."""
        self.set_identities([(row, name, cas)])

    def set_identities(self, changes: list) -> None:
        """``(row, name, cas)`` of several substances as one undo step (None keeps a value)."""
        from gcws.core.ident import Identification
        from gcws.ui.undo import IdentCommand, MultiCommand
        per_run: dict = {}                       # one command per run: each command snapshots the set
        for row, name, cas in changes:
            for key, i in (("source1", 0), ("source2", 1)):
                src = row.get(key)
                if src is None or i >= len(self.members) or self.members[i] not in self.ws.runs:
                    continue
                rid = self.members[i]
                st = self.ws.runs[rid]
                res = self.ws.result(rid, self.quant_signal())
                if res is None or not res.peaks:
                    continue
                j = min(range(len(res.peaks)), key=lambda q: abs(res.peaks[q].apex_rt - src["rt"]))
                peak = res.peaks[j]
                if abs(peak.apex_rt - src["rt"]) > 0.05:
                    continue
                old = st.ident_set(self.quant_signal()).for_peak(peak)
                ident = Identification(apex_rt=peak.apex_rt,
                                       name=name if name is not None else (old.name if old else ""),
                                       cas=cas if cas is not None else (old.cas if old else ""),
                                       score=old.score if old else None, status="Accepted (analyst)",
                                       formula=old.formula if old else "", library=old.library if old else "",
                                       hits=list(old.hits) if old else [], source="double determination",
                                       manual=True, istd=old.istd if old else "")
                what = "name" if name is not None else "CAS"
                value = name if name is not None else cas
                per_run.setdefault(rid, []).append((peak.apex_rt, ident,
                                                    f"{what} of peak {peak.apex_rt:.3f} = {value!r}"))
        cmds = [IdentCommand(self.ws, rid, self.quant_signal(), [(t, ident) for t, ident, _ in items],
                             items[0][2] if len(items) == 1 else f"{len(items)} names / CAS")
                for rid, items in per_run.items()]
        if not cmds:
            return
        if len(changes) == 1:
            _row, name, cas = changes[0]
            label = f"double determination: {'name' if name is not None else 'CAS'} = " \
                    f"{name if name is not None else cas}"
        else:
            label = f"double determination: {len(changes)} names / CAS changed"
        g = self.group()
        if g is None or not (g.get("signoff") and not g["signoff"].get("stale")):
            self._stack().push(MultiCommand(label, cmds))
            return
        stack = self._stack()                    # the names change what was accepted: one undo step for both
        stack.beginMacro(label)
        try:
            stack.push(MultiCommand(label, cmds))
            groups = copy.deepcopy(self.ws.replicate_groups)
            _stale(next(x for x in groups if x["id"] == g["id"]))
            self.set_groups(groups, label)
        finally:
            stack.endMacro()

    def _stack(self):
        return self.ws.undo_group.activeStack() or self.ws.project_undo

    def _row_index(self, visual_row: int):
        it = self.table.item(visual_row, 0)
        k = it.data(Qt.UserRole) if it is not None else None
        return k if k is not None and 0 <= k < len(self.rows) else None

    def _report_value(self, k: int, on: bool):
        """The stored report edit: on a red row always the analyst's answer (it decides the row);
        elsewhere None where ``on`` is what the default rule gives anyway."""
        if self.verdicts[k].level == "bad":
            return on
        return None if on == DV.default_report(self.base_rows[k], self.verdicts[k]) else on

    def _mark_rows(self, visual_rows: list, on: bool):
        """Enter / Backspace: the marked substances go into the report, or not."""
        self._mark_keys([self._row_index(r) for r in visual_rows], on)

    def _mark_keys(self, keys: list, on: bool):
        """The substances ``keys`` (indices into ``rows``) go into the report, or not (one undo step);
        an open red row is decided even where ``on`` is already its value."""
        changes = []
        for k in dict.fromkeys(keys):
            if k is None or not 0 <= k < len(self.rows):
                continue
            row = self.rows[k]
            if bool(row.get("report")) != on or DV.is_open(row, self.verdicts[k]):
                changes.append((row, "report", self._report_value(k, on)))
        self.set_edits(changes)

    def _toggle_rows(self, visual_rows: list):
        """Space on the Report box: every marked row switches."""
        self._toggle_keys([self._row_index(r) for r in visual_rows])

    def _toggle_keys(self, keys: list):
        """Into the report, unless every one of ``keys`` is reported already: then out of it."""
        keys = [k for k in keys if k is not None and 0 <= k < len(self.rows)]
        if keys:
            self._mark_keys(keys, not all(self.rows[k].get("report") for k in keys))

    def _delete_rows(self, visual_rows: list):
        """Delete: the marked substances leave the list, the counts and the report (one undo step);
        when all of them are deleted already, they are restored."""
        keys = [k for k in dict.fromkeys(self._row_index(r) for r in visual_rows) if k is not None]
        self.delete_keys(keys, not all(self.rows[k].get("deleted") for k in keys))

    def delete_keys(self, keys: list, deleted: bool = True) -> int:
        """Delete (or restore) the substances ``keys`` (indices into ``rows``); returns how many changed."""
        rows = [self.rows[k] for k in keys if 0 <= k < len(self.rows) and bool(self.rows[k].get("deleted")) != deleted]
        if not rows:
            return 0
        what = "deleted" if deleted else "restored"
        one = rows[0].get("name") or f"RT {rows[0]['rt']:.3f}"
        self.set_edits([(r, "deleted", True if deleted else None) for r in rows],
                       f"double determination: {one} {what}" if len(rows) == 1 else
                       f"double determination: {len(rows)} rows {what}")
        n = len(rows)
        self.ws.message.emit(f"{n} row{'s' if n != 1 else ''} {what}" +
                             (" · Ctrl+Z undoes · the Deleted chip shows them" if deleted else ""))
        return n

    def _report_clicked(self, index):
        """A click in a Report cell: switched once the click is over (the table is rebuilt then)."""
        k = index.data(Qt.UserRole)
        if k is not None:
            QTimer.singleShot(0, lambda: self._toggle_keys([k]))

    def _bulk_edit(self, cells: list):
        """Paste, Ctrl+D and the fill handle: many cells as one undo step."""
        from gcws.ui.docks.peak_table import parse_number
        edits, idents, bad = [], [], []
        for r, c, value in cells:
            k, field = self._row_index(r), FIELD_OF.get(c)
            if k is None or field is None:
                continue
            row = self.rows[k]
            if field == "report":
                edits.append((row, "report", self._report_value(k, bool(value))))
            elif field in ("name", "cas"):
                text = str(value).strip()
                if field == "name" and not text:
                    continue
                if text != (row.get(field) or ""):
                    idents.append((row, text, None) if field == "name" else (row, None, text))
            elif field == "comment":
                edits.append((row, "comment", str(value).strip() or None))
            else:
                text = str(value).strip()
                number = parse_number(text) if text else None
                if text and number is None:
                    bad.append(text)
                    continue
                edits.append((row, field, number))
        if bad:
            self.ws.message.emit(f"not a number, left out: {', '.join(sorted(set(bad))[:5])}")
        if not edits and not idents:
            return
        stack = self._stack()
        stack.beginMacro(f"double determination: {len(edits) + len(idents)} cells changed")
        try:
            self.set_identities(idents)
            self.set_edits(edits)
        finally:
            stack.endMacro()

    def _cell_edited(self, item):
        if self._filling:
            return
        field = FIELD_OF.get(item.column())
        k = item.data(Qt.UserRole)
        if field is None or k is None or not (0 <= k < len(self.rows)):
            return
        row = self.rows[k]
        if field == "report":
            self.set_edit(row, "report", self._report_value(k, item.checkState() == Qt.Checked))
            return
        text = item.text().strip()
        if field in ("name", "cas"):
            if text != (row.get(field) or ""):
                if field == "name" and not text:
                    self._reapply()
                    return
                self.set_identity(row, **{field: text})
            return
        if field == "comment":
            self.set_edit(row, "comment", text or None)
            return
        from gcws.ui.docks.peak_table import parse_number
        value = parse_number(text) if text else None
        if text and value is None:
            self.ws.message.emit(f"'{text}' is not a number")
            self._reapply()
            return
        if value is None or value == row.get(field):
            self.set_edit(row, field, None if value is None else value)
            return
        self.set_edit(row, field, value)

    def reset_row(self, row=None):
        row = row if row is not None else self._current_row()
        if row is None or row.get("edit_key") not in self.edits():
            return                               # nothing the analyst changed: nothing to reset
        self._sync_group(self.members)
        g = self.group()
        groups = copy.deepcopy(self.ws.replicate_groups)
        tg = next(x for x in groups if x["id"] == g["id"])
        tg.get(self.edits_key(), {}).pop(row["edit_key"], None)
        _stale(tg)
        self.set_groups(groups, f"double determination: reset {row.get('name') or row['rt']}")
        self.ws.log("Double determination changed", self._names(), f"{row.get('name')}: changes reset")

    def reset_all(self):
        g = self.group()
        if g is None or not g.get(self.edits_key()):
            return
        groups = copy.deepcopy(self.ws.replicate_groups)
        tg = next(x for x in groups if x["id"] == g["id"])
        tg[self.edits_key()] = {}
        _stale(tg)
        self.set_groups(groups, "double determination: all changes reset")
        self.ws.log("Double determination changed", self._names(), "all changes reset")

    def _show(self, rows, verdicts, problems):
        self.rows, self.verdicts = rows, verdicts
        labels = self.labels()
        if problems:
            self._update_chips()
            self.banner.setText("Cannot compare yet: " + "; ".join(problems))
            self.banner.setProperty("level", "warn")
        elif rows and rows[0].get("light"):
            self._show_lights(rows)
        else:
            self._show_summary(rows, labels)
        theme._repolish(self.banner)
        self.spec.setVisible(self.table_features is not None)      # classic pairing: no spectra to mirror
        self._frozen = None                                         # compared again: sorted afresh
        self._fill_table()
        self._update_report_note()
        self._draw_mirror()

    def _show_summary(self, rows, labels):
        """Chips and status line of the classic pairing."""
        live, verdicts = self._live()
        s = DV.summarize(live, verdicts, DV.limits(self.ws)[0], labels)
        self._update_chips()
        if len(self.members) != 2:
            self.banner.setText("Single determination: choose a partner B to compare.")
            self.banner.setProperty("level", "neutral")
            return
        level, text = s.level, s.text
        n_open, n_decided = self.decision_counts()
        if n_decided and not n_open:
            level, text = "ok", f"All decisions made ({n_decided} red decided by you): preview the report. " + text
        elif n_decided:
            text = f"Red: {n_open + n_decided} → {n_open} open · {n_decided} decided (Accept keeps the default for " \
                   "the open ones). " + text
        if s.mean_reldiff is not None:
            text += f" Mean difference {s.mean_reldiff:.1f} %."
        self.banner.setText(("✔  " if level == "ok" else "⚠  ") + text)
        self.banner.setProperty("level", level)

    def decision_counts(self) -> tuple[int, int]:
        """(red rows still open, red rows the analyst has decided); deleted rows do not count."""
        red = [r for r, v in zip(self.rows, self.verdicts) if v.level == "bad" and not r.get("deleted")]
        decided = sum(1 for r in red if r.get("decided"))
        return len(red) - decided, decided

    def _update_chips(self):
        """Count the rows of each colour on the chips; the chip of the current filter is marked."""
        lights = [DV.light_of(r, v) for r, v in zip(self.rows, self.verdicts)]
        n = {key: sum(1 for r, v, lt in zip(self.rows, self.verdicts, lights) if self._passes(key, r, v, lt))
             for key in FILTERS}
        self.counts = n
        if self.filter == "deleted" and not n["deleted"]:
            self.filter = "all"                      # every deleted row restored: the whole list again
        n_open, n_decided = self.decision_counts()
        for key, (label, level, tip) in FILTERS.items():
            text = f"{label} {n[key]}"
            if key == "red" and n_decided:
                text += f" · {n_open} open"
            theme.set_chip(self.chips[key], text if key != "deleted" or n[key] else "", level)
            self.chips[key].setProperty("selected", key == self.filter)
            theme._repolish(self.chips[key])
        lone = sum(1 for r, lt in zip(self.rows, lights) if lt == "red" and not r.get("deleted")
                   and str(r.get("status", "")).startswith("Artefact"))
        self.chips["red"].setToolTip(FILTERS["red"][2] + (f" {lone} of them found in one determination only."
                                                          if lone else ""))
        self._update_accept()
        self.summaryChanged.emit()

    def _update_accept(self) -> None:
        """Accept: possible once two determinations are compared; red rows still open keep their default."""
        g = self.group()
        signoff = (g or {}).get("signoff") or {}
        n_open = self.decision_counts()[0]
        ready = len(self.members) == 2 and bool(self.rows)
        if signoff and not signoff.get("stale"):
            left = signoff.get("open") or 0
            self.b_accept.setText(f"Accepted by {signoff.get('by', '?')} ✔")
            self.b_accept.setToolTip(f"Accepted on {str(signoff.get('at', '')).replace('T', ' ')}"
                                     + (f" with {left} red row{'s' if left != 1 else ''} left open (default)"
                                        if left else "")
                                     + "; any change of a value, a Report box or a name reopens it")
            self.b_accept.setEnabled(False)
            return
        self.b_accept.setText("Accept again" if signoff else "Accept double determination")
        self.b_accept.setEnabled(ready)
        if signoff:
            tip = f"Accepted by {signoff.get('by', '?')}, changed since. "
        else:
            tip = ""
        if not ready:
            tip += "Compare two determinations first"
        else:
            tip += ("Accept this double determination (saved with the project and in the audit trail). A pair "
                    "a workflow processed is accepted in Report² too; any other is listed there with its report")
            if n_open:
                tip += (f". {n_open} red row{'s are' if n_open != 1 else ' is'} still open: "
                        "they go into the report by the default rule (Only in A/B not reported, the others reported)")
        self.b_accept.setToolTip(tip)

    def accept(self) -> bool:
        """The analyst accepts the double determination: who and when go into the replicate group (undoable)
        and the audit trail; ``acceptRequested`` lets Report² accept its report too."""
        from datetime import datetime
        from gcws.core.audit import current_user
        if len(self.members) != 2 or not self.rows:
            return False
        self._sync_group(self.members)
        g = self.group()
        if g is None:
            return False
        n_open, decided = self.decision_counts()
        groups = copy.deepcopy(self.ws.replicate_groups)
        tg = next(x for x in groups if x["id"] == g["id"])
        tg["signoff"] = {"by": current_user(), "at": datetime.now().isoformat(timespec="seconds"), "open": n_open}
        live = self._live()[0]
        n_report = sum(1 for r in live if r.get("report"))
        self.set_groups(groups, f"double determination accepted: {self._names()}")
        self.ws.log("Double determination accepted", self._names(),
                    f"{n_report} of {len(live)} substances reported; {decided} red decided by the analyst"
                    + (f", {n_open} left open (default)" if n_open else ""))
        self.acceptRequested.emit(g["id"])
        return True

    @staticmethod
    def _passes(key: str, row: dict, v, light: str | None = None) -> bool:
        if row.get("deleted") or key == "deleted":
            return bool(row.get("deleted")) and key == "deleted"   # deleted rows only under their chip
        light = light or DV.light_of(row, v)
        if key == "all":
            return True
        if key == "check":
            return light in ("red", "yellow") or bool(row.get("edited"))
        return light == key

    def set_filter(self, key: str):
        """Show only the rows of one chip; the chip clicked again shows all rows."""
        self.filter = "all" if key == self.filter or key not in FILTERS else key
        self._update_chips()
        self._frozen = None
        self._fill_table()

    def _show_lights(self, rows):
        """Chips and status line of the feature pairing: how many rows are green, yellow, red, grey."""
        rows = [r for r in rows if not r.get("deleted")]
        n = {c: sum(1 for r in rows if r.get("light") == c) for c in ("green", "yellow", "red", "grey")}
        diffs = [r["reldiff"] for r in rows if r.get("reldiff") is not None and r.get("light") != "grey"]
        self._update_chips()
        notes = list(getattr(self.table_features, "notes", []) or [])
        auto = sum(1 for r in rows if r.get("gapfill"))
        n_open, n_decided = self.decision_counts()
        if n_open and n_decided:
            text = (f"Red: {n['red']} → {n_open} open · {n_decided} decided (F3 = next open; Accept keeps the "
                    f"default for the open ones); {n['yellow']} were made consistent automatically (yellow), "
                    f"{n['green']} are confirmed.")
            level = "bad"
        elif n_open:
            text = (f"{n['red']} of {len(rows)} substances need your decision (red; F3 = next; Accept keeps the "
                    f"default for the open ones); {n['yellow']} were made consistent automatically (yellow), "
                    f"{n['green']} are confirmed.")
            level = "bad"
        elif n_decided:
            text = (f"All decisions made ({n_decided} red decided by you): preview the report. "
                    f"{n['yellow']} made consistent automatically, {n['green']} confirmed.")
            level = "ok"
        elif n["yellow"]:
            text = f"Nothing to decide: {n['yellow']} made consistent automatically (yellow, a quick look), " \
                   f"{n['green']} confirmed."
            level = "warn"
        else:
            text, level = f"Double determination consistent: {n['green']} substances confirmed.", "ok"
        if auto:
            text += f" {auto} gap fill(s)."
        if diffs:
            text += f" Mean difference {sum(diffs) / len(diffs):.1f} %."
        drift = [x for x in notes if "drift" in x]
        if drift:
            text += " " + "; ".join(drift) + "."
        other = [x for x in notes if "drift" not in x and "time map" not in x]
        if other:
            text += " " + "; ".join(other)
        self.banner.setText(("✔  " if level == "ok" else "⚠  ") + text)
        self.banner.setProperty("level", level)
        props = self.boundary_proposals()
        self.b_bounds.setVisible(bool(props))
        self.b_bounds.setText(f"Harmonise boundaries ({len({p.rt for p in props})})")

    def headers(self) -> list[str]:
        """The title of every column (``COLUMN_KEYS``)."""
        labels = self.labels()
        unit = self.ws.quant_unit()
        out = ["", "Report", "RT [min]", "Substance", "CAS", f"Area {labels[0]}", f"Area {labels[1]}",
               f"{labels[0]} [{unit}]", f"{labels[1]} [{unit}]", f"Mean [{unit}]", "Diff. %", "Verdict", "Notes",
               "Comment", "Feature", "Similarity", f"Hit {labels[0]}", f"Hit {labels[1]}"]
        for other in QS.CONC_UNITS.values():
            out += [f"{labels[0]} [{other}]", f"{labels[1]} [{other}]", f"Mean [{other}]"]
        return out

    def _unit_cells(self, row: dict, units: list[str], settings) -> list[tuple]:
        """``(value, calculation)`` of the further-unit columns of ``row``: its A, B and mean
        (the analyst's edits included) converted from the mode's unit."""
        mode = self.ws.quant.get("mode", "nias_mgkg")
        if mode == "extraction":
            return self._extraction_cells(row, units)
        conv = {side: QS.from_mode_unit(mode, settings, row.get(field)) if units else None
                for side, field in UNIT_SIDES}
        return [(conv[side][unit], conv[side]["calc"].get(unit)) if unit in units else (None, None)
                for unit in QS.CONC_UNITS for side, _field in UNIT_SIDES]

    def _extraction_cells(self, row: dict, units: list[str]) -> list[tuple]:
        """Extraction: A and B converted with their own run's sample amount, the mean from both
        (without a dismissed outlier)."""
        from gcws.quant import extraction as EX
        from gcws.quant import units as U
        q, ids = self.ws.quant, list(self.members[:2])
        ids += [""] * (2 - len(ids))
        u1 = EX.unit(q)
        out = []
        for key in QS.CONC_UNITS:
            u = U.label_of(key)
            for (side, field), rid in zip(UNIT_SIDES, ids + [None]):
                if key not in units:
                    out.append((None, None))
                elif side == "mean":
                    v = EX.mean_in(q, ids, [row.get("c1"), row.get("c2")], row.get("mean"), u,
                                   int(row.get("dismissed") or 0))
                    out.append((v, f"Mean of {self.labels()[0]} and {self.labels()[1]} in {u}, each with its own "
                                   "sample amount" if v is not None else ""))
                else:
                    v = EX.convert_value(q, rid, row.get(field), u)
                    r = U.ratio(u1, u, EX.basis(q, rid))
                    out.append((v, U.missing(u, EX.basis(q, rid)) or
                                (f"{u} = {row.get(field):.6g} {u1} × {r:.6g}" if v is not None else "")))
        return out

    def _fill_table(self):
        from PySide6.QtGui import QFont
        from gcws.quant.nias_bridge import make_settings
        labels = self.labels()
        names = [self.ws.runs[m].name for m in self.members if m in self.ws.runs]
        headers = self.headers()
        feature_rows = bool(self.rows) and bool(self.rows[0].get("feature_id"))
        units = QS.unit_keys(self.ws.quant.get("mode", "nias_mgkg"), self.ws.quant)
        settings = make_settings(self.ws.quant.get("settings"))
        cur = self.table.currentItem()
        keep_cur = (cur.data(Qt.UserRole), cur.column()) if cur is not None else None
        keep_sel = {(i.data(Qt.UserRole), i.column()) for i in self.table.selectedItems()}
        scroll = (self.table.verticalScrollBar().value(), self.table.horizontalScrollBar().value())
        frozen = self._frozen if self._frozen is not None else {}
        self._filling = True
        self.table.setSortingEnabled(False)
        self.table.clear()
        self.table.setColumnCount(len(headers))
        self.table.setHorizontalHeaderLabels(headers)
        for i, n in enumerate(names[:2]):
            for c in (C_A1 + i, C_C1 + i) + tuple(range(C_UNIT0 + i, len(COLUMN_KEYS), len(UNIT_SIDES))):
                self.table.horizontalHeaderItem(c).setToolTip(n)
        self.table.horizontalHeaderItem(C_REPORT).setToolTip("Goes into the report (default: not for artefacts and "
                                                             "values below the reporting limit)")
        self.table.setRowCount(0)
        editable = set(FIELD_OF) - {C_REPORT}
        for k, (row, v) in enumerate(zip(self.rows, self.verdicts)):
            if not self._passes(self.filter, row, v):
                continue
            r = self.table.rowCount()
            self.table.insertRow(r)
            notes = "; ".join(x for x in (v.detail if row.get("light") else DV.english(row.get("review", "")),) if x)
            vals = [ICON.get(v.level, ""), None, row.get("rt"), row.get("name", ""), row.get("cas", ""),
                    row.get("a1"), row.get("a2"), row.get("c1"), row.get("c2"), row.get("mean"), row.get("reldiff"),
                    v.text, notes, row.get("comment", "")]
            if feature_rows:
                s1, s2 = row.get("source1") or {}, row.get("source2") or {}
                vals += [row.get("feature_id", ""), row.get("sim"),
                         s1.get("name", "") if s1 else "", s2.get("name", "") if s2 else ""]
            else:
                vals += [None] * len(FEATURE_COLUMNS)
            unit_cells = self._unit_cells(row, units, settings)
            vals += [v for v, _calc in unit_cells]
            edited = row.get("edited") or {}
            level, tip = v.level, v.detail
            if v.level == "bad" and row.get("decided"):
                level, vals[C_ICON] = "info", ICON["decided"]
                tip = "Decided by the analyst: " + (v.detail or v.text)
            n1 = ((row.get("source1") or {}).get("name") or "").strip()
            n2 = ((row.get("source2") or {}).get("name") or "").strip()
            hits_differ = feature_rows and n1 and n2 and n1.casefold() != n2.casefold()
            deleted = bool(row.get("deleted"))
            if deleted:
                tip = "Deleted by the analyst: not in the counts or the report. Right-click: Restore row."
            gone = row.get("dismissed") or 0
            # the dismissed determination's cells: area, concentration and its further units
            dismissed = ({C_A1 + gone - 1, C_C1 + gone - 1}
                         | set(range(C_UNIT0 + gone - 1, len(COLUMN_KEYS), len(UNIT_SIDES)))) if gone else set()
            for c, val in enumerate(vals):
                it = _SeverityItem() if c in (C_ICON, C_REPORT) else QTableWidgetItem()
                if c == C_REPORT and deleted:
                    it.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)      # no box: it is not reported
                elif c == C_REPORT:
                    it.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled | Qt.ItemIsSelectable)
                    it.setCheckState(Qt.Checked if row.get("report") else Qt.Unchecked)
                    it.setData(SORT_ROLE, (10000.0 if row.get("report") else 0.0) + float(row.get("rt") or 0.0))
                elif isinstance(val, float):
                    if c in (C_A1, C_A2):
                        it.setData(Qt.DisplayRole, int(round(val)))          # areas: whole counts
                    else:
                        decimals = UNIT_DECIMALS[UNIT_OF[c]] if c in UNIT_OF else {C_RT: 3, C_DIFF: 1, C_SIM: 2}.get(c, 4)
                        it.setData(Qt.DisplayRole, round(val, decimals))
                    it.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                else:
                    it.setText("" if val is None else str(val))
                if c not in editable and c != C_REPORT:
                    it.setFlags(it.flags() & ~Qt.ItemIsEditable)
                it.setData(Qt.UserRole, k)
                calc = unit_cells[c - C_UNIT0][1] if c in UNIT_OF and not deleted else None
                it.setToolTip(calc or tip)                     # a further unit: how it is calculated
                if c == C_ICON:
                    it.setData(OPEN_ROLE, DV.is_open(row, v))
                    sev = SEVERITY["decided"] if row.get("decided") and v.level == "bad" else SEVERITY.get(v.level, 4)
                    it.setData(SORT_ROLE, frozen.setdefault(k, sev * 10000.0 + float(row.get("rt") or 0.0)))
                if c == C_DIFF:
                    it.setData(LEVEL_ROLE, level)
                if c == C_NAME and hits_differ:
                    it.setData(Qt.DecorationRole, self.style().standardIcon(QStyle.SP_MessageBoxWarning))
                    it.setToolTip(f"The hits differ - {labels[0]}: {n1}  /  {labels[1]}: {n2}. Right-click for "
                                  "the candidate names.")
                if c in (C_ICON, C_VERDICT):
                    it.setBackground(theme.status_brush(level))
                    it.setForeground(QBrush(theme.status_color(level)))
                field = FIELD_OF.get(c)
                if field in edited or (c == C_COMMENT and row.get("comment")):
                    font = QFont()
                    font.setItalic(True)
                    it.setFont(font)
                    it.setData(EDITED_ROLE, True)
                    was = edited.get(field)
                    it.setToolTip(f"Changed by the analyst" + (f" (was {was:.4g})" if isinstance(was, float) else
                                                               (f" (was {'on' if was else 'off'})" if field == "report"
                                                                else "")))
                if deleted or c in dismissed:
                    font = it.font()
                    font.setStrikeOut(True)
                    it.setFont(font)
                    it.setForeground(QBrush(QColor(theme.FAINT)))
                if c in dismissed and not deleted:
                    it.setToolTip(f"{labels[gone - 1]} dismissed as an outlier: the result is {labels[2 - gone]}'s "
                                  f"concentration. Right-click: Use {labels[gone - 1]} again.")
                elif c == C_MEAN and gone and not deleted:
                    it.setToolTip(f"{labels[2 - gone]} only ({labels[gone - 1]} dismissed as an outlier)")
                self.table.setItem(r, c, it)
            if not row.get("report") and not deleted:
                for c in (C_RT, C_NAME, C_MEAN):
                    self.table.item(r, c).setForeground(QBrush(QColor(theme.FAINT)))
        self.table.resizeColumnsToContents()
        hh = self.table.horizontalHeader()
        hh.setSectionResizeMode(QHeaderView.Interactive)
        self.table.setColumnWidth(C_NAME, min(260, max(140, self.table.columnWidth(C_NAME))))
        hh.setStretchLastSection(True)
        self._apply_layout()
        self.table.setSortingEnabled(True)
        self.table.sortItems(*self._sort)
        self._frozen = frozen
        self._restore_selection(keep_cur, keep_sel)
        self.table.verticalScrollBar().setValue(scroll[0])          # an edit does not move the view
        self.table.horizontalScrollBar().setValue(scroll[1])
        self._filling = False

    # -- columns --------------------------------------------------------------------------------

    def _sort_changed(self, column: int, order) -> None:
        if not self._filling:
            self._sort = (column, order)
            if column == C_ICON and self._frozen:                 # sorted by severity again: as it is now
                self._frozen = None
                QTimer.singleShot(0, self._fill_table)

    # The column choice is remembered by column name (``COLUMN_KEYS``), never by index, so a column
    # added in a later version cannot shift the others: {"order": [names], "hidden": [names]}.

    @staticmethod
    def _layout() -> dict:
        """The remembered column choice; the index-based choice of an earlier version is taken over."""
        import json
        from PySide6.QtCore import QSettings
        s = QSettings()
        state = None
        try:
            state = json.loads(s.value(COLUMNS_SETTING) or "null")
        except (TypeError, ValueError):
            pass
        if not isinstance(state, dict):
            old = s.value("replicates/hidden_columns", None)
            if old is None:
                hidden = [COLUMN_KEYS[c] for c in HIDDEN_BY_DEFAULT]
            else:
                old = [old] if isinstance(old, str) and old else [] if isinstance(old, str) else old
                hidden = [COLUMN_KEYS[int(c)] for c in old if str(c).isdigit() and int(c) < C_UNIT0]
            state = {"order": [], "hidden": hidden}
        order = [k for k in state.get("order") or [] if k in COLUMN_KEYS]
        hidden = [k for k in state.get("hidden") or [] if k in COLUMN_KEYS]
        # columns that choice has never seen (new in this version) follow their default
        hidden += [COLUMN_KEYS[c] for c in HIDDEN_BY_DEFAULT if COLUMN_KEYS[c] not in order + hidden
                   and (order or c >= C_UNIT0)]
        return {"order": order + [k for k in COLUMN_KEYS if k not in order], "hidden": hidden}

    @staticmethod
    def _save_layout(state: dict) -> None:
        import json
        from PySide6.QtCore import QSettings
        QSettings().setValue(COLUMNS_SETTING, json.dumps(state))
        QSettings().remove("replicates/hidden_columns")

    @classmethod
    def hidden_columns(cls) -> set[int]:
        """The columns the analyst hid (indices)."""
        return {COLUMN_KEYS.index(k) for k in cls._layout()["hidden"]} - FIXED_COLUMNS

    def _forced_hidden(self) -> set[int]:
        """Columns without content in this comparison: the feature columns of the classic pairing,
        the units the quantification mode does not compute."""
        feature_rows = bool(self.rows) and bool(self.rows[0].get("feature_id"))
        units = QS.unit_keys(self.ws.quant.get("mode", "nias_mgkg"), self.ws.quant)
        return (set() if feature_rows else set(FEATURE_COLUMNS)) | {c for c, u in UNIT_OF.items() if u not in units}

    def _apply_layout(self) -> None:
        state = self._layout()
        hh = self.table.horizontalHeader()
        for pos, key in enumerate(state["order"]):
            logical = COLUMN_KEYS.index(key)
            if hh.visualIndex(logical) != pos:
                hh.moveSection(hh.visualIndex(logical), pos)
        hidden = self.hidden_columns() | self._forced_hidden()
        for c in range(self.table.columnCount()):
            self.table.setColumnHidden(c, c in hidden)

    def shown_keys(self) -> list[str]:
        """The names of the shown columns, left to right."""
        return [COLUMN_KEYS[c] for c in self.table.shown_columns()]

    def offered_keys(self) -> list[str]:
        """The columns the analyst can show or hide here (Columns... and the header's menu)."""
        forced = self._forced_hidden()
        return [k for c, k in enumerate(COLUMN_KEYS) if c not in forced and c not in (C_ICON, C_REPORT)]

    def set_shown(self, keys: list[str]) -> None:
        """Show ``keys`` in this order (the status and Report columns first, the substance always);
        remembered for the next start. Columns not offered now keep their choice."""
        keys = [k for k in dict.fromkeys(keys) if k in COLUMN_KEYS and k not in ("icon", "report")]
        if "name" not in keys:
            keys.insert(0, "name")
        shown = ["icon", "report"] + keys
        offered = set(self.offered_keys())
        old = self._layout()
        hidden = [k for k in COLUMN_KEYS if (k in offered and k not in shown)
                  or (k not in offered and k in old["hidden"])]
        order = shown + [k for k in old["order"] if k not in shown]
        self._save_layout({"order": order, "hidden": hidden})
        if self.table.columnCount():
            self._apply_layout()

    def set_column_hidden(self, column: int, hidden: bool) -> None:
        """Show or hide a column of the list; remembered for the next start."""
        if column in FIXED_COLUMNS:
            return
        state = self._layout()
        key = COLUMN_KEYS[column]
        state["hidden"] = [k for k in state["hidden"] if k != key] + ([key] if hidden else [])
        self._save_layout(state)
        self.table.setColumnHidden(column, hidden or column in self._forced_hidden())

    def column_dialog(self):
        """The column chooser of this list (not shown yet)."""
        from gcws.ui.dialogs.columns import ColumnChooserDialog
        headers = self.headers()
        offered = self.offered_keys()
        cols = [(k, headers[COLUMN_KEYS.index(k)], self._column_tip(k)) for k in offered]
        shown = [k for k in self.shown_keys() if k in offered] if self.table.columnCount() else \
            [k for k in self._layout()["order"] if k in offered and k not in self._layout()["hidden"]]
        defaults = [k for k in offered if COLUMN_KEYS.index(k) not in HIDDEN_BY_DEFAULT]
        return ColumnChooserDialog(cols, shown, defaults, self)

    @staticmethod
    def _column_tip(key: str) -> str:
        unit = key.rsplit("_", 1)[0]
        if unit in QS.CONC_UNITS:
            return f"Concentration in {QS.CONC_UNITS[unit]}, converted from the list's value (your changes " \
                   "included); hover a value for its calculation"
        return ""

    def choose_columns(self) -> None:
        dlg = self.column_dialog()
        if dlg.exec() == dlg.Accepted:
            self.set_shown(dlg.shown_keys())

    def _column_menu(self, pos) -> None:
        menu = QMenu(self)
        menu.addAction("Choose columns...", self.choose_columns)
        menu.addSeparator()
        hidden = self.hidden_columns()
        headers = self.headers()
        for key in self.offered_keys():
            c = COLUMN_KEYS.index(key)
            if c in FIXED_COLUMNS:
                continue
            a = menu.addAction(headers[c] or "")
            a.setCheckable(True)
            a.setChecked(c not in hidden)
            a.toggled.connect(lambda on, c=c: self.set_column_hidden(c, not on))
        menu.exec(self.table.horizontalHeader().mapToGlobal(pos))

    def _restore_selection(self, cur, sel):
        """After a rebuild the same substances and columns are marked again (no navigation)."""
        from PySide6.QtCore import QItemSelectionModel
        rowof = {self.table.item(r, 0).data(Qt.UserRole): r for r in range(self.table.rowCount())
                 if self.table.item(r, 0) is not None}
        self.table.blockSignals(True)
        try:
            for k, c in sel:
                if k in rowof and self.table.item(rowof[k], c) is not None:
                    self.table.item(rowof[k], c).setSelected(True)
            if cur is not None and cur[0] in rowof:
                idx = self.table.model().index(rowof[cur[0]], cur[1])
                self.table.selectionModel().setCurrentIndex(idx, QItemSelectionModel.NoUpdate)
        finally:
            self.table.blockSignals(False)

    # -- feature pairing ------------------------------------------------------------------------

    def features_mode(self) -> bool:
        from gcws.features import service as SV
        from gcws.features.model import PAIRING_FEATURES
        return self.ws.quant.get("mode") != "hs_screening" and SV.pairing(self.ws) == PAIRING_FEATURES

    def feature_of(self, row):
        t = self.table_features
        return t.by_id(row.get("feature_id")) if (t is not None and row and row.get("feature_id")) else None

    def edit_settings(self):
        from gcws.features import service as SV
        from gcws.ui.dialogs.feature_settings import FeatureSettingsDialog
        dlg = FeatureSettingsDialog(SV.settings(self.ws), self)
        if dlg.exec() != FeatureSettingsDialog.Accepted:
            return
        self.set_settings(dlg.settings())

    def set_settings(self, settings) -> None:
        """Store the feature settings (one undo step) and compare again."""
        q = copy.deepcopy(self.ws.quant)
        q["features"] = settings.to_dict()
        if q != self.ws.quant:
            self.ws.push_quant("double determination settings", q, "double determination (features)")
        self.compare(sync=False)

    def next_red(self):
        """F3: the next red row still open after the current one (from the top at the end)."""
        rows = [r for r in range(self.table.rowCount())
                if self.table.item(r, C_ICON) is not None and self.table.item(r, C_ICON).data(OPEN_ROLE)]
        if not rows:
            return
        cur = self.table.currentRow()
        nxt = next((r for r in rows if r > cur), rows[0])
        self.table.setCurrentCell(nxt, C_NAME)

    def boundary_proposals(self, row=None) -> list:
        t = self.table_features
        if t is None:
            return []
        feats = [self.feature_of(row)] if row is not None else t.features
        return [p for f in feats if f is not None for p in f.proposals if p.kind == "boundary"]

    def apply_boundaries(self, row=None) -> int:
        from gcws.features import service as SV
        props = self.boundary_proposals(row)
        if not props or self.table_features is None:
            return 0
        n = SV.apply(self.ws, self.table_features, props, stack=self._stack(),
                     label=f"double determination: {len({p.rt for p in props})} boundaries harmonised")
        self.compare(sync=False)
        return n

    def remove_gap_fill(self, row) -> bool:
        """Take the gap fill of this substance out again (one undo step)."""
        from gcws.features.model import GAPFILL, GAPFILL_OPTION
        from gcws.ui.undo import ManualEventsCommand, MultiCommand
        f = self.feature_of(row)
        if f is None:
            return False
        cmds = []
        for m in f.members:
            if m.origin != GAPFILL or m.peak is None:
                continue
            key = self.table_features.key
            st = self.ws.runs[m.run_id]
            keep = [e for e in st.events(key)
                    if not (e.option == GAPFILL_OPTION and e.t1 is not None
                            and min(e.t0, e.t1) - 1e-6 <= m.peak.rt <= max(e.t0, e.t1) + 1e-6)]
            if len(keep) != len(st.events(key)):
                cmds.append(ManualEventsCommand(self.ws, m.run_id, key, keep, f"{f.id}: gap fill removed"))
        if not cmds:
            return False
        self._stack().push(MultiCommand(f"double determination: gap fill of {f.id} removed", cmds))
        self.ws.log("Double determination (features)", self._names(), f"{f.id}: gap fill removed by the analyst")
        self.compare(sync=False)
        return True

    def _visual_row(self, row) -> int | None:
        """The table row showing ``row`` (None: filtered out)."""
        k = next((i for i, r in enumerate(self.rows) if r is row), None)
        return next((r for r in range(self.table.rowCount())
                     if self.table.item(r, 0) is not None and self.table.item(r, 0).data(Qt.UserRole) == k), None)

    def _marked_or(self, row) -> list[int]:
        """The marked table rows when ``row`` is one of them, else just ``row``'s."""
        r = self._visual_row(row)
        marked = self.table.selected_rows()
        return marked if r in marked else ([r] if r is not None else [])

    def edit_comment(self, row) -> None:
        from PySide6.QtWidgets import QInputDialog
        text, ok = QInputDialog.getText(self, "Comment", row.get("name") or f"RT {row['rt']:.3f}",
                                        text=row.get("comment", ""))
        if ok:
            self.set_edit(row, "comment", text.strip() or None)

    def copy_row(self, row) -> None:
        from PySide6.QtGui import QGuiApplication
        r = self._visual_row(row)
        if r is None:
            return
        cells = [self.table.cell_value(r, c) for c in self.table.shown_columns() if c != C_ICON]
        QGuiApplication.clipboard().setText("\t".join("yes" if v is True else "no" if v is False else str(v)
                                                       for v in cells))

    def row_actions(self, row) -> list:
        """``[(text, callable)]`` of the right-click menu of ``row``; ``(None, None)`` is a separator."""
        labels = self.labels()
        if row.get("deleted"):
            out = [("Restore row", lambda: self._delete_rows(self._marked_or(row)))]
        else:
            out = [("Report", lambda: self._mark_rows(self._marked_or(row), True)),
                   ("Not reported", lambda: self._mark_rows(self._marked_or(row), False)),
                   ("Delete row", lambda: self._delete_rows(self._marked_or(row)))]
            if row.get("edit_key") in self.edits():
                out.append(("Reset row", lambda: self.reset_row(row)))
            out.append(("Comment…", lambda: self.edit_comment(row)))
            if len(self.members) == 2 and row.get("c1") is not None and row.get("c2") is not None:
                n = row.get("dismissed")
                if n:
                    out.append((f"Use {labels[n - 1]} again (mean of both)",
                                lambda: self.set_edit(row, "dismiss", None)))
                else:
                    for i in (0, 1):
                        out.append((f"Dismiss {labels[i]} (outlier): result = {labels[1 - i]}",
                                    lambda n=i + 1: self.set_edit(row, "dismiss", n)))
        for key, i in (("source1", 0), ("source2", 1)):
            if row.get(key) is not None and i < len(self.members):
                out.append((f"Show in {labels[i]}", lambda b=bool(i): self._navigate(row, prefer_b=b)))
        out.append(("Copy row", lambda: self.copy_row(row)))
        f = self.feature_of(row)
        if f is None or row.get("deleted"):
            return out
        out.append((None, None))
        ident = f.identity
        if ident is not None and ident.case in ("C", "D"):
            for c in ident.candidates[:3]:
                out.append((f"Name: {c['name']}", lambda c=c: self.set_identities([(row, c["name"], c.get("cas") or "")])))
        if any(m.origin == "gapfill" for m in f.members):
            out.append(("Remove the gap fill", lambda: self.remove_gap_fill(row)))
        if self.boundary_proposals(row):
            out.append(("Harmonise the boundaries", lambda: self.apply_boundaries(row)))
        return out

    def _context_menu(self, pos):
        it = self.table.itemAt(pos)
        if it is None:
            return
        k = it.data(Qt.UserRole)
        if k is None or not (0 <= k < len(self.rows)):
            return
        actions = self.row_actions(self.rows[k])
        menu = QMenu(self)
        for text, fn in actions:
            if text is None:
                menu.addSeparator()
            else:
                menu.addAction(text, fn)
        menu.exec(self.table.viewport().mapToGlobal(pos))

    def _start_consensus_search(self):
        """Search the consensus spectra still unknown (worker thread), then compare again."""
        from gcws.features import service as SV
        from gcws.features.consensus import search_consensus
        t = self.table_features
        if t is None or self._search is not None or not SV.settings(self.ws).consensus_search:
            return
        needed = SV.consensus_needed(self.ws, t, SV.settings(self.ws))
        if not needed:
            return
        self._search_needed = needed
        self._search_group = SV.group_of(self.ws, t.members)
        self.banner.setText(self.banner.text() + f"  Searching {len(needed)} consensus spectra…")
        window = self.window()
        if hasattr(window, "begin_activity"):
            window.begin_activity(f"consensus:{id(self)}", f"Searching {len(needed)} spectra…")
        self._search = workers.submit(
            search_consensus, None, SV.search_method(self.ws, t.members), [f for _k, f in needed],
            on_done=self._consensus_found,
            on_error=lambda error: self._consensus_found(f"consensus search failed: {error.splitlines()[0]}"))

    def _consensus_found(self, note: str):
        from gcws.features import service as SV
        needed, self._search_needed = getattr(self, "_search_needed", []), []
        self._search = None
        window = self.window()
        if hasattr(window, "end_activity"):
            window.end_activity(f"consensus:{id(self)}")
        SV.store_consensus(self.ws, needed, note, group=getattr(self, "_search_group", None))
        if self.isVisible() and self.members:
            self.compare(sync=True)

    # -- spectra ---------------------------------------------------------------------------------

    def _draw_spectra(self):
        self.spec.clear()
        row = self._current_row()
        f = self.feature_of(row) if row is not None else None
        if f is None:
            return
        self.spec.addItem(pg.InfiniteLine(pos=0, angle=0, pen=pg.mkPen(theme.BORDER_STRONG)), ignoreBounds=True)
        for sign, m in zip((1, -1), f.members[:2]):
            if m.peak is None:
                continue
            spec = m.peak.spectrum if m.peak.spectrum is not None else m.peak.full_spectrum
            if spec is None or len(spec[0]) == 0:
                continue
            mz, ab = np.asarray(spec[0], float), np.asarray(spec[1], float)
            ab = ab / ab.max() * 100.0
            xs = np.repeat(mz, 2)
            ys = np.column_stack([np.zeros_like(ab), sign * ab]).ravel()
            st = self.ws.runs.get(m.run_id)
            color = st.color if st is not None else theme.ACCENT
            self.spec.plot(xs, ys, connect="pairs", pen=pg.mkPen(color, width=1.5))
            for i in np.argsort(-ab)[:4]:
                t = pg.TextItem(f"{int(mz[i])}", anchor=(0.5, 1.0 if sign > 0 else 0.0), color=theme.FAINT)
                t.setPos(mz[i], sign * ab[i])
                self.spec.addItem(t)
        sim = "–" if f.sim is None else f"{f.sim:.2f}"
        name = f.identity.name if f.identity else ""
        self.spec.setTitle(f"{f.id}  similarity {sim}  {name[:40]}", size="9pt")
        self.spec.setYRange(-110, 110, padding=0)

    # -- mirror plot -----------------------------------------------------------------------

    def _t0(self, st, sig) -> float:
        """Where the mirror plot starts: the integration start, else the solvent end."""
        from gcws.integration.autoparams import _integration_start
        t = _integration_start(sig.rt, self.ws.method_for(st, self.quant_signal()), self.ws.solvent_cut(st, self.quant_signal()))
        if t is None and self.quant_signal() == "TIC":
            return float(sig.rt[0])
        if t is None:
            from gcws.quant.nias_bridge import make_settings
            t = float(getattr(make_settings(self.ws.quant.get("settings")), "solvent_end", 0.0) or 0.0)
        return min(max(t, float(sig.rt[0])), float(sig.rt[-1]))

    def _trace(self, rid):
        """``(rt, signal above baseline, colour)`` of the FID from the integration start on.

        Both determinations keep their own signal units (no normalisation), so a peak that is
        twice as large in A also looks twice as large. The solvent front is left out and the
        slow baseline (solvent tail, column bleed) is removed, so peaks stand on zero."""
        st = self.ws.runs.get(rid)
        sig = st.run.signal(self.quant_signal()) if st is not None else None
        if sig is None or len(sig.rt) < 2:
            return None
        keep = sig.rt >= self._t0(st, sig)
        if keep.sum() < 2:
            return None
        from gcws.signal.envelope import envelope
        rt, y = sig.rt[keep], np.asarray(sig.y[keep], float)
        return rt, np.maximum(y - envelope(rt, y, BASELINE_WINDOW), 0.0), st.color

    def _draw_mirror(self):
        self.mirror.clear()
        self._traces, self._marks = [], []
        det = self.quant_signal()
        self.mirror.setLabel("bottom", f"RT ({det})", units="min")
        self.mirror.setLabel("left", f"A  ↑   {det}   ↓  B")
        self.mirror.addItem(self.cursor, ignoreBounds=True)
        traces = [self._trace(m) for m in self.members[:2]]
        for sign, tr in zip((1, -1), traces):
            if tr is None:
                continue
            rt, y, color = tr
            self.mirror.plot(rt, sign * y, pen=pg.mkPen(color, width=1.2), shadowPen=theme.glow_pen(color))
            self._traces.append((rt, sign * y))
        self.mirror.addItem(pg.InfiniteLine(pos=0, angle=0, pen=pg.mkPen(theme.BORDER_STRONG)), ignoreBounds=True)
        spots, self._marks = [], []
        self._signed = {}
        for i, tr in enumerate(traces):
            if tr is not None:
                self._signed[i] = (tr[0], (1 if i == 0 else -1) * tr[1])
        istd = self._istd_times()
        for k, (row, v) in enumerate(zip(self.rows, self.verdicts)):
            if row.get("deleted"):
                continue
            color = theme.status_color(v.level)
            for sign, key, tr in ((1, "source1", traces[0] if traces else None),
                                  (-1, "source2", traces[1] if len(traces) > 1 else None)):
                src = row.get(key)
                if src is None or tr is None or not tr[0][0] <= src["rt"] <= tr[0][-1]:
                    continue
                y = float(np.interp(src["rt"], tr[0], tr[1]))       # on the apex of the peak
                is_istd = any(abs(src["rt"] - t) < 0.02 for t in istd)
                self._marks.append((src["rt"], abs(y), is_istd))
                spots.append({"pos": (src["rt"], sign * y), "brush": pg.mkBrush(color),
                              "pen": pg.mkPen(theme.SURFACE, width=0.8), "size": 8,
                              "data": (k, row.get("name") or "")})
        if spots:
            dots = pg.ScatterPlotItem(spots=spots, hoverable=True,
                                      tip=lambda x, y, data: f"{data[1]}  RT {x:.3f} (click: select)")
            dots.sigClicked.connect(self._dot_clicked)
            self.mirror.addItem(dots, ignoreBounds=True)
        self.full_view()
        self._draw_bounds(self._current_row())

    def _dot_clicked(self, _item, points, *_):
        """A dot of the mirror plot selects its substance in the list (all rows shown if it was filtered out)."""
        if not points:
            return
        data = points[0].data()
        k = data[0] if isinstance(data, tuple) else None
        if k is None or not (0 <= k < len(self.rows)):
            return
        r = self._visual_row(self.rows[k])
        if r is None:
            self.set_filter("all")
            r = self._visual_row(self.rows[k])
        if r is not None:
            self.table.setCurrentCell(r, C_NAME)
            self.table.scrollToItem(self.table.item(r, C_NAME))

    def _source_peak(self, i: int, src):
        """The integrated peak of determination ``i`` at ``src`` (apex within 0.05 min), or None."""
        if src is None or i >= len(self.members) or self.members[i] not in self.ws.runs:
            return None
        res = self.ws.result(self.members[i], self.quant_signal())
        if res is None or not res.peaks:
            return None
        peak = min(res.peaks, key=lambda p: abs(p.apex_rt - src["rt"]))
        return peak if abs(peak.apex_rt - src["rt"]) <= 0.05 else None

    def _draw_bounds(self, row) -> None:
        """Shade the integrated peak of the selected substance in each determination, A above and B
        below the zero line, so different boundaries are seen before they are harmonised."""
        for item in self._bounds:
            self.mirror.removeItem(item)
        self._bounds = []
        if row is None:
            return
        for i, key in ((0, "source1"), (1, "source2")):
            peak = self._source_peak(i, row.get(key))
            if peak is None or i not in self._signed:
                continue
            rt, y = self._signed[i]
            m = (rt >= peak.start) & (rt <= peak.end)
            if m.sum() < 2:
                continue
            st = self.ws.runs.get(self.members[i])
            color = pg.mkColor(st.color if st is not None else theme.ACCENT)
            color.setAlpha(80)
            item = self.mirror.plot(rt[m], y[m], pen=None, fillLevel=0.0, brush=pg.mkBrush(color))
            item.setToolTip(f"{self.labels()[i]}: integrated {peak.start:.3f} - {peak.end:.3f} min")
            self._bounds.append(item)

    def full_view(self):
        """The whole integrated part of both determinations."""
        if not self._traces:
            return
        x0 = min(float(rt[0]) for rt, _ in self._traces)
        x1 = max(float(rt[-1]) for rt, _ in self._traces)
        self.mirror.setXRange(x0, x1, padding=0.01)
        self._fit_y()

    def _istd_times(self) -> list[float]:
        """FID apexes of the internal standards (from the quantification) in A and B."""
        out = []
        for rid in self.members[:2]:
            res = self.ws.result(rid, self.quant_signal())
            if res is None:
                continue
            for i, q in self.ws.quant_rows(rid, self.quant_signal()).items():
                if q.get("istd") and 0 <= i < len(res.peaks):
                    out.append(res.peaks[i].apex_rt)
        return out

    def _fit_y(self):
        """Symmetric intensity range from the largest peak inside the visible time window.

        Over a wide window (more than ``WIDE_VIEW`` min) a few very large peaks (internal
        standards, main components) would flatten everything else: the range then follows the
        ``WIDE_PERCENTILE`` of the substance peaks without the internal standards, and the
        largest peaks are clipped. Zooming in (a picked row) always shows the full peak."""
        if not self._traces:
            return
        (x0, x1), _ = self.mirror.getViewBox().viewRange()
        top = 0.0
        if x1 - x0 > WIDE_VIEW:
            heights = [h for t, h, is_istd in self._marks if x0 <= t <= x1 and not is_istd]
            top = float(np.percentile(heights, WIDE_PERCENTILE)) if len(heights) >= 5 else 0.0
        if not top:
            for rt, y in self._traces:
                m = (rt >= x0) & (rt <= x1)
                if m.any():
                    top = max(top, float(np.max(np.abs(y[m]))))
        top = top or 1.0
        self.mirror.setYRange(-1.12 * top, 1.12 * top, padding=0)

    # -- navigation ----------------------------------------------------------------------------

    def _current_row(self):
        it = self.table.currentItem()
        if it is None:
            items = self.table.selectedItems()
            it = items[0] if items else None
        if it is None:
            return None
        k = it.data(Qt.UserRole)
        return self.rows[k] if k is not None and 0 <= k < len(self.rows) else None

    def _row_selected(self):
        row = self._current_row()
        if row is None:
            return
        rt = row.get("rt")
        if rt is not None:
            self.cursor.setPos(rt)
            self.mirror.setXRange(rt - 0.4, rt + 0.4, padding=0)
            self._draw_bounds(row)
            self._navigate(row, prefer_b=False)
        self._draw_spectra()

    def _open_row(self, r, c):
        if c in FIELD_OF:                       # an editable cell: the double-click edits it
            return
        k = self.table.item(r, 0).data(Qt.UserRole)
        if k is not None:
            self._navigate(self.rows[k], prefer_b=c in (C_A2, C_C2))

    def _navigate(self, row, prefer_b=False):
        order = [("source2", 1), ("source1", 0)] if prefer_b else [("source1", 0), ("source2", 1)]
        for key, i in order:
            src = row.get(key)
            if src is None or i >= len(self.members):
                continue
            rid = self.members[i]
            self.ws.set_active(rid)
            if self.ws.signal_key != self.quant_signal():
                self.ws.set_signal_key(self.quant_signal())
            res = self.ws.result(rid, self.quant_signal())
            if res is not None and res.peaks:
                j = min(range(len(res.peaks)), key=lambda q: abs(res.peaks[q].apex_rt - src["rt"]))
                if abs(res.peaks[j].apex_rt - src["rt"]) < 0.05:
                    self.ws.select_peak(j)
                    self.ws.peakFocusRequested.emit(j)
            return

    def set_plots_visible(self, on: bool) -> None:
        """The chromatograms and spectra below the list (remembered)."""
        from PySide6.QtCore import QSettings
        self.plots.setVisible(on)
        QSettings().setValue("replicates/plots", bool(on))
        if self.b_plots.isChecked() != on:
            self.b_plots.setChecked(on)

    # -- output ------------------------------------------------------------------------------------

    def show_limit(self) -> None:
        """The limit of the quantification settings; marked when it is not the default."""
        value, default = DV.limits(self.ws)[0], DV.default_limit(self.ws)
        self.limit.blockSignals(True)
        self.limit.setValue(value)
        self.limit.blockSignals(False)
        changed = abs(value - default) > 1e-9
        self.limit.setProperty("changed", changed)
        theme.set_chip(self.limit_chip, f"report parameter · default {default:g} %" if changed else "report parameter",
                       "warn" if changed else "info")
        self.limit_chip.setToolTip("Duplicate difference limit of the report settings: every report of this "
                                   "project uses it" + (f". The default is {default:g} %." if changed else "."))

    def _limit_changed(self):
        from gcws.quant.nias_bridge import make_settings, settings_dict
        new = float(self.limit.value())
        if abs(new - DV.limits(self.ws)[0]) < 1e-9:
            return
        q = copy.deepcopy(self.ws.quant)
        if q.get("mode") == "hs_screening":
            q.setdefault("hs", {})["duplicate_max_reldiff"] = new
            self.ws.push_quant(f"HS duplicate difference limit = {new:g} %", q)
            self.ws.message.emit(f"Difference limit is a report parameter: every report now uses {new:g} %.")
            return
        s = make_settings(q.get("settings"))
        s.duplicate_max_reldiff = new
        q["settings"] = settings_dict(s)
        self.ws.push_quant(f"duplicate difference limit = {new:g} %", q)
        self.ws.message.emit(f"Difference limit is a report parameter: every report now uses {new:g} %.")

    def _report(self, kind, preview=False):
        self._sync_group(self.members)
        g = self.group()
        if g is not None:
            (self.previewRequested if preview else self.reportRequested).emit(kind, g["id"])

    def export_rows(self) -> list[list]:
        """Header and rows of the export: the fixed columns, then the further units shown in the list."""
        from gcws.quant.nias_bridge import make_settings
        labels = self.labels()
        unit = self.ws.quant_unit()
        extra = [c for c in self.table.shown_columns() if c in UNIT_OF]
        headers = self.headers()
        units = QS.unit_keys(self.ws.quant.get("mode", "nias_mgkg"), self.ws.quant)
        settings = make_settings(self.ws.quant.get("settings"))
        out = [["RT [min]", "Substance", "CAS", f"{labels[0]} [{unit}]", f"{labels[1]} [{unit}]",
                f"Mean [{unit}]", "Diff. %", "Verdict", "Explanation", "Report", "Changed by analyst", "Comment"]
               + [headers[c] for c in extra]]
        for row, v in zip(*self._live()):
            cells = self._unit_cells(row, units, settings) if extra else []
            out.append([row.get("rt"), row.get("name"), row.get("cas"), row.get("c1"), row.get("c2"),
                        row.get("mean"), row.get("reldiff"), v.text, v.detail, "yes" if row.get("report") else "no",
                        ", ".join(sorted(row.get("edited") or {})), row.get("comment", "")]
                       + [cells[c - C_UNIT0][0] for c in extra])
        return out

    def export(self):
        if not self.rows:
            return
        names = "_".join(self.ws.runs[m].name for m in self.members if m in self.ws.runs)[:80]
        path, _ = QFileDialog.getSaveFileName(self, "Export double determination", f"{names}_duplicate.xlsx",
                                              "Excel (*.xlsx)")
        if not path:
            return
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill
        from gcws.core.text import excel_safe
        wb = Workbook()
        sh = wb.active
        sh.title = "Double determination"
        labels = self.labels()
        sh.append([f"{labels[i]}: {self.ws.runs[m].name}" for i, m in enumerate(self.members) if m in self.ws.runs])
        sh.append(["Difference limit %", DV.limits(self.ws)[0]])
        sh.append([])
        head, *rows = self.export_rows()
        sh.append(head)
        fills = {lvl: PatternFill("solid", fgColor=theme.LEVELS[lvl][1].lstrip("#")) for lvl in theme.LEVELS}
        for values, v in zip(rows, self._live()[1]):
            sh.append([excel_safe(x) for x in values])
            sh.cell(sh.max_row, 8).fill = fills.get(v.level, fills["neutral"])
        wb.save(path)
        self.ws.message.emit(f"Double determination exported: {path}")

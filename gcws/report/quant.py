"""Quantification report of the extraction method (quant method): RT, Name, CAS, Qual, Conc. 1 and
Conc. 2 for a single or a double (N-fold) determination, with the calculation sheets behind it."""
from __future__ import annotations

from openpyxl import Workbook

from gcws.quant import units as U
from gcws.report import simple as S

TITLE = "GC – Quantification"


def method_text(m: dict, standards: list, detector: str, amounts: list, labels: list) -> str:
    """The method line: sample type and amount (each determination's when they differ), extract, standards."""
    what = "solid" if m["sample_type"] == "solid" else "foil"
    unit = "g" if m["sample_type"] == "solid" else "dm²"
    if len(set(amounts)) <= 1:
        amount = f"{amounts[0]:g} {unit}" if amounts and amounts[0] else "no amount"
    else:
        amount = ", ".join(f"{lab} {a:g} {unit}" if a else f"{lab} no amount" for lab, a in zip(labels, amounts))
    codes = ", ".join(str(s.get("code")) for s in standards if s.get("role") == "Quantification") or "none"
    return (f"Extraction of a {what} ({amount}) in {m['extract_volume_ml']:g} mL; internal standards "
            f"{codes} ({m['spike_ul']:g} µL spiked); {detector}")


def generate(job, progress=lambda _: None):
    from gcws.report import service as RS
    infos = [(getattr(s, "meta", None) or {}).get("extraction") for s in job.samples]
    if not job.samples or any(i is None for i in infos):
        raise ValueError("Select the Extraction (quant method) quantification before creating a Quantification "
                         "report")
    problems = [f"{n}: {i['problem'] or i['missing']}" for n, i in zip(job.names, infos)
                if i["problem"] or i["missing"] or not i["factor"]]
    if problems:
        raise ValueError("Resolve these before reporting:\n" + "\n".join(problems))
    m = infos[0]["method"]
    units = list(m["units"])
    u1 = units[0]
    if any(i["method"]["units"][0] != u1 for i in infos):
        raise ValueError("The determinations must have the same Conc. 1 unit")
    progress("1/3 combined rows")
    combined = RS.combined_rows(job)
    from gcws.core.text import clean_rows
    clean_rows(combined)
    ratio = lambda k, unit: U.ratio(u1, unit, infos[k]["basis"])
    rows, dets = S.report_rows(job, combined, units, ratio, m.get("reporting_limit") or 0.0)
    head = S.headers(units)
    decimals = [U.DECIMALS.get(u, 4) for u in units]
    detector = job.samples[0].meta.get("detector", "FID")
    labels = [chr(ord("A") + k) if len(job.samples) <= 26 else str(k + 1) for k in range(len(job.samples))]
    amounts = [i["basis"].get(U.AMOUNT_BASIS[m["sample_type"]]) for i in infos]
    method = method_text(m, job.samples[0].standards, detector, amounts, labels)
    subtitle = (f"{m['name'] + '; ' if m['name'] else ''}{'Single' if len(job.samples) == 1 else 'Double' if len(job.samples) == 2 else f'{len(job.samples)}-fold'}"
                f" determination; Conc. 1 in {units[0]}, Conc. 2 in {units[1]}")
    progress("2/3 workbook")
    wb = Workbook()
    sh = wb.active
    sh.title = "Result"
    S.result_sheet(sh, TITLE, subtitle, job.names, method, head, rows, decimals)
    det_sheet = wb.create_sheet("Determinations")
    det_sheet.append(S.determination_headers(labels, units))
    for d in dets:
        det_sheet.append(d)
    calc = wb.create_sheet("Calculation")
    calc.append(["Determination", "Sample", "Source", "Sample type", "Sample amount", "Amount unit",
                 "Extract volume mL", "Spiked volume µL", "Factor mg per area count", "Factor calculation",
                 "Reporting limit", "Blanks"])
    stds = wb.create_sheet("Standards")
    stds.append(["Determination", "Sample", "Code", "Name", "Stock conc. mg/mL", "Role", "RT min", "Area", "Status"])
    peaks = wb.create_sheet("Peak calculation")
    peaks.append(["Determination", "Sample", "RT min", "Name", "CAS", "Raw area", "Blank area", "Corrected area",
                  "Substance mg", f"Conc. 1 [{units[0]}]", f"Conc. 2 [{units[1]}]"])
    for k, (s, info) in enumerate(zip(job.samples, infos)):
        b = info["basis"]
        im = info["method"]
        key = U.AMOUNT_BASIS[im["sample_type"]]
        calc.append([labels[k], job.names[k], str(getattr(s, "path", "") or ""), U.SAMPLE_TYPES[im["sample_type"]],
                     b.get(key), "g" if key == "mass_g" else "dm²", b.get("volume_ml"), im["spike_ul"],
                     info["factor"], info["factor_text"], im["reporting_limit"] or None,
                     "; ".join(n for n in job.blank_names if n)])
        for d in s.standards:
            stds.append([labels[k], job.names[k], d.get("code"), d.get("name"), d.get("concentration"), d.get("role"),
                         d.get("fid_rt"), d.get("fid_area"), d.get("status")])
        r2 = ratio(k, units[1])
        for r in s.rows:
            d = r.derived
            c = d.get("extraction_conc")
            peaks.append([labels[k], job.names[k], r.rt, r.name, r.cas, d.get("raw_area"),
                          max(d.get("blank_area") or 0.0, d.get("blank_istd_area") or 0.0) or None, r.area,
                          d.get("extraction_mg"), c, c * r2 if (c is not None and r2 is not None) else None])
    audit = wb.create_sheet("Audit")
    audit.append(["Timestamp", "Run", "Action", "Before", "After", "Detail"])
    for record in job.audit:
        audit.append([record.timestamp, record.run, record.action, record.before, record.after, record.detail])
    S.plain_sheets(wb, (det_sheet, calc, stds, peaks, audit))
    job.target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(job.target)
    progress("3/3 Word document")
    warnings = list(getattr(job, "notes", None) or [])
    word = None
    try:
        S.write_word(job.word, TITLE, subtitle, job.names, head, rows, method, decimals)
        word = job.word
    except Exception as exc:  # noqa: BLE001 - the .xlsx is the report
        warnings.append(f"Word document not created: {exc}")
    edited = sum(1 for v in (job.overrides or {}).values() if any(v.get(f) is not None for f in ("c1", "c2", "mean")))
    if edited:
        warnings.append(f"{edited} value(s) set by the analyst in the double determination")
    reported = [{"name": r[1], "cas": r[2], "rt": r[0]} for r in rows]
    if job.record_seen:
        error = RS.record_seen(job.kind, reported, job.target, job.sample_key)
        if error:
            warnings.append(error)
    return RS.ReportResult(job.target, word, len(rows), warnings=warnings, reported=reported,
                           combined=RS.slim_rows(combined, job.overrides))

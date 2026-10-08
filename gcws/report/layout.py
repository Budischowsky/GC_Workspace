"""The red-band report layout of a :class:`~gcws.report.table.ReportTable` in Excel and Word (the NIAS look:
title band, subtitle, label/value header lines, bold header row, data, sum rows, footnotes and notes)."""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from gcws.report.table import ReportTable, fmt

RED = "BA0C2F"
#: usable table width (cm) of an A4 page with 1.75 cm side margins
PAGE_WIDTH = {"portrait": 17.5, "landscape": 26.2}
THIN = Side(style="thin", color="000000")


def _literal(cell, value):
    from gcws.core.text import excel_safe
    value = excel_safe(value) if isinstance(value, str) else value
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _number_format(decimals) -> str:
    if decimals is None:
        return "General"
    return "0." + "0" * decimals if decimals else "0"


def header_cells(n_cols: int, pairs: list) -> list:
    """``[(first column, last column, text, is label)]`` of one header line (1-based columns): each
    label/value pair gets an equal share of the columns, or the line is one text when they are too few."""
    k = len(pairs)
    if not k:
        return []
    if n_cols < 2 * k:
        text = "   ".join(f"{a} {b}".strip() for a, b in pairs)
        return [(1, n_cols, text, False)]
    width = n_cols // k
    out = []
    for j, (label, value) in enumerate(pairs):
        start = 1 + j * width
        end = n_cols if j == k - 1 else start + width - 1
        out.append((start, start, label, True))
        out.append((start + 1, end, value, False))
    return out


def result_sheet(sh, table: ReportTable) -> int:
    """Write ``table`` onto ``sh``; returns the header row's number."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    cols = table.columns
    last = max(len(cols), 2)
    sh.sheet_view.showGridLines = False
    font = lambda **kw: Font(name="Arial", size=8, color="000000", **kw)
    for r, text in ((1, table.title), (2, table.subtitle)):
        sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last)
        _literal(sh.cell(r, 1), text)
        sh.cell(r, 1).alignment = Alignment(horizontal="left", vertical="center")
    sh["A1"].fill = PatternFill("solid", fgColor=RED)
    sh["A1"].font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
    sh["A2"].font = font(bold=True, italic=True)
    row = 3
    for line in table.header_lines:
        for start, end, text, _label in header_cells(last, line):
            if end > start:
                sh.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
            _literal(sh.cell(row, start), text)
            sh.cell(row, start).font = font()
            sh.cell(row, start).alignment = Alignment(horizontal="left", vertical="center")
        sh.row_dimensions[row].height = 13
        row += 1
    head = row
    for c, col in enumerate(cols, 1):
        cell = sh.cell(head, c)
        _literal(cell, col.header)
        cell.font = font(bold=True)
        cell.border = Border(bottom=THIN)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    first = head + 1
    if not table.rows:
        sh.merge_cells(start_row=first, start_column=1, end_row=first, end_column=last)
        _literal(sh.cell(first, 1), table.empty_text)
        sh.cell(first, 1).font = font(italic=True)
        sh.cell(first, 1).alignment = Alignment(horizontal="left", vertical="center")
    sum_started = False
    for r, out in enumerate(table.rows, first):
        top = out.kind == "sum" and not sum_started
        sum_started = sum_started or out.kind == "sum"
        for c, (col, item) in enumerate(zip(cols, out.cells), 1):
            cell = sh.cell(r, c)
            v = item.value
            if item.marker:
                rich = CellRichText()
                if v not in (None, ""):
                    rich.append(f"{fmt(v, col.decimals)} ")
                rich.append(TextBlock(InlineFont(rFont="Arial", sz=8, color="000000", vertAlign="superscript"),
                                      f"({item.marker})"))
                cell.value = rich
            else:
                _literal(cell, v)
            cell.font = font(bold=item.bold or out.kind == "sum")
            numeric = col.numeric and isinstance(v, (int, float)) and not isinstance(v, bool)
            if numeric:
                cell.number_format = _number_format(col.decimals)
            cell.alignment = Alignment(vertical="top" if col.key == "name" else "center",
                                       wrap_text=col.key in ("name", "notes", "comment", "verdict", "footnote"),
                                       horizontal="right" if (numeric or item.marker) else "left")
            if top:
                cell.border = Border(top=THIN)
    end = first + max(len(table.rows), 1) - 1
    nxt = end + 2
    for k, text in enumerate(table.footnotes):
        marker = sh.cell(nxt, 1, f"({chr(ord('a') + k)})")
        marker.font = Font(name="Arial", size=8, vertAlign="superscript", color="000000")
        if last > 2:
            sh.merge_cells(start_row=nxt, start_column=2, end_row=nxt, end_column=last)
        _literal(sh.cell(nxt, 2), text)
        sh.cell(nxt, 2).font = font()
        sh.cell(nxt, 2).alignment = Alignment(wrap_text=True, vertical="top")
        nxt += 1
    if table.notes:
        nxt += 1 if table.footnotes else 0
        for text in table.notes:
            sh.merge_cells(start_row=nxt, start_column=1, end_row=nxt, end_column=last)
            _literal(sh.cell(nxt, 1), text)
            sh.cell(nxt, 1).font = font(italic=True)
            sh.cell(nxt, 1).alignment = Alignment(wrap_text=True, vertical="top")
            nxt += 1
    total = sum(c.width for c in cols) or 1.0
    scale = (120.0 if table.orientation == "landscape" else 90.0) / total
    for c, col in enumerate(cols, 1):
        sh.column_dimensions[get_column_letter(c)].width = max(4.0, col.width * min(scale, 1.4))
    sh.freeze_panes = f"A{first}"
    sh.print_title_rows = f"1:{head}"
    if table.rows:
        sh.auto_filter.ref = f"A{head}:{get_column_letter(last)}{end}"
    sh.page_setup.orientation = table.orientation
    sh.page_setup.paperSize = sh.PAPERSIZE_A4
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    return head


def plain_sheet(sh, rows: list) -> None:
    """A detail sheet: header row bold and frozen, text cells as text, Excel-safe strings."""
    for line in rows:
        sh.append([None if v is None else v for v in line])
    for line in sh.iter_rows():
        for cell in line:
            if isinstance(cell.value, str):
                _literal(cell, cell.value)
    for cell in sh[1] if sh.max_row else []:
        cell.font = Font(bold=True)
    sh.freeze_panes = "A2"
    for c in range(1, sh.max_column + 1):
        sh.column_dimensions[get_column_letter(c)].width = 18


# -- Word ---------------------------------------------------------------------------------------

def new_document(orientation: str = "portrait"):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    _page(doc.sections[0], orientation)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(8)
    normal.paragraph_format.space_after = Pt(0)
    return doc


def _page(section, orientation):
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm
    wide = orientation == "landscape"
    section.orientation = WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (Cm(29.7), Cm(21)) if wide else (Cm(21), Cm(29.7))
    section.left_margin = section.right_margin = Cm(1.75)
    section.top_margin = section.bottom_margin = Cm(1.5)


def _border(cell, side: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pr = cell._tc.get_or_add_tcPr()
    borders = pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        pr.append(borders)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "4")
    borders.append(el)


def _hyperlink(paragraph, text: str, url: str, bold: bool = False):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    if bold:
        props.append(OxmlElement("w:b"))
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def append_table(doc, table: ReportTable, *, new_page: bool = False) -> None:
    """``table`` as a Word table (header rows repeated on every page), then its footnotes and notes."""
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    if new_page:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        _page(section, table.orientation)
    cols = table.columns
    n = max(len(cols), 2)
    lines = [header_cells(n, line) for line in table.header_lines]
    widths = [c.width for c in cols] or [1.0, 1.0]
    while len(widths) < n:
        widths.append(widths[-1])
    page = PAGE_WIDTH.get(table.orientation, 17.5)
    width = lambda c: Cm(page * widths[c] / sum(widths))
    head_rows = 3 + len(lines)
    t = doc.add_table(rows=head_rows, cols=n)
    t.autofit = False
    for c in range(n):
        t.columns[c].width = width(c)
    for line in t.rows:
        for c, cell in enumerate(line.cells):
            cell.width = width(c)
    for r, text in ((0, table.title), (1, table.subtitle)):
        t.cell(r, 0).merge(t.cell(r, n - 1)).text = text or ""
    band = OxmlElement("w:shd")
    band.set(qn("w:fill"), RED)
    t.cell(0, 0)._tc.get_or_add_tcPr().append(band)
    for run in t.cell(0, 0).paragraphs[0].runs:
        run.bold, run.font.size = True, Pt(10)
        run.font.color.rgb = RGBColor.from_string("FFFFFF")
    for run in t.cell(1, 0).paragraphs[0].runs:
        run.bold = run.italic = True
    for i, cells in enumerate(lines, 2):
        for start, end, text, _label in cells:
            target = t.cell(i, start - 1) if end == start else t.cell(i, start - 1).merge(t.cell(i, end - 1))
            target.text = text or ""
    h = head_rows - 1
    for c, col in enumerate(cols):
        cell = t.cell(h, c)
        cell.text = col.header
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
        _border(cell, "bottom")
    for line in t.rows:
        line._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    if not table.rows:
        cells = t.add_row().cells
        merged = cells[0].merge(cells[n - 1])
        merged.text = table.empty_text or ""
        for run in merged.paragraphs[0].runs:
            run.italic = True
    sum_started = False
    for out in table.rows:
        cells = t.add_row().cells
        top = out.kind == "sum" and not sum_started
        sum_started = sum_started or out.kind == "sum"
        for c, (col, item) in enumerate(zip(cols, out.cells)):
            cells[c].width = width(c)
            p = cells[c].paragraphs[0]
            text = fmt(item.value, col.decimals)
            bold = item.bold or out.kind == "sum"
            if item.link and text:
                _hyperlink(p, text, item.link, bold)
            elif text:
                p.add_run(text).bold = bold
            if item.marker:
                run = p.add_run(("" if not text else " ") + f"({item.marker})")
                run.font.superscript = True
            numeric = col.numeric and isinstance(item.value, (int, float)) and not isinstance(item.value, bool)
            if numeric or item.marker:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if top:
                _border(cells[c], "top")
        t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    if table.footnotes or table.notes:
        doc.add_paragraph()
    for k, text in enumerate(table.footnotes):
        p = doc.add_paragraph()
        p.add_run(f"({chr(ord('a') + k)})").font.superscript = True
        p.add_run(" " + text)
    for text in table.notes:
        doc.add_paragraph().add_run(text).italic = True


def append_audit(doc, records: list, title: str = "Audit Trail") -> None:
    """A page with the audit records (timestamp, run, action, before, after, detail)."""
    from docx.enum.text import WD_BREAK
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph().add_run(title).bold = True
    if not records:
        doc.add_paragraph("No changes recorded.")
        return
    t = doc.add_table(rows=1, cols=6)
    for c, h in enumerate(("Timestamp", "Run", "Action", "Before", "After", "Detail")):
        t.cell(0, c).text = h
        for run in t.cell(0, c).paragraphs[0].runs:
            run.bold = True
    for rec in records:
        cells = t.add_row().cells
        for c, v in enumerate(rec[:6]):
            cells[c].text = "" if v is None else str(v)


def write_word(path, table: ReportTable, audit: list | None = None) -> None:
    doc = new_document(table.orientation)
    append_table(doc, table)
    if table.audit_page:
        append_audit(doc, audit or [])
    doc.save(path)

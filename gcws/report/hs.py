"""HS-Screening report: RT, Name, CAS, Qual, Conc. 1 and Conc. 2 (by default µg/dm² and mg/m²) for a single or
a double determination, in the layout of the Quantification report, with the HS calculation sheets."""
from __future__ import annotations

from openpyxl import Workbook


def generate(job, progress=lambda _: None):
    from gcws.report import service as RS
    from gcws.report import simple as S
    from gcws.quant.hs import UNITS, input_label, per_ug
    from gcws.quant import units as U
    if not job.samples or any(getattr(s, "mode", None) != "hs_screening" for s in job.samples):
        raise ValueError("Select HS-Screening quantification before creating an HS report")
    units = {s.meta.get("unit") for s in job.samples}
    if len(units) != 1 or not units.issubset(UNITS):
        raise ValueError("HS determinations must have the same result unit")
    if any(not s.mean_factor or any(r.derived.get("status") for r in s.rows) for s in job.samples):
        raise ValueError("Resolve HS standard and sample-normalization errors before reporting")
    unit = next(iter(units))
    report_units = list(job.samples[0].meta.get("report_units") or ("µg/dm²", "mg/m²"))
    missing = [f"{name}: enter the {input_label(u)} for {u}" for name, s in zip(job.names, job.samples)
               for u in report_units if per_ug(u, s.meta.get("inputs")) is None]
    if missing:
        raise ValueError("The HS report needs the sample amounts:\n" + "\n".join(dict.fromkeys(missing)))
    from gcws.quant.hs import METHODS
    method = job.samples[0].meta.get("method") or METHODS["internal"]
    combined = RS.combined_rows(job)
    ratio = lambda k, u: per_ug(u, job.samples[k].meta.get("inputs")) / per_ug(unit, job.samples[k].meta.get("inputs"))
    rows, dets = S.report_rows(job, combined, report_units, ratio)
    progress("HS-Screening: result table")
    head = S.headers(report_units)
    decimals = [U.DECIMALS.get(u, 4) for u in report_units]
    n = len(job.samples)
    wb = Workbook()
    sh = wb.active
    sh.title = "HS Result"
    title = "GC-MS – HS-Screening"
    subtitle = (f"HS-Screening; {'Single' if n == 1 else 'Double' if n == 2 else f'{n}-fold'} determination; "
                f"Conc. 1 in {report_units[0]}, Conc. 2 in {report_units[1]}")
    S.result_sheet(sh, title, subtitle, job.names, method, head, rows, decimals)
    labels = [chr(ord("A") + k) if n <= 26 else str(k + 1) for k in range(n)]
    det_sheet = wb.create_sheet("Determinations")
    det_sheet.append(S.determination_headers(labels, report_units))
    for d in dets:
        det_sheet.append(d)
    detail = wb.create_sheet("HS calculation")
    detail.append(["Sample", "Source", "Unit", "Sample area dm²", "Sample mass g", "µg per area count",
                   "Calculation", "Blank correction", "Blanks"])
    standards = wb.create_sheet("HS standards")
    standards.append(["Sample", "Code", "Name", "µg/HS", "RT min", "TIC area", "Active", "Status",
                      "Calibration runs (TIC area)"])
    peaks = wb.create_sheet("HS peak calculation")
    peaks.append(["Sample", "RT min", "Name", "CAS", "Raw TIC area", "Blank area", "Corrected area",
                  "µg/HS", unit] + [u for u in report_units if u != unit] + ["ISTD"])
    for s in job.samples:
        meta = s.meta
        detail.append([s.name, meta["source"], unit, meta["inputs"].get("area_dm2"), meta["inputs"].get("mass_g"),
                       s.mean_factor, meta["calculation"], meta["blank_correction"], "; ".join(meta["blanks"])])
        for d in s.standards:
            runs = "; ".join(f"{n}: {a:.6g}" if a is not None else f"{n}: not found" for n, a in d.get("areas") or [])
            standards.append([s.name, d["code"], d["name"], float(d["concentration"]), d["rt"], d["area"],
                              bool(d.get("quantify", True)), d["status"], runs])
        for r in s.rows:
            d = r.derived
            amount = d["amount_ug"]
            more = [amount * per_ug(u, meta.get("inputs")) if amount is not None else None
                    for u in report_units if u != unit]
            peaks.append([s.name, r.rt, r.name, r.cas, d["raw_area"], d["blank_area"], d["corr_area"],
                          amount, d["conc"]] + more + [d["istd"]])
    audit = wb.create_sheet("Audit")
    audit.append(["Timestamp", "Run", "Action", "Before", "After", "Detail"])
    for record in job.audit:
        audit.append([record.timestamp, record.run, record.action, record.before, record.after, record.detail])
    S.plain_sheets(wb, (det_sheet, detail, standards, peaks, audit))
    wb.save(job.target)
    S.write_word(job.word, title, subtitle, job.names, head, rows, method, decimals)
    reported = [{"name": r[1], "cas": r[2], "rt": r[0]} for r in rows]
    warnings = list(getattr(job, "notes", None) or [])
    if job.record_seen:
        error = RS.record_seen(job.kind, reported, job.target, job.sample_key)
        if error:
            warnings.append(error)
    return RS.ReportResult(job.target, job.word, len(rows), warnings=warnings, reported=reported,
                           combined=RS.slim_rows(combined, job.overrides))


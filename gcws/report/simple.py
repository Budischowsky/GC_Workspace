"""The six-column result report shared by the Quantification and the HS-Screening reports:
RT, Name, CAS, Qual, Conc. 1 and Conc. 2 (two units of the same result), as .xlsx and .docx.

A single determination reports its own values; a double (N-fold) determination the mean of its
determinations in each unit, the analyst's edits and a dismissed outlier included. Each determination
is converted with its own ratio (its own sample amount), so A and B of different sample amounts are
averaged in each unit correctly."""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RED = "BA0C2F"
WIDTHS = [8, 40, 14, 8, 13, 13]


def literal(cell, value):
    from gcws.core.text import excel_safe
    value = excel_safe(value)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def headers(units) -> list[str]:
    return ["RT (min)", "Name", "CAS-No.", "Qual"] + [f"Conc. {i} [{u}]" for i, u in enumerate(units, 1)]


def report_rows(job, combined, units, ratio, limit: float = 0.0):
    """``(rows, determinations)`` of the report.

    ``ratio(k, unit)``: the factor from determination k's concentration (``combined``'s unit) to ``unit``.
    ``rows``: [rt, name, cas, qual, conc 1, conc 2]; ``determinations``: [rt, name, cas, then per
    determination its value in each unit, "dismissed" / ""]. Rows without a result, and rows whose
    Conc. 1 is below ``limit`` (> 0), are left out."""
    from gcws.quant import extraction as EX
    rows, dets = [], []
    n = len(job.samples)
    for i, row in enumerate(combined):
        ov = job.overrides.get(i, {}) if job.overrides else {}
        cs = list(row.get("cs") or [row.get("c1"), row.get("c2")][:n])
        for k, f in enumerate(("c1", "c2")[:len(cs)]):
            if ov.get(f) is not None:
                cs[k] = ov[f]
        mean = ov.get("mean") if ov.get("mean") is not None else row.get("mean")
        dismissed = int(ov.get("dismissed") or 0)
        if mean is None:
            continue
        values = [EX.mean_in(None, list(range(len(cs))), cs, mean, u, dismissed, ratio=ratio) for u in units]
        if values[0] is None or (limit and values[0] < limit):
            continue
        sources = row.get("sources") or []
        scores = [s.get("score", s.get("quality")) for s in sources if s]
        scores = [float(s) for s in scores if s is not None]
        rows.append([row.get("rt"), row.get("name", "Unknown"), row.get("cas", ""),
                     min(scores) if scores else None] + values)
        det = [row.get("rt"), row.get("name", "Unknown"), row.get("cas", "")]
        for k, c in enumerate(cs):
            for u in units:
                r = ratio(k, u)
                det.append(c * r if (c is not None and r is not None) else None)
            det.append("dismissed" if dismissed == k + 1 else "")
        # the difference with the analyst's values (none with a dismissed determination)
        reldiff = ov["reldiff"] if "reldiff" in ov else row.get("reldiff")
        det += [reldiff, row.get("status", ""), row.get("review", "")]
        dets.append(det)
    return rows, dets


def determination_headers(labels, units) -> list[str]:
    out = ["RT (min)", "Name", "CAS-No."]
    for lab in labels:
        out += [f"{lab} [{u}]" for u in units] + [f"{lab} outlier"]
    return out + ["Relative difference [%]", "Status", "Review"]


def result_sheet(sh, title, subtitle, names, method, head, rows, decimals):
    """The red-banded result table (rows 1-4 title, sample and method; header row 5; data from row 6)."""
    last = len(head)
    sh.sheet_view.showGridLines = False
    for r, text in ((1, title), (2, subtitle)):
        sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last)
        literal(sh.cell(r, 1), text)
    sh["A1"].fill = PatternFill("solid", fgColor=RED)
    sh["A1"].font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
    sh["A2"].font = Font(name="Arial", size=8, bold=True, italic=True)
    for r, label, text in ((3, "Sample:", "; ".join(names)), (4, "Method:", method)):
        sh.cell(r, 1, label)
        sh.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last)
        literal(sh.cell(r, 2), text)
    for c, h in enumerate(head, 1):
        sh.cell(5, c, h)
    for r, values in enumerate(rows, 6):
        for c, value in enumerate(values, 1):
            literal(sh.cell(r, c), value)
    for cells in sh.iter_rows(min_row=3):
        for cell in cells:
            cell.font = Font(name="Arial", size=8, bold=cell.row == 5)
            cell.alignment = Alignment(vertical="center", horizontal="right" if cell.column >= 4 else "left",
                                       wrap_text=cell.column == 2)
            if cell.row == 5:
                cell.border = Border(bottom=Side(style="thin", color="000000"))
            if cell.row >= 6:
                if cell.column == 1:
                    cell.number_format = "0.000"
                elif cell.column == 4:
                    cell.number_format = "0"
                elif cell.column >= 5:
                    d = decimals[cell.column - 5]
                    cell.number_format = "0." + "0" * d if d else "0"
    for c, width in enumerate(WIDTHS[:last], 1):
        sh.column_dimensions[get_column_letter(c)].width = width * 1.1
    sh.freeze_panes = "A6"
    sh.print_title_rows = "1:5"
    sh.print_area = f"A1:{get_column_letter(last)}{max(6, sh.max_row)}"
    sh.page_setup.orientation = "portrait"
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True


def plain_sheets(wb, sheets):
    """Detail sheets: frozen header, text cells as text, Excel-safe strings."""
    from gcws.core.text import ILLEGAL, excel_safe
    for sheet in sheets:
        sheet.freeze_panes = "A2"
        for c in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(c)].width = 20
    for sheet in wb.worksheets:
        for line in sheet.iter_rows():
            for cell in line:
                if isinstance(cell.value, str):
                    if ILLEGAL.search(cell.value):
                        cell.value = excel_safe(cell.value)
                    cell.data_type = "s"


def write_word(path, title, subtitle, names, head, rows, method, decimals):
    """The result table as a portrait Word document (Arial 8, red title band, repeated header)."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.75)
    section.top_margin = section.bottom_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(8)
    normal.paragraph_format.space_after = Pt(0)
    cols = len(head)
    widths = WIDTHS[:cols]
    table = doc.add_table(rows=5, cols=cols)
    table.autofit = False
    width = lambda c: Cm(17.5 * widths[c] / sum(widths))
    for c in range(cols):
        table.columns[c].width = width(c)
    for row in table.rows:
        for c, cell in enumerate(row.cells):
            cell.width = width(c)
    for r, text in ((0, title), (1, subtitle)):
        table.cell(r, 0).merge(table.cell(r, cols - 1)).text = text
    band = OxmlElement("w:shd")
    band.set(qn("w:fill"), RED)
    table.cell(0, 0)._tc.get_or_add_tcPr().append(band)
    for run in table.cell(0, 0).paragraphs[0].runs:
        run.bold, run.font.size = True, Pt(10)
        run.font.color.rgb = RGBColor.from_string("FFFFFF")
    for r, label, value in ((2, "Sample:", "; ".join(names)), (3, "Method:", method)):
        table.cell(r, 0).text = label
        table.cell(r, 1).merge(table.cell(r, cols - 1)).text = value
    for c, h in enumerate(head):
        table.cell(4, c).text = h
        borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        borders.append(bottom)
        table.cell(4, c)._tc.get_or_add_tcPr().append(borders)
        for run in table.cell(4, c).paragraphs[0].runs:
            run.bold = True
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for values in rows:
        cells = table.add_row().cells
        for c, value in enumerate(values):
            cells[c].width = width(c)
            if value is None:
                text = ""
            elif c == 0:
                text = f"{value:.2f}"
            elif c == 3:
                text = f"{value:.0f}"
            elif c >= 4:
                text = f"{value:.{decimals[c - 4]}f}"
            else:
                text = str(value)
            cells[c].text = text
            if c >= 3:
                cells[c].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    doc.save(path)

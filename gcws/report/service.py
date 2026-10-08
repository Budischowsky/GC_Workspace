"""NIAS, Fingerprint and Total extraction reports from the workspace.

The chain is the NIAS one, unchanged in its later steps:
  1. the determinations of a replicate group as ``gc_model.Session``
  2. the combined rows (single, duplicate or N-fold)
  3. the intermediate workbook (``gc_export``)
     (a double determination: only the substances the analyst lets through, with the values
     the analyst set written into the workbook as numbers)
  4. ``process_workbook`` of the NIAS main script -> report .xlsx
  5. ``create_combined_word`` -> report .docx
  6. the reported substances into the unknown register ("already reported")
"""
from __future__ import annotations

import re
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Optional

KINDS = {"nias": "NIAS Report", "fingerprint": "Fingerprint Report",
         "total_extraction": "Total Extraction Report", "hs_screening": "HS-Screening Report",
         "quant": "Quantification Report", "template": "Template Report"}
SUFFIXES = {"nias": "_NIAS_Report", "fingerprint": "_Fingerprint_Report",
            "total_extraction": "_Total_Extraction_Report", "hs_screening": "_HS_Screening_Report",
            "quant": "_Quantification_Report", "template": "_Template_Report"}
SEEN_TYPES = {"nias": "NIAS", "fingerprint": "Fingerprint", "total_extraction": "Total extraction",
              "hs_screening": "HS-Screening", "quant": "Quantification", "template": "Template"}
#: quantification modes with their own report (every other mode: the NIAS, Fingerprint and Total extraction reports)
MODE_KINDS = {"hs_screening": "hs_screening", "extraction": "quant"}


def default_kind(quant: dict) -> str:
    """The report that fits the quantification mode: HS-Screening, Quantification (extraction method) or NIAS."""
    return MODE_KINDS.get((quant or {}).get("mode"), "nias")


def kind_fits(kind: str, quant: dict) -> bool:
    """Whether a ``kind`` report can be made in the quantification mode of ``quant`` (the Template report:
    in every mode)."""
    if kind == "template":
        return True
    own = MODE_KINDS.get((quant or {}).get("mode"))
    return kind == own if own else kind not in MODE_KINDS.values()


def preview_kind(quant: dict) -> str:
    """The report the panels' Report preview buttons show: the method's template when it asks for it,
    else the report of the quantification mode."""
    from gcws.report import template as TP
    tpl = TP.of(quant)
    return "template" if tpl is not None and tpl["extras"]["default_report"] else default_kind(quant)


def suffix(kind: str, quant: Optional[dict] = None, template: Optional[dict] = None) -> str:
    """The end of the report's file name: the template's own ending, else ``_<template name>_Report``
    (``_<name>_Template_Report`` where that would be a fixed report's), else the kind's."""
    if kind != "template":
        return SUFFIXES[kind]
    from gcws.report import template as TP
    tpl = TP.normalise(template) if template is not None else TP.of(quant)
    if tpl is None:
        return SUFFIXES[kind]
    if tpl["extras"]["file_suffix"]:
        own = tpl["extras"]["file_suffix"]
        return own if own.startswith(("_", "-", " ")) else "_" + own
    name = re.sub(r"[^\w.-]+", "_", tpl["name"]).strip("_.")
    if not name:
        return SUFFIXES[kind]
    out = f"_{name}_Report"
    return out.replace("_Report", "_Template_Report") if out in SUFFIXES.values() else out


@dataclass
class ReportJob:
    kind: str
    samples: list                  # NiasSample, in determination order
    names: list[str]
    settings: object
    target: Path                   # report .xlsx
    word: Path
    cas_path: Optional[Path]
    migration: dict = field(default_factory=dict)
    blank_names: tuple = ("", "")
    audit: list = field(default_factory=list)       # AuditRecord
    policy: str = "all"
    batch_target: Optional[Path] = None
    keep_middle: Optional[Path] = None
    sample_key: str = ""
    record_seen: bool = True
    ri_options: Optional[dict] = None
    edits: dict = field(default_factory=dict)       # analyst edits of the double determination
    notes: list = field(default_factory=list)       # warnings known before the report is made (e.g. no blank)
    overrides: dict = field(default_factory=dict)   # filled by combined_rows: row position -> values
    features: object = None                         # FeatureTable: pair by features (None: AutoLib's pairing)
    template: Optional[dict] = None                 # Template report: the template ...
    table: object = None                            # ... and its table (gcws.report.table), built with the job


@dataclass
class ReportResult:
    target: Path
    word: Optional[Path]
    rows: int
    batch: Optional[Path] = None
    middle: Optional[Path] = None
    warnings: list = field(default_factory=list)
    reported: list = field(default_factory=list)
    #: what the NIAS main script found (``process_workbook``'s result: SML exceedances, unidentified count, ...)
    summary: dict = field(default_factory=dict)
    #: the merged rows of the determinations (rt, name, cas, mean, c1, c2, reldiff, status, id_status, review)
    combined: list = field(default_factory=list)


def report_stem(names: list[str]) -> str:
    for n in names:
        m = re.search(r"(?<!\d)(\d{8})(?!\d)", n)
        if m:
            return m.group(1)
    return re.sub(r'[<>:"/\\|?*]+', "_", names[0]) if names else "Report"


def build_session(job: ReportJob):
    import gc_model as M
    session = M.Session()
    for i, sample in enumerate(job.samples, 1):
        sample.label = f"{i:02d}" if len(job.samples) > 9 else str(i)
        session.add_sample(sample)
    for rec in job.audit:
        session.audit.append(M.EditRecord(rec.timestamp, rec.run, 0, 0, rec.action,
                                          rec.before, (rec.after + ("  " + rec.detail if rec.detail else "")).strip()))
    return session


def combined_rows(job: ReportJob):
    """The merged rows of the determinations. For a double determination only the substances
    the analyst reports (default: AutoLib's rule, no artefacts and nothing below the reporting
    limit); ``job.overrides`` receives the values the analyst set."""
    from gcws.quant.replicates import combine, engine_peaks
    tol = float(getattr(job.settings, "rt_tolerance", 0.035) or 0.035)
    lists = [engine_peaks(s) for s in job.samples]
    job.overrides = {}
    if job.features is not None and len(job.samples) >= 2:
        return feature_rows(job, lists, tol)
    rows = combine(lists, tol, job.policy)
    if len(job.samples) == 2 and job.policy == "all":
        from gcws.quant import duplicate_view as DV
        limit = DV.reldiff_limit(job.settings)
        rl = 0.0
        if job.kind == "nias":
            import gc_duplicate as GD
            rl = getattr(job.settings, "reporting_limit", None)
            rl = float(rl if rl is not None else GD.DEFAULT_REPORTING_LIMIT)
        rows, job.overrides = DV.rows_for_report(rows, job.edits or {}, limit, rl, tol)
    return rows


def feature_rows(job: ReportJob, lists: list, tol: float):
    """The combined rows of the feature double determination (``gcws.features.combine``); for two
    determinations the analyst's report choices and numbers as in :func:`combined_rows`."""
    from gcws.features import combine as FC
    from gcws.quant import duplicate_view as DV
    limit = DV.reldiff_limit(job.settings)
    rl = 0.0
    if job.kind == "nias":
        import gc_duplicate as GD
        rl = getattr(job.settings, "reporting_limit", None)
        rl = float(rl if rl is not None else GD.DEFAULT_REPORTING_LIMIT)
    rows = FC.rows(job.features, lists, limit, rl, job.policy)
    if len(job.samples) == 2 and job.policy == "all":
        rows, job.overrides = DV.rows_for_report(rows, job.edits or {}, limit, rl, tol)
    return rows


def _metadata_save(main, metadata: dict):
    """``save(workbook, path)`` for the workbook writers: the migration metadata written into the
    loaded workbook and saved as ``write_migration_metadata_to_workbook`` saves it (recalculation
    flags, normalised package), the same file without reading it back first."""
    def save(workbook, path):
        main._apply_migration_metadata(workbook, main.validate_migration_metadata(metadata))
        main._save_recalculating(workbook, Path(path))
    return save


def apply_overrides(middle: Path, overrides: dict) -> int:
    """Write the analyst's values into the intermediate NIAS workbook as numbers (the NIAS main
    script takes a saved number before its formula fallback). Returns the rows changed."""
    if not overrides:
        return 0
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    wb = load_workbook(middle)
    if "Doppelbestimmung" not in wb.sheetnames:
        return 0
    sh = wb["Doppelbestimmung"]
    col = {str(c.value): c.column for c in sh[1] if c.value is not None}
    cells = {"a1": "Area 1", "c1": "Concentration 1 [mg/kg]", "a2": "Area 2", "c2": "Concentration 2 [mg/kg]",
             "mean": "mg/kg (mean)"}
    note = Comment("set by the analyst in GC Workspace (double determination)", "GC Workspace")
    for i, values in overrides.items():
        r = 2 + int(i)
        for field, header in cells.items():
            if header in col and values.get(field) is not None:
                c = sh.cell(r, col[header])
                c.value = float(values[field])
                c.comment = note
        rd = values.get("reldiff")
        if "Relative difference [%]" in col:
            sh.cell(r, col["Relative difference [%]"]).value = None if rd is None else float(rd) / 100.0
        if "Review" in col:
            c = sh.cell(r, col["Review"])
            c.value = "; ".join(x for x in (str(c.value or ""), "values set by the analyst") if x)
    wb.save(middle)
    return len(overrides)


def generate(job: ReportJob, progress: Callable[[str], None] = lambda s: None) -> ReportResult:
    if job.kind == "template":
        # the table was built with the job: writing it never touches the workspace's samples
        from gcws.report.template_report import generate as generate_template
        return generate_template(job, progress)
    clean_samples(job.samples)
    if job.kind == "hs_screening":
        from gcws.report.hs import generate as generate_hs
        return generate_hs(job, progress)
    if job.kind == "quant":
        from gcws.report.quant import generate as generate_quant
        return generate_quant(job, progress)
    if any(getattr(s, "mode", None) == "hs_screening" for s in job.samples):
        raise ValueError("HS samples require the HS-Screening report")
    import gc_export
    from gcws.report.legacy_api import main_script
    warnings = list(job.notes)
    session = build_session(job)
    combined = combined_rows(job)
    from gcws.core.text import clean_rows
    clean_rows(combined)
    tmp = Path(tempfile.mkdtemp(prefix="gcws_report_"))
    try:
        middle = tmp / f"{job.target.stem}_intermediate.xlsx"
        progress("1/4 intermediate workbook")
        audit = None
        main = main_script()
        # the report metadata go into the workbook before its one save (no load and save again); with
        # the analyst's overrides, written into the saved file, they follow those as before
        with_metadata = _metadata_save(main, job.migration) if job.migration and job.kind == "nias" else None
        if job.kind == "nias":
            gc_export.run_nias_duplicate(session, middle, job.settings,
                                         blank_path=Path(job.blank_names[0]) if job.blank_names[0] else None,
                                         blank_istd_path=Path(job.blank_names[1]) if job.blank_names[1] else None,
                                         combined=combined, ri_options=job.ri_options,
                                         save=None if job.overrides else with_metadata)
            n = apply_overrides(middle, job.overrides)
            if n:
                warnings.append(f"{n} value(s) set by the analyst in the double determination")
        else:
            audit = gc_export.write_fingerprint_workbook(session, middle, job.kind, job.settings, combined=combined,
                                                         ri_options=job.ri_options)
            if any(any(v.get(k) is not None for k in ("a1", "a2", "c1", "c2")) for v in job.overrides.values()):
                warnings.append("Values edited in the double determination are used in the NIAS report only")
        progress("2/4 report (NIAS main script)")
        if with_metadata is not None and job.overrides:
            main.write_migration_metadata_to_workbook(middle, job.migration)
        job.target.parent.mkdir(parents=True, exist_ok=True)
        summary = main.process_workbook(middle, job.cas_path if job.kind == "nias" else None, job.target)
        progress("3/4 Word document")
        word = None
        try:
            word = Path(main.create_combined_word([job.target], job.word))
        except Exception as exc:  # noqa: BLE001 - the .xlsx is the report
            warnings.append(f"Word document not created: {exc}")
        batch = None
        if job.kind == "nias" and job.batch_target is not None:
            try:
                gc_export.write_batch_workbook(session, job.batch_target, job.settings,
                                               sample_name=job.target.stem.replace(SUFFIXES["nias"], ""),
                                               save=with_metadata)
                batch = job.batch_target
            except Exception as exc:  # noqa: BLE001
                warnings.append(f"Batch workbook not written: {exc}")
        progress("4/4 register")
        rows = reported_rows(job, combined, audit)
        if job.record_seen:
            err = record_seen(job.kind, rows, job.target, job.sample_key)
            if err:
                warnings.append(err)
        kept = None
        if job.keep_middle is not None:
            kept = job.keep_middle
            shutil.copy2(middle, kept)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)               # also when the report failed
    return ReportResult(job.target, word, len(rows), batch, kept, warnings, rows, json_safe(summary or {}),
                        slim_rows(combined, job.overrides))


def clean_samples(samples) -> None:
    """Remove the characters Excel refuses from the names of the determinations' peak rows
    (some library files carry control characters in their names)."""
    from gcws.core.text import ILLEGAL, clean_name
    for sample in samples or []:
        for row in getattr(sample, "rows", None) or []:
            for attr in ("name", "cas", "ref", "mark"):
                v = getattr(row, attr, None)
                if isinstance(v, str) and ILLEGAL.search(v):
                    setattr(row, attr, clean_name(v))


def json_safe(value):
    """``value`` with paths, sets and other objects as JSON-compatible values."""
    import json
    return json.loads(json.dumps(value, default=lambda o: sorted(o) if isinstance(o, (set, frozenset)) else str(o)))


SLIM_KEYS = ("rt", "name", "cas", "mean", "c1", "c2", "reldiff", "status", "id_status", "review", "n")


def slim_rows(combined, overrides: dict | None = None) -> list[dict]:
    """The merged rows as plain values, with the analyst's values where set."""
    out = []
    for i, r in enumerate(combined or []):
        d = {k: r.get(k) for k in SLIM_KEYS}
        for k, v in ((overrides or {}).get(i) or {}).items():
            if k in d and v is not None:
                d[k] = v
        out.append(json_safe(d))
    return out


def record_seen(kind: str, rows: list, target: Path, sample_key: str = "") -> str:
    """Count the reported substances in the register ('already reported')."""
    try:
        import gc_seen
        con = gc_seen.open_db()
        try:
            gc_seen.record(con, rows, report_id=str(Path(target).resolve()).casefold(),
                           report_type=SEEN_TYPES[kind], sample_key=sample_key)
        finally:
            con.close()
    except Exception as exc:  # noqa: BLE001 - the report is already on disk
        return f"Register not updated: {exc}"
    return ""


def reported_rows(job: ReportJob, combined, audit) -> list:
    if audit is not None and isinstance(audit.get("rows"), list):
        return list(audit["rows"])
    limit = float(getattr(job.settings, "reporting_limit", 0.01) or 0.01)
    out = []
    for i, r in enumerate(combined):
        mean = (job.overrides.get(i) or {}).get("mean", r.get("mean"))
        if mean is None or mean < limit:
            continue
        out.append({"name": r.get("name", ""), "cas": r.get("cas", ""), "mean_mgkg": mean,
                    "rt": r.get("rt"), "status": r.get("status", "")})
    return out


# -- preview -------------------------------------------------------------------

WD_FORMAT_PDF = 17


def docx_to_pdf(docx: Path, pdf: Path) -> Path:
    """Microsoft Word over COM (the layout engine that will open the file)."""
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(Path(docx).resolve()), ReadOnly=True, AddToRecentFiles=False, Visible=False)
        try:
            doc.SaveAs2(str(Path(pdf).resolve()), FileFormat=WD_FORMAT_PDF)
        finally:
            doc.Close(False)
    finally:
        if word is not None:
            try:
                word.Quit(False)
            except Exception:  # noqa: BLE001
                pass
        pythoncom.CoUninitialize()
    if not Path(pdf).is_file():
        raise RuntimeError("Word did not write a PDF")
    return Path(pdf)


def render_pages(pdf: Path, scale: float = 1.5) -> list:
    """``QImage`` of every page (pypdfium2; needs numpy only, not Pillow)."""
    import numpy as np
    import pypdfium2
    from PySide6.QtGui import QImage
    doc = pypdfium2.PdfDocument(str(pdf))
    pages = []
    try:
        for i in range(len(doc)):
            bitmap = doc[i].render(scale=scale, rev_byteorder=True)      # RGB byte order
            arr = np.ascontiguousarray(bitmap.to_numpy()[:, :, :3])
            h, w = arr.shape[:2]
            pages.append(QImage(arr.data, w, h, 3 * w, QImage.Format_RGB888).copy())
    finally:
        doc.close()
    return pages

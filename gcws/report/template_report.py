"""Template report: the analyst's template applied to a single or N-fold determination, as .xlsx and .docx.

The table is built on the GUI thread when the job is made (``assemble.build_job``); generating only
writes it, so the worker never touches the workspace. The workbook carries the table as JSON on a very
hidden sheet, from which :func:`combined_word` makes one Word document of several reports (batch)."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from gcws.report import layout as L
from gcws.report.table import ReportTable

PAYLOAD_SHEET = "_gcws_template"
#: characters per payload cell (Excel holds 32 767)
CHUNK = 30000


def _audit_rows(job) -> list:
    return [[r.timestamp, r.run, r.action, r.before, r.after, r.detail] for r in (job.audit or [])]


def _payload(wb, table: ReportTable, audit: list) -> None:
    sh = wb.create_sheet(PAYLOAD_SHEET)
    text = json.dumps({"table": table.to_dict(), "audit": audit}, ensure_ascii=False, default=str)
    for i in range(0, len(text), CHUNK):
        sh.cell(i // CHUNK + 1, 1).value = text[i:i + CHUNK]
        sh.cell(i // CHUNK + 1, 1).data_type = "s"
    sh.sheet_state = "veryHidden"


def is_template_workbook(path) -> bool:
    try:
        wb = load_workbook(path, read_only=True)
    except Exception:  # noqa: BLE001
        return False
    try:
        return PAYLOAD_SHEET in wb.sheetnames
    finally:
        wb.close()


def read_payload(path) -> tuple[ReportTable, list]:
    """``(table, audit rows)`` of a Template report workbook."""
    wb = load_workbook(path, read_only=True)
    try:
        if PAYLOAD_SHEET not in wb.sheetnames:
            raise ValueError(f"{Path(path).name} is not a Template report")
        text = "".join(str(row[0] or "") for row in wb[PAYLOAD_SHEET].iter_rows(values_only=True))
    finally:
        wb.close()
    data = json.loads(text)
    return ReportTable.from_dict(data["table"]), data.get("audit") or []


def generate(job, progress=lambda _: None):
    from gcws.report import service as RS
    table = job.table if isinstance(job.table, ReportTable) else ReportTable.from_dict(job.table or {})
    if not table.columns:
        raise ValueError("The report template has no column that this quantification can fill")
    progress("1/2 workbook")
    audit = _audit_rows(job)
    wb = Workbook()
    sh = wb.active
    sh.title = "Result"
    L.result_sheet(sh, table)
    if table.sheets.get("determinations") and len(table.determinations) > 1:
        L.plain_sheet(wb.create_sheet("Determinations"), table.determinations)
    if table.sheets.get("calculation"):
        L.plain_sheet(wb.create_sheet("Calculation"), table.calculation)
        if len(table.standards) > 1:
            L.plain_sheet(wb.create_sheet("Standards"), table.standards)
    if table.sheets.get("audit"):
        L.plain_sheet(wb.create_sheet("Audit"), [["Timestamp", "Run", "Action", "Before", "After", "Detail"]] + audit)
    _payload(wb, table, audit)
    job.target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(job.target)
    warnings = list(getattr(job, "notes", None) or []) + list(table.warnings)
    word = None
    if table.word:
        progress("2/2 Word document")
        try:
            L.write_word(job.word, table, audit)
            word = job.word
        except Exception as exc:  # noqa: BLE001 - the .xlsx is the report
            warnings.append(f"Word document not created: {exc}")
    if job.record_seen and table.reported:
        error = RS.record_seen(job.kind, table.reported, job.target, job.sample_key)
        if error:
            warnings.append(error)
    return RS.ReportResult(job.target, word, len(table.reported), warnings=list(dict.fromkeys(warnings)),
                           reported=table.reported, summary=dict(table.summary), combined=table.combined)


def combined_word(paths: list, target) -> Path:
    """One Word document of several Template report workbooks (each on its own page), audit pages last."""
    tables = [read_payload(p) for p in paths]
    if not tables:
        raise ValueError("No Template report to combine")
    doc = L.new_document(tables[0][0].orientation)
    audit = []
    for k, (table, rows) in enumerate(tables):
        L.append_table(doc, table, new_page=k > 0)
        audit += rows
    if any(t.audit_page for t, _ in tables):
        L.append_audit(doc, audit)
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    return target

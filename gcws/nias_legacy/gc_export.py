#!/usr/bin/env python3
"""Export the workspace session into the existing report pipeline.

Deviation from the spec, deliberate
-----------------------------------
The spec's seam 1 hands ``list[Peak]`` straight into the report engines, which
would require splitting ``process()`` -- one ~130-line function that both
evaluates and writes its workbook -- into a pure ``evaluate()`` and an I/O
wrapper.

This module instead writes the sheets in exactly the layout the engine already
reads, and hands it that file. The evaluation, the clustering, the QC gates and
the sheet formatting then run unchanged, so results are identical to today's by
construction rather than by inspection. The refactor stays available and is the
better long-term shape; it is simply not worth putting between the analyst and a
correct report right now.

The grid remains the source of truth: the workbook is a generated intermediate,
not something anyone edits.

Scope, since spec v3.2 §VIII.4
------------------------------
This module is the NIAS export plus the vocabulary both tools share
(``load_main_script``, ``_write_audit_sheet``, ``excluded_rows``, the header
styles). The raw workbook and the DIN SPEC 91521 run moved to ``gc_dinspec``;
nothing here imports it, so deleting that file leaves this module intact.
"""

from __future__ import annotations

import importlib.util
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Optional

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import gc_model as M
# Shared with the workspace sheet: the grid and the workbook must parse
# numbers and count register sightings the same way.
from gc_duplicate import _num, _seen_lookup

AUDIT_SHEET = "Manuelle Änderungen"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


# --------------------------------------------------------------------------
# Shared workbook vocabulary
# --------------------------------------------------------------------------

def excluded_rows(sample: M.Sample) -> list[M.PeakRow]:
    return [r for r in sample.rows if r.area is None]


def _write_audit_sheet(wb: Workbook, session: M.Session) -> None:
    """Every manual change, so a report can be traced back to what was edited."""
    ws = wb.create_sheet(AUDIT_SHEET)
    ws.append(["Zeitpunkt", "Injektion", "Peak", "Feld", "vorher", "nachher"])
    for record in session.audit:
        ws.append([record.timestamp, record.sample, record.peak_no,
                   record.field, record.old, record.new])

    dropped = [(label, r) for label in session.labels
               for r in excluded_rows(session.samples[label])]
    if dropped:
        ws.append([])
        ws.append(["Nicht exportierte Peaks: Bibliothekstreffer ohne integrierten "
                   "TIC-Peak (keine Fläche, daher nicht quantifizierbar):"])
        ws.append(["Injektion", "Peak", "RT", "SI", "Name", "CAS #"])
        for label, row in dropped:
            ws.append([label, row.peak_no, row.rt, row.si, row.name,
                       M.display_cas(row.cas)])

    reviews = session.open_reviews()
    if reviews:
        ws.append([])
        ws.append(["Unsichere Identifikationen (Quality unterhalb der Schwelle), "
                   "nicht bestätigt:"])
        ws.append(["Injektion", "Peak", "RT", "Quality", "Status", "Name", "CAS #"])
        for label, row in reviews:
            ws.append([label, row.peak_no, row.rt, row.si, row.id_status,
                       row.name, M.display_cas(row.cas)])

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for index, width in enumerate((20, 11, 8, 14, 30, 30), start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


# --------------------------------------------------------------------------
# Handing over to the existing engine
# --------------------------------------------------------------------------

_MAIN_SCRIPT = "NIAS Reporting v27.py"


def load_main_script(script_path: Optional[Path] = None):
    """Import the main NIAS script as a module.

    Its filename contains spaces, so it cannot be imported by name. ``main()``
    is guarded, so importing it has no side effects beyond definitions.
    """
    path = Path(script_path) if script_path else Path(__file__).resolve().parent / _MAIN_SCRIPT
    if not path.is_file():
        raise FileNotFoundError(f"Hauptskript nicht gefunden: {path}")
    spec = importlib.util.spec_from_file_location("nias_main", path)
    if spec is None or spec.loader is None:
        raise ImportError(f"{path.name} konnte nicht geladen werden")
    module = importlib.util.module_from_spec(spec)
    sys.modules.setdefault("nias_main", module)
    spec.loader.exec_module(module)
    return module


# --------------------------------------------------------------------------
# Unknown register
# --------------------------------------------------------------------------

def write_unknowns(db_path: Path, confirmed: Iterable[dict[str, Any]], *,
                   context_window: float = 0.6) -> dict[str, int]:
    """Write confirmed unknowns into the SQLite register.

    Called only after explicit confirmation. Each sighting carries its measured
    spectrum and the chromatogram around the peak, so the entry can be redrawn
    later without the original ``.D``.
    """
    import gc_register as R

    confirmed = list(confirmed)
    if not confirmed:
        return {"entries": 0, "sightings": 0, "spectra": 0}

    con = R.connect(Path(db_path))
    written = {"entries": 0, "sightings": 0, "spectra": 0}
    try:
        R.create_schema(con)
        with R.writing(con):
            next_id = int(R.meta_get(con, "next_id", 1))
            for item in confirmed:
                ions = R.significant_ions(item.get("spectrum") or [])
                if not ions:
                    continue
                # Identity is the strongest few ions, not the whole significant
                # set: keying on all of them would fragment the register.
                mz_key, ranked = R.identity_key(ions)

                # Match on ``identity_key``, not ``mz_key``: a migrated entry
                # carries its *full* canonical list in ``mz_key``, so keying on
                # the strongest four would never find it and would file a
                # duplicate alongside it. ``identity_key`` is the strongest four
                # for both origins. It is not UNIQUE - a collision means two
                # entries answer to one key, and new sightings land on the first
                # (spec v2.1 SS V.7 step 4); the register tab shows them.
                found = con.execute(
                    "SELECT entry_id FROM entries WHERE identity_key = ? "
                    "ORDER BY entry_id LIMIT 1", (mz_key,)
                ).fetchone()
                if found:
                    entry_id = found["entry_id"]
                else:
                    entry_id = next_id
                    next_id += 1
                    con.execute(
                        "INSERT INTO entries (entry_id, unknown_id, mz_key, "
                        "identity_key, label, "
                        "canonical_mz, ranked_mz, base_peak, status, search_blob) "
                        "VALUES (?,?,?,?,?,?,?,?,'offen',?)",
                        (entry_id, f"UNK-{entry_id:04d}", mz_key, mz_key,
                         f"unknown (m/z {'/'.join(str(m) for m in ranked[:4])})",
                         mz_key, "/".join(str(m) for m in ranked),
                         ranked[0] if ranked else None,
                         f"unk-{entry_id:04d} {item.get('name_raw', '')}".casefold()),
                    )
                    con.executemany(
                        "INSERT OR IGNORE INTO entry_ions(entry_id, mz, rank) "
                        "VALUES (?,?,?)",
                        [(entry_id, mz, rank) for mz, _, rank in ions])
                    written["entries"] += 1

                # A hand-filed unknown may arrive with the analyst's own name,
                # CAS and note (the register window and the spectrum panel's
                # "Im Unbekannten-Register speichern"). Only non-empty values
                # are written: a second sighting arriving from the batch export
                # must never blank out what somebody typed.
                named = {k: str(item.get(k) or "").strip()
                         for k in ("assigned_name", "assigned_cas", "note", "status")}
                named = {k: v for k, v in named.items() if v}
                if named:
                    con.execute(
                        "UPDATE entries SET "
                        + ", ".join(f"{k} = ?" for k in named)
                        + ", search_blob = search_blob || ' ' || ? "
                        "WHERE entry_id = ?",
                        [*named.values(),
                         " ".join(named.values()).casefold(), entry_id])

                rt = item.get("rt")
                cur = con.execute(
                    """INSERT OR IGNORE INTO sightings (
                        entry_id, sample, sample_key, sample_name, report_type,
                        report_type_key, rt, rt_key, conc_kg, area_pct,
                        match_pct, analyst, simulant, source_file, date_text,
                        date_iso, ranked_mz, name_raw)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (entry_id, item.get("sample", ""),
                     str(item.get("sample", "")).casefold(),
                     item.get("sample_name", ""), item.get("report_type", ""),
                     str(item.get("report_type", "")).casefold(),
                     rt, round(rt, 4) if rt is not None else None,
                     item.get("conc_kg"), item.get("area_pct"), item.get("si"),
                     item.get("analyst", ""), item.get("simulant", ""),
                     item.get("source_file", ""), item.get("date", ""),
                     R._iso_date(item.get("date")),
                     "/".join(str(m) for m, _, _ in ions),
                     item.get("name_raw", "")),
                )
                # INSERT OR IGNORE leaves lastrowid pointing at the previous
                # successful insert when the row is deduplicated, so it cannot
                # be trusted here -- look the sighting up instead.
                if cur.rowcount != 1:
                    continue
                found_sighting = con.execute(
                    "SELECT sighting_id FROM sightings WHERE entry_id = ? "
                    "AND sample_key = ? AND report_type_key = ? "
                    "AND rt_key IS ?",
                    (entry_id, str(item.get("sample", "")).casefold(),
                     str(item.get("report_type", "")).casefold(),
                     round(rt, 4) if rt is not None else None),
                ).fetchone()
                if found_sighting is None:
                    continue
                sighting_id = found_sighting["sighting_id"]
                written["sightings"] += 1

                spectrum_id = R.attach_spectrum(
                    con, entry_id, sighting_id, item.get("spectrum") or [],
                    source_path=item.get("source_file", ""),
                    apex_scan=item.get("apex_scan"), bg_scan=item.get("bg_scan"),
                    bounds_rule=item.get("bounds_rule", ""), rt=rt,
                    kind=str(item.get("kind") or "measured"))
                written["spectra"] += 1

                # What the analyst looked at is the entry's spectrum, above.
                # When that was not the raw apex scan -- a background was
                # subtracted, or the deconvoluted component was on screen --
                # the untouched measurement is stored beside it under its own
                # ``kind``, so the register can always show where the decision
                # came from. ``spectra`` counts spectra, so it counts both.
                for extra in item.get("extra_spectra") or ():
                    points = extra.get("spectrum") or []
                    if not points:
                        continue
                    R.attach_spectrum(
                        con, entry_id, sighting_id, points,
                        source_path=item.get("source_file", ""),
                        apex_scan=item.get("apex_scan"),
                        bg_scan=extra.get("bg_scan"),
                        bounds_rule=item.get("bounds_rule", ""), rt=rt,
                        kind=str(extra.get("kind") or "measured"))
                    written["spectra"] += 1

                tic = item.get("tic")
                if tic:
                    R.attach_tic_slice(con, sighting_id, spectrum_id,
                                       tic[0], tic[1], apex_rt=rt)

                # The counts and the two dates are what the register window
                # lists, so they are maintained here rather than recomputed on
                # every read: one UPDATE per written sighting against a table
                # scan per displayed row.
                con.execute(
                    """UPDATE entries SET
                         n_sightings = (SELECT COUNT(*) FROM sightings
                                         WHERE entry_id = ?),
                         n_samples   = (SELECT COUNT(DISTINCT sample_key)
                                          FROM sightings WHERE entry_id = ?),
                         first_seen_iso = (SELECT MIN(date_iso) FROM sightings
                                            WHERE entry_id = ? AND date_iso IS NOT NULL),
                         last_seen_iso  = (SELECT MAX(date_iso) FROM sightings
                                            WHERE entry_id = ? AND date_iso IS NOT NULL)
                       WHERE entry_id = ?""",
                    (entry_id, entry_id, entry_id, entry_id, entry_id))
                # The two shown dates come from the *earliest* and *latest*
                # sighting by ``date_iso``, never from the row that happens to
                # be written last: a sighting filed today for a measurement of
                # last March must not move ``last_seen`` forward. Spec defect 6
                # of §V.7 is the same mistake made in the JSON register, where
                # the ``dd.mm.yyyy`` strings were sorted as text.
                con.execute(
                    """UPDATE entries SET
                         first_seen = COALESCE((SELECT s.date_text FROM sightings s
                                                 WHERE s.entry_id = ?
                                                   AND s.date_iso IS NOT NULL
                                                 ORDER BY s.date_iso, s.sighting_id
                                                 LIMIT 1), first_seen),
                         last_seen  = COALESCE((SELECT s.date_text FROM sightings s
                                                 WHERE s.entry_id = ?
                                                   AND s.date_iso IS NOT NULL
                                                 ORDER BY s.date_iso DESC,
                                                          s.sighting_id DESC
                                                 LIMIT 1), last_seen)
                       WHERE entry_id = ?""",
                    (entry_id, entry_id, entry_id))
                # An unparseable date leaves ``date_iso`` empty; then the text
                # is still better than a blank column.
                stamp = str(item.get("date") or "").strip()
                if stamp:
                    con.execute(
                        "UPDATE entries SET "
                        "first_seen = CASE WHEN first_seen = '' THEN ? ELSE first_seen END, "
                        "last_seen  = CASE WHEN last_seen  = '' THEN ? ELSE last_seen  END "
                        "WHERE entry_id = ?", (stamp, stamp, entry_id))
                sample_text = str(item.get("sample") or "").strip()
                if sample_text:
                    con.execute(
                        "INSERT OR IGNORE INTO entry_samples(entry_id, sample) "
                        "VALUES (?,?)", (entry_id, sample_text))

            R.meta_set(con, "next_id", next_id)
            R.meta_set(con, "updated_at",
                       datetime.now().isoformat(timespec="seconds"))
            R.meta_set(con, "schema_version", R.SCHEMA_VERSION)
    finally:
        con.close()
    return written


# --------------------------------------------------------------------------
# NIAS Doppelbestimmung
# --------------------------------------------------------------------------
#
# AutoLib's make_duplicate_workbook re-reads RESULTS.CSV, so calling it would
# silently discard everything the analyst corrected in the grid. The workbook is
# therefore written from the session, while the *evaluation* -- standard
# detection, blank correction, the duplicate merge -- still comes from AutoLib
# through gc_fid. The sheet layout follows the reference workbook.

#: Header of the "seen before" column (SS VII.5). **Appended**, never inserted:
#: the sheet is addressed positionally by the number-format loops below, by the
#: conditional formats and by whatever the analyst built on top of it, so a
#: column that exists today may not change its letter.
SEEN_BEFORE_HEADER = "Times reported before"

#: Header of the retention-index column (SS VII.9). Unlike the one above this
#: one *does* move columns when ``report_ri`` is on -- see
#: :func:`_result_headers` for why that is a deliberate, opt-in exception.
RI_HEADER = "RI"

NIAS_DUPLICATE_HEADERS = (
    "RT (mean)", "Name", "CAS", "mg/kg (mean)", "Area 1",
    "Concentration 1 [mg/kg]", "Area 2", "Concentration 2 [mg/kg]",
    "Duplicate status", "Identification status", "Review",
    "Relative difference [%]", SEEN_BEFORE_HEADER,
)

#: The v3.0 columns are **appended**, never inserted: the sheet is read
#: positionally by number-format loops here and by whatever the analyst has
#: built on top of it, so an existing column may not change its letter
#: (SS VI.14).
NIAS_DETAIL_HEADERS = (
    "RT [min]", "Name", "CAS", "Quality", "Blank-corrected FID area", "mg/dm2",
    "mg/kg", "Identification status", "Match status", "Review", "FID peak",
    "PBM peak", "MS RT [min]", "Corrected MS RT [min]", "Raw FID area",
    "Blank area", "Blank+ISTD area", "Blank correction", "ISTD protected",
    "Quantification factor",
    "FID von [min]", "FID bis [min]", "Integration", "S/N", "LOQ [mg/kg]",
    "Reinheit [%]", "RI",
)

#: ``Manuell_pruefen`` gains one column, so a reviewer sees *why* a settled
#: duplicate is on the sheet: because somebody moved its integration. In v3.1 it
#: inherits ``Times reported before`` from the result sheet (SS VII.5), again
#: appended last.
NIAS_MANUAL_HEADERS = (
    "RT (mean)", "Name", "CAS", "mg/kg (mean)", "Concentration 1 [mg/kg]",
    "Concentration 2 [mg/kg]", "Relative difference [%]", "Duplicate status",
    "Identification status", "Review", "Integration", SEEN_BEFORE_HEADER,
)

#: The GLP record of every integration action (SS VI.14). One row per action,
#: projected from the ``EditRecord``s the session already holds.
NIAS_INTEGRATION_HEADERS = (
    "Zeitpunkt", "Bestimmung", "Peak", "Aktion", "Grenzen alt", "Grenzen neu",
    "Fläche alt", "Fläche neu", "mg/kg alt", "mg/kg neu", "Basislinie",
    "Benutzer",
)

NIAS_STANDARD_HEADERS = (
    "Name", "Target RT [min]", "Role", "FID peak", "FID RT [min]", "FID area",
    "Concentration [mg/mL]", "Factor", "Deviation from mean [%]", "Status",
    # Appended last (v3.2) so every column letter the formulas use stays put.
    "Area edited", "ISTD",
)

#: Fill for a standard area the analyst set by hand. Same blue as the grid's
#: "edited" colour, so the workbook and the workspace agree visually.
EDITED_FILL = PatternFill("solid", fgColor="D6E4F7")

#: One row per PBM peak, three hits wide -- the layout AutoLib's own
#: ``make_duplicate_workbook`` writes, so the sheet stays recognisable; the peak
#: and retention-time headers are spelled as in ``NIAS_DETAIL_HEADERS``.
NIAS_PBM_HEADERS = (
    "PBM peak", "MS RT [min]", "Area %",
    "Hit 1", "CAS 1", "Q1", "Hit 2", "CAS 2", "Q2", "Hit 3", "CAS 3", "Q3",
)

NIAS_BLANK_HEADERS = (
    "Determination", "FID peak", "RT [min]", "Name", "Raw FID area",
    "Blank area", "Blank+ISTD area", "Corrected FID area", "Correction status",
    "ISTD protected", "Corrected?",
)

#: The columns of "Doppelbestimmung" that belong to the second determination.
#: An analysis with a single determination writes them empty and hides them
#: rather than dropping them, so the sheet has one layout and a consumer that
#: reads by header name never learns a second one (spec v2.1 SS V.6).
SINGLE_DETERMINATION_HIDDEN = ("G", "H", "L")

#: Duplicate status of a row whose analysis is a single determination. AutoLib
#: writes the value (``AutoLib.SINGLE_DETERMINATION_STATUS``); it is repeated
#: here rather than imported because the engine is loaded lazily and by path,
#: and a test asserts the two stay equal.
SINGLE_DETERMINATION_STATUS = "Einzelbestimmung"


#: Colours of the two conditional-formatting rules on ``Doppelbestimmung``.
#: They replace the static fills this writer used until v3.0: a fill cannot
#: react to an area the analyst edits in Excel afterwards, and the whole point
#: of the sheet is that it recalculates (SS VI.22).
BELOW_LIMIT_FILL = PatternFill("solid", fgColor="D9D9D9")
BELOW_LIMIT_FONT = Font(color="7F7F7F")
REVIEW_FILL = PatternFill("solid", fgColor="FFC7CE")
REVIEW_FONT = Font(color="9C0006")

#: Green, the workbook's convention for "this cell is a formula".
FORMULA_FONT = Font(color="008000")

#: Blue, the workbook's convention for "change this value" (Parameter!B).
PARAMETER_FONT = Font(color="0000FF")


def _letter(headers: tuple[str, ...], header: str) -> str:
    """The column letter of ``header`` in a sheet with ``headers``.

    Derived, never hardcoded. This writer's ``Bestimmung_n`` has no
    ``Subtracted area`` column, so its quantification factor sits one column
    left of where AutoLib's does -- and SS VI.14 forbids inserting the missing
    column to make the letters line up, because an existing column may not move.
    """
    return get_column_letter(headers.index(header) + 1)


def _parameter_ref(key: str) -> str:
    """Absolute reference to one editable value on the ``Parameter`` sheet.

    Read out of ``gc_fid.PARAMETER_LAYOUT`` rather than written down, so the
    formulas follow the sheet if a parameter is ever added to it.
    """
    import gc_fid

    for index, (attr, _label, _unit) in enumerate(gc_fid.PARAMETER_LAYOUT):
        if attr == key:
            return f"Parameter!$B${index + 2}"
    raise KeyError(key)


# --------------------------------------------------------------------------
# v3.1: retention index, duplicate difference limit, "seen before"
# --------------------------------------------------------------------------

def _ri_options(session: Optional[M.Session] = None,
                settings: Any = None,
                override: Optional[dict] = None) -> tuple[bool, bool]:
    """``(report_ri, replace_rt)`` for the session's method (SS VII.9).

    The ladder store (``ladders.json``) belongs to the workspace, which resolves
    the two flags for the method actually used and hands them down. This writer
    only *reads* them, and it reads them from wherever the caller happened to
    park them, because a report must never depend on the analyst having opened
    the alkane dialog first. Both default to ``False``, which is the v3.0
    behaviour: a session without RI options exports exactly as it does today.

    ``override`` is the dict ``GCWorkspace.ri_options`` holds; passing it is the
    explicit route and outranks anything parked on the session or the settings.
    """
    if isinstance(override, dict):
        report = bool(override.get("report_ri", False))
        return report, bool(override.get("replace_rt", False)) and report
    for holder in (settings, session):
        if holder is None:
            continue
        options = getattr(holder, "ri_options", None)
        if options is None and isinstance(holder, dict):
            options = holder.get("ri_options")
        if isinstance(options, dict):
            report = bool(options.get("report_ri", False))
            # ``RT durch RI ersetzen`` is meaningless on its own: the dialog
            # forces it off with its parent checkbox and so does this.
            return report, bool(options.get("replace_rt", False)) and report
        if hasattr(holder, "report_ri"):
            report = bool(getattr(holder, "report_ri"))
            return report, bool(getattr(holder, "replace_rt", False)) and report
    return False, False


def _result_headers(base: tuple[str, ...], report_ri: bool,
                    replace_rt: bool) -> tuple[str, ...]:
    """``base`` with the RI column of SS VII.9 applied.

    Three cases, and only the middle one moves a column:

    * ``report_ri = False`` -- ``base`` unchanged, byte for byte;
    * ``report_ri, not replace_rt`` -- ``RI`` inserted directly right of the
      retention time. This is the one place in the workbook where a column
      shifts, and it does so only because the analyst ticked a box; every
      formula, width, number format and conditional format on the two result
      sheets is derived from the header tuple precisely so that it may;
    * ``replace_rt`` -- the first column keeps its position and is headed
      ``RI``. Nothing moves, and the measured retention time survives in the
      detail sheets, which are the raw-data documentation.
    """
    if not report_ri:
        return base
    if replace_rt:
        return (RI_HEADER,) + base[1:]
    return base[:1] + (RI_HEADER,) + base[1:]


def _sources(item: dict[str, Any]) -> list:
    """GCWS-PATCH: every determination's source dict (N-fold replicates)."""
    if item.get("sources") is not None:
        return list(item["sources"])
    return [item.get("source1"), item.get("source2")]


def _merged_ri(item: dict[str, Any]) -> Optional[int]:
    """The retention index of a merged row, as an integer (SS VII.9).

    The mean of both determinations; with one side missing the other is used,
    with neither the cell stays empty. Kovats indices are reported without
    decimals, so the mean is rounded here rather than formatted away in Excel --
    a cell that *looks* like an integer but sorts as 1043.5 is a trap.
    """
    values = [_num((src or {}).get("ri")) for src in _sources(item)]  # GCWS-PATCH: N-fold
    present = [v for v in values if v is not None]
    if not present:
        return None
    return int(round(sum(present) / len(present)))


def _merged_number(item: dict[str, Any], key: str) -> Optional[float]:
    """Mean of ``source1[key]`` and ``source2[key]``, or the side that has one."""
    values = [_num((src or {}).get(key)) for src in _sources(item)]  # GCWS-PATCH: N-fold
    present = [v for v in values if v is not None]
    if not present:
        return None
    return sum(present) / len(present)


def _seen_before(item: dict[str, Any], seen) -> Optional[int]:
    """``Times reported before`` for one merged row, or ``None``.

    ``seen`` is the lookup :func:`_seen_lookup` built. A row already carrying
    ``seen_before`` -- which is what ``gc_duplicate.build`` puts there when the
    workspace passed it the counts -- wins, so the workbook and the grid can
    never disagree about a number the analyst is looking at in both.
    """
    direct = item.get("seen_before")
    if direct is not None:
        return int(direct)
    if seen is None:
        return None
    return seen(item.get("cas"), item.get("name"))


def _exceeds_reldiff(item: dict[str, Any], limit: float,
                     reporting_limit: Optional[float] = None) -> bool:
    """SS VII.4's flag for one merged row, delegated to ``gc_duplicate``.

    The rule lives in ``gc_duplicate`` and is not reimplemented here; this
    wrapper only adapts AutoLib's merged dict into the ``DuplicateRow`` the rule
    takes. The literal fallback exists so that this module stays usable against
    a ``gc_duplicate`` that predates SS VII.4 -- it is the contract's own
    sentence in code and must be deleted, not maintained, once every checkout
    has the rule.

    ``reporting_limit`` is passed through because ``exceeds_reldiff`` enforces
    the ``below_limit`` precedence itself and defaults that limit to 0.01 mg/kg.
    An analyst who *lowers* the reporting limit reports rows between their limit
    and 0.01, and those rows must be flaggable -- the workbook's amber rule has
    no such floor either, and the sheet and the audit payload may not disagree
    (agent RULES flagged exactly this in ``.v31_reports\\rules.md``, SS 4).
    """
    import gc_duplicate

    row = gc_duplicate.DuplicateRow(item=item)
    rule = getattr(gc_duplicate, "exceeds_reldiff", None)
    if rule is not None:
        if reporting_limit is None:
            return bool(rule(row, limit))
        try:
            return bool(rule(row, limit, reporting_limit=reporting_limit))
        except TypeError:                 # pragma: no cover - pre-RULES module
            return bool(rule(row, limit))
    c1, c2, reldiff = row.c1, row.c2, row.reldiff
    return c1 is not None and c2 is not None and reldiff is not None \
        and reldiff > limit


def _pairs_above_limit(combined: list[dict[str, Any]], limit: float,
                       reporting_limit: float) -> int:
    """How many reported pairs disagree by more than ``limit`` percent.

    ``below_limit`` keeps its precedence exactly as it does in the workbook's
    conditional formats: a result under the reporting limit is not reported and
    is therefore not flagged either (SS VII.4).
    """
    import gc_duplicate

    count = 0
    for item in combined:
        row = gc_duplicate.DuplicateRow(item=item)
        if gc_duplicate.below_limit(row, reporting_limit):
            continue
        if _exceeds_reldiff(item, limit, reporting_limit):
            count += 1
    return count


def duplicate_limit(settings: Any) -> float:
    """The duplicate difference limit in percent, from the settings (SS VII.4).

    A normal parameter since v3.1, so it comes off ``AutoLib.Settings`` like
    every other one; ``gc_duplicate.DEFAULT_MAX_RELDIFF`` is the fallback for a
    settings object written before the parameter existed.
    """
    import gc_duplicate

    value = _num(getattr(settings, "duplicate_max_reldiff", None))
    return gc_duplicate.DEFAULT_MAX_RELDIFF if value is None else value


def duplicate_audit_fields(combined: list[dict[str, Any]], settings: Any
                           ) -> dict[str, Any]:
    """The two SS VII.4 fields for the report's audit payload.

    ``{"Duplicate difference limit": 30.0, "Pairs above limit": 3}``. Returned
    rather than written: the NIAS report's payload is assembled in the main
    script, the fingerprint one in :func:`write_fingerprint_workbook`, and both
    want the same two numbers computed the same way.
    """
    limit = duplicate_limit(settings)
    reporting_limit = float(getattr(settings, "reporting_limit", 0.01) or 0.01)
    return {
        "Duplicate difference limit": limit,
        "Pairs above limit": _pairs_above_limit(combined, limit,
                                                reporting_limit),
    }


def _single_determination_rows(sample: M.Sample) -> list[dict[str, Any]]:
    """The Doppelbestimmung rows of an analysis with one determination.

    AutoLib owns the row shape, so the rows are built there rather than here.
    ``gc_fid.combine`` cannot be used because it needs two samples; what it
    would do to each row -- turn it into an engine peak dict -- is all that is
    needed, and that is what ``_as_engine_peak`` does.
    """
    import gc_fid

    engine = gc_fid.engine()
    to_peak = getattr(gc_fid, "_as_engine_peak")
    return engine.combine_determinations([to_peak(r) for r in gc_fid.report_rows(sample)],
                                         None, 0.0)


def run_nias_duplicate(session: M.Session, output_path: Path, settings, *,
                       blank_path: Optional[Path] = None,
                       blank_istd_path: Optional[Path] = None,
                       seen_counts: Optional[dict] = None,
                       ri_options: Optional[dict] = None,
                       combined: Optional[list] = None,
                       save=None) -> Path:
    """Write the NIAS workbook from the edited session.

    GCWS-PATCH: ``save(workbook, path)`` replaces the plain ``workbook.save`` (the caller adds
    the report metadata before the one save instead of loading the file again).

    One determination is a complete analysis (spec v2.1 SS V.6). The layout does
    not change for it: ``Area 2``, ``Concentration 2 [mg/kg]`` and
    ``Relative difference [%]`` are written empty and hidden, ``Duplicate
    status`` reads ``Einzelbestimmung`` and ``mg/kg (mean)`` is the one measured
    value. Only ``Bestimmung_2``, ``Standards_2`` and ``PBM_2`` are left out,
    because an empty detail sheet documents nothing.

    ``seen_counts`` is the ``gc_seen.counts`` mapping the workspace already
    fetched once per rebuild (SS VII.5). Keyword-only and optional, so every
    existing call site keeps working; without it the ``Times reported before``
    column is written empty rather than the register being opened here -- this
    module owns no database handle, for the same reason ``gc_duplicate`` does
    not.

    ``ri_options`` is ``GCWorkspace.ri_options`` -- ``{"report_ri": bool,
    "replace_rt": bool}`` for the method actually used (SS VII.9). Also
    keyword-only and optional: omitted, the flags are looked for on the settings
    and then on the session, and a caller that has none exports exactly what
    v3.0 exported.
    """
    import gc_fid

    labels = session.labels
    if not labels:
        raise ValueError("Es ist keine Bestimmung geladen.")
    # GCWS-PATCH: every determination (N-fold replicates) and precomputed
    # combined rows; two determinations without ``combined`` behave as before.
    samples = [session.samples[label] for label in (labels if combined is not None else labels[:2])]
    single = len(samples) == 1
    if combined is None:
        combined = (_single_determination_rows(samples[0]) if single
                    else gc_fid.combine(samples[0], samples[1]))

    wb = Workbook()
    wb.properties.title = ("NIAS Einzelbestimmung" if single
                           else "NIAS Doppelbestimmung" if len(samples) == 2
                           else f"NIAS {len(samples)}-fach-Bestimmung")
    # The sheets are created in their final order first and filled afterwards:
    # every formula on "Doppelbestimmung" and "Manuell_pruefen" addresses a row
    # of a detail sheet, so those have to be written before the two result
    # sheets can be -- but the tab order the analyst sees must not change.
    ws = wb.active
    ws.title = "Doppelbestimmung"
    wsr = wb.create_sheet("Manuell_pruefen")
    _write_parameter_sheet(wb.create_sheet("Parameter"), settings, session,
                           blank_path, blank_istd_path)
    detail_rows: dict[int, dict[int, int]] = {}
    for n, sample in enumerate(samples, 1):
        detail_rows[n] = _write_detail_sheet(
            wb.create_sheet(f"Bestimmung_{n}"), sample, n)
        _write_standards_sheet(wb.create_sheet(f"Standards_{n}"), sample)
        _write_pbm_sheet(wb.create_sheet(f"PBM_{n}"), sample)

    limit = float(getattr(settings, "reporting_limit", 0.01) or 0.01)
    report_ri, replace_rt = _ri_options(session, settings, ri_options)
    seen = _seen_lookup(seen_counts)
    _write_duplicate_sheet(ws, combined, detail_rows, single=single, limit=limit,
                           report_ri=report_ri, replace_rt=replace_rt,
                           seen=seen, reldiff_limit=duplicate_limit(settings),
                           n_determinations=len(samples))
    _write_manual_sheet(wsr, combined, detail_rows,
                        report_ri=report_ri, replace_rt=replace_rt, seen=seen)
    _write_blank_sheet(wb.create_sheet("Blankkorrektur"), samples)
    _write_integration_sheet(wb.create_sheet("Integration"), session)
    _write_audit_sheet(wb, session)

    output_path = Path(output_path)
    (save or _plain_save)(wb, output_path)      # GCWS-PATCH: one save with the caller's additions
    return output_path


def _plain_save(wb, path: Path) -> None:
    """GCWS-PATCH: the writers' default save."""
    wb.save(path)


# --------------------------------------------------------------------------
# The batch workbook, filled from the workspace (v3.2)
# --------------------------------------------------------------------------
#
# ``Doppelbestimmung Batch auswerten`` writes ``<Probe>_Doppelbestimmung.xlsx``
# through ``AutoLib.make_duplicate_workbook``. That function re-reads
# RESULTS.CSV, so it cannot carry what the analyst changed in the workspace.
# ``write_batch_workbook`` writes the *same layout* from the session instead --
# Parameter rows at their fixed positions, ``Bestimmung_n`` with the ISTD
# column V, ``ISTD_n`` and the linked result sheets -- so someone without the
# workspace can follow the whole evaluation in Excel, and NIAS Report can
# process the file again like any batch workbook.

#: Detail columns of ``AutoLib.make_duplicate_workbook``, in its order. Two
#: columns are appended after the ISTD column V so no formula letter moves.
BATCH_DETAIL_HEADERS = (
    "RT [min]", "Name", "CAS", "Quality", "Blank-corrected FID area",
    "mg/dm2", "mg/kg", "Identification status", "Match status", "Review",
    "FID peak", "PBM peak", "MS RT [min]", "Corrected MS RT [min]",
    "Raw FID area", "Blank area", "Blank+ISTD area", "Subtracted area",
    "Blank correction", "ISTD protected", "Quantification factor", "ISTD",
    "Integration", "Manuell geändert",
)
BATCH_RESULT_HEADERS = (
    "RT (mean)", "Name", "CAS", "mg/kg (mean)", "Concentration 1 [mg/kg]",
    "Concentration 2 [mg/kg]", "Relative difference [%]", "Duplicate status",
    "Identification status", "Review",
)
#: AutoLib's fixed Parameter positions (``Parameter!$B$9`` and following).
BATCH_PARAM_REF = {"cell_area": "Parameter!$B$9", "coverage": "Parameter!$B$10",
                   "ov_ratio": "Parameter!$B$11", "is_amount": "Parameter!$B$12"}
#: AutoLib keeps the mean in row 7 and the factor in row 8 of ``ISTD_n`` for
#: its four standards; a longer list moves both down.
BATCH_ISTD_MEAN_ROW = 7
BATCH_ISTD_FACTOR_ROW = 8
BATCH_ISTD_COLUMN = "V"
#: Columns of ``Doppelbestimmung``/``Manuell_pruefen`` hidden for one
#: determination (``AutoLib.SINGLE_DETERMINATION_HIDDEN_COLUMNS``).
BATCH_SINGLE_HIDDEN = ("F", "G")
#: Workspace field -> column of ``Bestimmung_n``, for the "edited" colour.
_BATCH_EDITED_COLUMNS = {"rt": 1, "name": 2, "cas": 3, "si": 4, "area": 5,
                         "raw_area": 15, "blank_area": 16}


def batch_workbook_name(stem: str) -> str:
    """``<stem>_Doppelbestimmung.xlsx``, the name the batch gives the file."""
    import gc_fid

    safe = getattr(gc_fid.engine(), "_safe_output_stem", None)
    clean = safe(stem) if callable(safe) else re.sub(r'[<>:"/\\|?*]+', "_", stem)
    return f"{clean or 'Probe'}_Doppelbestimmung.xlsx"


def _batch_istd_rows(n_istds: int) -> tuple[int, int, int]:
    """``(mean row, factor row, mode row)`` of ``ISTD_n`` for ``n_istds``."""
    mean_row = max(BATCH_ISTD_MEAN_ROW, n_istds + 3)
    return mean_row, mean_row + 1, mean_row + 2


def write_batch_workbook(session: M.Session, output_path: Path, settings, *,
                         blank_path: Optional[Path] = None,
                         blank_istd_path: Optional[Path] = None,
                         sample_name: str = "",
                         save=None) -> Path:
    """The ``Doppelbestimmung Batch auswerten`` workbook, from the edited session.

    GCWS-PATCH: ``save(workbook, path)`` as in :func:`run_nias_duplicate`.

    Same sheets, same positions and the same formula chain as
    ``AutoLib.make_duplicate_workbook``: ``Bestimmung_n!V`` marks the ISTD
    rows, ``ISTD_n`` looks their areas up, forms the mean and the factor, and
    every concentration on ``Bestimmung_n``/``Doppelbestimmung`` follows from
    it. What differs is the input: the rows, areas, names and ISTDs are the
    workspace's, including every manual change -- and ``ISTD_n`` lists the
    workspace's ISTDs with their typed concentrations, the ``Quantify`` flag
    and the mean/reference mode instead of the fixed IS1..IS4. Edited cells are
    coloured, and ``Integration`` and ``Manuelle Änderungen`` document the
    changes.
    """
    import gc_fid

    labels = session.labels
    if not labels:
        raise ValueError("Es ist keine Bestimmung geladen.")
    samples = [session.samples[label] for label in labels[:2]]
    single = len(samples) == 1
    combined = (_single_determination_rows(samples[0]) if single
                else gc_fid.combine(samples[0], samples[1]))
    defs = list(getattr(session, "istd_defs", None)
                or getattr(samples[0], "istd_defs", None) or [])
    options = (getattr(session, "istd_options", None)
               or getattr(samples[0], "istd_options", None)
               or gc_fid.DEFAULT_ISTD_OPTIONS)
    ref = BATCH_PARAM_REF

    wb = Workbook()
    wb.properties.title = sample_name or "NIAS Doppelbestimmung"
    ws = wb.active
    ws.title = "Doppelbestimmung"
    wsr = wb.create_sheet("Manuell_pruefen")
    wsp = wb.create_sheet("Parameter")

    delays = [s.meta.get("delay", "") for s in samples]
    istd_note = "→ ISTD_n, Spalte B (im GC Workspace festgelegt)"
    parameters = [
        ("Sample name", sample_name, "used for result file and workbook identification"),
        ("Reporting limit", getattr(settings, "reporting_limit", 0.01),
         "mg/kg; applied at workbook generation"),
        ("Solvent end", settings.solvent_end, "min"),
        ("Quality threshold", settings.quality_limit, ""),
        ("Duplicate RT tolerance", settings.rt_tolerance, "min"),
        ("FID-MS delay determination 1", delays[0], "min"),
        ("FID-MS delay determination 2", "" if single else delays[1], "min"),
        ("Cell area", settings.cell_area_dm2, "dm2"),
        ("Coverage", settings.coverage, ""),
        ("O/V ratio", settings.ov_ratio, ""),
        ("Internal standard amount", settings.is_amount, ""),
        ("FC17 concentration", istd_note, "mg/mL"),
        ("BBP-d4 concentration", istd_note, "mg/mL"),
        ("DnNP-d4 concentration", istd_note, "mg/mL"),
        ("Blank RT tolerance", settings.blank_rt_tolerance,
         "min; one-to-one FID apex matching"),
        ("Blank file", str(blank_path) if blank_path else "not used", ""),
        ("Blank+ISTD file", str(blank_istd_path) if blank_istd_path else "not used", ""),
        ("Blank FID-MS delay", "", "min"),
        ("Blank+ISTD FID-MS delay", "", "min"),
        ("Note", "Im GC Workspace bearbeitete Auswertung. Die ISTD-Zeilen sind in "
                 "Spalte 'ISTD' von 'Bestimmung_1'/'Bestimmung_2' markiert, ihre "
                 "Konzentrationen stehen in 'ISTD_1'/'ISTD_2'. Blaue Zellen sind "
                 "änderbar; Konzentrationen, Mittelwerte und relative Differenzen "
                 "rechnen sich neu. Hellblau hinterlegt: im Workspace geändert.", ""),
    ]
    wsp.append(["Parameter", "Value", "Unit / note"])
    for row in parameters:
        wsp.append(list(row))

    detail_rows: dict[int, dict[int, int]] = {}
    istd_sheets: list[str] = []
    for number, sample in enumerate(samples, 1):
        wd = wb.create_sheet(f"Bestimmung_{number}")
        istd_name = f"ISTD_{number}"
        wi = wb.create_sheet(istd_name)
        istd_sheets.append(istd_name)
        mean_row, factor_row, mode_row = _batch_istd_rows(len(defs))
        factor_cell = f"{istd_name}!$B${factor_row}"
        code_of = {s.get("row_id"): s.get("code") for s in
                   getattr(sample, "standards", []) or []
                   if s.get("row_id") is not None and s.get("code")}

        wd.append(list(BATCH_DETAIL_HEADERS))
        rows: dict[int, int] = {}
        for row in gc_fid.report_rows(sample):
            d = row.derived
            r = wd.max_row + 1
            rows[row.row_id] = r
            wd.append([
                row.rt, row.name, M.display_cas(row.cas), row.si, row.area,
                f'=IF(U{r}="","",E{r}*U{r})',
                f'=IF(F{r}="","",F{r}*{ref["ov_ratio"]})',
                d.get("id_status", ""), d.get("match_status", ""),
                d.get("review", ""), d.get("fid_peak"), d.get("pbm_peak"),
                d.get("ms_rt"), d.get("corrected_ms_rt"), d.get("raw_area"),
                d.get("blank_area"), d.get("blank_istd_area"),
                d.get("subtracted_area"), d.get("blank_status", ""),
                d.get("protected_standard", "") or "Nein",
                f'=IF({factor_cell}="","",{factor_cell})',
                code_of.get(row.row_id),
                row.integration_label,
                ", ".join(sorted(row.edited)),
            ])
            for field in row.edited:
                column = _BATCH_EDITED_COLUMNS.get(field)
                if column:
                    wd.cell(r, column).fill = EDITED_FILL
            if row.integration_origin != M.ORIGIN_CHEMSTATION:
                wd.cell(r, 23).fill = EDITED_FILL
        detail_rows[number] = rows
        last_row = max(wd.max_row, 2)
        _write_batch_istd_sheet(wi, sample, defs, options, number, last_row)

        if wd.max_row > 1:
            validation = DataValidation(
                type="list", formula1=f"={istd_name}!$A$2:$A${max(len(defs), 1) + 1}",
                allow_blank=True)
            validation.error = "Bitte einen ISTD aus der Liste in ISTD_n wählen."
            validation.errorTitle = "ISTD"
            validation.prompt = "Diese Zeile als ISTD markieren"
            validation.promptTitle = "ISTD"
            wd.add_data_validation(validation)
            validation.add(f"{BATCH_ISTD_COLUMN}2:{BATCH_ISTD_COLUMN}{wd.max_row}")
        _write_pbm_sheet(wb.create_sheet(f"PBM_{number}"), sample)

    _write_batch_blank_sheet(wb.create_sheet("Blankkorrektur"), samples)

    settled = ("Valid duplicate", SINGLE_DETERMINATION_STATUS)
    ws.append(list(BATCH_RESULT_HEADERS))
    wsr.append(list(BATCH_RESULT_HEADERS))
    for item in combined:
        s1, s2 = item.get("source1"), item.get("source2")
        r1 = detail_rows[1].get(s1.get("row_id")) if s1 else None
        r2 = (detail_rows.get(2, {}).get(s2.get("row_id"))
              if s2 and not single else None)
        c1 = f"=Bestimmung_1!G{r1}" if r1 else None
        c2 = f"=Bestimmung_2!G{r2}" if r2 else None

        def result_row(target, c1=c1, c2=c2, item=item):
            r = target.max_row + 1
            if single:
                mean, rel = (f"=E{r}" if c1 else None), None
            else:
                both = bool(c1 and c2)
                mean = f'=IFERROR(AVERAGE(E{r}:F{r}),"")' if both else None
                rel = f'=IFERROR(ABS(E{r}-F{r})/D{r}*100,"")' if both else None
            return [item.get("rt"), item.get("name"), item.get("cas"), mean, c1, c2,
                    rel, item.get("status", ""), item.get("id_status", ""),
                    item.get("review", "")]

        ws.append(result_row(ws))
        if (not str(item.get("status", "")).startswith(settled)
                or item.get("id_status") != "Accepted" or item.get("review")):
            wsr.append(result_row(wsr))
            for cell in ws[ws.max_row]:
                cell.fill = REVIEW_FILL
                cell.font = REVIEW_FONT

    _write_integration_sheet(wb.create_sheet("Integration"), session)
    _write_audit_sheet(wb, session)
    _style_batch_workbook(wb, samples, istd_sheets, single)
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:                        # pragma: no cover - old openpyxl
        pass
    output_path = Path(output_path)
    (save or _plain_save)(wb, output_path)      # GCWS-PATCH: one save with the caller's additions
    return output_path


def _write_batch_istd_sheet(wi, sample: M.Sample, defs: list, options: dict,
                            number: int, last_row: int) -> None:
    """``ISTD_n`` for the workspace's ISTDs, with the workspace's mode.

    One row per ISTD: label, typed concentration (blue), the FID area looked
    up from the row marked with the label in ``Bestimmung_n!V`` -- or the area
    typed in the workspace when the ISTD has no peak there -- a note and
    ``Quantify``. Below: the mean (or the reference ISTD), the factor and the
    ``Quantification mode`` row NIAS Report reads.
    """
    import gc_fid

    ref = BATCH_PARAM_REF
    wi.append(["Name", "c [mg/mL]", "FID area", "Note", "Quantify", "Substance"])
    label_range = f"Bestimmung_{number}!${BATCH_ISTD_COLUMN}$2:${BATCH_ISTD_COLUMN}${last_row}"
    area_range = f"Bestimmung_{number}!$E$2:$E${last_row}"
    binding = {str(s.get("code")): s for s in getattr(sample, "standards", []) or []
               if s.get("code")}
    for d in defs:
        r = wi.max_row + 1
        std = binding.get(str(d["code"]), {})
        manual = std.get("manual_area")
        if std.get("row_id") is not None:
            note = (f"FID peak {std.get('fid_peak')} bei RT "
                    f"{_minutes(std.get('fid_rt'))} min")
            area: Any = f'=IFERROR(INDEX({area_range},MATCH(A{r},{label_range},0)),"")'
        elif manual is not None:
            note, area = "Fläche im GC Workspace eingetragen", manual
        else:
            note = "nicht gefunden"
            area = f'=IFERROR(INDEX({area_range},MATCH(A{r},{label_range},0)),"")'
        if d.get("quantify", True) and d.get("concentration") is None:
            note += "; Konzentration fehlt"
        if not d.get("quantify", True):
            note += "; nicht in der Berechnung (QC)"
        wi.append([d["code"], d.get("concentration"), area, note,
                   "Ja" if d.get("quantify", True) else "Nein", d.get("name", "")])

    first, last = 2, max(len(defs), 1) + 1
    mean_row, factor_row, mode_row = _batch_istd_rows(len(defs))
    rng = lambda col: f"{col}{first}:{col}{last}"
    if options.get("use_mean_area", True):
        wi.cell(mean_row, 1, "Mean (Quantify = Ja)")
        wi.cell(mean_row, 2, f'=IFERROR(AVERAGEIFS({rng("B")},{rng("C")},">0",'
                             f'{rng("E")},"Ja",{rng("B")},">0"),"")')
        wi.cell(mean_row, 3, f'=IFERROR(AVERAGEIFS({rng("C")},{rng("C")},">0",'
                             f'{rng("E")},"Ja",{rng("B")},">0"),"")')
        wi.cell(mean_row, 4, "Mittelwert über die ISTDs mit Konzentration und Fläche")
        mode, code = "mean ISTD area", ""
    else:
        # The ISTD the workspace actually quantified with in this determination.
        code = next((s.get("code") for s in getattr(sample, "standards", []) or []
                     if s.get("status") == gc_fid.STATUS_REFERENCE), None)
        code = code or gc_fid.reference_code(defs, options)
        row = next((first + i for i, d in enumerate(defs) if d["code"] == code), first)
        wi.cell(mean_row, 1, f"Reference {code}")
        wi.cell(mean_row, 2, f"=B{row}")
        wi.cell(mean_row, 3, f"=C{row}")
        wi.cell(mean_row, 4, "Ein Referenz-ISTD quantifiziert alle Peaks")
        mode = "reference"
    wi.cell(factor_row, 1, "Quantification factor")
    wi.cell(factor_row, 2, (f'=IFERROR(B{mean_row}*{ref["is_amount"]}/1000/C{mean_row}/'
                            f'{ref["cell_area"]}/{ref["coverage"]},"")'))
    wi.cell(factor_row, 4, f"mg/dm2 je Flächeneinheit; gilt für jede Zeile von "
                           f"Bestimmung_{number}")
    wi.cell(mode_row, 1, "Quantification mode")
    wi.cell(mode_row, 2, mode)
    wi.cell(mode_row, 3, code)


def _write_batch_blank_sheet(ws, samples: list) -> None:
    """AutoLib's ``Blankkorrektur`` sheet, from the workspace rows."""
    import gc_fid

    ws.append(["Determination", "FID peak", "RT [min]", "Name", "Raw FID area",
               "Blank area", "Blank+ISTD area", "Subtracted area",
               "Corrected FID area", "Correction status", "ISTD protected",
               "Corrected?"])
    for number, sample in enumerate(samples, 1):
        for row in gc_fid.report_rows(sample):
            d = row.derived
            subtracted = d.get("subtracted_area")
            ws.append([number, d.get("fid_peak"), row.rt, row.name,
                       d.get("raw_area"), d.get("blank_area"),
                       d.get("blank_istd_area"), subtracted, row.area,
                       d.get("blank_status", ""),
                       d.get("protected_standard", "") or "Nein",
                       "Ja" if subtracted else "Nein"])


def _style_batch_workbook(wb, samples: list, istd_sheets: list[str],
                          single: bool) -> None:
    """AutoLib's look: blue header, green formulas, blue inputs, tables."""
    detail_sheets = {f"Bestimmung_{n}" for n in range(1, len(samples) + 1)}
    for sheet in wb.worksheets:
        if sheet.title in ("Integration", AUDIT_SHEET):
            continue                          # already styled by their writers
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
        for col in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(r, col).value or "")
                      for r in range(1, min(sheet.max_row, 200) + 1)]
            sheet.column_dimensions[get_column_letter(col)].width = min(
                max(max(map(len, values)) + 2, 10), 45)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.font = FORMULA_FONT
    wsp = wb["Parameter"]
    for r in range(2, 16):
        if not str(wsp.cell(r, 2).value or "").startswith("→"):
            wsp.cell(r, 2).font = PARAMETER_FONT
    for name in istd_sheets:
        wi = wb[name]
        for row in wi.iter_rows(min_row=2):
            if row[4].value in ("Ja", "Nein"):
                for cell in (row[0], row[1], row[4]):
                    cell.font = PARAMETER_FONT
                if not isinstance(row[2].value, str):
                    row[2].font = PARAMETER_FONT
            row[2].number_format = "0"
        wi.column_dimensions["A"].width = 22
        wi.column_dimensions["D"].width = 50
        wi.column_dimensions["F"].width = 34
    for name in detail_sheets:
        wd = wb[name]
        wd.column_dimensions[BATCH_ISTD_COLUMN].width = 12
        for cell in wd[BATCH_ISTD_COLUMN][1:]:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
        for col in ("F", "G"):
            for cell in wd[col][1:]:
                cell.number_format = "0.000000"
        for cell in wd["U"][1:]:
            cell.number_format = "0.000E+00"
        if wd.max_row > 1:
            wd.auto_filter.ref = f"A1:{get_column_letter(wd.max_column)}{wd.max_row}"
    for sheet in (wb["Doppelbestimmung"], wb["Manuell_pruefen"]):
        for cell in sheet["A"][1:]:
            cell.number_format = "0.0000"
        for col in ("D", "E", "F"):
            for cell in sheet[col][1:]:
                cell.number_format = "0.000000"
        for cell in sheet["G"][1:]:
            cell.number_format = "0.0"
        if sheet.max_row > 1:
            sheet.auto_filter.ref = (f"A1:{get_column_letter(sheet.max_column)}"
                                     f"{sheet.max_row}")
        if single:
            for col in BATCH_SINGLE_HIDDEN:
                sheet.column_dimensions[col].hidden = True
    for row in wb["Manuell_pruefen"].iter_rows(min_row=2):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="FCE4D6")


def _style_header(ws, width_map: dict[int, int], autofilter: bool = True) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    if autofilter and ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for index, width in width_map.items():
        ws.column_dimensions[get_column_letter(index)].width = width


#: Column widths of the two result sheets, by header rather than by index, so
#: that the RI column of SS VII.9 cannot shift a width onto the wrong column.
_RESULT_WIDTHS = {
    "RT (mean)": 12, RI_HEADER: 9, "Name": 44, "CAS": 22, "mg/kg (mean)": 14,
    "Area 1": 14, "Concentration 1 [mg/kg]": 20, "Area 2": 14,
    "Concentration 2 [mg/kg]": 20, "Duplicate status": 32,
    "Identification status": 22, "Review": 44, "Relative difference [%]": 18,
    "Integration": 18, SEEN_BEFORE_HEADER: 18,
}


def _widths(headers: tuple[str, ...],
            table: Optional[dict[str, int]] = None) -> dict[int, int]:
    """``{column index: width}`` for ``headers``, derived from the names.

    ``table`` defaults to the two result sheets' widths; the fingerprint sheet
    passes its own (:data:`_FINGERPRINT_WIDTHS`), because ``RT (min)`` and
    ``CAS`` are narrower there than on the Doppelbestimmung sheet.
    """
    lookup = _RESULT_WIDTHS if table is None else table
    return {index: lookup[header]
            for index, header in enumerate(headers, 1)
            if header in lookup}


def _write_duplicate_sheet(ws, combined: list[dict[str, Any]],
                           detail_rows: dict[int, dict[int, int]],
                           single: bool = False,
                           limit: float = 0.01,
                           report_ri: bool = False,
                           replace_rt: bool = False,
                           seen=None,
                           reldiff_limit: Optional[float] = None,
                           n_determinations: int = 2) -> None:
    """The result sheet, linked rather than snapshotted (SS VI.22).

    GCWS-PATCH: with three or more determinations the extra ``Area k`` /
    ``Concentration k`` pairs and ``SD`` are appended after the last column (no
    existing letter moves), and ``mg/kg (mean)`` / ``Relative difference`` are
    written as values -- the mean of all determinations and the RSD --
    because the report reads column D and an AVERAGE over two columns would be
    wrong. The per-determination concentrations stay linked formulas.

    Every number here is a formula over ``Bestimmung_n`` and ``Parameter``:

    ``Area n``  ``=Bestimmung_n!E{d}``            the blank-corrected FID area
    ``Conc n``  ``=Area*Bestimmung_n!{F}{d}*O/V`` area x factor x O/V ratio
    ``mean``    ``=IF(COUNT(F,H)=0,"",AVERAGE(F,H))``
    ``rel.``    ``=IFERROR(ABS(F-H)/AVERAGE(F,H),"")``

    so an area corrected in Excel after the export recalculates the
    concentration, the mean, the relative difference *and* the row colour --
    the same chain the workspace runs in Python. A side the merge did not find
    leaves its two cells empty; the relative difference is written only when
    both sides exist, because ``ABS(F-H)`` would read an empty cell as zero and
    report a settled artefact as 100 % apart.

    Every column letter in every formula is derived from the header tuple, not
    written down: with ``report_ri`` on and ``replace_rt`` off the RI column of
    SS VII.9 sits second and pushes the whole sheet one letter right. With the
    RI options off the tuple is ``NIAS_DUPLICATE_HEADERS`` and every letter is
    the one v3.0 wrote.
    """
    headers = _result_headers(NIAS_DUPLICATE_HEADERS, report_ri, replace_rt)
    extra = list(range(3, n_determinations + 1)) if n_determinations > 2 else []
    if extra:
        headers = tuple(headers) + tuple(h for k in extra for h in (f"Area {k}", f"Concentration {k} [mg/kg]")) \
            + ("SD [mg/kg]",)
    area_letter = _letter(NIAS_DETAIL_HEADERS, "Blank-corrected FID area")
    factor_letter = _letter(NIAS_DETAIL_HEADERS, "Quantification factor")
    ov_ref = _parameter_ref("ov_ratio")

    col = {header: _letter(headers, header) for header in headers}
    a1, c1 = col["Area 1"], col["Concentration 1 [mg/kg]"]
    a2, c2 = col["Area 2"], col["Concentration 2 [mg/kg]"]

    ws.append(list(headers))
    pairs = [(1, a1, c1), (2, a2, c2)] + [(k, col[f"Area {k}"], col[f"Concentration {k} [mg/kg]"])
                                          for k in extra]
    for item in combined:
        r = ws.max_row + 1
        cells: dict[str, Any] = {}
        sources = _sources(item)
        for n, column, conc_column in pairs:
            source = (sources[n - 1] if n - 1 < len(sources) else None) or {}
            detail = detail_rows.get(n, {}).get(source.get("row_id"))
            if not source or detail is None:
                continue
            sheet = f"Bestimmung_{n}"
            cells[column] = f"={sheet}!{area_letter}{detail}"
            cells[conc_column] = (f"={column}{r}*{sheet}!{factor_letter}{detail}"
                                  f"*{ov_ref}")
        # SS V.6: the mean of one determination *is* that determination.
        # AVERAGE over the blank second column would give the same number, but
        # a reference says so, and it keeps following column F if the
        # concentration there is ever overridden by hand.
        if not cells:
            mean = None
        elif single:
            mean = f"={c1}{r}" if c1 in cells else None
        else:
            mean = f'=IF(COUNT({c1}{r},{c2}{r})=0,"",AVERAGE({c1}{r},{c2}{r}))'
        reldiff = (f'=IFERROR(ABS({c1}{r}-{c2}{r})/AVERAGE({c1}{r},{c2}{r}),"")'
                   if c1 in cells and c2 in cells else None)
        if extra:
            mean = _num(item.get("mean"))
            rsd = _num(item.get("rsd"))
            reldiff = rsd / 100.0 if rsd is not None else None
        ri = _merged_ri(item)
        values = {
            "RT (mean)": item.get("rt"), "Name": item.get("name", ""),
            "CAS": item.get("cas", ""), "mg/kg (mean)": mean,
            "Area 1": cells.get(a1), "Concentration 1 [mg/kg]": cells.get(c1),
            "Area 2": cells.get(a2), "Concentration 2 [mg/kg]": cells.get(c2),
            "Duplicate status": item.get("status", ""),
            "Identification status": item.get("id_status", ""),
            "Review": item.get("review", ""),
            "Relative difference [%]": reldiff,
            SEEN_BEFORE_HEADER: _seen_before(item, seen),
            # SS VII.9 case 2 needs no special case: ``_result_headers`` has
            # already renamed the first column, so the index simply lands where
            # the retention time used to be. The measured RT is not lost -- it
            # stays in the detail sheets, the raw-data documentation.
            RI_HEADER: ri,
            "SD [mg/kg]": _num(item.get("sd")),
        }
        for k in extra:
            values[f"Area {k}"] = cells.get(col[f"Area {k}"])
            values[f"Concentration {k} [mg/kg]"] = cells.get(col[f"Concentration {k} [mg/kg]"])
        ws.append([values.get(header) for header in headers])
    _style_header(ws, _widths(headers))
    if single:
        # Written, sized and headed like always -- only not shown, because for
        # one determination they can never hold anything (SS V.6).
        for header in ("Area 2", "Concentration 2 [mg/kg]",
                       "Relative difference [%]"):
            ws.column_dimensions[col[header]].hidden = True
    formats = {
        # Four decimals, not three: a merged RT is the mean of two and lands on
        # a half-digit (11.3815), which three would round away.
        "RT (mean)": "0.0000",
        "mg/kg (mean)": "0.000000",
        "Concentration 1 [mg/kg]": "0.000000",
        "Concentration 2 [mg/kg]": "0.000000",
        "Area 1": "#,##0", "Area 2": "#,##0", "SD [mg/kg]": "0.000000",
        **{f"Area {k}": "#,##0" for k in extra},
        **{f"Concentration {k} [mg/kg]": "0.000000" for k in extra},
        # A fraction, formatted as a percentage -- the reference workbook's own
        # convention for this column, and the reason it carries no "*100".
        "Relative difference [%]": "0.0%",
        # Both integers: a Kovats index is reported without decimals (SS VII.9)
        # and a report count cannot be fractional (SS VII.5).
        RI_HEADER: "0", SEEN_BEFORE_HEADER: "0",
    }
    index_format = {index: formats[header]
                    for index, header in enumerate(headers)
                    if header in formats}
    for row in ws.iter_rows(min_row=2):
        for index, fmt in index_format.items():
            row[index].number_format = fmt
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.font = FORMULA_FONT
    _add_duplicate_rules(ws, limit, headers=headers,
                         reldiff_limit=reldiff_limit)


def _add_duplicate_rules(ws, limit: float, headers: tuple[str, ...] = (),
                         reldiff_limit: Optional[float] = None) -> None:
    """The three rules that colour the sheet, in ``gc_duplicate``'s wording.

    Conditional formatting rather than fills, so the colour follows an edited
    area. Grey is ``stopIfTrue`` and comes first: a result under the reporting
    limit is not reported, so flagging it for review would send a reviewer after
    a number nobody will publish -- and for the same reason it also outranks the
    duplicate-difference flag of SS VII.4, which sits between grey and the
    review rule in the same amber the review rule uses.

    **Unit deviation from the spec, deliberate.** SS VII.4 writes the amber rule
    as ``$L2 > Parameter!$B$n``. Column ``Relative difference [%]`` holds a
    *fraction* (0.41) shown through a ``0.0%`` format, while the parameter is a
    percentage (30.0), so that comparison can never be true. The rule is written
    with the conversion -- ``$L2*100 > Parameter!$B$n`` -- because a rule that
    never fires would fail SS VII.16's own check ("a pair at 41 % is painted").
    """
    if ws.max_row < 2:
        return
    headers = headers or NIAS_DUPLICATE_HEADERS
    col = {header: _letter(headers, header) for header in headers}
    mean = col["mg/kg (mean)"]
    c1, c2 = col["Concentration 1 [mg/kg]"], col["Concentration 2 [mg/kg]"]
    rel = col["Relative difference [%]"]
    status, id_status = col["Duplicate status"], col["Identification status"]
    review_col = col["Review"]
    data = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"

    grey = FormulaRule(
        formula=[f"AND(COUNT(${mean}2,${c1}2,${c2}2)>0,"
                 f"IF(ISNUMBER(${mean}2),${mean}2,MAX(${c1}2,${c2}2))"
                 f"<{limit:g})"],
        fill=BELOW_LIMIT_FILL, font=BELOW_LIMIT_FONT, stopIfTrue=True)
    grey.priority = 1
    ws.conditional_formatting.add(data, grey)

    # SS VII.4: the limit is a normal parameter, so the rule reads it off the
    # Parameter sheet through _parameter_ref instead of baking the number in --
    # changing 30 to 50 in Excel then repaints the sheet without a re-export.
    # ``reldiff_limit`` is the fallback for a gc_fid that predates the
    # parameter: the number is then written into the formula, which still
    # paints correctly but no longer follows an edit.
    priority = 2
    try:
        limit_term = _parameter_ref("duplicate_max_reldiff")
    except KeyError:                          # pragma: no cover - old gc_fid
        limit_term = None if reldiff_limit is None else f"{reldiff_limit:g}"
    if limit_term is not None:
        amber = FormulaRule(
            formula=[f"AND(ISNUMBER(${rel}2),${rel}2*100>{limit_term})"],
            fill=REVIEW_FILL, font=REVIEW_FONT, stopIfTrue=True)
        amber.priority = priority
        ws.conditional_formatting.add(data, amber)
        priority += 1

    # Both settled statuses may carry a ", also in Blank" suffix, so the test is
    # on the leading characters and not on equality (SS V.6).
    import gc_duplicate

    settled = ",".join(
        f'LEFT(${status}2,{len(s)})<>"{s}"'
        for s in gc_duplicate.SETTLED_STATUSES)
    review = FormulaRule(
        formula=[f'OR(AND({settled}),${id_status}2<>"Accepted",'
                 f'${review_col}2<>"")'],
        fill=REVIEW_FILL, font=REVIEW_FONT, stopIfTrue=True)
    review.priority = priority
    ws.conditional_formatting.add(data, review)


def _integration_label(item: dict[str, Any]) -> str:
    """The integration state of a merged row, or "" when nobody touched it.

    A Doppelbestimmung row stands for up to two determinations, so both sides
    are asked. ``Automatisch`` is never written: the column exists to make a
    *touched* integration visible, and filling it for every untouched row would
    bury exactly the rows it is there to surface.
    """
    labels = []
    for key in ("source1", "source2"):
        source = item.get(key) or {}
        if source.get("integration_touched") or (
                source.get("integration_origin", M.ORIGIN_CHEMSTATION)
                in M.TOUCHED_ORIGINS):
            label = source.get("integration_label") or M.INTEGRATION_LABELS.get(
                source.get("integration_origin"), "")
            if label and label not in labels:
                labels.append(label)
    return " / ".join(labels)


def _write_manual_sheet(ws, combined: list[dict[str, Any]],
                        detail_rows: dict[int, dict[int, int]],
                        report_ri: bool = False, replace_rt: bool = False,
                        seen=None) -> None:
    """Rows needing a decision: conflicts, artefacts, uncertain identifications.

    Since v3.0 a row whose integration was changed by hand is on this sheet too,
    however clean its duplicate status and its identification are (SS VI.14): a
    reviewer must see that somebody moved the bounds, and an accepted, valid
    duplicate would otherwise never appear.

    Inherits both v3.1 columns from the result sheet: ``Times reported before``
    appended last (SS VII.5) and the RI column of SS VII.9, in the same three
    shapes -- a reviewer comparing the two sheets must not have to learn a
    second layout.
    """
    headers = _result_headers(NIAS_MANUAL_HEADERS, report_ri, replace_rt)
    mgkg_letter = _letter(NIAS_DETAIL_HEADERS, "mg/kg")
    col = {header: _letter(headers, header) for header in headers}
    mean_col = col["mg/kg (mean)"]
    c1, c2 = col["Concentration 1 [mg/kg]"], col["Concentration 2 [mg/kg]"]

    ws.append(list(headers))
    # A single determination raises no duplicate question, so it is as settled
    # as a valid duplicate and must not flood this sheet (SS V.6).
    settled = ("Valid duplicate", SINGLE_DETERMINATION_STATUS)
    for item in combined:
        status = str(item.get("status", ""))
        id_status = str(item.get("id_status", ""))
        integration = _integration_label(item)
        if id_status == "Accepted" and status in settled and not integration:
            continue
        r = ws.max_row + 1
        # Linked to the same detail rows the result sheet uses, so a correction
        # made in Excel reaches both sheets instead of only one of them.
        conc = []
        for n in (1, 2):
            source = item.get(f"source{n}") or {}
            detail = detail_rows.get(n, {}).get(source.get("row_id"))
            conc.append(None if not source or detail is None
                        else f"=Bestimmung_{n}!{mgkg_letter}{detail}")
        mean = f"=AVERAGE({c1}{r}:{c2}{r})" if any(conc) else None
        reldiff = (f'=IFERROR(ABS({c1}{r}-{c2}{r})/{mean_col}{r}*100,"")'
                   if all(conc) else None)
        values = {
            "RT (mean)": item.get("rt"), "Name": item.get("name", ""),
            "CAS": item.get("cas", ""), "mg/kg (mean)": mean,
            "Concentration 1 [mg/kg]": conc[0],
            "Concentration 2 [mg/kg]": conc[1],
            "Relative difference [%]": reldiff, "Duplicate status": status,
            "Identification status": id_status,
            "Review": item.get("review", ""), "Integration": integration,
            SEEN_BEFORE_HEADER: _seen_before(item, seen),
            RI_HEADER: _merged_ri(item),
        }
        ws.append([values.get(header) for header in headers])
    _style_header(ws, _widths(headers))
    integration_index = headers.index("Integration")
    formats = {
        "mg/kg (mean)": "0.000000",
        "Concentration 1 [mg/kg]": "0.000000",
        "Concentration 2 [mg/kg]": "0.000000",
        # Percent points here, not a fraction: this column carries the "*100"
        # in its own formula, unlike the result sheet's.
        "Relative difference [%]": "0.0",
        RI_HEADER: "0", SEEN_BEFORE_HEADER: "0",
    }
    index_format = {index: formats[header]
                    for index, header in enumerate(headers)
                    if header in formats}
    for row in ws.iter_rows(min_row=2):
        # The integration column is the reason a settled row is here at all.
        if row[integration_index].value:
            row[integration_index].fill = EDITED_FILL
        for index, fmt in index_format.items():
            row[index].number_format = fmt
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.font = FORMULA_FONT


def _write_parameter_sheet(ws, settings, session: M.Session,
                           blank_path, blank_istd_path) -> None:
    import gc_fid

    labels = session.labels
    meta: dict[str, Any] = {
        "sample_name": Path(session.samples[labels[0]].path).name if labels else "",
        "blank_file": str(blank_path or ""),
        "blank_istd_file": str(blank_istd_path or ""),
    }
    for n, key in enumerate(("delay_1", "delay_2")):
        if n < len(labels):
            meta[key] = session.samples[labels[n]].meta.get("delay", "")

    ws.append(["Parameter", "Wert", "Unit / note"])
    for label, value, unit, _key in gc_fid.parameter_rows(settings, meta):
        ws.append([label, value, unit])
    ws.append(["Hinweis",
               "Blau markierte Werte sind änderbar; Konzentrationen, "
               "Mittelwerte und relative Differenzen rechnen sich automatisch "
               "neu.", ""])
    # Positional sheet: the row order is the contract, so no autofilter.
    _style_header(ws, {1: 34, 2: 46, 3: 52}, autofilter=False)
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].font = Font(bold=True, color="404040")
    # Blue is the reference workbook's convention for "change this". Every one
    # of these cells is referenced absolutely by a formula on another sheet, so
    # editing one requantifies the whole workbook -- and nothing else on this
    # sheet does anything when it is edited.
    selectable = getattr(session, "istd_defs", None) is not None
    for index, (attr, _label, _unit) in enumerate(gc_fid.PARAMETER_LAYOUT):
        if selectable and attr in gc_fid.ISTD_CONCENTRATION_ATTRS:
            # v3.2: the rows keep their position -- ``_parameter_ref`` counts
            # on it -- but the concentrations now live on Standards_n.
            ws.cell(index + 2, 2).value = "→ Standards_n (Spalte G)"
            continue
        if attr:
            ws.cell(index + 2, 2).font = PARAMETER_FONT


def _write_detail_sheet(ws, sample: M.Sample,
                        number: int = 1) -> dict[int, int]:
    """One determination, with its quantification chain left as formulas.

    ``mg/dm2 = area x factor`` and ``mg/kg = mg/dm2 x O/V ratio``, and the
    factor is the one cell on ``Standards_n`` that the three quantification
    standards average into. Writing the three as formulas is what lets an
    analyst correct an area, an ISTD area or the O/V ratio in Excel and see
    every dependent number follow -- here, on ``Doppelbestimmung`` and on
    ``Manuell_pruefen`` (SS VI.22).

    Returns ``{row_id: sheet row}``, which is how the result sheets address
    these rows.
    """
    area_letter = _letter(NIAS_DETAIL_HEADERS, "Blank-corrected FID area")
    dm2_letter = _letter(NIAS_DETAIL_HEADERS, "mg/dm2")
    factor_letter = _letter(NIAS_DETAIL_HEADERS, "Quantification factor")
    ov_ref = _parameter_ref("ov_ratio")
    factor_ref = f"=Standards_{number}!$H${_standards_mean_row(sample)}"

    rows: dict[int, int] = {}
    ws.append(list(NIAS_DETAIL_HEADERS))
    for row in sample.rows:
        if row.integration_origin == M.ORIGIN_DISABLED:
            continue
        d = row.derived
        r = ws.max_row + 1
        rows[row.row_id] = r
        ws.append([
            row.rt, row.name, M.display_cas(row.cas), row.si, row.area,
            f"={area_letter}{r}*{factor_letter}{r}",
            f"={dm2_letter}{r}*{ov_ref}", d.get("id_status", ""),
            d.get("match_status", ""), d.get("review", ""), d.get("fid_peak"),
            d.get("pbm_peak"), d.get("ms_rt"), d.get("corrected_ms_rt"),
            d.get("raw_area"), d.get("blank_area"), d.get("blank_istd_area"),
            d.get("blank_status", ""), d.get("protected_standard", ""),
            factor_ref,
            # v3.0, appended so no existing column changes its letter:
            row.fid_start, row.fid_end, row.integration_label, row.sn,
            d.get("loq_mgkg"), row.purity, row.ri,
        ])
    _style_header(ws, {1: 11, 2: 44, 3: 20, 4: 9, 5: 22, 6: 14, 7: 12, 8: 20,
                       9: 14, 10: 40, 11: 11, 12: 11, 13: 13, 14: 18, 15: 16,
                       16: 13, 17: 17, 18: 30, 19: 15, 20: 20,
                       21: 14, 22: 14, 23: 16, 24: 9, 25: 14, 26: 14, 27: 9})
    for r in ws.iter_rows(min_row=2):
        r[0].number_format = "0.000"
        for index in (4, 14, 15, 16):
            r[index].number_format = "#,##0"
        r[5].number_format = "0.000000"
        r[6].number_format = "0.0000"
        for index in (12, 13):
            r[index].number_format = "0.0000"
        for index in (20, 21):
            r[index].number_format = "0.000"
        r[23].number_format = "0.0"
        r[24].number_format = "0.0000"
        for index in (25, 26):
            r[index].number_format = "0"
        # An integration nobody touched reads "Automatisch"; the ones that were
        # touched are what the reviewer is looking for, so they are marked.
        if r[22].value and r[22].value != M.INTEGRATION_LABELS[M.ORIGIN_CHEMSTATION]:
            r[22].fill = EDITED_FILL
        for cell in r:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.font = FORMULA_FONT
    return rows


def _standards_mean_row(sample: M.Sample) -> int:
    """Row of ``Standards_n`` that carries the mean quantification factor.

    Header, one row per standard, one blank row, then the mean -- the layout
    :func:`_write_standards_sheet` writes and the one AutoLib writes. Computed
    rather than remembered, because ``Bestimmung_n`` has to reference the cell
    before that sheet exists.
    """
    return len(sample.meta.get("standards") or []) + 3


def _write_standards_sheet(ws, sample: M.Sample) -> None:
    """The quantification standards and the factor they average into.

    ``factor_i = c_i * ISTD amount / 1000 / area_i / cell area / coverage``
    (SS 11.3), written as a formula against the ``Parameter`` cells, so that
    correcting a standard's area or a concentration in Excel requantifies the
    whole determination -- which is what it already does in the workspace.

    Only the quantification standards get one. The QC standard has no
    concentration and never gets a factor (SS 11.2), and ``conc_attr`` is
    exactly the field that distinguishes the two.

    With the selectable ISTDs (v3.2) the concentrations are typed input cells
    and the mean row follows the workspace's mode: ``mean c x amount / 1000 /
    mean area / cell / coverage`` over the quantifying ISTDs, or the reference
    ISTD's own factor.
    """
    import gc_fid

    is_amount = _parameter_ref("is_amount")
    cell_area = _parameter_ref("cell_area_dm2")
    coverage = _parameter_ref("coverage")
    mean_row = _standards_mean_row(sample)

    ws.append(list(NIAS_STANDARD_HEADERS))
    standards = sample.meta.get("standards") or []
    selectable = getattr(sample, "istd_defs", None) is not None
    options = getattr(sample, "istd_options", None) or {}
    quantification: list[int] = []
    codes: dict[int, str] = {}
    inputs: list[int] = []
    for std in standards:
        r = ws.max_row + 1
        concentration, factor, deviation = (
            std.get("concentration"), std.get("factor"), std.get("deviation"))
        attr = std.get("conc_attr")
        if attr:
            concentration = f"={_parameter_ref(attr)}"
            if std.get("fid_area"):
                factor = f"=G{r}*{is_amount}/1000/F{r}/{cell_area}/{coverage}"
                deviation = f'=IFERROR((H{r}/$H${mean_row}-1)*100,"")'
                quantification.append(r)
        elif selectable:
            # v3.2: the concentration is the analyst's typed number, written
            # as a blue input cell so correcting it in Excel requantifies.
            inputs.append(r)
            if (std.get("role") == gc_fid.ROLE_QUANTIFICATION
                    and std.get("fid_area") and _num(concentration)):
                factor = f"=G{r}*{is_amount}/1000/F{r}/{cell_area}/{coverage}"
                deviation = f'=IFERROR((H{r}/$H${mean_row}-1)*100,"")'
                quantification.append(r)
                codes[r] = str(std.get("code") or "")
        ws.append([std.get("name", ""), std.get("target_rt"), std.get("role", ""),
                   std.get("fid_peak"), std.get("fid_rt"), std.get("fid_area"),
                   concentration, factor, deviation, std.get("status", ""),
                   "Ja" if (std.get("edited") or std.get("manual_area")) else "",
                   std.get("code") or ""])
    ws.append([])
    label = "Mean quantification factor"
    if selectable and quantification:
        if options.get("use_mean_area", True):
            g = ",".join(f"G{r}" for r in quantification)
            f = ",".join(f"F{r}" for r in quantification)
            mean = (f"=AVERAGE({g})*{is_amount}/1000/AVERAGE({f})"
                    f"/{cell_area}/{coverage}")
            label = "Quantification factor (mean ISTD area)"
        else:
            ref = gc_fid.reference_code(
                [{"code": c, "quantify": True} for c in codes.values()], options)
            ref_row = next(r for r, c in codes.items() if c == ref)
            mean = f"=H{ref_row}"
            label = f"Quantification factor (reference {ref})"
    elif quantification and not selectable:
        mean = "=AVERAGE(" + ",".join(f"H{r}" for r in quantification) + ")"
    else:
        mean = sample.meta.get("mean_factor")
    ws.append(["", "", "", "", "", "", label, mean, "", "", "", ""])
    _style_header(ws, {1: 32, 2: 16, 3: 16, 4: 11, 5: 14, 6: 15, 7: 22, 8: 16,
                       9: 22, 10: 18, 11: 13, 12: 8}, autofilter=False)
    for r in ws.iter_rows(min_row=2, max_row=1 + len(standards)):
        r[1].number_format = "0.000"
        r[4].number_format = "0.000"
        r[5].number_format = "#,##0"
        r[7].number_format = "0.000000E+00"
        r[8].number_format = "+0.00;-0.00"
        # A standard drifting far from the mean makes every result suspect.
        # Reachable only while the deviation is still a number: once the column
        # is a formula, openpyxl cannot evaluate it, and the mean factor's own
        # ">20 %" guard already runs in the workspace before the export.
        if isinstance(r[8].value, (int, float)) and abs(r[8].value) > 20:
            r[8].fill = PatternFill("solid", fgColor="FCE4D6")
        if str(r[9].value) not in ("Found", ""):
            r[9].fill = PatternFill("solid", fgColor="F4CCCC")
        # A hand-set standard area rewrites the whole determination's
        # quantification factor, so the sheet has to show it (spec v2.1 SS V.1).
        if r[10].value:
            r[5].fill = EDITED_FILL
        if r[0].row in inputs:
            r[6].font = PARAMETER_FONT
        for cell in r:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.font = FORMULA_FONT
    mean_cell = ws.cell(_standards_mean_row(sample), 8)
    mean_cell.number_format = "0.000000E+00"
    if isinstance(mean_cell.value, str) and mean_cell.value.startswith("="):
        mean_cell.font = FORMULA_FONT


def _pbm_entries(sample: M.Sample) -> list[dict[str, Any]]:
    """``{peak, rt, area_pct, hits}`` per PBM peak of one determination.

    The determination's own table is used when it is there, because it is the
    complete library report: a hit the FID-tuned integrator merged into its
    neighbour never reaches a row. The per-row fallback covers a sample that was
    not loaded through :func:`gc_fid.load_determination` -- fewer peaks, but the
    hits are still the ones the row was identified from.
    """
    table = sample.meta.get("pbm_hits")
    if table:
        return list(table)

    seen: dict[Any, dict[str, Any]] = {}
    for row in sample.rows:
        peak = row.derived.get("pbm_peak")
        if peak is None or peak in seen:
            continue
        seen[peak] = {
            "peak": peak,
            "rt": row.derived.get("ms_rt"),
            "area_pct": row.pbm_area_pct,
            "hits": row.library_hits,
        }
    return [seen[k] for k in sorted(seen)]


def _write_pbm_sheet(ws, sample: M.Sample) -> None:
    """The library report behind one determination: every PBM peak, all hits.

    Identification uses only hit 1 against the quality threshold, with hits 2
    and 3 deciding class inference. Both are invisible in ``Bestimmung_n``,
    which shows the resolved name alone -- this sheet is what an analyst checks
    a weak or class-inferred identification against.
    """
    ws.append(list(NIAS_PBM_HEADERS))
    for entry in _pbm_entries(sample):
        hits = list(entry.get("hits") or [])
        values: list[Any] = [entry.get("peak"), entry.get("rt"),
                             entry.get("area_pct")]
        for index in range(3):
            name, cas, qual = hits[index] if index < len(hits) else ("", "", None)
            values += [name, M.display_cas(cas) or "", qual]
        ws.append(values)

    _style_header(ws, {1: 11, 2: 13, 3: 10, 4: 44, 5: 20, 6: 7, 7: 44, 8: 20,
                       9: 7, 10: 44, 11: 20, 12: 7})
    for r in ws.iter_rows(min_row=2):
        r[1].number_format = "0.0000"
        r[2].number_format = "0.0000"


# --------------------------------------------------------------------------
# The Integration sheet -- a projection of the audit trail (SS VI.14)
# --------------------------------------------------------------------------
#
# Deliberately not a second set of books. Every row on this sheet comes from an
# ``EditRecord`` the session already holds; nothing is recorded here that is not
# in ``Manuelle Änderungen`` too. A command that creates, splits, merges or
# disables a row therefore has to append its own record -- ``gc_model``'s
# ``integration_record()`` is the one way in -- and the sheet then has it
# without the exporter needing to know what a command is.


def current_user() -> str:
    """The ``Benutzer`` column: who was logged in when the workbook was written.

    The audit trail records *when* and *what*, never *who* -- the workspace is a
    single-user desktop tool and has no accounts. The OS login is the honest
    answer available; it identifies the workstation session that produced the
    report, which is what a GLP reviewer needs to follow up.
    """
    import getpass
    try:
        return getpass.getuser()
    except Exception:
        return ""


def _group_audit(session: M.Session) -> list[list[M.EditRecord]]:
    """Consecutive records of one row at one instant -- i.e. one action.

    A command moves ``fid_start``, ``fid_end`` and ``raw_area`` in one gesture
    and writes a record for each. Grouping by (timestamp, determination, row)
    turns those back into the single action the analyst performed, which is what
    "one row per action" means.
    """
    groups: list[list[M.EditRecord]] = []
    key = None
    for record in session.audit:
        current = (record.timestamp, record.sample, record.row_id)
        if current != key:
            groups.append([])
            key = current
        groups[-1].append(record)
    return groups


def _action_label(group: list[M.EditRecord], row: Optional[M.PeakRow]) -> str:
    """What the analyst did, in the words the workspace uses for it."""
    fields = {r.field for r in group}
    for record in group:
        if record.field == "integration_origin":
            return M.INTEGRATION_LABELS.get(str(record.new), str(record.new))
    if fields & {"fid_start", "fid_end"}:
        return "Grenzen geändert"
    if "fid_baseline" in fields:
        return "Basislinie geändert"
    return (row.integration_label if row is not None else "")


def _bounds_text(start: Any, end: Any) -> str:
    if start is None and end is None:
        return ""
    return f"{_minutes(start)} – {_minutes(end)}"


def _minutes(value: Any) -> str:
    return "" if value is None else f"{float(value):.3f}"


def _write_integration_sheet(ws, session: M.Session) -> int:
    """One row per integration action. Returns how many were written."""
    ws.append(list(NIAS_INTEGRATION_HEADERS))
    user = current_user()
    written = 0

    for group in _group_audit(session):
        if not {r.field for r in group} & M.INTEGRATION_AUDIT_FIELDS:
            continue                       # an ordinary cell edit, not this sheet
        first = group[0]
        sample = session.samples.get(first.sample)
        row = sample.row(first.row_id) if sample is not None else None

        # "Old" is the value before the action, "new" the value after it. A
        # field the action did not touch has one value, not two, and the row's
        # current one is that value.
        old: dict[str, Any] = {}
        new: dict[str, Any] = {}
        for record in group:
            old.setdefault(record.field, record.old)
            new[record.field] = record.new
        for field in ("fid_start", "fid_end", "area"):
            if field not in old and row is not None:
                old[field] = new[field] = getattr(row, field, None)

        area_new = new.get("area", old.get("area"))
        ws.append([
            first.timestamp,
            first.sample,
            first.peak_no,
            _action_label(group, row),
            _bounds_text(old.get("fid_start"), old.get("fid_end")),
            _bounds_text(new.get("fid_start"), new.get("fid_end")),
            old.get("area"),
            area_new,
            _concentration(sample, old.get("area")),
            _concentration(sample, area_new),
            (new.get("fid_baseline")
             or (row.fid_baseline if row is not None else "")),
            user,
        ])
        written += 1

    _style_header(ws, {1: 20, 2: 16, 3: 8, 4: 22, 5: 20, 6: 20, 7: 16, 8: 16,
                       9: 14, 10: 14, 11: 14, 12: 18})
    for r in ws.iter_rows(min_row=2):
        for index in (6, 7):
            r[index].number_format = "#,##0"
        for index in (8, 9):
            r[index].number_format = "0.0000"
    return written


def _concentration(sample: Optional[M.Sample], area: Any) -> Optional[float]:
    """``mg/kg`` for an area through the determination's own factor chain.

    The audit trail records areas, not concentrations -- a concentration is
    derived and was never edited. Reconstructing it here rather than storing it
    keeps the sheet a projection. The factor used is the determination's
    *current* one, so if the action moved a standard it moved this number too;
    that is the honest reading, because the old concentration no longer exists.
    """
    if area is None or sample is None:
        return None
    factor = getattr(sample, "mean_factor", None)
    if not factor:
        return None
    ov_ratio = getattr(getattr(sample, "settings", None), "ov_ratio", 1.0) or 1.0
    return float(area) * factor * ov_ratio


def _write_blank_sheet(ws, samples) -> None:
    ws.append(list(NIAS_BLANK_HEADERS))
    for n, sample in enumerate(samples, 1):
        for row in sample.rows:
            d = row.derived
            ws.append([n, d.get("fid_peak"), row.rt, row.name, d.get("raw_area"),
                       d.get("blank_area"), d.get("blank_istd_area"), row.area,
                       d.get("blank_status", ""), d.get("protected_standard", ""),
                       "Ja" if d.get("subtracted_area") else "Nein"])
    _style_header(ws, {1: 13, 2: 11, 3: 11, 4: 44, 5: 16, 6: 13, 7: 17, 8: 20,
                       9: 32, 10: 15, 11: 12})
    for r in ws.iter_rows(min_row=2):
        r[2].number_format = "0.000"
        for index in (4, 5, 6, 7):
            r[index].number_format = "#,##0"


# --------------------------------------------------------------------------
# Fingerprint / Total extraction (SS VII.10)
# --------------------------------------------------------------------------
#
# The report chain of SS VII.10 is: session -> intermediate workbook -> report.
# This is the first arrow for the two fingerprint-shaped reports. The
# intermediate is deliberately *the shape the main script already reads*:
# ``is_fingerprint_workbook`` matches on normalised header names in row 1 of a
# sheet named ``Fingerprint``, so a header that is one character off routes the
# workbook silently down the Doppelbestimmung branch and produces a wrong
# report rather than an error. The headers below are therefore written once,
# here, and asserted against the real predicate by the v3.1 test.

FINGERPRINT_SHEET = "Fingerprint"

#: Exactly what ``is_fingerprint_workbook`` requires, in AutoLib's own order.
#: ``normalized_header`` reduces these to ``rtmin``, ``name``, ``cas``,
#: ``qualitymatch``, ``pbmarea``, ``imblank`` -- the required set.
FINGERPRINT_HEADERS = (
    "RT (min)", "Name", "CAS", "Quality match", "PBM Area %", "Im Blank",
)

#: The one extra column of a total-extraction workbook (SS VII.10). Appended,
#: so a total-extraction workbook is still a valid fingerprint workbook and
#: ``is_fingerprint_workbook`` keeps recognising it.
TOTAL_EXTRACTION_HEADER = "c [µg/L]"

#: ``kind`` values :func:`write_fingerprint_workbook` accepts, mapped to the
#: ``report_type`` string SS VII.5 stores in the register.
FINGERPRINT_KINDS = {
    "fingerprint": "Fingerprint",
    "total_extraction": "Total extraction",
}

#: Suffix AutoLib's ``combine_determinations`` appends to the duplicate status
#: of a peak whose area was reduced by a blank subtraction. The ``Im Blank``
#: column is read off this rather than recomputed, because the status is what
#: the analyst reviewed and signed off (SS VII.10).
BLANK_SUFFIX = ", also in Blank"

#: Column widths of the Fingerprint sheet, by header rather than by index --
#: same reason as :data:`_RESULT_WIDTHS`: once SS VII.9 may insert an ``RI``
#: column, an index-keyed map silently puts every width one column to the left.
#: The six fingerprint widths are the ones the pre-v3.1 writer used, unchanged.
_FINGERPRINT_WIDTHS = {
    "RT (min)": 11, RI_HEADER: 9, "Name": 44, "CAS": 20, "Quality match": 13,
    "PBM Area %": 13, "Im Blank": 10, TOTAL_EXTRACTION_HEADER: 13,
}

#: Number formats of the Fingerprint sheet, likewise by header. A header that
#: is not listed keeps openpyxl's ``General`` -- which is what ``Name``, ``CAS``
#: and ``Im Blank`` had before v3.1 and still have.
#:
#: ``RI`` is ``"0"`` and the value written is an ``int``: a Kovats index is
#: reported without decimals, and a cell that merely *looks* integral but sorts
#: as 1043.5 is a trap (same rule as the two result sheets).
_FINGERPRINT_FORMATS = {
    "RT (min)": "0.0000", RI_HEADER: "0", "Quality match": "0",
    "PBM Area %": "0.0000", TOTAL_EXTRACTION_HEADER: "0.00",
}


def fingerprint_headers(kind: str, report_ri: bool = False,
                        replace_rt: bool = False) -> tuple[str, ...]:
    """The header row for ``kind``. Raises ``ValueError`` on an unknown kind.

    SS VII.9 reaches this sheet too (analyst's decision on the open point of the
    v3.1 integration report). The three cases are exactly the ones
    :func:`_result_headers` implements for the two result sheets:

    * ``report_ri = False`` -- the pre-v3.1 tuple, byte for byte. Both flags
      default to ``False``, so every existing caller keeps that tuple;
    * ``report_ri`` alone -- ``RI`` directly right of ``RT (min)``;
    * ``replace_rt`` -- the first column keeps its position and is headed
      ``RI``. The Fingerprint sheet has no detail sheets, so this is the one
      place where turning the option on **loses** the retention time from the
      report. That is what ``RT durch RI ersetzen`` means here.

    ``c [µg/L]`` is appended after the RI column in every case, so a
    total-extraction workbook keeps its concentration column last (SS VII.10)
    and ``is_total_extraction_workbook`` keeps discriminating on it.
    """
    if kind not in FINGERPRINT_KINDS:
        raise ValueError(
            f"Unbekannte Berichtsart: {kind!r}. Erlaubt sind "
            f"{' und '.join(sorted(FINGERPRINT_KINDS))}.")
    report_ri = bool(report_ri)
    base = _result_headers(FINGERPRINT_HEADERS, report_ri,
                           bool(replace_rt) and report_ri)
    if kind == "total_extraction":
        return base + (TOTAL_EXTRACTION_HEADER,)
    return base


def istd_concentration_ugl(session: Optional[M.Session] = None,
                           settings: Any = None) -> float:
    """The ISTD concentration in the extract, in µg/L (SS VII.10).

    Three sources, in order:

    1. an explicit ``istd_conc_ugl`` on the settings or on the session -- what
       the analyst typed into ``ISTD-Konzentration im Extrakt``. A hand-entered
       number always wins, so a measured concentration can be used directly;
    2. the **derivation** from the Parameter sheet,
       :func:`gc_fid.derived_istd_concentration_ugl`: the three quantification
       standards' concentrations, the ISTD amount and the extract volume give
       ``mean(fc17, bbp, dnnp) [mg/mL] * is_amount [µL] * 1000 /
       extract_volume_ml [mL]``, i.e. **823.33 µg/L** with the defaults;
    3. ``gc_model.IS_CONCENTRATION_UG_L`` only when neither holder carries the
       five settings the derivation needs -- a bare ``Session`` with no
       settings anywhere, which is the DIN SPEC shape.

    This used to stop at 3, with the DIN SPEC constant as the
    unconditional default and the whole number marked **[Annahme] 1 of
    SS VII.15**. It no longer is: the analyst has stated that the NIAS
    total-extraction extract is 10 mL, spiked with the NIAS ISTD, and 10 mL is
    what ``extract_volume_ml`` carries. The two paths therefore quantify with
    two different numbers on purpose -- DIN SPEC still reads
    ``IS_CONCENTRATION_UG_L`` (``gc_model.Sample._recalculate_derived``), NIAS
    derives its own, 12.35x lower. Changing one no longer changes the other,
    which is the point.

    ``ov_ratio`` is not part of this and must not become part of it: O/V only
    applies to a report with an SML, and Total Extraction has none.
    """
    import gc_fid

    for holder in (settings, session):
        if holder is None:
            continue
        for attr in ("istd_conc_ugl", "istd_concentration_ugl"):
            value = _num(getattr(holder, attr, None))
            if value is not None:
                return value
    # ``session.settings`` is tried too: a caller that passes only the session
    # still gets the derivation if the session knows its settings, rather than
    # silently dropping to the DIN SPEC constant.
    # v3.2: the typed concentrations of the session's ISTD list, when it has one.
    defs = getattr(session, "istd_defs", None)
    options = getattr(session, "istd_options", None)
    for holder in (settings, session, getattr(session, "settings", None)):
        if holder is None:
            continue
        derived = gc_fid.derived_istd_concentration_ugl(holder, defs, options)
        if derived is not None:
            return float(derived)
    return float(M.IS_CONCENTRATION_UG_L)


def _quantification_istd_area(sample: M.Sample) -> Optional[float]:
    """Mean FID area of the three quantification ISTDs of one determination.

    ``conc_attr`` is exactly the field that separates the three quantification
    standards from the QC standard -- the same test ``_write_standards_sheet``
    uses for the factor -- so the mean is over the same three areas the
    quantification factor already averages.
    """
    if getattr(sample, "istd_defs", None) is not None:
        # v3.2: the same area the factor is built from -- the mean over the
        # quantifying ISTDs with a concentration, or the reference ISTD's.
        return _num(sample.meta.get("istd_mean_area"))
    areas = [
        _num(std.get("fid_area"))
        for std in (sample.meta.get("standards") or [])
        if std.get("conc_attr")
    ]
    areas = [a for a in areas if a]
    if not areas:
        return None
    return sum(areas) / len(areas)


def _sample_name(session: M.Session) -> str:
    """The name the report's title band shows, from the first determination."""
    labels = session.labels
    if not labels:
        return ""
    sample = session.samples[labels[0]]
    return str(sample.meta.get("sample") or Path(sample.path).stem)


def _merged_si(item: dict[str, Any]) -> Optional[int]:
    """``Quality match`` of a merged row: the merged ``si``, as an integer.

    A quality match is a whole-number percentage in every sheet this workbook
    talks to, so the mean of the two determinations is rounded rather than
    carried at full precision into a column formatted ``0``.
    """
    value = _merged_number(item, "quality")
    return None if value is None else int(round(value))


def _pbm_area_reader(sample: M.Sample):
    """A ``source dict -> Optional[float]`` reader for ``PBM Area %``.

    The percentage the fingerprint report is built on is the *library search's*
    own ``Area %``, and on a determination loaded through
    :func:`gc_fid.load_determination` it lives on the PBM table
    (``sample.meta["pbm_hits"]``, keyed by PBM peak number), **not** on the
    ``PeakRow``: ``PeakRow.pbm_area_pct`` is only filled on the DIN SPEC path,
    where the rows *are* the MS peaks. Reading the row alone therefore left the
    column empty on every real NIAS determination, which is why the lookup goes
    through :func:`_pbm_entries` -- the same table ``PBM_n`` is written from,
    so the report and the sheet quote one number.

    Two fallbacks, in order: the row's own ``pbm_area_pct`` (DIN SPEC rows and
    samples not built by ``load_determination``), then the FID area percentage,
    because a semi-quantitative column carrying the FID's own percentage is
    still a percentage of the same chromatogram -- and the report renormalises.
    """
    table = {entry.get("peak"): _num(entry.get("area_pct"))
             for entry in _pbm_entries(sample)}
    by_id = {row.row_id: row for row in sample.rows}

    def read(source: dict[str, Any]) -> Optional[float]:
        peak = source.get("pbm_peak")
        if peak is not None:
            value = table.get(peak)
            if value is not None:
                return value
        row = by_id.get(source.get("row_id"))
        if row is None:
            return None
        value = _num(row.pbm_area_pct)
        return _num(row.area_pct) if value is None else value

    return read


def _merged_pbm_area(item: dict[str, Any], readers: dict[int, Any]
                     ) -> Optional[float]:
    """``PBM Area %`` of a merged row, meaned over the determinations present."""
    values: list[float] = []
    for n, source in enumerate(_sources(item), 1):  # GCWS-PATCH: N-fold
        source = source or {}
        if not source:
            continue
        reader = readers.get(n)
        value = None if reader is None else reader(source)
        if value is not None:
            values.append(value)
    if not values:
        return None
    return sum(values) / len(values)


def _extract_concentration(item: dict[str, Any],
                           istd_area: dict[int, Optional[float]],
                           conc_ugl: float) -> Optional[float]:
    """``c [µg/L]`` of a merged row (SS VII.10, assumption 1 of SS VII.15).

    Per determination ``area_corr / mean(ISTD areas) * istd_conc``, then the
    mean of the sides that produced a number. ``area_corr`` is the
    blank-corrected FID area -- the same ``area`` the ``Area n`` column of the
    Doppelbestimmung sheet links to, so the two reports quantify off one number.
    """
    values: list[float] = []
    for n, source in enumerate(_sources(item), 1):  # GCWS-PATCH: N-fold
        source = source or {}
        area = _num(source.get("area"))
        reference = istd_area.get(n)
        if area is None or not reference:
            continue
        values.append(area / reference * conc_ugl)
    if not values:
        return None
    return sum(values) / len(values)


def write_fingerprint_workbook(session: M.Session, path: Path, kind: str,
                               settings, *,
                               ri_options: Optional[dict] = None,
                               combined: Optional[list] = None
                               ) -> dict[str, Any]:
    """The intermediate workbook of the fingerprint and total-extraction reports.

    One sheet, ``Fingerprint``, with exactly the headers
    ``is_fingerprint_workbook`` requires, plus ``RI`` when SS VII.9 asks for it
    and ``c [µg/L]`` for ``total_extraction``. The rows are the merged rows of
    SS VI.22 that reach
    the report: a mean is present and it is at or above the reporting limit.
    An artefact -- one determination only, no mean -- is not a result and does
    not appear.

    Returns the audit payload. ``payload["rows"]`` is the list of rows that
    actually reached the sheet, which is what SS VII.10 step 6 hands to
    ``gc_seen.record``: a substance the writer dropped was never reported, and
    counting it would make ``Schon berichtet`` lie.

    µg/L for ``total_extraction``::

        c = area_corr / mean(area of the three quantification ISTDs) * istd_conc

    computed per determination and then averaged, so each side is divided by
    *its own* ISTD level -- which is the whole point of an internal standard.
    No renormalisation to 100 %: a concentration column that sums to a measured
    total is the point (SS VII.15, assumption 7).

    ``ri_options`` is the explicit route for the ``{"report_ri", "replace_rt"}``
    dict of SS VII.9 -- the same keyword :func:`run_nias_duplicate` takes, and it
    outranks the copy the workspace parks on the session. Omitted, the flags are
    read off the settings and then off the session, so a caller written before
    this change exports exactly what it exported before: with ``report_ri`` off
    the sheet is unchanged, cell for cell, including its widths and number
    formats.
    """
    report_ri, replace_rt = _ri_options(session, settings, ri_options)
    headers = fingerprint_headers(kind, report_ri, replace_rt)
    labels = session.labels
    if not labels:
        raise ValueError("Es ist keine Bestimmung geladen.")

    import gc_fid

    # GCWS-PATCH: N-fold replicates with precomputed combined rows
    samples = [session.samples[label] for label in (labels if combined is not None else labels[:2])]
    single = len(samples) == 1
    if combined is None:
        combined = (_single_determination_rows(samples[0]) if single
                    else gc_fid.combine(samples[0], samples[1]))

    reporting_limit = float(getattr(settings, "reporting_limit", 0.01) or 0.01)
    conc_ugl = istd_concentration_ugl(session, settings)
    # {determination number: reader} -- the PBM area percentage is a library
    # number that never travelled into AutoLib's merged dict, so it is read back
    # off the determination's own PBM table (see _pbm_area_reader).
    pbm_readers = {n: _pbm_area_reader(sample)
                   for n, sample in enumerate(samples, 1)}
    istd_area = {n: _quantification_istd_area(sample)
                 for n, sample in enumerate(samples, 1)}

    rows: list[dict[str, Any]] = []
    for item in combined:
        mean = _num(item.get("mean"))
        if mean is None or mean < reporting_limit:
            continue
        row: dict[str, Any] = {
            "rt": _num(item.get("rt")),
            "name": item.get("name") or "",
            "cas": item.get("cas") or "",
            "match": _merged_si(item),
            "area_pct": _merged_pbm_area(item, pbm_readers),
            "blank": ("Ja" if BLANK_SUFFIX in str(item.get("status", ""))
                      else "Nein"),
            "mean_mgkg": mean,
            "status": item.get("status", ""),
            "id_status": item.get("id_status", ""),
        }
        if report_ri:
            # Only when the column exists: with SS VII.9 off, payload["rows"]
            # is the mapping it has always been, key for key, because SS VII.10
            # step 6 hands these rows straight to ``gc_seen.record``.
            row["ri"] = _merged_ri(item)
        if kind == "total_extraction":
            row["conc_ugl"] = _extract_concentration(item, istd_area, conc_ugl)
        rows.append(row)

    wb = Workbook()
    # The report's title band reads the sample name off the workbook title when
    # there is no Parameter sheet, stripping exactly this suffix -- so it is
    # spelled the way AutoLib's own fingerprint workbook spells it.
    sample_name = _sample_name(session)
    wb.properties.title = f"{sample_name} - Fingerprint Screening"
    ws = wb.active
    ws.title = FINGERPRINT_SHEET
    ws.append(list(headers))
    # One row builder, driven by the header tuple, so that inserting RI cannot
    # put a value in the wrong column. ``RI`` is written as an ``int`` or not at
    # all: ``_merged_ri`` returns ``None`` when neither determination has an
    # index, and ``None`` is an empty cell -- never a zero, which would read as
    # a measured index of 0.
    by_header = {
        "RT (min)": "rt", "Name": "name", "CAS": "cas",
        "Quality match": "match", "PBM Area %": "area_pct",
        "Im Blank": "blank", TOTAL_EXTRACTION_HEADER: "conc_ugl",
        RI_HEADER: "ri",
    }
    for row in rows:
        ws.append([row.get(by_header[header]) for header in headers])
    _style_header(ws, _widths(headers, _FINGERPRINT_WIDTHS))
    # Number formats by header for the same reason as the widths. With
    # ``report_ri`` off this sets exactly the four columns the pre-v3.1 writer
    # set, in the same places, and leaves the other three ``General``.
    formats = {index: _FINGERPRINT_FORMATS[header]
               for index, header in enumerate(headers, 1)
               if header in _FINGERPRINT_FORMATS}
    for r in ws.iter_rows(min_row=2):
        for index, number_format in formats.items():
            r[index - 1].number_format = number_format

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)

    fields: dict[str, Any] = {
        "Source format": FINGERPRINT_KINDS[kind],
        "Sample": sample_name,
        "Output file": str(path.resolve()),
        "Reported rows": len(rows),
        "Reporting limit": reporting_limit,
        "Determinations": len(samples),
    }
    fields.update(duplicate_audit_fields(combined, settings))
    if kind == "total_extraction":
        fields["Quantity basis"] = "ISTD; concentration in extract"
        fields["ISTD concentration [ug/L]"] = conc_ugl
        total = [row["conc_ugl"] for row in rows if row["conc_ugl"] is not None]
        fields["Sum c [ug/L]"] = round(sum(total), 4) if total else None
    else:
        fields["Quantity basis"] = "PBM Area %; semi-quantitative"

    return {
        "path": str(path),
        "kind": kind,
        "report_type": FINGERPRINT_KINDS[kind],
        "sheet": FINGERPRINT_SHEET,
        "headers": list(headers),
        "sample": sample_name,
        "rows": rows,
        "fields": fields,
    }
